"""Revision delivery consumers; offline Office doubles are not calculation evidence."""
from copy import deepcopy
import hashlib
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree as ET
import zipfile

import openpyxl
import pytest

from ai_sow_lite import cli, office, workbook
from ai_sow_lite.contracts import canonical_json_bytes, file_sha256
from .support.clarify import (clarify_case, delivered_baseline, check_edits,
                              edit_draft, complexity_draft, confirm_plan)
from .support.cli import run_request
from .support.fixtures import read_json, write_json


@pytest.fixture
def offline_office(monkeypatch):
    """Only replace the external engine: empty string caches, no amount calculation.

    The actual projector, seal, prepared verifier and CLI consumers all run.
    This fixture cannot supply evidence about real Office formula evaluation.
    """
    calls = []

    def recalculate(source, destination):
        calls.append(source)
        raw = destination.with_name('sow.office-raw.xlsx')
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with zipfile.ZipFile(source) as src, zipfile.ZipFile(raw, 'w') as dst:
            for info in src.infolist():
                content = src.read(info.filename)
                if info.filename.startswith('xl/worksheets/') and info.filename.endswith('.xml'):
                    root = ET.fromstring(content)
                    for cell in root.iter(ns + 'c'):
                        if cell.find(ns + 'f') is not None:
                            cell.set('t', 'str')
                            value = cell.find(ns + 'v')
                            if value is None:
                                value = ET.SubElement(cell, ns + 'v')
                            value.text = None
                    content = ET.tostring(root, encoding='utf-8')
                dst.writestr(info, content)
        report = workbook.seal_office_output(source, raw, destination)
        engine = dict(name='LibreOffice', version='LibreOffice offline-test-double',
                      binary_sha256='0' * 64, platform='Darwin', arguments=office.ARGUMENTS)
        return dict(office_identity='sha256:' + hashlib.sha256(canonical_json_bytes(engine)).hexdigest(),
                    engine=engine, input_hash=file_sha256(source), raw_hash=file_sha256(raw),
                    final_hash=file_sha256(destination), exit_code=0, elapsed_ms=0, verification=report)

    monkeypatch.setattr(office, 'recalculate', recalculate)
    return calls


def invoke(case, operation, payload):
    return cli.execute(dict(protocol_version='1.0', request_id=case['request_id'],
                            project_path=str(case['project']), operation=operation, payload=payload))


def baseline(case):
    return case['project'] / '.ai-sow-lite/versions' / case['current']['version_id']


def render(case, result, *, real=False):
    payload = dict(candidate_path=result['candidate_ref']['path'],
                   check_path=result['check_ref']['path'], expected_current=case['current'])
    reply = (run_request(case['project'], case['request_id'], 'render', payload) if real
             else invoke(case, 'render', payload))
    assert reply['ok'], reply
    output = reply['result']
    prepared = read_json(case['project'] / output['prepared_ref']['path'])
    assert workbook.verify_prepared(case['project'], prepared) == dict(diagnostics=[])
    return output


def collision_draft(case, kind):
    old = read_json(baseline(case) / 'model.json')
    # These valid UUIDs sort before the baseline's Story 19 and Task 31.
    story = dict(deepcopy(old['stories'][0]), id='00000000-0000-4000-8000-000000000001')
    story['acs'] = [dict(ac, id=str(uuid4())) for ac in story['acs']]
    task = dict(deepcopy(old['tasks'][0]), id='00000000-0000-4000-8000-000000000002')
    if kind == 'story':
        task['story_id'] = story['id']
        task['name'] = '新增故事的独立查询任务'
        additions = [('stories', story), ('tasks', task)]
    else:
        additions = [('tasks', task)]
    draft = edit_draft(case, [dict(op='add', collection=c, object_id=obj['id'], field=None, value=obj)
                              for c, obj in additions])
    draft['change_summary'] = '增加明确的同名独立范围，保留已有对象。'
    return draft, story, task


@pytest.mark.parametrize('kind', ['story', 'task'])
def test_revision_new_lower_uuid_cannot_steal_existing_alias(clarify_case, offline_office, kind):
    case = clarify_case
    before = {p.name: p.read_bytes() for p in baseline(case).iterdir() if p.is_file()}
    original = read_json(baseline(case) / 'projection.json')
    draft, story, task = collision_draft(case, kind)
    checked = check_edits(case, draft)
    assert checked['ok'], checked
    output = render(case, checked['result'])
    projection = read_json(case['project'] / output['projection_ref']['path'])
    names = {o['object_id']: o['display_name'] for o in projection['objects']}
    for obj in original['objects']:
        if obj['kind'] in ('story', 'task'):
            assert names[obj['object_id']] == obj['display_name']
    newcomer = story if kind == 'story' else task
    original_id = case['ids']['S-01' if kind == 'story' else 'T-01']
    assert names[newcomer['id']] != names[original_id]
    with_book = openpyxl.load_workbook(case['project'] / output['workbook_ref']['path'])
    try:
        task_row = next(o['rows'][0] for o in projection['objects'] if o['object_id'] == task['id'])
        assert with_book['02-任务清单'].cell(task_row, 1).value == names[task['story_id']]
        assert with_book['02-任务清单'].cell(task_row, 2).value == names[task['id']]
    finally:
        with_book.close()
    assert {p.name: p.read_bytes() for p in baseline(case).iterdir() if p.is_file()} == before


@pytest.mark.parametrize('subset', [False, True], ids=['full', 'selected-subset'])
def test_revision_summary_describes_only_the_bound_actual_changes(clarify_case, offline_office, subset):
    case = clarify_case
    edits = [dict(op='replace', collection='stories', object_id=case['ids'][identity],
                  field='notes', value=value) for identity, value in
             [('S-01', '保留资料查询的全部验收范围'), ('S-02', '暂缓的公共客户端备注')]]
    draft = edit_draft(case, edits)
    draft['change_summary'] = '调整资料查询与公共客户端两项备注。'
    checked = check_edits(case, draft)
    assert checked['ok'], checked
    if subset:
        selected = dict(deepcopy(draft), subset_of=checked['result']['plan_ref'])
        selected['edits'] = selected['edits'][:1]
        checked = check_edits(case, selected)
        assert checked['ok'], checked
    # A plausible adjacent filename is not the plan bound by the check record.
    write_json(case['project'] / f'.ai-sow-lite/work/clarify/{case["request_id"]}/plan.json',
               dict(change_summary='无关文件不得采用', changes=[]))
    output = render(case, checked['result'])
    summary = (case['project'] / output['summary_ref']['path']).read_text()
    assert '首版' not in summary
    assert '修改版' in summary and output['version_id'] in summary
    assert '保留资料查询的全部验收范围' in summary
    assert case['ids']['S-01'] in summary and 'notes' in summary
    assert '无关文件不得采用' not in summary
    if subset:
        assert '暂缓的公共客户端备注' not in summary
        assert draft['change_summary'] not in summary
        assert case['ids']['S-02'] not in summary
    else:
        assert '暂缓的公共客户端备注' in summary
        assert draft['change_summary'] in summary


def test_no_change_uses_existing_delivery_without_office(clarify_case, offline_office):
    case = clarify_case
    before = {p.name: p.read_bytes() for p in baseline(case).iterdir() if p.is_file()}
    checked = check_edits(case, edit_draft(case, []))
    assert checked['ok'] and checked['result']['no_change'], checked
    result = checked['result']
    reply = invoke(case, 'render', dict(candidate_path=result['candidate_ref']['path'],
        check_path=result['check_ref']['path'], expected_current=case['current']))
    assert not reply['ok'] and reply['diagnostics'][0]['code'] == 'CANDIDATE_INVALID'
    assert not offline_office
    assert {p.name: p.read_bytes() for p in baseline(case).iterdir() if p.is_file()} == before


def notes_result(case):
    checked = check_edits(case, edit_draft(case, [dict(op='replace', collection='stories',
        object_id=case['ids']['S-01'], field='notes', value='仅补充交付说明')]))
    assert checked['ok'], checked
    return checked['result']


def legacy_attempt(case, result, prepared_ref=None):
    from ai_sow_lite.contracts import semantic_digest
    payload = dict(candidate_path=result['candidate_ref']['path'],
                   check_path=result['check_ref']['path'], expected_current=case['current'])
    signature = semantic_digest(dict(check=read_json(case['project'] / payload['check_path']),
        payload=payload, projector_version='lite-projection-v1', engine=office.selection_fingerprint(),
        implementation_version='lite-render-v3'))
    area = case['project'] / f'.ai-sow-lite/work/clarify/{case["request_id"]}'
    write_json(area / 'render-attempt.json', dict(signature=signature, prepared_ref=prepared_ref,
                                                implementation_version='lite-render-v3'))
    return area


@pytest.mark.parametrize('legacy', [True, False], ids=['old-success', 'new-package-tamper'])
def test_only_recorded_old_success_can_reuse_legacy_summary(clarify_case, offline_office, legacy):
    from ai_sow_lite.project import file_ref
    case = clarify_case
    result = notes_result(case)
    output = render(case, result)
    prepared_path = case['project'] / output['prepared_ref']['path']
    prepared = read_json(prepared_path)
    summary_path = case['project'] / output['summary_ref']['path']
    old_text = (baseline(case) / 'summary.md').read_text()
    summary_path.write_text(old_text.replace(case['current']['version_id'], output['version_id']))
    for ref in prepared['files']:
        if Path(ref['path']).name == 'summary.md':
            ref['sha256'] = file_sha256(summary_path)
    write_json(prepared_path, prepared)
    if legacy:
        # A pre-v4 capture has only the request-wide attempt receipt.
        (case['project'] / result['candidate_ref']['path']).with_name('render-attempt.json').unlink(missing_ok=True)
        area = legacy_attempt(case, result, file_ref(case['project'], prepared_path))
        before = (area / 'checkpoint.json').read_bytes()
        assert workbook.verify_prepared(case['project'], prepared) == dict(diagnostics=[])
        repeated = render(case, result)
        assert repeated['version_id'] == output['version_id']
        assert len(offline_office) == 1
        assert (area / 'checkpoint.json').read_bytes() == before
    else:
        assert workbook.verify_prepared(case['project'], prepared)['diagnostics']


def test_new_renderer_can_retry_v3_failure_once_without_resetting_budget(clarify_case, offline_office):
    case = clarify_case
    result = notes_result(case)
    area = legacy_attempt(case, result)
    checkpoint = read_json(area / 'checkpoint.json')
    checkpoint['repair_batches'] = 1
    write_json(area / 'checkpoint.json', checkpoint)
    output = render(case, result)
    checkpoint = read_json(area / 'checkpoint.json')
    assert checkpoint['repair_batches'] == 2 and checkpoint['operation_retries']['render'] == 1
    assert render(case, result)['version_id'] == output['version_id']
    assert len(offline_office) == 1


@pytest.mark.parametrize('prior_retries', [0, 1], ids=['remaining-retry', 'exhausted-retry'])
def test_legacy_failed_receipt_keeps_budget_when_engine_changes(clarify_case, monkeypatch, prior_retries):
    from ai_sow_lite.project import StorageError
    case = clarify_case
    result = notes_result(case)
    payload = dict(candidate_path=result['candidate_ref']['path'],
                   check_path=result['check_ref']['path'], expected_current=case['current'])
    calls, environment = [], []

    def fail(source, destination):
        calls.append(source)
        raise StorageError('CALCULATION_FAILED', '受控持续失败；没有 Office 进程。')

    monkeypatch.setattr(office, 'recalculate', fail)
    monkeypatch.setattr(office, 'selection_fingerprint', lambda: list(environment))
    # Real consumer failures consume the budget; do not manufacture checkpoint counts.
    for number in range(prior_retries + 1):
        environment[:] = [f'engine-{number}']
        assert invoke(case, 'render', payload)['diagnostics'][0]['code'] == 'CALCULATION_FAILED'
    # Model the pre-v4 storage format: global signature, no candidate identity.
    area = legacy_attempt(case, result)
    (case['project'] / result['candidate_ref']['path']).with_name('render-attempt.json').unlink()
    checkpoint = read_json(area / 'checkpoint.json')
    assert checkpoint['operation_retries'].get('render', 0) == prior_retries
    old_receipt = (area / 'render-attempt.json').read_bytes()
    environment[:] = ['changed-after-upgrade']
    response = invoke(case, 'render', payload)
    assert response['diagnostics'][0]['code'] == ('LOOP_LIMIT_REACHED' if prior_retries else 'CALCULATION_FAILED')
    assert len(calls) == 2
    current_checkpoint = read_json(area / 'checkpoint.json')
    assert current_checkpoint['repair_batches'] == 1
    assert current_checkpoint['operation_retries']['render'] == 1
    if prior_retries:
        assert current_checkpoint == checkpoint
    assert (area / 'render-attempt.json').read_bytes() == old_receipt
    assert read_json(case['project'] / '.ai-sow-lite/current.json') == case['current']
    environment[:] = ['another-engine-change']
    assert invoke(case, 'render', payload)['diagnostics'][0]['code'] == 'LOOP_LIMIT_REACHED'
    assert len(calls) == 2


def test_normal_r1_r2_subset_previews_reuse_each_slot_without_repair_cost(clarify_case, offline_office):
    case = clarify_case
    draft = edit_draft(case, [dict(op='replace', collection='stories', object_id=case['ids'][identity],
        field='notes', value=value) for identity, value in [('S-01', '首稿说明'), ('S-02', '公共客户端说明')]])
    first = check_edits(case, draft)
    assert first['ok'], first
    r1 = render(case, first['result'])
    draft['revision'] = 2
    draft['edits'][0]['value'] = '明确修订后的说明'
    second = check_edits(case, draft)
    assert second['ok'], second
    r2 = render(case, second['result'])
    selected = dict(deepcopy(draft), subset_of=second['result']['plan_ref'])
    selected['edits'] = selected['edits'][:1]
    subset = check_edits(case, selected)
    assert subset['ok'], subset
    r_subset = render(case, subset['result'])
    assert len({r1['version_id'], r2['version_id'], r_subset['version_id']}) == 3
    # Returning to any immutable slot reuses its complete package, not only the latest.
    for result, output in [(first['result'], r1), (second['result'], r2), (subset['result'], r_subset),
                           (first['result'], r1)]:
        assert render(case, result)['version_id'] == output['version_id']
    assert len(offline_office) == 3
    checkpoint = read_json(case['project'] / f'.ai-sow-lite/work/clarify/{case["request_id"]}/checkpoint.json')
    assert checkpoint['repair_batches'] == 0 and checkpoint['operation_retries'].get('render', 0) == 0


def test_failure_retry_cap_is_shared_across_preview_slots(clarify_case, monkeypatch):
    from ai_sow_lite.project import StorageError
    case = clarify_case
    draft = edit_draft(case, [dict(op='replace', collection='stories', object_id=case['ids']['S-01'],
                                  field='notes', value='首稿说明')])
    checked = check_edits(case, draft)
    assert checked['ok'], checked
    calls = []
    environment = ['original-engine']

    def fail(source, destination):
        calls.append(source)
        raise StorageError('CALCULATION_FAILED', '受控外部引擎失败；没有 Office 进程。')

    monkeypatch.setattr(office, 'recalculate', fail)
    monkeypatch.setattr(office, 'selection_fingerprint', lambda: list(environment))

    def attempt(result):
        return invoke(case, 'render', dict(candidate_path=result['candidate_ref']['path'],
            check_path=result['check_ref']['path'], expected_current=case['current']))

    assert attempt(checked['result'])['diagnostics'][0]['code'] == 'CALCULATION_FAILED'
    assert attempt(checked['result'])['diagnostics'][0]['code'] == 'LOOP_LIMIT_REACHED'
    environment[:] = ['corrected-engine']
    assert attempt(checked['result'])['diagnostics'][0]['code'] == 'CALCULATION_FAILED'
    draft['revision'] = 2
    draft['edits'][0]['value'] = '正常的第二轮说明'
    second = check_edits(case, draft)
    assert second['ok'], second
    assert attempt(second['result'])['diagnostics'][0]['code'] == 'CALCULATION_FAILED'
    environment[:] = ['another-engine-change']
    assert attempt(second['result'])['diagnostics'][0]['code'] == 'LOOP_LIMIT_REACHED'
    assert attempt(checked['result'])['diagnostics'][0]['code'] == 'LOOP_LIMIT_REACHED'
    assert len(calls) == 3
    checkpoint = read_json(case['project'] / f'.ai-sow-lite/work/clarify/{case["request_id"]}/checkpoint.json')
    assert checkpoint['repair_batches'] == 1 and checkpoint['operation_retries']['render'] == 1


@pytest.mark.parametrize('change', ['rename-story', 'rename-task', 'remove-task'])
def test_changed_or_removed_names_do_not_reserve_retired_aliases(clarify_case, offline_office, change):
    case = clarify_case
    kind = 'story' if change == 'rename-story' else 'task'
    draft, story, task = collision_draft(case, kind)
    identity = case['ids']['S-01' if kind == 'story' else 'T-01']
    collection, field = ('stories', 'title') if kind == 'story' else ('tasks', 'name')
    if change.startswith('rename'):
        draft['edits'].append(dict(op='replace', collection=collection, object_id=identity,
                                   field=field, value='明确改名后的原有对象'))
    else:
        lineage = dict(from_version_id=case['current']['version_id'], from_ids=[identity], to_ids=[],
                       reason='明确退出原任务，保留历史身份。', evidence_refs=[case['ids']['EV-P-B1']])
        key = canonical_json_bytes([lineage['from_version_id'], lineage['from_ids']]).decode()
        draft['edits'].extend([dict(op='remove', collection=collection, object_id=identity, field=None),
            dict(op='add', collection='lineage', object_id=key, field=None, value=lineage)])
    checked = check_edits(case, draft)
    assert checked['ok'], checked
    output = render(case, checked['result'])
    projection = read_json(case['project'] / output['projection_ref']['path'])
    names = {obj['object_id']: obj['display_name'] for obj in projection['objects']}
    newcomer = story if kind == 'story' else task
    assert names[newcomer['id']] == newcomer[field]
    if change.startswith('rename'):
        assert names[identity] == '明确改名后的原有对象'
    else:
        assert identity not in names
        original = read_json(baseline(case) / 'model.json')
        assert identity in {obj['id'] for obj in original[collection]}


def test_generate_summary_still_describes_first_version(tmp_path, offline_office):
    from .support.excel import prepare_case
    from ai_sow_lite.contracts import semantic_digest
    seed, payload, _ = prepare_case(tmp_path / 'project', render=False)
    case = dict(project=seed.project, request_id=seed.request_id)
    reply = invoke(case, 'render', payload)
    assert reply['ok'], reply
    output = reply['result']
    text = (seed.project / output['summary_ref']['path']).read_text()
    assert '本次生成首版候选' in text and '修改版' not in text and '本次变化' not in text
    prepared = read_json(seed.project / output['prepared_ref']['path'])
    assert workbook.verify_prepared(seed.project, prepared) == dict(diagnostics=[])
    # Generate v3 used the same projection/text contract. A successful exact
    # v3 receipt must keep its bytes and budget when the implementation advances.
    attempt = read_json(seed.file('render-attempt.json'))
    attempt['implementation_version'] = 'lite-render-v3'
    attempt['signature'] = semantic_digest(dict(check=read_json(seed.project / payload['check_path']),
        payload=payload, projector_version='lite-projection-v1', engine=office.selection_fingerprint(),
        implementation_version='lite-render-v3'))
    write_json(seed.file('render-attempt.json'), attempt)
    checkpoint = read_json(seed.file('checkpoint.json'))
    checkpoint['repair_batches'] = 1
    write_json(seed.file('checkpoint.json'), checkpoint)
    before = {ref['path']: (seed.project / ref['path']).read_bytes() for ref in
              [output['prepared_ref'], *prepared['files']]}
    repeated = invoke(case, 'render', payload)
    assert repeated['ok'], repeated
    assert repeated['result']['prepared_ref'] == output['prepared_ref']
    assert repeated['result']['workbook_ref'] == output['workbook_ref']
    assert repeated['result']['version_id'] == output['version_id']
    assert len(offline_office) == 1
    assert read_json(seed.file('checkpoint.json')) == checkpoint
    assert all((seed.project / path).read_bytes() == raw for path, raw in before.items())


@pytest.mark.office
def test_real_cli_revision_preserves_aliases_and_delivers_bound_change_summary(clarify_case):
    """One isolated real Office calculation covers both workbook review defects."""
    case = clarify_case
    old = baseline(case)
    before = {p.name: p.read_bytes() for p in old.iterdir() if p.is_file()}
    template = case['project'] / '.ai-sow-lite/template' / workbook.TEMPLATE_HASH / 'sow-template.xlsx'
    assert file_sha256(template) == '6abc55d44bc66476a60c2251e18c0dfdb66709e07539c246dfdec3a0373f5332'
    draft = complexity_draft(case, 'S')
    collision, story, task = collision_draft(case, 'story')
    task['name'] = read_json(old / 'model.json')['tasks'][0]['name']
    draft['edits'].extend(collision['edits'])
    draft['change_summary'] += '增加同名独立 Story 和 Task。'
    checked = check_edits(case, draft)
    assert checked['ok'], checked
    result = checked['result']
    confirmed, _ = confirm_plan(case, result)
    checked = run_request(case['project'], case['request_id'], 'check', dict(
        candidate_path=result['candidate_ref']['path'], plan_path=confirmed, scope='full'))
    assert checked['ok'], checked
    result['check_ref'] = checked['result']['check_ref']
    output = render(case, result, real=True)
    projection = read_json(case['project'] / output['projection_ref']['path'])
    names = {obj['object_id']: obj['display_name'] for obj in projection['objects']}
    for obj in read_json(old / 'projection.json')['objects']:
        if obj['kind'] in ('story', 'task'):
            assert names[obj['object_id']] == obj['display_name']
    assert names[story['id']] != names[case['ids']['S-01']]
    assert names[task['id']] != names[case['ids']['T-01']]
    summary = (case['project'] / output['summary_ref']['path']).read_text()
    assert '本次生成修改版候选' in summary and '本次生成首版候选' not in summary
    assert draft['change_summary'] in summary and '"M" → "S"' in summary
    assert story['id'] in summary and task['id'] in summary and case['ids']['T-06'] in summary
    applied = run_request(case['project'], case['request_id'], 'apply', dict(entrypoint='clarify',
        prepared_path=output['prepared_ref']['path'], plan_path=confirmed, expected_current=case['current']))
    assert applied['ok'], applied
    destination = case['project'] / '.ai-sow-lite/versions' / output['version_id']
    assert (destination / 'summary.md').read_bytes() == (case['project'] / output['summary_ref']['path']).read_bytes()
    assert {p.name: p.read_bytes() for p in old.iterdir() if p.is_file()} == before
    assert file_sha256(template) == workbook.TEMPLATE_HASH
    print('REVISION_REAL_OFFICE', dict(version_id=output['version_id'],
        workbook_sha256=output['workbook_ref']['sha256'], template_sha256=file_sha256(template)))
