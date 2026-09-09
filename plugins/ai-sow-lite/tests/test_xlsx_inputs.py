"""Public CLI reads real sparse workbooks; semantics remain agent-owned."""
from pathlib import Path
from uuid import uuid4
import hashlib

from .support.cli import run_request
from .support.fixtures import read_json
from .test_inputs import sources_payload, inspect

HISTORY = Path(__file__).parent / 'fixtures/history'


def ingest_history(tmp_path):
    project, request = tmp_path / 'project', str(uuid4())
    payload = sources_payload(HISTORY / 'sparse-history.xlsx', project_type='existing')
    payload['sources'][0].update(material_types=['prior-sow'], uses=['as-is'])
    response = run_request(project, request, 'ingest', payload)
    assert response['ok'], response
    entry = response['result']['input_refs'][0]
    reading = read_json(project / response['result']['reading_refs'][0]['path'])
    return project, request, entry, reading


def region(project, request, entry, reading, sheet='历史范围', address='A1:C4', **kwargs):
    return inspect(project, request, 'regions', dict(input_version_id=entry['input_version_id'],
                   locator=dict(kind='xlsx_range', sheet=sheet, range=address, read_id=reading['read_id'])), **kwargs)


def test_real_xlsx_ingest_directory_and_exact_region(tmp_path):
    project, request, entry, reading = ingest_history(tmp_path)
    assert entry['format'] == 'xlsx' and 'encoding' not in entry
    assert (project / entry['relative_path']).read_bytes() == (HISTORY / 'sparse-history.xlsx').read_bytes()
    directory = inspect(project, request, 'regions', {'input_version_id': entry['input_version_id']}, limit=1)
    assert directory['ok'], directory
    assert directory['result']['matched_count'] == 4
    assert directory['result']['items'][0]['sheet'] == '历史范围'
    assert directory['result']['items'][0]['merges'] == ['A2:A3']
    assert directory['result']['next_cursor']
    response = region(project, request, entry, reading)
    assert response['ok'], response
    result = response['result']
    cells = {cell['address']: cell for cell in result['items']}
    assert cells['A2']['value'] == {'type': 'string', 'value': '订单查询'}
    assert cells['A3']['merged_range'] == 'A2:A3'
    assert cells['A3']['merge_anchor'] == 'A2'
    assert cells['C4']['value'] == {'type': 'blank', 'value': None}
    ref = result['coverage']['excerpt_ref']
    assert hashlib.sha256((project / ref['path']).read_bytes()).hexdigest() == result['coverage']['excerpt_hash']
    assert result['coverage']['locator']['range'] == 'A1:C4'
    assert result['coverage']['locator']['read_id'] != reading['read_id']


def test_typed_values_formula_cache_and_hidden_comments(tmp_path):
    project, request, entry, reading = ingest_history(tmp_path)
    result = region(project, request, entry, reading, '类型观察', 'A1:H2')['result']
    cells = {cell['address']: cell for cell in result['items']}
    expected = {'A1': ('number', '12.5'), 'B1': ('boolean', True), 'C1': ('date', '2026-01-02'),
                'D1': ('datetime', '2026-01-02T03:04:05'), 'E1': ('time', '03:04:05'),
                'F1': ('error', '#DIV/0!'), 'A2': ('duration', '86430')}
    for address, (kind, value) in expected.items():
        assert cells[address]['value'] == dict(type=kind, value=value)
    assert cells['G1']['formula']['text'] == '=1+2'
    assert cells['G1']['cache'] == dict(present=True, value=dict(type='number', value='3'))
    assert cells['H1']['cache'] == dict(present=False, value=None)
    assert cells['D2']['cache'] == dict(present=True, value=dict(type='number', value='0'))
    assert cells['E2']['cache'] == dict(present=True, value=dict(type='boolean', value=False))
    assert cells['G1']['value'] == dict(type='blank', value=None)
    assert cells['C1']['raw_value'] == dict(type='number', value='46024')
    result = region(project, request, entry, reading)['result']
    cells = {cell['address']: cell for cell in result['items']}
    assert cells['B3']['hidden_row'] is True and cells['C2']['hidden_column'] is True
    assert '不含订单明细' in cells['B3']['comment']['text']
    assert result['coverage']['metadata']['tables'] == [{'name': 'HistoryTable', 'range': 'A1:C8'}]
    notes = region(project, request, entry, reading, '附注', 'A1:A2')['result']
    assert notes['coverage']['metadata']['state'] == 'hidden'


def rewrite_xlsx(path, *, replace=None, extras=None):
    from io import BytesIO
    from zipfile import ZipFile, ZIP_DEFLATED
    output = BytesIO()
    with ZipFile(path) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            raw = source.read(item.filename)
            if replace and item.filename in replace:
                raw = replace[item.filename](raw)
            target.writestr(item, raw)
        for name, raw in (extras or {}).items():
            target.writestr(name, raw)
    path.write_bytes(output.getvalue())


def test_unread_surfaces_are_reported_without_failing_readable_cells(tmp_path):
    import shutil
    source = tmp_path / 'objects.xlsx'
    shutil.copyfile(HISTORY / 'sparse-history.xlsx', source)
    rewrite_xlsx(source, extras={'xl/drawings/drawing9.xml': b'<drawing>unread</drawing>',
                                 'xl/embeddings/object1.bin': b'unread'})
    response = run_request(tmp_path / 'project', str(uuid4()), 'ingest', sources_payload(source))
    assert response['ok'], response
    entry = response['result']['input_refs'][0]
    directory = inspect(tmp_path / 'project', str(uuid4()), 'regions', {'input_version_id': entry['input_version_id']})
    limits = directory['result']['coverage']['limitations']
    assert any(x['part'] == 'xl/drawings/drawing9.xml' for x in limits)
    assert any(x['part'] == 'xl/embeddings/object1.bin' for x in limits)
    assert all(x['coverage'] == 'unread' for x in limits)


import pytest


@pytest.mark.parametrize('mutation', ['dimensions', 'actual-cell', 'members', 'compression', 'long-text'])
def test_xlsx_limits_never_return_a_partial_success(tmp_path, mutation):
    import shutil
    source = tmp_path / 'bounded.xlsx'
    shutil.copyfile(HISTORY / 'sparse-history.xlsx', source)
    if mutation == 'dimensions':
        rewrite_xlsx(source, replace={'xl/worksheets/sheet1.xml': lambda raw: raw.replace(b'A1:C8', b'A1:XFD1048576')})
    elif mutation == 'actual-cell':
        rewrite_xlsx(source, replace={'xl/worksheets/sheet1.xml': lambda raw: raw.replace(b'r="B2"', b'r="XFD1048576"')})
    elif mutation == 'members':
        rewrite_xlsx(source, extras={f'extra/{n}': b'x' for n in range(2049)})
    elif mutation == 'compression':
        rewrite_xlsx(source, extras={'extra/huge': b'0' * 3000000})
    else:
        rewrite_xlsx(source, replace={'xl/worksheets/sheet1.xml': lambda raw: raw.replace(b'<t>', b'<t>' + b'x' * 32768, 1)})
    response = run_request(tmp_path / 'project', str(uuid4()), 'ingest', sources_payload(source))
    assert not response['ok'], response
    assert response['diagnostics'][0]['code'] == 'RESULT_TOO_LARGE'
    assert response['result']['reading_refs'] == []
    assert '上限' in response['diagnostics'][0]['message']


def test_huge_unicode_cell_has_bounded_representation_and_finite_page(tmp_path):
    from openpyxl import Workbook
    source = tmp_path / 'huge.xlsx'
    book = Workbook()
    book.active['A1'] = '中' * 32767
    book.active['B1'] = 'next'
    book.save(source)
    response = run_request(tmp_path / 'project', str(uuid4()), 'ingest', sources_payload(source))
    assert response['ok'], response
    project = tmp_path / 'project'
    entry = response['result']['input_refs'][0]
    reading = read_json(project / response['result']['reading_refs'][0]['path'])
    result = region(project, str(uuid4()), entry, reading, 'Sheet', 'A1:B1')['result']
    assert result['next_cursor'] is None and result['returned_count'] == 2
    assert result['items'][0]['representation'] == 'attachment'
    assert result['items'][0]['address'] == 'A1'
    assert result['items'][1]['value']['value'] == 'next'
    excerpt = read_json(project / result['coverage']['excerpt_ref']['path'])
    assert excerpt['cells'][0]['value']['value'] == '中' * 32767


def xlsx_analysis_case(tmp_path, mutation=None):
    import json
    from .support.fixtures import build_ingested_case, write_json
    case = build_ingested_case(tmp_path / 'project')
    response = run_request(case.project, case.request_id, 'ingest', sources_payload(HISTORY / 'sparse-history.xlsx'))
    assert response['ok'], response
    entry = response['result']['input_refs'][0]
    reading = read_json(case.project / response['result']['reading_refs'][0]['path'])
    result = region(case.project, case.request_id, entry, reading)['result']
    source = dict(input_version_id=entry['input_version_id'], locator=result['coverage']['locator'],
                  excerpt_hash=result['coverage']['excerpt_hash'])
    analysis = read_json(case.file('analysis.json'))
    old_evidence, new_evidence = analysis['evidence'][0]['id'], str(uuid4())
    old_topic, new_topic = analysis['topics'][0]['topic_version_id'], str(uuid4())
    for name in ['analysis.json', 'candidate.json', 'model.json', 'pending-items.json', 'decisions.json']:
        path = case.file(name)
        data = json.loads(json.dumps(read_json(path)).replace(old_evidence, new_evidence).replace(old_topic, new_topic))
        write_json(path, data)
    analysis = read_json(case.file('analysis.json'))
    analysis['evidence'][0]['source_refs'] = [source]
    topic = analysis['topics'][0]
    topic['input_version_ids'].append(entry['input_version_id'])
    topic['covered_regions'].append(source.copy())
    topic['historical_items'] = [dict(id=str(uuid4()), label='订单查询', description='原文仅有范围及工作说明', evidence_refs=[new_evidence])]
    if mutation == 'hash':
        analysis['evidence'][0]['source_refs'][0]['excerpt_hash'] = '0' * 64
    elif mutation == 'range':
        source['locator'] = dict(source['locator'], range='B1:C4')
    elif mutation == 'read-id':
        source['locator'] = dict(source['locator'], read_id=str(uuid4()))
    elif mutation == 'original':
        (case.project / entry['relative_path']).write_bytes(b'corrupted')
    elif mutation == 'missing':
        (case.project / entry['relative_path']).unlink()
    elif mutation == 'attachment':
        (case.project / result['coverage']['excerpt_ref']['path']).write_bytes(b'{}')
    elif mutation == 'forged-attachment':
        from ai_sow_lite.contracts import canonical_json_bytes
        from ai_sow_lite.project import file_ref
        ref = result['coverage']['reading_ref']
        path = case.project / ref['path']
        data = read_json(path)
        attachment = case.project / data['excerpts'][0]['file_ref']['path']
        altered = read_json(attachment)
        altered['cells'][0]['value']['value'] = 'forged'
        attachment.write_bytes(canonical_json_bytes(altered))
        data['excerpts'][0]['file_ref'] = file_ref(case.project, attachment)
        write_json(path, data)
        for cache in (case.project / entry['relative_path']).parent.glob('readings/*.json'):
            if read_json(cache)['path'] == ref['path']:
                write_json(cache, file_ref(case.project, path))
        source['excerpt_hash'] = data['excerpts'][0]['file_ref']['sha256']
    write_json(case.file('analysis.json'), analysis)
    candidate = read_json(case.candidate_path)
    candidate['input_version_ids'].append(entry['input_version_id'])
    write_json(case.candidate_path, candidate)
    response = run_request(case.project, case.request_id, 'ingest', dict(kind='analysis', entrypoint='generate',
                           analysis_path=case.file('analysis.json').relative_to(case.project).as_posix()))
    return case, response, result


def test_xlsx_ac_and_sparse_history_evidence_survive_full_candidate_check(tmp_path):
    case, response, region_result = xlsx_analysis_case(tmp_path)
    assert response['ok'], response
    checked = run_request(case.project, case.request_id, 'check', dict(candidate_path=case.candidate_path.relative_to(case.project).as_posix(),
                          scope='full', plan_path=None))
    assert checked['ok'], checked
    assert checked['result']['valid_for_render'] is True
    report = read_json(case.project / checked['result']['check_ref']['path'])
    assert region_result['coverage']['excerpt_ref'] in report['dependencies']
    assert region_result['coverage']['reading_ref'] in report['dependencies']


@pytest.mark.parametrize('mutation', ['hash', 'range', 'read-id', 'original', 'missing', 'attachment', 'forged-attachment'])
def test_xlsx_analysis_rejects_invalid_source_before_publication(tmp_path, mutation):
    case, response, _ = xlsx_analysis_case(tmp_path, mutation)
    assert not response['ok'], response
    assert 'OPERATION_UNSUPPORTED' not in {d['code'] for d in response['diagnostics']}
    analysis = read_json(case.file('analysis.json'))
    assert not (case.project / '.ai-sow-lite/analysis/topics' / analysis['topics'][0]['topic_version_id']).exists()


def test_selection_and_multiuse_reuse_physical_reading_without_old_analysis(tmp_path):
    project, request, entry, reading = ingest_history(tmp_path)
    first = region(project, request, entry, reading)['result']
    same = region(project, request, entry, reading)['result']
    assert same['coverage']['reading_ref'] == first['coverage']['reading_ref']
    different = region(project, request, entry, reading, address='B2:C3')['result']
    assert different['coverage']['locator']['read_id'] != first['coverage']['locator']['read_id']
    payload = sources_payload(HISTORY / 'sparse-history.xlsx', project_type='existing')
    payload['sources'][0].update(material_types=['prior-sow', 'hld'], uses=['as-is', 'to-be-architecture'],
                                use_regions=[dict(material_type='hld', use='to-be-architecture', locators=[first['coverage']['locator']])])
    added = run_request(project, request, 'ingest', payload)
    assert added['ok'], added
    assert added['result']['input_refs'][0]['input_version_id'] == entry['input_version_id']
    assert read_json(project / added['result']['reading_refs'][0]['path'])['read_id'] == reading['read_id']
    topics = inspect(project, request, 'topics', {'uses': ['to-be-architecture']})
    assert topics['ok'] and topics['result']['items'] == []


def test_history_query_scope_changes_on_new_member_even_after_zero_match(tmp_path):
    from .support.fixtures import write_json
    project, request, entry, reading = ingest_history(tmp_path)
    selected = region(project, request, entry, reading)['result']['coverage']
    evidence_id = str(uuid4())
    source = dict(input_version_id=entry['input_version_id'], locator=selected['locator'], excerpt_hash=selected['excerpt_hash'])
    evidence = dict(id=evidence_id, kind='statement', text='原文历史范围', source_refs=[source], basis_refs=[], limitations='仅原文粒度')
    topic = dict(topic_id=str(uuid4()), topic_version_id=str(uuid4()), title='历史候选导航', uses=['as-is'], input_version_ids=[entry['input_version_id']],
                 covered_regions=[source], uncovered_regions=[], evidence_refs=[evidence_id], related_object_ids=[],
                 external_responsibilities='', limitations='', conclusion='保留两个部分重叠范围。', historical_items=[
                     dict(id=str(uuid4()), label='订单查询', description='当前订单', evidence_refs=[evidence_id]),
                     dict(id=str(uuid4()), label='订单查询', description='归档订单；不与当前订单合并', evidence_refs=[evidence_id])])
    path = project / f'.ai-sow-lite/work/generate/{request}/history.json'
    write_json(path, dict(schema_version='1.0', topics=[topic], evidence=[evidence], observations=[]))
    registered = run_request(project, request, 'ingest', dict(kind='analysis', entrypoint='generate', analysis_path=path.relative_to(project).as_posix()))
    assert registered['ok'], registered
    selector = {'uses': ['as-is'], 'historical_label': '订单'}
    first = inspect(project, request, 'topics', selector, limit=1)
    assert first['ok'], first
    assert first['result']['matched_count'] == 2
    assert first['result']['items'][0]['historical_item']['description'] == '当前订单'
    zero_selector = {'uses': ['as-is'], 'historical_label': '付款'}
    zero = inspect(project, request, 'topics', zero_selector)['result']
    assert zero['items'] == [] and zero['coverage']['scope_digest']
    before = zero['coverage']['scope_digest']
    extra = tmp_path / 'extra.md'
    extra.write_text('付款范围\n', encoding='utf-8')
    payload = sources_payload(extra, project_type='existing')
    payload['sources'][0].update(uses=['as-is'], material_types=['prior-sow'])
    assert run_request(project, request, 'ingest', payload)['ok']
    stale = inspect(project, request, 'topics', selector, limit=1, cursor=first['result']['next_cursor'])
    assert not stale['ok'] and stale['diagnostics'][0]['code'] == 'VERSION_INCOMPATIBLE'
    fresh = inspect(project, request, 'topics', zero_selector)['result']
    assert fresh['coverage']['scope_digest'] != before
    assert fresh['items'] == []  # no semantic history invented from newly ingested text
    assert len(fresh['coverage']['source_members']) == 2
    assert fresh['coverage']['selector'] == zero_selector


def test_region_cursor_pins_selection_and_refuses_zero_progress(tmp_path):
    import base64
    from ai_sow_lite.contracts import strict_json_loads, canonical_json_bytes
    project, request, entry, reading = ingest_history(tmp_path)
    first = region(project, request, entry, reading, limit=1)['result']
    changed = region(project, request, entry, reading, address='B1:C4', cursor=first['next_cursor'])
    assert not changed['ok'] and changed['diagnostics'][0]['code'] == 'VERSION_INCOMPATIBLE'
    decoded = strict_json_loads(base64.urlsafe_b64decode(first['next_cursor']))
    decoded['offset'] = 0
    bad_cursor = base64.urlsafe_b64encode(canonical_json_bytes(decoded)).decode()
    bad = region(project, request, entry, reading, cursor=bad_cursor)
    assert not bad['ok'] and bad['diagnostics'][0]['code'] == 'VERSION_INCOMPATIBLE'
    cursor, addresses = None, []
    for _ in range(20):
        result = region(project, request, entry, reading, limit=1, cursor=cursor)['result']
        addresses.extend(cell['address'] for cell in result['items'])
        assert result['next_cursor'] != cursor or cursor is None
        cursor = result['next_cursor']
        if cursor is None:
            break
    assert addresses == [f'{col}{row}' for row in range(1, 5) for col in 'ABC']
    assert cursor is None


@pytest.mark.parametrize('damage', ['options', 'version', 'selection'])
def test_reading_identity_is_checked_before_reuse(tmp_path, damage):
    from .support.fixtures import write_json
    project, request, entry, reading = ingest_history(tmp_path)
    first = region(project, request, entry, reading)['result']
    path = project / first['coverage']['reading_ref']['path']
    record = read_json(path)
    if damage == 'options':
        record['options']['data_only'] = True
    elif damage == 'version':
        record['adapter_version'] = 'different-reader'
    else:
        record['selection']['range'] = 'B1:C4'
    write_json(path, record)
    observed = region(project, request, entry, reading)
    assert not observed['ok'] and observed['diagnostics'][0]['code'] == 'EVIDENCE_MISSING'


def test_utf16_xml_entity_declaration_is_not_expanded(tmp_path):
    import shutil
    source = tmp_path / 'entity.xlsx'
    shutil.copyfile(HISTORY / 'sparse-history.xlsx', source)
    def entity(raw):
        text = raw.decode('utf-8')
        text = '<!DOCTYPE worksheet [<!ENTITY hidden "silently expanded">]>' + text.replace('订单查询', '&hidden;')
        return text.encode('utf-16')
    rewrite_xlsx(source, replace={'xl/worksheets/sheet1.xml': entity})
    response = run_request(tmp_path / 'project', str(uuid4()), 'ingest', sources_payload(source))
    assert not response['ok'], response
    assert response['diagnostics'][0]['code'] == 'INPUT_UNAVAILABLE'
    assert response['result']['reading_refs'] == []


def _review_book(tmp_path, configure):
    from openpyxl import Workbook
    source = tmp_path / 'review.xlsx'
    book = Workbook()
    configure(book.active)
    book.save(source)
    book.close()
    return source


def _review_ingest(tmp_path, source):
    project, request = tmp_path / 'project', str(uuid4())
    response = run_request(project, request, 'ingest', sources_payload(source))
    assert response['ok'], response
    entry = response['result']['input_refs'][0]
    reading = read_json(project / response['result']['reading_refs'][0]['path'])
    return project, request, entry, reading


@pytest.mark.parametrize('surface', ['metadata', 'limitations'])
def test_review_f1_unbounded_region_context_is_explicitly_rejected(tmp_path, surface):
    def configure(sheet):
        sheet['A1'] = 'scope'
        if surface == 'metadata':
            for row in range(1, 20001):
                sheet.row_dimensions[row].hidden = True
    source = _review_book(tmp_path, configure)
    if surface == 'limitations':
        rewrite_xlsx(source, extras={f'xl/embeddings/object{n}.bin': b'unread' for n in range(500)})
    project, request, entry, reading = _review_ingest(tmp_path, source)
    response = region(project, request, entry, reading, 'Sheet', 'A1', limit=1)
    assert not response['ok'], response
    assert response['diagnostics'][0]['code'] == 'RESULT_TOO_LARGE'
    assert 'coverage' not in response['result'] and 'items' not in response['result']  # observation may still report the diagnostic


def test_review_f1_page_budget_includes_metadata_and_preserves_exact_continuation(tmp_path):
    from ai_sow_lite.contracts import canonical_json_bytes
    def configure(sheet):
        for row in range(1, 31):
            sheet.cell(row, 1, str(row) + ':' + 'x' * 600)
        for row in range(1, 7001):
            sheet.row_dimensions[row].hidden = True
    source = _review_book(tmp_path, configure)
    project, request, entry, reading = _review_ingest(tmp_path, source)
    cursor, addresses, pages = None, [], 0
    for _ in range(30):
        response = region(project, request, entry, reading, 'Sheet', 'A1:A30', limit=100, cursor=cursor)
        assert response['ok'], response
        result = response['result']
        assert len(canonical_json_bytes(result)) <= 64 * 1024
        assert result['coverage']['metadata']['hidden_rows'] == list(range(1, 7001))
        assert result['returned_count'] > 0
        addresses.extend(cell['address'] for cell in result['items'])
        pages += 1
        cursor = result['next_cursor']
        if cursor is None:
            break
    assert cursor is None and pages > 1
    assert addresses == [f'A{row}' for row in range(1, 31)]


@pytest.mark.parametrize('storage', ['inline', 'shared'])
@pytest.mark.parametrize('rich', [False, True])
def test_review_f2_phonetic_annotations_never_become_cell_body(tmp_path, storage, rich):
    from xml.etree import ElementTree as ET
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    source = _review_book(tmp_path, lambda sheet: setattr(sheet['A1'], 'value', '東京'))
    body = ET.Element(ns + ('is' if storage == 'inline' else 'si'))
    if rich:
        for text in ['東', '京']:
            ET.SubElement(ET.SubElement(body, ns + 'r'), ns + 't').text = text
    else:
        ET.SubElement(body, ns + 't').text = '東京'
    ET.SubElement(ET.SubElement(body, ns + 'rPh', sb='0', eb='2'), ns + 't').text = 'とうきょう'
    ET.SubElement(body, ns + 'phoneticPr', fontId='0')
    def replace(raw):
        root = ET.fromstring(raw)
        cell = root.find('.//' + ns + 'c')
        for child in list(cell):
            cell.remove(child)
        if storage == 'inline':
            cell.append(body)
        else:
            cell.set('t', 's')
            ET.SubElement(cell, ns + 'v').text = '0'
        return ET.tostring(root, encoding='utf-8')
    extras = {}
    replacements = {'xl/worksheets/sheet1.xml': replace}
    if storage == 'shared':
        shared = ET.Element(ns + 'sst', count='1', uniqueCount='1')
        shared.append(body)
        extras['xl/sharedStrings.xml'] = ET.tostring(shared, encoding='utf-8')
        def content_types(raw):
            root = ET.fromstring(raw)
            ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                          PartName='/xl/sharedStrings.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
            return ET.tostring(root, encoding='utf-8')
        replacements['[Content_Types].xml'] = content_types
    rewrite_xlsx(source, replace=replacements, extras=extras)
    project, request, entry, reading = _review_ingest(tmp_path, source)
    response = region(project, request, entry, reading, 'Sheet', 'A1')
    assert response['ok'], response
    cell = response['result']['items'][0]
    assert cell['value'] == {'type': 'string', 'value': '東京'}
    assert cell['raw_value'] == {'type': 'string', 'value': '東京'}
    excerpt = read_json(project / response['result']['coverage']['excerpt_ref']['path'])
    assert excerpt['cells'][0]['value'] == {'type': 'string', 'value': '東京'}


@pytest.mark.parametrize('dimension', ['absent', 'undersized'])
def test_review_f4_sparse_merge_extents_are_readable_within_known_bounds(tmp_path, dimension):
    from xml.etree import ElementTree as ET
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    def configure(sheet):
        sheet['A1'] = 'scope'
        sheet.merge_cells('A1:C3')
    source = _review_book(tmp_path, configure)
    def replace(raw):
        root = ET.fromstring(raw)
        node = root.find(ns + 'dimension')
        if dimension == 'absent':
            root.remove(node)
        else:
            node.set('ref', 'A1')
        for row in root.findall(ns + 'sheetData/' + ns + 'row'):
            for cell in list(row):
                if cell.attrib.get('r') != 'A1':
                    row.remove(cell)
        return ET.tostring(root, encoding='utf-8')
    rewrite_xlsx(source, replace={'xl/worksheets/sheet1.xml': replace})
    project, request, entry, reading = _review_ingest(tmp_path, source)
    directory = inspect(project, request, 'regions', {'input_version_id': entry['input_version_id']})
    assert directory['result']['items'][0]['used_range'] == 'A1:C3'
    response = region(project, request, entry, reading, 'Sheet', 'A1:C3')
    assert response['ok'], response
    cells = response['result']['items']
    assert len(cells) == 9
    assert cells[-1]['address'] == 'C3' and cells[-1]['merge_anchor'] == 'A1'
    assert cells[-1]['value'] == {'type': 'blank', 'value': None}
    outside = region(project, request, entry, reading, 'Sheet', 'A1:D3')
    assert not outside['ok'] and outside['diagnostics'][0]['code'] == 'EVIDENCE_MISSING'
