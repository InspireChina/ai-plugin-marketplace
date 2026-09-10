"""Real Lite CLI delivery from a fresh locked plugin copy; no installed old plugin."""
import argparse
import atexit
from collections import Counter
import hashlib
import json
import mimetypes
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

PLUGIN = Path(__file__).resolve().parents[2]


def install_read_audit():
    """Audit Python opens in the worker and CLI children, after interpreter startup.

    This is test instrumentation, not an OS sandbox or an Office syscall trace.
    Report only category counts; never source contents or absolute read paths.
    """
    workspace = Path(os.environ['LITE_SMOKE_WORKSPACE']).resolve()
    roots = [('plugin', PLUGIN), ('project', workspace / '项目 with spaces'),
             ('temporary', workspace), ('python', Path(sys.base_prefix).resolve())]
    engines = {Path(p).expanduser().resolve() for p in [os.environ.get('AI_SOW_LITE_OFFICE_BIN'),
               shutil.which('soffice'), shutil.which('libreoffice')] if p}
    system_files = {Path(p).resolve() for p in mimetypes.knownfiles}
    counts = Counter()
    report = dict(read_counts=counts, violations=0, office_conversions=0, pid=os.getpid(),
                  invocation_id=os.environ.get('LITE_SMOKE_INVOCATION'),
                  operation=os.environ.get('LITE_SMOKE_OPERATION'))

    def audit(event, args):
        if event == 'subprocess.Popen':
            arguments = args[1]
            if isinstance(arguments, (list, tuple)) and '--convert-to' in arguments:
                report['office_conversions'] += 1
        if event != 'open' or isinstance(args[0], int):
            return
        if args[2] & os.O_ACCMODE == os.O_WRONLY:
            return
        path = Path(os.fsdecode(args[0])).resolve()
        category = next((name for name, root in roots if path.is_relative_to(root)), None)
        if category is None and path in engines:
            category = 'office_binary'
        if category is None and path == Path(os.devnull).resolve():
            category = 'null_device'
        if category is None and path in system_files:
            category = 'system_mime_config'
        if category is None:
            report['violations'] += 1
            report.setdefault('denied_paths', []).append(str(path))  # Retained failure workspace only.
            raise PermissionError('Copy smoke refused a Python read outside its declared roots')
        counts[category] += 1

    def save():
        identity = report['invocation_id'] or str(os.getpid())
        (workspace / 'audit' / f'{identity}.json').write_text(json.dumps(report), encoding='utf-8')

    sys.addaudithook(audit)
    atexit.register(save)


def verify_delivery(case, prepared, applied, *, confirmation_path=None, previous_files=None):
    """Read delivered files and independently enumerate adopted dependency bytes."""
    from tests.support.fixtures import read_json
    project = case.project
    candidate = read_json(case.candidate_path)
    clarify = candidate['entrypoint'] == 'clarify'
    assert clarify == (confirmation_path is not None)

    def bound(ref):
        path = project / ref['path']
        assert path.resolve().is_relative_to(project)
        assert hashlib.sha256(path.read_bytes()).hexdigest() == ref['sha256'], ref['path']
        return path

    manifest_path = bound(applied['manifest_ref'])
    manifest = read_json(manifest_path)
    version = prepared['version_id']
    assert applied['applied_version'] == applied['current_version'] == manifest['version_id'] == version
    assert read_json(project / '.ai-sow-lite/current.json') == dict(
        version_id=version, manifest_hash=applied['manifest_ref']['sha256'])
    files = {Path(ref['path']).name: bound(ref) for ref in manifest['files']}
    expected_files = {'model.json', 'pending-items.json', 'decisions.json', 'projection.json', 'sow.xlsx',
            'summary.md', 'pending-items.md', 'details.md', 'verification.json', 'input-records.json'}
    if clarify:
        expected_files.update({'plan.json', 'shown-plan.json', 'candidate.json', 'prepared.json', 'confirmation.json'})
    assert expected_files == files.keys()
    for name, key in [('model.json', 'model_path'), ('pending-items.json', 'pending_items_path'),
                      ('decisions.json', 'decisions_path')]:
        assert files[name].read_bytes() == (project / candidate[key]).read_bytes()
    assert files['sow.xlsx'].read_bytes() == bound(prepared['workbook_ref']).read_bytes()
    assert version in files['summary.md'].read_text(encoding='utf-8')
    assert version in files['pending-items.md'].read_text(encoding='utf-8')
    assert version in files['details.md'].read_text(encoding='utf-8')
    assert prepared['pending_count'] == 1
    projection = read_json(files['projection.json'])
    assert projection['version_id'] == version
    assert projection['workbook_hash'] == applied['workbook_ref']['sha256']
    assert manifest['verification_ref']['path'].endswith(f'/{version}/verification.json')
    bound(manifest['verification_ref'])
    verification = read_json(files['verification.json'])
    receipt = verification['office']
    assert verification['version_id'] == version and receipt['exit_code'] == 0
    assert receipt['final_hash'] == applied['workbook_ref']['sha256']
    assert receipt['office_identity'] == prepared['office_identity'] == manifest['office_identity']
    assert receipt['engine']['name'] == 'LibreOffice' and receipt['elapsed_ms'] > 0

    # Enumerate references from genuine registration files, not the generated check report.
    expected = {}
    counts = Counter(identity=0, originals=0, readings=0, analysis=0, template=0, observations=0, history=0)

    def expect(ref, category):
        bound(ref)
        if ref['path'] not in expected:
            counts[category] += 1
        expected[ref['path']] = ref['sha256']

    identity = project / '.ai-sow-lite/project.json'
    expect(dict(path='.ai-sow-lite/project.json', sha256=hashlib.sha256(identity.read_bytes()).hexdigest()), 'identity')
    records = read_json(files['input-records.json'])['items']
    assert {entry['input_version_id'] for entry in records} == set(candidate['input_version_ids'])
    if clarify:
        from ai_sow_lite.contracts import semantic_digest
        plan = read_json(files['plan.json'])
        proof = read_json(files['confirmation.json'])
        assert files['plan.json'].read_bytes() == (project / confirmation_path).read_bytes()
        assert files['shown-plan.json'].read_bytes() == bound(plan['confirmation']['shown_plan_ref']).read_bytes()
        assert files['candidate.json'].read_bytes() == case.candidate_path.read_bytes()
        assert files['prepared.json'].read_bytes() == bound(prepared['prepared_ref']).read_bytes()
        assert bound(proof['plan_ref']) == files['plan.json']
        assert bound(proof['shown_plan_ref']) == files['shown-plan.json']
        assert proof['source_ref'] == plan['confirmation']['input_ref']
        assert proof['selected_changes'] == plan['changes']
        content = {key: plan[key] for key in ('plan_id', 'revision', 'base_version_id', 'changes', 'read_set',
            'write_set', 'read_boundary', 'conditions', 'unresolved_items', 'change_summary')}
        assert proof['digest'] == semantic_digest(content) == plan['confirmation']['digest']
        assert proof['input_record']['input_version_id'] == proof['source_ref']['input_version_id']
        source = proof['input_record']
        expect(dict(path=source['relative_path'], sha256=source['content_hash']), 'originals')
        assert manifest['base_version_id'] == read_json(bound(prepared['prepared_ref']))['expected_current']['version_id']
        assert previous_files
        for relative, raw in previous_files.items():
            expect(dict(path=str(relative), sha256=hashlib.sha256(raw).hexdigest()), 'history')
            if str(relative).endswith('/input-records.json'):
                records.extend(json.loads(raw)['items'])
            elif str(relative).endswith('/confirmation.json'):
                source = json.loads(raw)['input_record']
                expect(dict(path=source['relative_path'], sha256=source['content_hash']), 'originals')
    else:
        assert len(records) == 3
        assert manifest['base_version_id'] is None
    for entry in records:
        expect(dict(path=entry['relative_path'], sha256=entry['content_hash']), 'originals')
        reading_ref = read_json((project / entry['relative_path']).parent / 'reading-ref.json')
        expect(reading_ref, 'readings')
        for excerpt in read_json(bound(reading_ref))['excerpts']:
            expect(excerpt['file_ref'], 'readings')
    assert manifest['input_version_ids'] == candidate['input_version_ids']
    assert manifest['topic_version_ids'] == candidate['topic_version_ids']
    assert manifest['evidence_ids'] == candidate['evidence_ids']
    topics = set(candidate['topic_version_ids'])
    for relative, raw in (previous_files or {}).items():
        if str(relative).endswith('/manifest.json'):
            topics.update(json.loads(raw)['topic_version_ids'])
    for topic in topics:
        path = project / '.ai-sow-lite/analysis/topics' / topic / 'analysis.json'
        expect(dict(path=path.relative_to(project).as_posix(),
                    sha256=hashlib.sha256(path.read_bytes()).hexdigest()), 'analysis')
        registration = read_json(path.parent / 'registration-ref.json')
        expect(registration, 'analysis')
        assert read_json(path)['observations'] == []  # Observation registration is I4.
    assert read_json(files['model.json'])['lineage'] == []
    template = project / '.ai-sow-lite/template' / case.template_hash / 'sow-template.xlsx'
    expect(dict(path=template.relative_to(project).as_posix(), sha256=case.template_hash), 'template')
    assert hashlib.sha256((PLUGIN / 'assets/sow-template.xlsx').read_bytes()).hexdigest() == case.template_hash
    assert {ref['path']: ref['sha256'] for ref in manifest['dependencies']} == expected, 'Delivered dependencies differ from registered sources and prior files'
    for ref in manifest['dependencies']:
        bound(ref)

    # Independent input/structure/cache checks; do not calculate or gate monetary totals.
    import openpyxl
    workbook = openpyxl.load_workbook(files['sow.xlsx'])
    cached = openpyxl.load_workbook(files['sow.xlsx'], data_only=True)
    try:
        assert workbook.sheetnames == ['01-需求故事', '02-任务清单', '03-工作量汇总', '90-估算标准', '04-待确认事项', '05-完整说明']
        assert workbook['01-需求故事']['C5'].value == '资料查询'
        assert workbook['02-任务清单']['E10'].value == 'M'
        assert workbook['02-任务清单']['F5'].value is None
        assert workbook['02-任务清单']['J5'].data_type == 'f'
        assert cached['02-任务清单']['J5'].value is not None
        assert all(workbook[name].protection.sheet for name in workbook.sheetnames[:4])
        for name in workbook.sheetnames[4:]:
            assert workbook[name].sheet_state == 'visible'
            assert not workbook[name].protection.sheet
            assert all(cell.data_type != 'f' for row in workbook[name] for cell in row)
        assert sum(len(sheet.tables) for sheet in workbook) == 5
    finally:
        workbook.close()
        cached.close()
    return dict(template_hash=case.template_hash, dependency_counts=dict(counts),
                confirmation='registered-source-bound' if clarify else None,
                pending_count=prepared['pending_count'], details=True, office=receipt['engine'],
                office_elapsed_ms=receipt['elapsed_ms'])


def run_clarify_delivery(original):
    """Read the actual delivery, then confirm one finite notes edit as test controller."""
    from copy import deepcopy
    from uuid import uuid4
    from ai_sow_lite.project import ensure_request
    from tests.support.clarify import prepare_notes_change
    from tests.support.cli import run_request
    from tests.support.fixtures import Case, read_json
    project = original.project
    before = {p.relative_to(project).as_posix(): p.read_bytes()
              for p in (project / '.ai-sow-lite/versions').rglob('*') if p.is_file()}
    current = read_json(project / '.ai-sow-lite/current.json')
    base = project / '.ai-sow-lite/versions' / current['version_id']
    # Actual offline files are read before discussion and the controller's confirmation.
    old_model = read_json(base / 'model.json')
    assert '资料查询' in (base / 'details.md').read_text(encoding='utf-8')
    request_id = str(uuid4())
    ensure_request(project, request_id, 'clarify')
    case = dict(project=project, ids=original.ids, request_id=request_id, current=current)
    notes = '沿用当前验收范围，保留复核记录'
    prepared = prepare_notes_change(case, notes)
    reused = run_request(project, request_id, 'render', dict(candidate_path=case['result']['candidate_ref']['path'],
        check_path=case['result']['check_ref']['path'], expected_current=current))
    assert reused['ok'] and reused['result']['prepared_ref'] == prepared['prepared_ref'], reused
    payload = dict(entrypoint='clarify', prepared_path=case['prepared_ref']['path'],
                   plan_path=case['confirmed_path'], expected_current=current)
    applied = run_request(project, request_id, 'apply', payload)
    assert applied['ok'], applied
    consumer = Case(project, request_id, project / case['result']['candidate_ref']['path'],
                    original.template_hash, original.ids)
    result = verify_delivery(consumer, prepared, applied['result'],
                             confirmation_path=case['confirmed_path'], previous_files=before)
    latest = project / '.ai-sow-lite/versions' / applied['result']['applied_version']
    expected = deepcopy(old_model)
    story = next(s for s in expected['stories'] if s['id'] == original.ids['S-01'])
    story['notes'] = notes
    assert read_json(latest / 'model.json') == expected
    for name in ('pending-items.json', 'decisions.json'):
        assert (latest / name).read_bytes() == (base / name).read_bytes()
    import openpyxl
    workbook = openpyxl.load_workbook(latest / 'sow.xlsx', read_only=True)
    try:
        assert workbook['01-需求故事']['E5'].value.splitlines()[0] == notes
    finally:
        workbook.close()
    pointer = (project / '.ai-sow-lite/current.json').read_bytes()
    repeated = run_request(project, request_id, 'apply', payload)
    assert repeated['ok'] and repeated['result']['idempotent'], repeated
    recovered = run_request(project, request_id, 'recover', dict(target_request_id=request_id))
    assert recovered['ok'] and recovered['result']['applied_version'] == applied['result']['applied_version'], recovered
    assert (project / '.ai-sow-lite/current.json').read_bytes() == pointer
    assert all((project / relative).read_bytes() == raw for relative, raw in before.items())
    return dict(result, base_version=current['version_id'], applied_version=applied['result']['applied_version'],
                preview_reused=True, idempotent=True, recovered='applied', unrelated_bytes_preserved=True)


def run_delivery(project):
    from tests.support.cli import run_request
    from tests.support.fixtures import build_ingested_case, prepare_case, read_json
    case = build_ingested_case(project)
    prepared = prepare_case(case)
    assert not (project / '.ai-sow-lite/current.json').exists()
    payload = dict(entrypoint='generate', prepared_path=prepared['prepared_ref']['path'],
                   expected_current=None, plan_path=None)
    first = run_request(project, case.request_id, 'apply', payload)
    assert first['ok'], first
    result = verify_delivery(case, prepared, first['result'])
    before = {p.relative_to(project): p.read_bytes() for p in (project / '.ai-sow-lite/versions').rglob('*') if p.is_file()}
    current = (project / '.ai-sow-lite/current.json').read_bytes()
    again = run_request(project, case.request_id, 'apply', payload)
    assert again['ok'] and again['result']['idempotent'], again
    assert again['result']['applied_version'] == first['result']['applied_version']
    recovered = run_request(project, case.request_id, 'recover', dict(target_request_id=case.request_id))
    assert recovered['ok'] and recovered['result']['state'] == 'applied', recovered
    assert recovered['result']['applied_version'] == first['result']['applied_version']
    observed = run_request(project, case.request_id, 'inspect', dict(
        view='telemetry', selector=dict(request_id=case.request_id), limit=100, cursor=None))
    assert observed['ok'], observed
    report_ref = observed['result']['report_ref']
    report_path = project / report_ref['path']
    assert hashlib.sha256(report_path.read_bytes()).hexdigest() == report_ref['sha256']
    metrics = read_json(report_path)['metrics']
    tokens = next(m for m in metrics if m['name'] == 'total_tokens')
    duration = next(m for m in metrics if m['name'] == 'tool_duration_ns')
    assert tokens['value'] is None and tokens['coverage'] == 'unknown'
    assert duration['value'] > 0
    assert before == {p.relative_to(project): p.read_bytes() for p in (project / '.ai-sow-lite/versions').rglob('*') if p.is_file()}
    assert (project / '.ai-sow-lite/current.json').read_bytes() == current
    clarification = run_clarify_delivery(case)
    latest = (project / '.ai-sow-lite/current.json').read_bytes()
    old_recovered = run_request(project, case.request_id, 'recover', dict(target_request_id=case.request_id))
    old_repeated = run_request(project, case.request_id, 'apply', payload)
    assert old_recovered['ok'] and old_recovered['result']['applied_version'] == first['result']['applied_version']
    assert old_repeated['ok'] and old_repeated['result']['idempotent']
    assert old_repeated['result']['current_version'] == clarification['applied_version']
    assert (project / '.ai-sow-lite/current.json').read_bytes() == latest
    return dict(result, idempotent=True, recovered='applied', tool_duration_ns=duration['value'], total_tokens=None,
                clarify=clarification, older_request_preserves_current=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--copy-plugin', action='store_true', help='在独立副本创建锁定环境并验收交付')
    parser.add_argument('--worker', action='store_true', help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.worker:
        assert Path(sys.prefix).resolve() == PLUGIN / '.venv'
        assert Path.cwd().resolve() != PLUGIN
        print(json.dumps(run_delivery(Path(os.environ['LITE_SMOKE_WORKSPACE']) / '项目 with spaces'), ensure_ascii=False))
        return 0
    workspace = Path(tempfile.mkdtemp(prefix='lite-copy-smoke-')).resolve()
    try:
        plugin = workspace / '插件 with spaces' if args.copy_plugin else PLUGIN
        if args.copy_plugin:
            shutil.copytree(PLUGIN, plugin, ignore=shutil.ignore_patterns('.venv', '__pycache__', '*.pyc',
                            '.pytest_cache', '.runtime', '.git'))
        environment = dict(os.environ)
        for name in ('PYTHONPATH', 'PYTHONHOME', 'VIRTUAL_ENV', 'UV_PROJECT_ENVIRONMENT'):
            environment.pop(name, None)
        temporary = workspace / 'temporary'
        temporary.mkdir()
        environment.update(TMPDIR=str(temporary), TMP=str(temporary), TEMP=str(temporary))
        uv = shutil.which('uv')
        assert uv, 'uv is required for the locked smoke environment'
        synced = subprocess.run([uv, 'sync', '--project', str(plugin), '--locked'], cwd=workspace,
                                env=environment, capture_output=True, text=True, timeout=120)
        assert synced.returncode == 0, synced.stderr
        audit = workspace / 'audit'
        audit.mkdir()
        (audit / 'sitecustomize.py').write_text(
            'from tests.support.smoke_plugin import install_read_audit\ninstall_read_audit()\n', encoding='utf-8')
        environment.update(PYTHONPATH=os.pathsep.join(map(str, [audit, plugin / 'runtime', plugin])),
                           PYTHONDONTWRITEBYTECODE='1', PYTHONNOUSERSITE='1', LITE_SMOKE_WORKSPACE=str(workspace))
        python = plugin / '.venv' / ('Scripts/python.exe' if os.name == 'nt' else 'bin/python')
        # Import beside this script: the coordinator need not have plugin runtime on sys.path.
        from process_audit import reconcile_audits, run_process
        process = run_process([str(python), str(plugin / 'tests/support/smoke_plugin.py'), '--worker'],
                              operation='worker', cwd=workspace, env=environment,
                              capture_output=True, text=True, encoding='utf-8', timeout=180)
        assert process.returncode == 0, (process.stdout, process.stderr)
        assert not process.stderr, process.stderr
        delivery = json.loads(process.stdout)
        observed = reconcile_audits(workspace)
        assert observed['violations'] == 0
        assert observed['office_conversions'] == 2  # apply/retry/recover must never recalculate.
        shutil.rmtree(workspace)
        print(json.dumps(dict(ok=True, copied=args.copy_plugin, locked_environment=True, cleaned=True,
                              delivery=delivery, audit=observed), ensure_ascii=False, indent=2))
        return 0
    except Exception as error:
        print(json.dumps(dict(ok=False, retained_path=str(workspace), error=str(error)), ensure_ascii=False))
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
