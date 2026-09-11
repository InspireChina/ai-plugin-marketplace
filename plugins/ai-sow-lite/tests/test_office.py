"""Real Office tests are marked; child-process fixtures only test failure boundaries."""
import importlib
import os
from pathlib import Path
import shutil
import sys
import time
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from ai_sow_lite.project import StorageError
from .test_workbook import project, module as workbook


def office():
    try:
        return importlib.import_module('ai_sow_lite.office')
    except ModuleNotFoundError:
        pytest.fail('I1.3 Office adapter implementation is missing')


def mutate_zip(source, destination, change):
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(destination,'w') as dst:
        for name in src.namelist():
            dst.writestr(name,change(name,src.read(name)))


def _engine_available():
    """Use the adapter's own discovery: Windows never puts LibreOffice on PATH."""
    try:
        office().discover_engine()
        return True
    except StorageError:
        return False


@pytest.mark.office
def test_real_office_normalizes_only_known_omissions_then_readonly_seals(tmp_path):
    if not _engine_available(): pytest.skip('Real Office engine unavailable')
    project(tmp_path)
    source=tmp_path/'projected.xlsx'; before=source.read_bytes()
    target=tmp_path/'sow.xlsx'
    result=office().recalculate(source,target)
    assert source.read_bytes()==before
    assert result['exit_code']==0 and result['engine']['name']=='LibreOffice'
    assert result['raw_hash'] and result['final_hash']
    assert result['engine']['version'].startswith('LibreOffice ')
    assert str(tmp_path) not in str(result)
    raw=tmp_path/'sow.office-raw.xlsx'
    assert raw.is_file()
    assert not list(tmp_path.glob('.office-*'))
    inventory=workbook().audit_workbook(target,source)
    assert inventory['formula_count']==1304
    assert inventory['cache_count']==1304
    # Original amount formula is preserved; no numeric completeness expectations.
    w=openpyxl.load_workbook(target); cached=openpyxl.load_workbook(target,data_only=True)
    assert w['02-任务清单']['J5'].value=='=IF(OR(NOT(ISNUMBER($H5)),NOT(ISNUMBER($I5))),"",ROUND($H5*$I5,1))'
    assert cached['01-需求故事']['C5'].value=='资料查询'
    assert '资料查询页及公共客户端接入' in cached['01-需求故事']['H5'].value
    assert w['02-任务清单'].tables['TaskTable'].tableColumns[7].calculatedColumnFormula is not None
    assert str(w['02-任务清单'].data_validations.dataValidation[0].sqref)=='A5:A1048576'
    assert workbook().formula_cache_inventory(raw)==workbook().formula_cache_inventory(target)
    saved=target.read_bytes(); workbook().audit_workbook(target,source)
    assert target.read_bytes()==saved


@pytest.mark.office
@pytest.mark.parametrize('damage',['formula','cache','table','protection','validation_rule','validation_range','table_formula'])
def test_raw_damage_cannot_be_repaired_by_compatibility_adapter(tmp_path,damage):
    if not _engine_available(): pytest.skip('Real Office engine unavailable')
    project(tmp_path); office().recalculate(tmp_path/'projected.xlsx',tmp_path/'sow.xlsx')
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    def change(name,raw):
        if name=='xl/worksheets/sheet1.xml':
            root=ET.fromstring(raw)
            if damage=='formula': root.find('.//m:c[@r="I5"]/m:f',ns).text='999'
            if damage=='cache':
                c=root.find('.//m:c[@r="I5"]',ns); c.remove(c.find('m:v',ns))
            if damage=='protection': root.remove(root.find('m:sheetProtection',ns))
            if damage=='validation_rule': root.find('.//m:dataValidation/m:formula1',ns).text='1=1'
            if damage=='validation_range': root.find('.//m:dataValidation',ns).set('sqref','C5:C1063')
            return ET.tostring(root)
        if name.startswith('xl/tables/'):
            root=ET.fromstring(raw)
            if root.get('name')=='SOWStoryTable':
                if damage=='table': root.set('ref','A4:J63')
                if damage=='table_formula':
                    col=root.find('m:tableColumns',ns)[5]
                    ET.SubElement(col,'{'+ns['m']+'}calculatedColumnFormula').text='999'
                return ET.tostring(root)
        return raw
    raw=tmp_path/'sow.office-raw.xlsx'; bad=tmp_path/'bad.xlsx'
    mutate_zip(raw,bad,change); before=bad.read_bytes()
    with pytest.raises(StorageError,match='WORKBOOK_INVALID'):
        workbook().seal_office_output(tmp_path/'projected.xlsx',bad,tmp_path/'rejected.xlsx')
    assert bad.read_bytes()==before
    assert not (tmp_path/'rejected.xlsx').exists()


def test_missing_engine_does_not_write_output(tmp_path,monkeypatch):
    adapter=office(); monkeypatch.delenv('AI_SOW_LITE_OFFICE_BIN',raising=False)
    monkeypatch.setattr(adapter.shutil,'which',lambda _:None)
    # Windows also probes the default install location; absence must be total.
    for variable in ('ProgramFiles','ProgramFiles(x86)'): monkeypatch.delenv(variable,raising=False)
    project(tmp_path)
    with pytest.raises(StorageError,match='OFFICE_ENGINE_UNAVAILABLE'):
        adapter.recalculate(tmp_path/'projected.xlsx',tmp_path/'sow.xlsx')
    assert not (tmp_path/'sow.xlsx').exists()


@pytest.mark.skipif(os.name=='nt',reason='os.kill(pid,0) liveness probe is POSIX-only')
def test_timeout_cleans_only_owned_process_tree(tmp_path):
    adapter=office()
    child_pid=tmp_path/'child.pid'
    code='import subprocess,time,pathlib,sys; p=subprocess.Popen([sys.executable,"-c","import time; time.sleep(30)"]); pathlib.Path(sys.argv[1]).write_text(str(p.pid)); time.sleep(30)'
    start=time.monotonic()
    with pytest.raises(StorageError,match='CALCULATION_FAILED'):
        adapter._run([sys.executable,'-c',code,str(child_pid)],timeout=.4,environment=os.environ.copy())
    assert time.monotonic()-start<5
    pid=int(child_pid.read_text())
    for _ in range(30):
        try: os.kill(pid,0)
        except ProcessLookupError: break
        time.sleep(.05)
    else: pytest.fail('Owned child survived timeout')


@pytest.mark.skipif(os.name=='nt',reason='Shebang + chmod fake engine is POSIX-only')
def test_exit_zero_without_output_is_failure_and_cleans_profile(tmp_path,monkeypatch):
    adapter=office(); project(tmp_path)
    for variable in ('ProgramFiles','ProgramFiles(x86)'): monkeypatch.delenv(variable,raising=False)
    script=tmp_path/'fake-office'
    script.write_text('#!'+sys.executable+'\nimport sys\nif "--version" in sys.argv: print("LibreOffice 0.0 test")\n')
    script.chmod(0o700)
    monkeypatch.setenv('AI_SOW_LITE_OFFICE_BIN',str(script))
    with pytest.raises(StorageError,match='CALCULATION_FAILED'):
        adapter.recalculate(tmp_path/'projected.xlsx',tmp_path/'sow.xlsx')
    assert not list(tmp_path.glob('.office-*'))
    assert not (tmp_path/'sow.xlsx').exists()


@pytest.mark.office
@pytest.mark.parametrize('story_count,task_count,last_s,last_t,formulas',[(0,0,64,204,1304),(1,1,64,204,1304),(60,200,64,204,1304),(61,201,65,205,1314)])
def test_real_office_capacity_inputs_array_refs_and_cached_storage(tmp_path,story_count,task_count,last_s,last_t,formulas):
    if not _engine_available(): pytest.skip('Real Office engine unavailable')
    from .test_workbook import bundle
    from copy import deepcopy
    from uuid import uuid4
    model,pending,decisions=bundle(); s0=model['stories'][0];t0=model['tasks'][0]
    model['stories']=[dict(deepcopy(s0),id=str(uuid4()),title=f'故事 {i}') for i in range(story_count)]
    for story in model['stories']:
        for ac in story['acs']:ac['id']=str(uuid4())
    model['tasks']=[dict(deepcopy(t0),id=str(uuid4()),name=f'任务 {i}',story_id=model['stories'][i%story_count]['id']) for i in range(task_count)]
    model['dependencies']=[]
    literals=['=SUM(A1:A9)','+原文','-原文','@原文',"'真实前缀",'#N/A']
    for task,text in zip(model['tasks'],literals):task['notes']=text
    project(tmp_path,model,dict(schema_version='1.0',items=[]),decisions)
    result=office().recalculate(tmp_path/'projected.xlsx',tmp_path/'sow.xlsx')
    assert result['verification']['formula_count']==formulas
    w=openpyxl.load_workbook(tmp_path/'sow.xlsx'); cached=openpyxl.load_workbook(tmp_path/'sow.xlsx',data_only=True)
    assert w['01-需求故事'].tables['SOWStoryTable'].ref==f'A4:J{last_s}'
    assert w['02-任务清单'].tables['TaskTable'].ref==f'A4:L{last_t}'
    assert w['01-需求故事'][f'H{last_s}'].value.ref==f'H{last_s}'
    assert w['02-任务清单'][f'L{last_t}'].protection.locked
    by_name={w['02-任务清单'].cell(r,2).value:r for r in range(5,last_t+1)}
    for index,text in enumerate(literals[:task_count]):
        row=by_name[f'任务 {index}']; cell=w['02-任务清单'].cell(row,7)
        assert cell.data_type=='s' and cell.value==text
        assert cached['02-任务清单'].cell(row,7).value==text
    assert workbook().formula_cache_inventory(tmp_path/'sow.office-raw.xlsx')==workbook().formula_cache_inventory(tmp_path/'sow.xlsx')


def test_style_tint_damage_cannot_pass_metadata_reread(tmp_path):
    project(tmp_path)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    from copy import copy
    fill=copy(w['01-需求故事']['A4'].fill)
    fill.fgColor.tint=.5
    w['01-需求故事']['A4'].fill=fill
    w.save(tmp_path/'tinted.xlsx')
    with pytest.raises(StorageError,match='WORKBOOK_INVALID'):
        workbook().audit_workbook(tmp_path/'tinted.xlsx',tmp_path/'projected.xlsx',caches=False)


@pytest.mark.office
@pytest.mark.parametrize('damage',['hidden_row','font_name','font_scheme','font_vertical','font_charset','collapsed_row'])
def test_raw_display_damage_is_rejected_without_restoration(tmp_path,damage):
    if not _engine_available(): pytest.skip('Real Office engine unavailable')
    project(tmp_path)
    office().recalculate(tmp_path/'projected.xlsx',tmp_path/'sow.xlsx')
    raw=tmp_path/'sow.office-raw.xlsx'; bad=tmp_path/'unreadable.xlsx'
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    def change(name,payload):
        if name=='xl/worksheets/sheet1.xml' and damage in ('hidden_row','collapsed_row'):
            root=ET.fromstring(payload)
            root.find(f'{ns}sheetData/{ns}row[@r="5"]').set('hidden' if damage=='hidden_row' else 'collapsed','1')
            return ET.tostring(root)
        if name=='xl/styles.xml' and damage.startswith('font_'):
            root=ET.fromstring(payload)
            tag,value={'font_name':('name','Wingdings'),'font_scheme':('scheme','major'),
                       'font_vertical':('vertAlign','superscript'),'font_charset':('charset','2')}[damage]
            for font in root.find(ns+'fonts'):
                node=font.find(ns+tag)
                if node is None: node=ET.SubElement(font,ns+tag)
                node.set('val',value)
            return ET.tostring(root)
        return payload
    mutate_zip(raw,bad,change);before=bad.read_bytes()
    with pytest.raises(StorageError,match='WORKBOOK_INVALID'):
        workbook().seal_office_output(tmp_path/'projected.xlsx',bad,tmp_path/'rejected.xlsx')
    assert bad.read_bytes()==before
    assert not (tmp_path/'rejected.xlsx').exists()


@pytest.mark.office
@pytest.mark.parametrize('damage',['business_pair','other_a2_pair','a2_italic','a2_text','a2_theme'])
def test_instructional_font_fallback_cannot_admit_other_display_changes(tmp_path,damage):
    if not _engine_available(): pytest.skip('Real Office engine unavailable')
    project(tmp_path)
    office().recalculate(tmp_path/'projected.xlsx',tmp_path/'sow.xlsx')
    raw=tmp_path/'sow.office-raw.xlsx';bad=tmp_path/'bad-fallback.xlsx'
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(raw) as z:
        sheet=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        cell=sheet.find(f'.//{ns}c[@r="'+('C5' if damage=='business_pair' else 'A2')+'"]')
        styles=ET.fromstring(z.read('xl/styles.xml'))
        font_id=int(styles.find(ns+'cellXfs')[int(cell.get('s'))].get('fontId'))
        a2=sheet.find(f'.//{ns}c[@r="A2"]')
        assert a2.get('t')=='s'
        string_index=int(a2.find(ns+'v').text)
    def change(name,payload):
        if name=='xl/styles.xml' and damage in ('business_pair','other_a2_pair','a2_italic'):
            root=ET.fromstring(payload);font=root.find(ns+'fonts')[font_id]
            if damage=='a2_italic': ET.SubElement(font,ns+'i').set('val','1')
            else: font.find(ns+'name').set('val','Arial Unicode MS' if damage=='business_pair' else 'Microsoft YaHei')
            return ET.tostring(root)
        if name=='xl/sharedStrings.xml' and damage=='a2_text':
            root=ET.fromstring(payload);root[string_index].find('.//'+ns+'t').text='已改写的模板说明'
            return ET.tostring(root)
        if name=='xl/theme/theme1.xml' and damage=='a2_theme':
            root=ET.fromstring(payload);a='{http://schemas.openxmlformats.org/drawingml/2006/main}'
            root.find('.//'+a+'minorFont/'+a+'font[@script="Hans"]').set('typeface','Wingdings')
            return ET.tostring(root)
        return payload
    mutate_zip(raw,bad,change)
    with pytest.raises(StorageError,match='WORKBOOK_INVALID'):
        workbook().seal_office_output(tmp_path/'projected.xlsx',bad,tmp_path/'rejected.xlsx')
    assert not (tmp_path/'rejected.xlsx').exists()
