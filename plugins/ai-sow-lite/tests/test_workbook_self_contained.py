"""Single-file delivery: no appendix dependency or stale pending flags."""
from copy import deepcopy
from uuid import uuid4

import openpyxl
import pytest
from openpyxl.worksheet.hyperlink import Hyperlink

from ai_sow_lite.project import StorageError
from ai_sow_lite.workbook import audit_workbook
from .test_workbook import bundle, project


def assert_internal_delivery(book):
    assert book.sheetnames==['01-需求故事','02-任务清单','03-工作量汇总','90-估算标准']
    for sheet in book:
        assert sheet.sheet_state=='visible'
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    assert 'details.md' not in cell.value and 'pending-items.md' not in cell.value
                if cell.hyperlink:
                    assert cell.hyperlink.target is None


def test_xlsx_alone_preserves_long_ac_notes_and_pending(tmp_path):
    model,pending,decisions=bundle()
    raw='验收原文😀\n'*8+'最终验收条件。'
    model['stories'][0]['acs'][0]['text']=raw
    model['tasks'][0]['notes']='=不是公式\n'+'备注全文。'*15
    pending['items'][0]['question']='需要决定的真实问题。'*2
    pending['items'][0]['current_handling']='采用M但记录数量仍未知。'*2
    project(tmp_path,model,pending,decisions)
    for path in tmp_path.glob('*.md'):path.unlink()
    book=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert_internal_delivery(book)
    assert raw in book['01-需求故事']['D5'].value
    assert model['stories'][0]['acs'][1]['text'] in book['01-需求故事']['D5'].value
    assert model['tasks'][0]['notes']==book['02-任务清单']['G5'].value
    assert pending['items'][0]['question'] in book['02-任务清单']['G10'].value
    assert pending['items'][0]['current_handling'] in book['02-任务清单']['G10'].value


def test_overlapping_targets_show_one_question_per_row_and_all_ac_positions(tmp_path):
    m,p,d=bundle();story=m['stories'][0]
    p['items'][0]['targets']=[dict(object_id=o['id'],field=None) for o in [story,*story['acs'],m['features'][0]]]
    project(tmp_path,m,p,d)
    book=openpyxl.load_workbook(tmp_path/'projected.xlsx');note=book['01-需求故事']['E5'].value
    assert note.count(p['items'][0]['question'])==1
    assert '第 1 条 AC' in note and '第 2 条 AC' in note
    assert_internal_delivery(book)


def test_multiple_questions_and_authored_exception_are_all_inline(tmp_path):
    m,p,d=bundle();story=m['stories'][0];story['notes']='生产切换由客户执行。'
    p['items'][0]['targets']=[dict(object_id=story['id'],field='notes')]
    second=dict(deepcopy(p['items'][0]),id=str(uuid4()),question='另一项范围是否包含？')
    p['items'].append(second)
    project(tmp_path,m,p,d)
    book=openpyxl.load_workbook(tmp_path/'projected.xlsx');note=book['01-需求故事']['E5'].value
    assert note.startswith('待确认：')
    assert all(item['question'] in note and item['current_handling'] in note for item in p['items'])
    assert story['notes'] in note
    assert_internal_delivery(book)


@pytest.mark.parametrize('prefix,repetitions,overflow', [
    ('\n'*3,10,False), (' '*80,10,False),
    ('\n'*8,150,True), (' '*800,150,True)])
def test_leading_spaces_and_lines_are_preserved_inline(tmp_path,prefix,repetitions,overflow):
    m,p,d=bundle();raw=prefix+'正文。'*repetitions;m['tasks'][0]['notes']=raw
    if overflow:
        with pytest.raises(StorageError, match='WORKBOOK_LAYOUT_OVERFLOW'):
            project(tmp_path,m,p,d)
        assert m['tasks'][0]['notes']==raw
        assert not (tmp_path/'projected.xlsx').exists()
        return
    project(tmp_path,m,p,d)
    book=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert book['02-任务清单']['G5'].value==raw
    audit_workbook(tmp_path/'projected.xlsx',tmp_path/'projected.xlsx',caches=False)


@pytest.mark.parametrize('location',["'01-需求故事'!D99999","'01-需求故事'!D3","'不存在的说明表'!D5"])
def test_audit_rejects_empty_internal_target_even_when_expected_matches(tmp_path,location):
    project(tmp_path)
    book=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    book['01-需求故事']['D5'].hyperlink=Hyperlink(ref='D5',location=location)
    damaged=tmp_path/'damaged.xlsx';book.save(damaged)
    with pytest.raises(StorageError,match='WORKBOOK_INVALID'):audit_workbook(damaged,damaged,caches=False)


@pytest.mark.office
def test_real_office_preserves_inline_content_pending_status_and_literal_text(tmp_path):
    from ai_sow_lite import office
    m,p,d=bundle()
    for obj in [*m['stories'],*m['tasks']]:obj['notes']=''
    m['tasks'][0]['notes']='无需待确认；=这仍是文字。'
    second=dict(deepcopy(p['items'][0]),id=str(uuid4()),question='该验收边界是否包含？',
                targets=[dict(object_id=m['stories'][1]['acs'][0]['id'],field='text')])
    p['items'].append(second)
    project(tmp_path,m,p,d)
    office.recalculate(tmp_path/'projected.xlsx',tmp_path/'sow.xlsx')
    book=openpyxl.load_workbook(tmp_path/'sow.xlsx')
    cached=openpyxl.load_workbook(tmp_path/'sow.xlsx',data_only=True)
    assert_internal_delivery(book)
    assert cached['02-任务清单']['L10'].value=='待确认'
    assert cached['01-需求故事']['J6'].value=='待确认'
    assert cached['02-任务清单']['L5'].value=='通过'
    assert cached['01-需求故事']['J5'].value=='通过'
    assert cached['02-任务清单']['L11'].value=='通过'
    assert cached['01-需求故事']['J64'].value is None
    assert cached['02-任务清单']['E10'].value=='M'
    assert cached['02-任务清单']['J10'].value is not None
    assert book['02-任务清单']['L10'].protection.locked
    assert book['01-需求故事']['J6'].protection.locked
    # Same accepted model, now all decisions resolved: normal notes and status.
    p['items']=[]
    output=tmp_path/'resolved';project(output,m,p,d)
    office.recalculate(output/'projected.xlsx',output/'sow.xlsx')
    resolved=openpyxl.load_workbook(output/'sow.xlsx',data_only=True)
    assert resolved['02-任务清单']['G10'].value is None
    assert resolved['02-任务清单']['L10'].value=='通过'
    assert resolved['01-需求故事']['J6'].value=='通过'
