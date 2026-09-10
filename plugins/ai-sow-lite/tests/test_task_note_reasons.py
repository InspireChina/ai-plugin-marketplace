"""Task remarks expose the reasons needed to review non-default classifications."""
from copy import deepcopy

import openpyxl
import pytest

from ai_sow_lite.project import StorageError
from .test_workbook import bundle, project


def test_task_remarks_include_non_new_mode_and_non_m_complexity_reasons(tmp_path):
    model, pending, decisions = bundle()
    task = model['tasks'][0]
    task.update(work_mode='调整', complexity='S', notes='')
    basis = task['classification_basis'][0]
    task['classification_basis'] = [
        dict(basis, fields=['work_mode'], rationale='沿用既有查询 API，本期只增加一个筛选条件。'),
        dict(basis, fields=['complexity'], rationale='单对象精确查询，无新增权限分支，按标准归 S。'),
        dict(basis, fields=['work_type_name'], rationale='这是查询 API 类型。'),
    ]
    before = deepcopy(model)
    projection = project(tmp_path, model, pending, decisions)
    book = openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert book['02-任务清单']['G5'].value == (
        '工作方式（调整）判断原因：沿用既有查询 API，本期只增加一个筛选条件。\n'
        '复杂度（S）判断原因：单对象精确查询，无新增权限分支，按标准归 S。')
    book.close()
    record = next(obj for obj in projection['objects'] if obj['object_id'] == task['id'])
    assert next(field for field in record['fields'] if field['field'] == 'classification_basis')['cells'] == [
        {'sheet': '02-任务清单', 'cell': 'G5'}]
    assert model == before


def test_shared_reason_is_written_once_with_both_applicable_labels(tmp_path):
    model, pending, decisions = bundle()
    task = model['tasks'][0]
    task.update(work_mode='接入复用', complexity='L', notes='')
    basis = task['classification_basis'][0]
    task['classification_basis'] = [
        dict(basis, fields=['work_mode'], rationale='复用现有能力，但需适配多条已明确的复杂业务分支。'),
        dict(basis, fields=['complexity', 'work_mode'], rationale='复用现有能力，但需适配多条已明确的复杂业务分支。'),
    ]
    project(tmp_path, model, pending, decisions)
    book = openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert book['02-任务清单']['G5'].value == (
        '工作方式（接入复用）、复杂度（L）判断原因：复用现有能力，但需适配多条已明确的复杂业务分支。')
    book.close()


@pytest.mark.parametrize('mode,complexity,expected', [
    ('新建', 'M', None),
    (None, 'M', None),
    ('新建', 'S', '复杂度（S）判断原因：已有材料支持该分类。'),
    ('新建', 'L', '复杂度（L）判断原因：已有材料支持该分类。'),
    ('调整', 'M', '工作方式（调整）判断原因：已有材料支持该分类。'),
    ('接入复用', 'M', '工作方式（接入复用）判断原因：已有材料支持该分类。'),
])
def test_only_applicable_non_default_reasons_are_displayed(tmp_path, mode, complexity, expected):
    model, pending, decisions = bundle()
    task = model['tasks'][0]
    task.update(work_mode=mode, complexity=complexity, notes='')
    task['classification_basis'][0]['rationale'] = '已有材料支持该分类。'
    project(tmp_path, model, pending, decisions)
    book = openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert book['02-任务清单']['G5'].value == expected
    book.close()


@pytest.mark.parametrize('rationale', ['😀' * 17000, '原因\x00正文'])
def test_unwritable_reason_points_to_classification_basis(tmp_path, rationale):
    model, pending, decisions = bundle()
    task = model['tasks'][0]
    task['complexity'] = 'S'
    task['classification_basis'][0]['rationale'] = rationale
    with pytest.raises(StorageError) as caught:
        project(tmp_path, model, pending, decisions)
    target = caught.value.diagnostics[0]['target']
    assert target['object_id'] == task['id'] and target['field'] == 'classification_basis'
    assert not (tmp_path/'projected.xlsx').exists()


@pytest.mark.office
def test_pending_and_reasons_coexist_and_only_open_questions_change_status(tmp_path):
    from ai_sow_lite import office
    model, pending, decisions = bundle()
    for obj in model['tasks']:
        obj['notes'] = ''
    task = model['tasks'][0]
    task.update(work_mode='调整', complexity='S')
    task['classification_basis'][0]['rationale'] = '基于既有页面调整一个展示字段。'
    pending['items'][0].update(question='是否保留现有字段名称？', current_handling='暂保留现有名称。',
        targets=[dict(object_id=task['id'], field='name')])
    project(tmp_path, model, pending, decisions)
    office.recalculate(tmp_path/'projected.xlsx', tmp_path/'sow.xlsx')
    book = openpyxl.load_workbook(tmp_path/'sow.xlsx', data_only=True)
    assert book['02-任务清单']['G5'].value == (
        '待确认：是否保留现有字段名称？\n当前处理：暂保留现有名称。\n'
        '工作方式（调整）、复杂度（S）判断原因：基于既有页面调整一个展示字段。')
    assert book['02-任务清单']['L5'].value == '待确认'
    assert '复杂度（S）判断原因：' in book['02-任务清单']['G6'].value
    assert book['02-任务清单']['L6'].value == '通过'
    assert book['02-任务清单']['G7'].value is None
    assert book['02-任务清单']['L7'].value == '通过'
    book.close()
