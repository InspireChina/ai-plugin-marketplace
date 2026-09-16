"""The original tables alone carry acceptance and unresolved user decisions."""
from copy import copy, deepcopy
from uuid import uuid4

import openpyxl
import pytest

from ai_sow_lite.project import StorageError
from .test_workbook import bundle, project


def test_original_four_sheets_complete_acs_and_normally_empty_notes(tmp_path):
    m,p,d=bundle()
    for obj in [*m['stories'],*m['tasks']]: obj['notes']=''
    for task in m['tasks']: task.update(work_mode='新建', complexity='M')
    p['items']=[]
    result=project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert w.sheetnames==['01-需求故事','02-任务清单','03-工作量汇总','90-估算标准']
    assert w['01-需求故事']['D5'].value=='\n'.join(f'{i}. {ac["text"]}' for i,ac in enumerate(m['stories'][0]['acs'],1))
    assert all(w['01-需求故事'][f'E{r}'].value is None for r in range(5,9))
    assert all(w['02-任务清单'][f'G{r}'].value is None for r in range(5,12))
    assert not any(c.hyperlink for s in list(w)[:2] for row in s.iter_rows(min_row=5) for c in row)
    assert result['details']==[] and not (tmp_path/'details.md').exists()


def test_pending_question_and_current_handling_are_inline_without_ids_or_paths(tmp_path):
    m,p,d=bundle()
    m['tasks'][-2]['notes']=''
    project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    text=w['02-任务清单']['G10'].value
    item=p['items'][0]
    assert text.startswith('待确认：')
    assert item['question'] in text and item['current_handling'] in text
    assert item['id'] not in text and '.md' not in text and '见 ' not in text
    assert w['02-任务清单']['L10'].data_type=='f'


@pytest.mark.parametrize('target_kind',['epic','feature','story','ac','task'])
def test_pending_is_on_its_own_rows_and_not_unrelated_rows(tmp_path,target_kind):
    m,p,d=bundle()
    for o in [*m['stories'],*m['tasks']]:o['notes']=''
    s=m['stories'][0]
    obj={'epic':m['epics'][0],'feature':m['features'][0],'story':s,'ac':s['acs'][1],'task':m['tasks'][0]}[target_kind]
    p['items'][0].update(question='该边界是否包含？',current_handling='等待答复。',targets=[dict(object_id=obj['id'],field=None)])
    project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    sheet,col=('02-任务清单','G') if target_kind=='task' else ('01-需求故事','E')
    note=w[sheet][f'{col}5'].value
    assert note.startswith('待确认：') and '该边界是否包含？' in note
    if target_kind=='ac':assert '第 2 条 AC' in note
    assert w['01-需求故事']['E8'].value is None
    assert w['02-任务清单']['G10'].value is None


@pytest.mark.parametrize('collection',['epics','features'])
def test_undivided_scope_has_an_original_table_row_without_inventing_a_story(tmp_path,collection):
    m,p,d=bundle(); original=deepcopy(m)
    parent=dict(deepcopy(m[collection][0]),id=str(uuid4()),title='尚待明确的数据移交范围')
    m[collection].append(parent)
    p['items'][0].update(question='哪些数据需要移交？',current_handling='未拆明，未估算。',unestimated_work=True,
                          targets=[dict(object_id=parent['id'],field=None)])
    result=project(tmp_path,m,p,d)
    record=next(o for o in result['objects'] if o['object_id']==parent['id'])
    assert len(record['rows'])==1
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx');row=record['rows'][0]
    assert w['01-需求故事'][f'C{row}'].value is None
    assert w['01-需求故事'][f'E{row}'].value.startswith('待确认：')
    assert '哪些数据需要移交？' in w['01-需求故事'][f'E{row}'].value
    assert m['stories']==original['stories'] and m['tasks']==original['tasks']
    assert result['pending_items'][0]['targets'][0]['cells']


def test_closed_questions_do_not_keep_rows_pending_and_history_stays_in_project(tmp_path):
    m,p,d=bundle()
    for obj in [*m['stories'],*m['tasks']]:obj['notes']=''
    p['items'][0].update(status='resolved',resolution=dict(decision_id=d['items'][0]['id'],request_id=str(uuid4()),summary='确认沿用M。'))
    project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert w['02-任务清单']['G10'].value is None
    assert p['items'][0]['question'] in (tmp_path/'pending-items.md').read_text()
    assert '已采用答复' in (tmp_path/'pending-items.md').read_text()


def test_many_short_acs_keep_original_text_without_ac_count_gate(tmp_path):
    m,p,d=bundle()
    m['stories'][0]['acs']=[dict(id=str(uuid4()),text=f'{n}项验收。',evidence_refs=m['stories'][0]['evidence_refs']) for n in range(20)]
    project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert w['01-需求故事']['D5'].value=='\n'.join(f'{i}. {ac["text"]}' for i,ac in enumerate(m['stories'][0]['acs'],1))
    assert w['01-需求故事'].row_dimensions[5].height<=409
    assert not (tmp_path/'details.md').exists()


def test_text_beyond_excel_capacity_returns_specific_target_without_partial_xlsx(tmp_path):
    m,p,d=bundle(); m['stories'][0]['acs'][0]['text']='😀'*17000
    with pytest.raises(StorageError) as caught:project(tmp_path,m,p,d)
    assert caught.value.diagnostics[0]['code']=='WORKBOOK_INVALID'
    assert caught.value.diagnostics[0]['target']['object_id']==m['stories'][0]['id']
    assert caught.value.diagnostics[0]['target']['field']=='acs'
    assert not (tmp_path/'projected.xlsx').exists()


@pytest.mark.parametrize('filename', ['sow-template-before-inline.xlsx', 'sow-template-before-task-reasons.xlsx'])
def test_old_project_template_is_unchanged_but_gets_current_inline_status_formulas(tmp_path, filename):
    import hashlib
    from .support.fixtures import FIXTURES,PLUGIN
    from ai_sow_lite.workbook import project_workbook,formula_text
    old=FIXTURES/'workbook'/filename;original=old.read_bytes()
    m,p,d=bundle()
    result=project_workbook(old,m,p,d,str(uuid4()),tmp_path)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    before=openpyxl.load_workbook(old);current=openpyxl.load_workbook(PLUGIN/'assets/sow-template.xlsx')
    assert result['template_hash']==hashlib.sha256(original).hexdigest()
    for sn,col in [('01-需求故事','J'),('02-任务清单','L')]:
        assert w[sn][f'{col}5'].value==current[sn][f'{col}5'].value
        assert w[sn]['A2'].value==current[sn]['A2'].value
        assert next(c for c in w[sn].tables.values()).tableColumns[-1].calculatedColumnFormula.text==w[sn][f'{col}5'].value[1:]
    for sheet in before:
        for row in sheet:
            for cell in row:
                if cell.data_type=='f' and not ((sheet.title,cell.column_letter) in [('01-需求故事','J'),('02-任务清单','L')]):
                    assert formula_text(w[sheet.title][cell.coordinate].value)==formula_text(cell.value)
    assert old.read_bytes()==original


def test_crlf_capacity_is_checked_after_display_newline_normalization(tmp_path):
    m,p,d=bundle();raw='字\r\n'*6
    m['tasks'][0]['notes']=raw
    project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert w['02-任务清单']['G5'].value==raw.replace('\r\n','\n')


@pytest.mark.parametrize('field',['question','current_handling'])
def test_oversize_scope_note_points_to_an_existing_pending_field(tmp_path,field):
    m,p,d=bundle();orphan=dict(deepcopy(m['epics'][0]),id=str(uuid4()),title='未拆明范围')
    m['epics'].append(orphan)
    item=p['items'][0];item.update(targets=[dict(object_id=orphan['id'],field=None)],unestimated_work=True)
    item[field]='字段正文。'*6600
    with pytest.raises(StorageError) as caught:project(tmp_path,m,p,d)
    target=caught.value.diagnostics[0]['target']
    assert target['object_id']==item['id'] and target['field']==field
    assert not (tmp_path/'projected.xlsx').exists()


@pytest.mark.parametrize('collection,field',[('stories','title'),('tasks','name')])
def test_oversize_alias_original_points_to_the_authored_name(tmp_path,collection,field):
    m,p,d=bundle();obj=m[collection][0];obj[field]='长名称'*11000
    with pytest.raises(StorageError) as caught:project(tmp_path,m,p,d)
    assert caught.value.diagnostics[0]['target']['object_id']==obj['id']
    assert caught.value.diagnostics[0]['target']['field']==field


@pytest.mark.parametrize('target', ['acs', 'notes', 'tasks'])
def test_visible_layout_overflow_names_its_original_object_and_cell(tmp_path, target):
    m,p,d=bundle()
    story=m['stories'][0]
    p['items']=[]
    if target == 'acs':
        story['acs'][0]['text']='一行验收。\n'*60
        identity,field,sheet,cell=story['id'],'acs','01-需求故事','D5'
    elif target == 'notes':
        m['tasks'][0]['notes']='一行必要说明。\n'*60
        identity,field,sheet,cell=m['tasks'][0]['id'],'notes','02-任务清单','G5'
    else:
        m['tasks']=[dict(deepcopy(m['tasks'][0]),id=str(uuid4()),name=f'实际独立任务{n}') for n in range(40)]
        identity,field,sheet,cell=story['id'],None,'01-需求故事','H5'
    original=deepcopy((m,p,d))
    with pytest.raises(StorageError) as caught:
        project(tmp_path,m,p,d)
    diag=caught.value.diagnostics[0]
    assert diag['code']=='WORKBOOK_LAYOUT_OVERFLOW'
    assert diag['target']['object_id']==identity and diag['target']['field']==field
    assert sheet+'!'+cell in diag['message']
    assert (m,p,d)==original
    assert not (tmp_path/'projected.xlsx').exists()


def _multiple_layout_overflows(model, pending):
    """Each note component fits alone; their actual G-column composition does not."""
    model['stories'][0]['acs'][0]['text']='一行验收。\n'*60
    model['stories'][1]['notes']='一行必要说明。\n'*60
    tasks=model['tasks'][:3]
    for task in tasks:
        task.update(work_mode='调整',complexity='S',notes='\n'.join(['必要补充。']*4))
        task['classification_basis'][0]['rationale']='\n'.join(['沿用既有实例并调整既定查询条件。']*4)
    pending['items'][0].update(question='\n'.join(['确认适用边界？']*4),
        current_handling='\n'.join(['暂按当前处理。']*4),
        targets=[dict(object_id=task['id'],field='work_mode') for task in tasks])
    prototype=model['tasks'][-1]
    model['tasks'].extend(dict(deepcopy(prototype),id=str(uuid4()),name=f'独立发布成果{index}',notes='',complexity='M')
                         for index in range(12))
    return [
        (model['stories'][0]['id'],'acs','01-需求故事!D5'),
        (model['stories'][1]['id'],'notes','01-需求故事!E6'),
        *[(task['id'],'classification_basis',f'02-任务清单!G{row}') for row,task in enumerate(tasks,5)],
        (model['stories'][-1]['id'],None,'01-需求故事!H9'),
    ]


def _assert_layout_targets(diagnostics, expected):
    assert len(diagnostics)==len(expected)
    for identity,field,cell in expected:
        matching=[d for d in diagnostics if d['target']['object_id']==identity and d['target']['field']==field]
        assert len(matching)==1,diagnostics
        assert matching[0]['code']=='WORKBOOK_LAYOUT_OVERFLOW'
        assert cell in matching[0]['message']


def test_projection_reports_all_composed_note_and_story_overflows_in_one_pass(tmp_path):
    model,pending,decisions=bundle()
    expected=_multiple_layout_overflows(model,pending)
    before=deepcopy((model,pending,decisions))
    with pytest.raises(StorageError) as caught:
        project(tmp_path,model,pending,decisions)
    _assert_layout_targets(caught.value.diagnostics,expected)
    assert (model,pending,decisions)==before
    assert not (tmp_path/'projected.xlsx').exists()


def test_full_check_reports_all_layout_overflows_without_office_or_export(contract_case,monkeypatch):
    from ai_sow_lite import office,workbook
    from ai_sow_lite.cli import execute
    from .support.fixtures import read_json,write_json
    case=contract_case
    model,pending=read_json(case.file('model.json')),read_json(case.file('pending-items.json'))
    expected=_multiple_layout_overflows(model,pending)
    write_json(case.file('model.json'),model);write_json(case.file('pending-items.json'),pending)
    before={name:case.file(name).read_bytes() for name in ('model.json','pending-items.json','candidate.json')}
    def unexpected(*args,**kwargs):
        pytest.fail('Full layout check must not select/start Office or export a workbook')
    monkeypatch.setattr(office,'recalculate',unexpected)
    monkeypatch.setattr(office,'selection_fingerprint',unexpected)
    monkeypatch.setattr(workbook,'project_workbook',unexpected)
    payload=dict(candidate_path=case.candidate_path.relative_to(case.project).as_posix(),scope='full',plan_path=None)
    request=dict(protocol_version='1.0',request_id=case.request_id,project_path=str(case.project),operation='check',payload=payload)
    result=execute(request)
    assert not result['ok'] and not result['result']['valid_for_render']
    _assert_layout_targets(result['diagnostics'],expected)
    saved=read_json(case.project/result['result']['check_ref']['path'])
    assert saved['diagnostics']==result['diagnostics']
    rendered=execute(dict(request,operation='render',payload=dict(candidate_path=payload['candidate_path'],
        check_path=result['result']['check_ref']['path'],expected_current=None)))
    assert not rendered['ok']
    _assert_layout_targets(rendered['diagnostics'],expected)
    assert before=={name:case.file(name).read_bytes() for name in before}
    assert not case.file('render-attempt.json').exists()
    assert not (case.project/'.ai-sow-lite/current.json').exists()
    assert not list(case.candidate_path.parent.glob('render-*'))
    payload['scope']='slice'
    assert execute(request)['ok']


@pytest.mark.parametrize('case', ['long_acs', 'overflow_acs', 'padding_boundary_acs', 'above_padding_boundary_acs'])
def test_real_consumer_acceptance_layout_preserves_text_or_diagnoses_overflow(tmp_path, case):
    from .support.fixtures import FIXTURES,read_json,PLUGIN
    from ai_sow_lite.workbook import preflight_layout
    m,p,d=bundle()
    before=(PLUGIN/'assets/sow-template.xlsx').read_bytes()
    texts=read_json(FIXTURES/'workbook/ac-layout.json')[case]
    story=m['stories'][0]
    story['acs']=[dict(id=str(uuid4()),text=text,evidence_refs=story['evidence_refs']) for text in texts]
    expected='\n'.join(f'{i}. {text}' for i,text in enumerate(texts,1))
    diagnostics=preflight_layout(PLUGIN/'assets/sow-template.xlsx',m,p,d)
    if case in ('overflow_acs', 'above_padding_boundary_acs'):
        _assert_layout_targets(diagnostics,[(story['id'],'acs','01-需求故事!D5')])
        with pytest.raises(StorageError) as caught: project(tmp_path,m,p,d)
        assert caught.value.diagnostics[0]['code']=='WORKBOOK_LAYOUT_OVERFLOW'
        assert caught.value.diagnostics[0]['target']['object_id']==story['id']
        assert not (tmp_path/'projected.xlsx').exists()
    else:
        assert diagnostics==[]
        project(tmp_path,m,p,d)
        w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
        cell=w['01-需求故事']['D5']
        assert cell.value==expected and cell.alignment.vertical=='top' and cell.alignment.wrap_text
        assert w['01-需求故事'].column_dimensions['D'].width==88
        assert w['01-需求故事'].row_dimensions[5].height<=409
        if case=='padding_boundary_acs':
            assert w['01-需求故事'].row_dimensions[5].height==409
        assert copy(cell.font)==copy(openpyxl.load_workbook(PLUGIN/'assets/sow-template.xlsx')['01-需求故事']['D5'].font)
        assert w['01-需求故事']['F5'].protection.locked
    assert (PLUGIN/'assets/sow-template.xlsx').read_bytes()==before


def test_layout_preflight_uses_extended_row_styles_without_expanding_formulas(tmp_path,monkeypatch):
    from ai_sow_lite import workbook
    from .support.fixtures import PLUGIN
    model,pending,decisions=bundle()
    prototype,task=deepcopy(model['stories'][0]),deepcopy(model['tasks'][0])
    model['stories']=[dict(deepcopy(prototype),id=str(uuid4()),title=f'独立故事{index}') for index in range(61)]
    model['tasks']=[dict(deepcopy(task),id=str(uuid4()),name=f'独立任务{index}',notes='',
                         story_id=model['stories'][index%61]['id']) for index in range(201)]
    model['stories'][-1]['acs'][0]['text']='一行验收。\n'*60
    # Tasks are projected by their parent Story, so the last Task belongs to the
    # last Story and appears after all 200 others, at the first expanded row.
    model['tasks'][-1]['story_id']=model['stories'][-1]['id']
    model['tasks'][-1]['notes']='一行必要说明。\n'*60
    pending['items']=[]
    expected=[(model['stories'][-1]['id'],'acs','01-需求故事!D65'),
              (model['tasks'][-1]['id'],'notes','02-任务清单!G205')]
    with monkeypatch.context() as patched:
        def unexpected(*args,**kwargs):
            pytest.fail('Preflight must not expand tables or formulas')
        patched.setattr(workbook,'_extend',unexpected)
        diagnostics=workbook.preflight_layout(PLUGIN/'assets/sow-template.xlsx',model,pending,decisions)
    _assert_layout_targets(diagnostics,expected)
    with pytest.raises(StorageError) as caught:project(tmp_path,model,pending,decisions)
    assert caught.value.diagnostics==diagnostics


def test_layout_failure_does_not_reach_office_or_apply_and_keeps_candidate(tmp_path,monkeypatch):
    from .support.excel import prepare_case
    from .support.fixtures import read_json,write_json
    from .support.cli import run_request
    from ai_sow_lite import office
    case,payload,_=prepare_case(tmp_path/'project',render=False)
    model=read_json(case.file('model.json'))
    model['stories'][0]['acs'][0]['text']='验收正文仍在候选中。\n'*60
    write_json(case.file('model.json'),model)
    before=case.file('model.json').read_bytes()
    checked=run_request(case.project,case.request_id,'check',dict(candidate_path=payload['candidate_path'],scope='full',plan_path=None))
    assert not checked['ok'] and not checked['result']['valid_for_render']
    assert checked['diagnostics'][0]['code']=='WORKBOOK_LAYOUT_OVERFLOW'
    payload['check_path']=checked['result']['check_ref']['path']
    monkeypatch.setattr(office,'recalculate',lambda *args: pytest.fail('Overflow must not invoke Office'))
    from ai_sow_lite.cli import execute
    result=execute(dict(protocol_version='1.0',request_id=case.request_id,project_path=str(case.project),operation='render',payload=payload))
    assert not result['ok'] and result['diagnostics'][0]['code']=='WORKBOOK_LAYOUT_OVERFLOW'
    assert case.file('model.json').read_bytes()==before
    assert not (case.project/'.ai-sow-lite/current.json').exists()
    assert not list(case.file('render-attempt.json').parent.glob('render-*/prepared.json'))
