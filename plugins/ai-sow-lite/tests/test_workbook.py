"""Projection expectations use hand-inspected template cells, never output builders."""
from copy import deepcopy
import hashlib
import importlib
from pathlib import Path
from uuid import uuid4
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from .support.fixtures import PLUGIN, FIXTURES, read_json


def module():
    try:
        return importlib.import_module('ai_sow_lite.workbook')
    except ModuleNotFoundError:
        pytest.fail('I1.3 workbook implementation is missing')


def bundle():
    return [read_json(FIXTURES / 'generate' / name) for name in
            ('model.json', 'pending-items.json', 'decisions.json')]


def project(tmp_path, model=None, pending=None, decisions=None):
    m, p, d = bundle()
    return module().project_workbook(PLUGIN / 'assets/sow-template.xlsx',
        model if model is not None else m, pending if pending is not None else p,
        decisions if decisions is not None else d, str(uuid4()), tmp_path)


@pytest.mark.parametrize('text', ['=SUM(A1:A9)', '+原文', '-原文', '@原文', "'真实前缀", '#N/A'])
def test_literal_task_name_is_not_a_formula(tmp_path, text):
    cell = openpyxl.Workbook().active['B5']
    module().write_literal(cell, text)
    assert cell.data_type == 's'
    assert cell.value == text
    path = tmp_path / 'literal.xlsx'
    cell.parent.parent.save(path)
    reopened = openpyxl.load_workbook(path).active['B5']
    assert (reopened.data_type, reopened.value) == ('s', text)
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        assert root.find('.//m:c[@r="B5"]/m:f', ns) is None


def test_representative_exact_columns_pending_and_unchanged_template(tmp_path):
    before = (PLUGIN / 'assets/sow-template.xlsx').read_bytes()
    result = project(tmp_path)
    book = openpyxl.load_workbook(tmp_path / 'projected.xlsx')
    s, t = book['01-需求故事'], book['02-任务清单']
    assert [s.cell(5, c).value for c in range(1, 5)] == ['信息查询', '资料查询', '资料查询',
        '1. 依约定条件查询并显示结果、空结果或失败。\n2. 本页失败提示一致，支持手动重试，迟到响应不覆盖新查询。']
    assert [t.cell(5, c).value for c in range(1, 7)] == ['资料查询', '资料查询页及公共客户端接入', '信息展示与查询页面', '新建', 'M', None]
    assert t['E10'].value == 'M'
    assert t['G10'].value.startswith('待确认：')
    assert (s['E5'].value or '') == bundle()[0]['stories'][0]['notes']
    assert isinstance(s['H5'].value, ArrayFormula) and s['H5'].value.ref == 'H5'
    assert t['J5'].value == '=IF(OR(NOT(ISNUMBER($H5)),NOT(ISNUMBER($I5))),"",ROUND($H5*$I5,1))'
    assert book['90-估算标准']['Q4'].value == 'SIT适用'
    assert book['90-估算标准']['R4'].value == 'UAT适用'
    assert book['03-工作量汇总']['B5'].value == '=SUM(TaskTable[任务人天])'
    assert book.sheetnames == ['01-需求故事', '02-任务清单', '03-工作量汇总', '90-估算标准']
    assert result['details'] == [] and not (tmp_path/'details.md').exists()
    assert result['objects'][0]['kind'] == 'epic'
    ac = next(o for o in result['objects'] if o['object_id'].endswith('000024'))
    assert ac['fields'] == [{'field': 'text', 'cells': [{'sheet': '01-需求故事', 'cell': 'D5', 'entry': 1}]}]
    assert result['pending_items'][0]['targets'][0]['cells'] == [{'sheet': '02-任务清单', 'cell': 'E10'}]
    text = (tmp_path / 'pending-items.md').read_text()
    assert result['version_id'] in text and '条数未知' in text
    assert (PLUGIN / 'assets/sow-template.xlsx').read_bytes() == before
    assert hashlib.sha256(before).hexdigest() == '28d23be2b50e6abcb3abd81e97ebba3e9e696a8994e05dd9bda72987a1c8ab9e'


@pytest.mark.parametrize('stories,tasks,last_s,last_t', [(0,0,64,204),(1,1,64,204),(60,200,64,204),(61,201,65,205)])
def test_capacity_keeps_reserved_rows_and_extends_original_prototypes(tmp_path, stories, tasks, last_s, last_t):
    m, p, d = bundle()
    s0, t0 = m['stories'][0], m['tasks'][0]
    m['stories'] = [dict(deepcopy(s0), id=str(uuid4()), title=f'故事 {i}') for i in range(stories)]
    for s in m['stories']:
        for ac in s['acs']: ac['id'] = str(uuid4())
    m['tasks'] = [dict(deepcopy(t0), id=str(uuid4()), name=f'任务 {i}', story_id=m['stories'][i % stories]['id']) for i in range(tasks)]
    m['dependencies'] = []
    project(tmp_path, m, dict(schema_version='1.0', items=[]), d)
    w = openpyxl.load_workbook(tmp_path / 'projected.xlsx')
    for sn,tn,end,col in [('01-需求故事','SOWStoryTable',last_s,'J'),('02-任务清单','TaskTable',last_t,'L')]:
        s=w[sn]; table=s.tables[tn]
        assert table.ref == f'A4:{col}{end}'
        assert table.autoFilter.ref.replace('$','') == table.ref
        assert s.freeze_panes == 'A5' and s.print_title_rows == '$1:$4'
        assert s[f'{col}{end}'].data_type == 'f'
        assert s[f'{col}{end}'].protection.locked
        assert not s[f'A{end}'].protection.locked
        assert s[f'{col}{end}']._style == s[f'{col}5']._style
        assert s.protection.sheet
    assert w['01-需求故事'][f'H{last_s}'].value.ref == f'H{last_s}'
    assert w['01-需求故事'].data_validations.dataValidation[0].sqref == 'C5:C1048576'
    assert w['02-任务清单']['A205'].value == ('故事 60' if tasks == 201 else None)
    if not stories:
        assert w['01-需求故事']['C5'].value is None


def test_aliases_are_stable_unique_case_unicode_and_criteria_safe():
    items = [{'id': str(uuid4()), 'title': n} for n in ['ABC','abc','é','e\u0301','查询~*?','>3','=1','😀'*61,'正常']]
    aliases = module().allocate_aliases(items, 'title', 'S')
    assert aliases[items[-1]['id']] == '正常'
    for obj in items[4:8]:
        assert aliases[obj['id']].startswith('S-')
    assert len({module().name_key(v) for v in aliases.values()}) == len(items)
    assert all(len(v.encode('utf-16-le')) // 2 <= 120 for v in aliases.values())
    assert module().allocate_aliases(list(reversed(items)), 'title', 'S') == aliases
    newcomer = {'id':str(uuid4()), 'title':'ABC'}
    previous = {o['id']: {'original': o['title'], 'display_name': aliases[o['id']]} for o in items}
    updated = module().allocate_aliases([newcomer, *items], 'title', 'S', previous)
    assert all(updated[k] == v for k,v in aliases.items())


def test_long_text_stays_inline_with_literal_notes_and_alias_original(tmp_path):
    m,p,d=bundle(); original='验收原文😀\n' * 1200
    m['stories'][0]['title']='查询*?~订单'
    m['stories'][0]['acs'][0]['text']=original
    m['tasks'][0]['notes']='=SUM(A1:A9)'
    result=project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx'); s=w['01-需求故事']; t=w['02-任务清单']
    assert original in s['D5'].value
    assert '原名：查询*?~订单' in s['E5'].value
    assert t['G5'].value=='=SUM(A1:A9)' and t['G5'].data_type=='s'
    assert result['details']==[] and not (tmp_path/'details.md').exists()
    assert s['H5'].data_type=='f'


def test_invalid_xml_literal_and_changed_template_are_diagnosed(tmp_path):
    from ai_sow_lite.project import StorageError
    with pytest.raises(StorageError,match='WORKBOOK_INVALID'):
        module().write_literal(openpyxl.Workbook().active['A1'], 'bad\x00text')
    changed=tmp_path/'changed.xlsx'; w=openpyxl.load_workbook(PLUGIN/'assets/sow-template.xlsx')
    w['02-任务清单']['J5']='=999'; w.save(changed)
    m,p,d=bundle()
    with pytest.raises(StorageError,match='VERSION_INCOMPATIBLE'):
        module().project_workbook(changed,m,p,d,str(uuid4()),tmp_path/'out')


@pytest.mark.office
def test_public_packet_real_render_apply_recover_and_no_repeat_office(tmp_path,monkeypatch):
    import shutil
    if not shutil.which('soffice') and not shutil.which('libreoffice'): pytest.skip('Real Office engine unavailable')
    from .support.excel import prepare_case
    from ai_sow_lite import cli,office
    from ai_sow_lite.contracts import schema_validator
    from ai_sow_lite.contracts import semantic_digest
    from .support.fixtures import write_json
    case,payload,result=prepare_case(tmp_path/'project')
    assert result['ok'],result
    output=result['result']
    prepared=read_json(case.project/output['prepared_ref']['path'])
    projection=read_json(case.project/output['projection_ref']['path'])
    assert not list(schema_validator('artifacts','projection').iter_errors(projection))
    assert projection['projector_version']=='lite-projection-v1'
    assert 'projection_version' not in projection
    assert projection['version_id']==prepared['version_id']==output['version_id']
    assert projection['workbook_hash']==output['workbook_ref']['sha256']
    assert output['pending_count']==1 and output['details_ref'] is None
    assert not (case.project/'.ai-sow-lite/current.json').exists()
    # Prior successful text signatures reuse the same genuinely calculated bytes.
    attempt=read_json(case.file('render-attempt.json'))
    def forbidden(*args,**kwargs): pytest.fail('Prepared reuse/apply must not invoke Office')
    monkeypatch.setattr(office,'recalculate',forbidden)
    def invoke(op,p):
        return cli.execute(dict(protocol_version='1.0',request_id=case.request_id,project_path=str(case.project),operation=op,payload=p))
    for implementation in (None,'lite-render-v2'):
        signature_inputs=dict(check=read_json(case.project/payload['check_path']),payload=payload,
            projector_version='lite-projection-v1',engine=office.selection_fingerprint())
        if implementation is None:
            attempt.pop('implementation_version',None)
        else:
            attempt['implementation_version']=implementation
            signature_inputs['implementation_version']=implementation
        attempt['signature']=semantic_digest(signature_inputs)
        write_json(case.file('render-attempt.json'),attempt)
        repeated=invoke('render',payload)
        assert repeated['ok'] and repeated['result']==output
    before={p:p.read_bytes() for p in case.project.rglob('*') if p.is_file() and ('inputs' in p.parts or 'analysis' in p.parts)}
    applied_payload=dict(entrypoint='generate',prepared_path=output['prepared_ref']['path'],expected_current=None,plan_path=None)
    applied=invoke('apply',applied_payload)
    assert applied['ok'],applied
    assert applied['result']['applied_version']==output['version_id']
    assert invoke('apply',applied_payload)['result']['idempotent']
    recovered=invoke('recover',dict(target_request_id=case.request_id))
    assert recovered['ok'] and recovered['result']['state']=='applied'
    assert all(p.read_bytes()==raw for p,raw in before.items())
    manifest=read_json(case.project/applied['result']['manifest_ref']['path'])
    assert manifest['projection_version']==projection['projector_version']
    final=case.project/applied['result']['workbook_ref']['path']
    assert final.read_bytes()==(case.project/output['workbook_ref']['path']).read_bytes()


@pytest.mark.parametrize('change',['model','candidate','source','check_scope','check_flag'])
def test_render_refuses_changed_or_slice_check_before_office(tmp_path,monkeypatch,change):
    from .support.fixtures import build_ingested_case,write_json
    from .support.cli import run_request
    from ai_sow_lite import cli
    from ai_sow_lite import office
    case=build_ingested_case(tmp_path/'project')
    candidate=read_json(case.candidate_path)
    payload=dict(candidate_path=case.candidate_path.relative_to(case.project).as_posix(),scope='slice' if change=='check_scope' else 'full',plan_path=None)
    checked=run_request(case.project,case.request_id,'check',payload)
    if change=='model':
        model=read_json(case.file('model.json'));model['tasks'][0]['notes']+='修改';write_json(case.file('model.json'),model)
    elif change=='candidate':
        case.candidate_path.write_bytes(case.candidate_path.read_bytes()+b'\n')
    elif change=='source':
        next((case.project/'.ai-sow-lite/inputs/originals').glob('*/*.md')).write_text('changed',encoding='utf-8')
    elif change=='check_flag':
        path=case.project/checked['result']['check_ref']['path']; check=read_json(path);check['candidate_digest']='json-v1:'+'a'*64;write_json(path,check)
    def forbidden(*a,**k): pytest.fail('Invalid check reached Office')
    monkeypatch.setattr(office,'recalculate',forbidden)
    result=cli.execute(dict(protocol_version='1.0',request_id=case.request_id,project_path=str(case.project),operation='render',
        payload=dict(candidate_path=payload['candidate_path'],check_path=checked['result']['check_ref']['path'],expected_current=None)))
    assert not result['ok'] and result['diagnostics'][0]['code']=='CANDIDATE_INVALID',result
    assert not (case.project/'.ai-sow-lite/current.json').exists()


@pytest.mark.office
@pytest.mark.parametrize('target',['sow.xlsx','pending-items.md','projection.json','model.json','verification.json'])
def test_tampered_prepared_never_activates_or_recalculates(tmp_path,monkeypatch,target):
    import shutil
    if not shutil.which('soffice') and not shutil.which('libreoffice'): pytest.skip('Real Office engine unavailable')
    from .support.excel import prepare_case
    from ai_sow_lite import cli,office
    case,payload,result=prepare_case(tmp_path/'project'); assert result['ok'],result
    prepared_path=case.project/result['result']['prepared_ref']['path']
    path=prepared_path.parent/target;path.write_bytes(path.read_bytes()+b' ')
    def forbidden(*a,**k):pytest.fail('Tampered prepared output reached Office')
    monkeypatch.setattr(office,'recalculate',forbidden)
    applied=cli.execute(dict(protocol_version='1.0',request_id=case.request_id,project_path=str(case.project),operation='apply',
        payload=dict(entrypoint='generate',prepared_path=prepared_path.relative_to(case.project).as_posix(),expected_current=None,plan_path=None)))
    assert not applied['ok'],applied
    assert not (case.project/'.ai-sow-lite/current.json').exists()


@pytest.mark.office
def test_candidate_mutation_during_office_invalidates_preparation_and_preserves_inputs(tmp_path,monkeypatch):
    import shutil
    if not shutil.which('soffice') and not shutil.which('libreoffice'): pytest.skip('Real Office engine unavailable')
    from .support.excel import prepare_case
    from .support.fixtures import write_json
    from ai_sow_lite import cli,office
    case,payload,_=prepare_case(tmp_path/'project',render=False)
    actual=office.recalculate
    def changed(source,destination):
        receipt=actual(source,destination)
        m=read_json(case.file('model.json'));m['tasks'][0]['notes']+='新的业务内容';write_json(case.file('model.json'),m)
        return receipt
    monkeypatch.setattr(office,'recalculate',changed)
    result=cli.execute(dict(protocol_version='1.0',request_id=case.request_id,project_path=str(case.project),operation='render',payload=payload))
    assert not result['ok'] and result['diagnostics'][0]['code']=='CANDIDATE_INVALID',result
    assert not list(case.candidate_path.parent.glob('render-*/prepared.json'))
    assert list(case.candidate_path.parent.glob('render-*/sow.office-raw.xlsx'))
    assert not (case.project/'.ai-sow-lite/current.json').exists()


def test_loop_budget_rejects_identical_failure_and_allows_one_changed_retry(tmp_path,monkeypatch):
    from .support.excel import prepare_case
    from .support.fixtures import write_json
    from ai_sow_lite import cli,office
    from ai_sow_lite.project import StorageError
    case,payload,_=prepare_case(tmp_path/'project',render=False)
    calls=[]
    def fail(source,destination):
        calls.append(str(source)); raise StorageError('CALCULATION_FAILED','受控引擎故障')
    monkeypatch.setattr(office,'recalculate',fail)
    def render():return cli.execute(dict(protocol_version='1.0',request_id=case.request_id,project_path=str(case.project),operation='render',payload=payload))
    assert render()['diagnostics'][0]['code']=='CALCULATION_FAILED'
    assert render()['diagnostics'][0]['code']=='LOOP_LIMIT_REACHED'
    assert len(calls)==1
    from .support.cli import run_request
    for index in range(2):
        m=read_json(case.file('model.json'));m['tasks'][0]['notes']+=f'具体修正 {index}';write_json(case.file('model.json'),m)
        checked=run_request(case.project,case.request_id,'check',dict(candidate_path=payload['candidate_path'],scope='full',plan_path=None))
        payload['check_path']=checked['result']['check_ref']['path']
        assert render()['diagnostics'][0]['code']==('CALCULATION_FAILED' if index==0 else 'LOOP_LIMIT_REACHED')
    assert len(calls)==2
    checkpoint=read_json(case.file('checkpoint.json'))
    assert checkpoint['repair_batches']==1 and checkpoint['operation_retries']['render']==1


def test_registered_xlsx_and_judgment_sources_render_before_office(tmp_path,monkeypatch):
    from .test_xlsx_inputs import xlsx_analysis_case
    from .support.fixtures import write_json
    from .support.cli import run_request
    from ai_sow_lite import office
    case,registered,_=xlsx_analysis_case(tmp_path)
    assert registered['ok'],registered
    analysis=read_json(case.file('analysis.json'))
    xlsx=analysis['evidence'][0]
    text=next(e for e in analysis['evidence'] if e['source_refs'] and e['source_refs'][0]['locator']['kind']=='text_lines')
    judgment=dict(id=str(uuid4()),kind='judgment',text='综合已登记依据。',source_refs=[],
                  basis_refs=[xlsx['id'],text['id']],limitations='')
    analysis['evidence'].append(judgment)
    for topic in analysis['topics']: topic['topic_version_id']=str(uuid4())
    analysis['topics'][0]['evidence_refs'].append(judgment['id'])
    write_json(case.file('label-analysis.json'),analysis)
    registered=run_request(case.project,case.request_id,'ingest',dict(kind='analysis',entrypoint='generate',
        analysis_path=case.file('label-analysis.json').relative_to(case.project).as_posix()))
    assert registered['ok'],registered
    candidate=read_json(case.candidate_path)
    candidate.update(evidence_ids=registered['result']['evidence_ids'],topic_version_ids=registered['result']['topic_version_ids'])
    write_json(case.candidate_path,candidate)
    pending=read_json(case.file('pending-items.json'))
    pending['items'][0]['evidence_refs']=[xlsx['id'],text['id'],judgment['id']]
    write_json(case.file('pending-items.json'),pending)
    checked=run_request(case.project,case.request_id,'check',dict(candidate_path=case.candidate_path.relative_to(case.project).as_posix(),scope='full',plan_path=None))
    assert checked['ok'] and checked['result']['valid_for_render'],checked
    before={case.file(n):case.file(n).read_bytes() for n in ('candidate.json','model.json','pending-items.json','decisions.json')}
    class OfficeBoundaryReached(Exception): pass
    def stop(source,destination):
        note=source.with_name('pending-items.md').read_text()
        expected_xlsx=f'sparse-history.xlsx；历史范围!A1:C4；依据 {xlsx["id"]}\n'
        ref=text['source_refs'][0]; loc=ref['locator']
        index=read_json(case.project/'.ai-sow-lite/inputs/index.json')
        filename=next(Path(i['relative_path']).name for i in index['items'] if i['input_version_id']==ref['input_version_id'])
        expected_text=f'{filename}；第 {loc["start_line"]}—{loc["end_line"]} 行；依据 {text["id"]}\n'
        assert f'依据：{xlsx["id"]}\n'+expected_xlsx in note
        assert f'依据：{text["id"]}\n'+expected_text in note
        assert f'依据：{judgment["id"]}\n'+expected_xlsx+expected_text in note
        book=openpyxl.load_workbook(source); assert book['01-需求故事']['C5'].value; book.close()
        for name in ('model.json','pending-items.json','decisions.json'):
            assert source.with_name(name).read_bytes()==before[case.file(name)]
        raise OfficeBoundaryReached
    monkeypatch.setattr(office,'recalculate',stop)
    with pytest.raises(OfficeBoundaryReached):
        module().render_candidate(case.project,case.request_id,dict(candidate_path=case.candidate_path.relative_to(case.project).as_posix(),
            check_path=checked['result']['check_ref']['path'],expected_current=None))
    assert all(p.read_bytes()==raw for p,raw in before.items())
    assert not (case.project/'.ai-sow-lite/current.json').exists()


@pytest.mark.parametrize('previous_version', [None, 'lite-render-v2'])
def test_renderer_implementation_fix_retries_legacy_failure_without_reset(tmp_path,monkeypatch,previous_version):
    from .support.excel import prepare_case
    from .support.fixtures import write_json
    from ai_sow_lite import office
    from ai_sow_lite.contracts import semantic_digest
    from ai_sow_lite.project import StorageError
    case,payload,_=prepare_case(tmp_path/'project',render=False)
    check=read_json(case.project/payload['check_path'])
    # Previous signatures, including the implementation before prototype labels.
    signature_inputs=dict(check=check,payload=payload,projector_version='lite-projection-v1',engine=office.selection_fingerprint())
    if previous_version is not None: signature_inputs['implementation_version']=previous_version
    legacy=semantic_digest(signature_inputs)
    attempt=case.file('render-attempt.json');write_json(attempt,dict(signature=legacy,prepared_ref=None))
    checkpoint=read_json(case.file('checkpoint.json'));checkpoint['repair_batches']=1
    write_json(case.file('checkpoint.json'),checkpoint)
    preserved=case.file('render-previous-failure');preserved.mkdir()
    for name in ('model.json','pending-items.json','decisions.json'):
        (preserved/name).write_bytes(case.file(name).read_bytes())
    before={p:p.read_bytes() for p in preserved.iterdir()}
    class OfficeBoundaryReached(Exception): pass
    def stop(source,destination):
        assert source.exists()
        raise OfficeBoundaryReached
    monkeypatch.setattr(office,'recalculate',stop)
    with pytest.raises(OfficeBoundaryReached): module().render_candidate(case.project,case.request_id,payload)
    assert read_json(attempt)['signature']!=legacy
    checkpoint=read_json(case.file('checkpoint.json'))
    assert checkpoint['repair_batches']==2 and checkpoint['operation_retries']['render']==1
    with pytest.raises(StorageError) as caught: module().render_candidate(case.project,case.request_id,payload)
    assert caught.value.diagnostics[0]['code']=='LOOP_LIMIT_REACHED'
    assert all(p.read_bytes()==raw for p,raw in before.items())
    assert not (case.project/'.ai-sow-lite/current.json').exists()


def test_crlf_is_displayed_inline_with_normalized_excel_newlines(tmp_path):
    model,pending,decisions=bundle(); original='第一行\r\n第二行\r\n'
    model['tasks'][0]['notes']=original
    project(tmp_path,model,pending,decisions)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert w['02-任务清单']['G5'].value == original.replace('\r\n','\n')
    assert not (tmp_path/'details.md').exists()


@pytest.mark.office
def test_checked_partial_classification_and_unestimated_targets_keep_original_formulas(tmp_path):
    import shutil
    if not shutil.which('soffice') and not shutil.which('libreoffice'): pytest.skip('Real Office engine unavailable')
    from .support.excel import prepare_case
    from .support.fixtures import write_json
    from .support.cli import run_request
    case,payload,_=prepare_case(tmp_path/'project',render=False)
    m,p,d=[read_json(case.file(n)) for n in ('model.json','pending-items.json','decisions.json')]
    task=m['tasks'][0];task.update(work_type_name=None,work_mode=None,not_applicable_fields=[])
    task['classification_basis'][0].update(standard_id=None,fields=['complexity'])
    def question(target,field,unestimated=False):
        return dict(id=str(uuid4()),revision=1,question='请确认已有资料中未确定的业务事项。',
            targets=[dict(object_id=target,field=field)],evidence_refs=task['evidence_refs'],
            current_handling='保留已知值；未知分类留空，未拆明工作列待确认。',unestimated_work=unestimated,status='open',resolution=None)
    for field in ('work_type_name','work_mode','integration_type'):p['items'].append(question(task['id'],field))
    integration=m['tasks'][2];integration.update(work_type_name='跨系统业务交互集成',integration_type=None,not_applicable_fields=[])
    integration['classification_basis'][0]['standard_id']='IN-INTEGRATION'
    p['items'].append(question(integration['id'],'integration_type'))
    # A known new-build value and open instance question coexist.
    p['items'].append(question(m['tasks'][1]['id'],'work_mode'))
    orphan=dict(deepcopy(m['epics'][0]),id=str(uuid4()),title='尚未拆明的父级工作');m['epics'].append(orphan)
    row_story=dict(deepcopy(m['stories'][0]),id=str(uuid4()),title='有行尚缺任务的业务工作')
    for ac in row_story['acs']:ac['id']=str(uuid4())
    m['stories'].append(row_story)
    gap=question(orphan['id'],None,True);gap['targets'] += [dict(object_id=row_story['id'],field=None),dict(object_id=m['stories'][0]['id'],field=None)]
    p['items'].append(gap)
    write_json(case.file('model.json'),m);write_json(case.file('pending-items.json'),p)
    checked=run_request(case.project,case.request_id,'check',dict(candidate_path=payload['candidate_path'],scope='full',plan_path=None)); assert checked['ok'],checked
    payload['check_path']=checked['result']['check_ref']['path']
    result=run_request(case.project,case.request_id,'render',payload);assert result['ok'],result
    out=result['result'];w=openpyxl.load_workbook(case.project/out['workbook_ref']['path'])
    assert [w['02-任务清单'].cell(5,c).value for c in (3,4,5,6)]==[None,None,'M',None]
    assert w['02-任务清单']['D6'].value=='新建'
    assert w['02-任务清单']['C7'].value=='跨系统业务交互集成' and w['02-任务清单']['F7'].value is None
    assert w['01-需求故事']['C6'].value=='有行尚缺任务的业务工作'
    assert w['01-需求故事']['H6'].data_type=='f'
    projection=read_json(case.project/out['projection_ref']['path'])
    gap_mapping=next(x for x in projection['pending_items'] if x['pending_item_id']==gap['id'])
    assert gap_mapping['targets'][0]['cells'] and gap_mapping['targets'][1]['cells']
    note=(case.project/out['pending_items_ref']['path']).read_text()
    assert '无工作簿行' not in note and '未拆明业务工作：是' in note and gap['id'] in note
    cached=openpyxl.load_workbook(case.project/out['workbook_ref']['path'],data_only=True)
    scope_row=next(o['rows'][0] for o in projection['objects'] if o['object_id']==orphan['id'])
    assert cached['01-需求故事'][f'J{scope_row}'].value=='待确认'
    assert all(cached['01-需求故事'][f'{col}{scope_row}'].value is None for col in ('C','D','I'))
    assert cached['02-任务清单']['L5'].value==cached['02-任务清单']['L7'].value=='待确认'
    assert out['pending_count']==7
    assert (case.project/out['prepared_ref']['path']).parent.joinpath('model.json').read_bytes()==case.file('model.json').read_bytes()


@pytest.mark.office
def test_real_registered_no_gap_evidence_exports_empty_reserved_inputs(tmp_path):
    import shutil
    if not shutil.which('soffice') and not shutil.which('libreoffice'): pytest.skip('Real Office engine unavailable')
    from .support.fixtures import build_ingested_case,write_json
    from .support.cli import run_request
    from ai_sow_lite.contracts import file_sha256
    case=build_ingested_case(tmp_path/'project')
    scope='范围确认：既有系统已覆盖本期全部要求，没有新增、调整、接入或未拆明的业务工作。\n'
    source=case.file('no-gap.md');source.write_text(scope,encoding='utf-8')
    ingested=run_request(case.project,case.request_id,'ingest',dict(kind='sources',entrypoint='generate',project_type='new',sources=[
        dict(source_path=str(source),input_id=None,material_types=['answer'],uses=['to-be-scope'],use_regions=[])]));assert ingested['ok'],ingested
    input_id=ingested['result']['input_refs'][0]['input_version_id'];evidence_id=str(uuid4());topic_id=str(uuid4());topic_version=str(uuid4())
    ref=dict(input_version_id=input_id,locator=dict(kind='text_lines',start_line=1,end_line=1),excerpt_hash=file_sha256(source))
    topic=read_json(case.file('analysis.json'))['topics'][0]
    topic.update(topic_id=topic_id,topic_version_id=topic_version,title='已确认无本期 gap',input_version_ids=[input_id],uses=['to-be-scope'],
                 covered_regions=[ref],uncovered_regions=[],evidence_refs=[evidence_id],related_object_ids=[],external_responsibilities='',limitations='',conclusion=scope,historical_items=[])
    analysis=dict(schema_version='1.0',evidence=[dict(id=evidence_id,kind='statement',text=scope,source_refs=[ref],basis_refs=[],limitations='')],topics=[topic],observations=[])
    write_json(case.file('no-gap-analysis.json'),analysis)
    registered=run_request(case.project,case.request_id,'ingest',dict(kind='analysis',entrypoint='generate',analysis_path=case.file('no-gap-analysis.json').relative_to(case.project).as_posix()));assert registered['ok'],registered
    candidate=read_json(case.candidate_path);candidate.update(input_version_ids=[input_id],topic_version_ids=[topic_version],evidence_ids=[evidence_id])
    write_json(case.candidate_path,candidate)
    write_json(case.file('model.json'),dict(schema_version='1.0',epics=[],features=[],stories=[],tasks=[],dependencies=[],lineage=[]))
    for name in ('pending-items.json','decisions.json'):write_json(case.file(name),dict(schema_version='1.0',items=[]))
    relative=case.candidate_path.relative_to(case.project).as_posix()
    checked=run_request(case.project,case.request_id,'check',dict(candidate_path=relative,scope='full',plan_path=None));assert checked['ok'],checked
    result=run_request(case.project,case.request_id,'render',dict(candidate_path=relative,check_path=checked['result']['check_ref']['path'],expected_current=None));assert result['ok'],result
    out=result['result'];w=openpyxl.load_workbook(case.project/out['workbook_ref']['path'])
    assert w['01-需求故事']['C5'].value is None and w['02-任务清单']['B5'].value is None
    assert w['01-需求故事'].tables['SOWStoryTable'].ref=='A4:J64'
    assert out['pending_count']==0
    assert '无待确认事项' in (case.project/out['pending_items_ref']['path']).read_text()


def test_numeric_criteria_names_receive_stable_identity_aliases():
    names=['1','01','+1','-1','1.0','1e0','TRUE','#N/A']
    items=[dict(id=str(uuid4()),title=name) for name in names]
    aliases=module().allocate_aliases(items,'title','S')
    assert all(value.startswith('S-') for value in aliases.values())
    assert len(set(aliases.values()))==len(items)


@pytest.mark.office
def test_apply_rechecks_original_sources_after_delivery_verifier(tmp_path,monkeypatch):
    import shutil
    if not shutil.which('soffice') and not shutil.which('libreoffice'): pytest.skip('Real Office engine unavailable')
    from .support.excel import prepare_case
    from ai_sow_lite import project as storage
    case,payload,result=prepare_case(tmp_path/'project');assert result['ok'],result
    verifier=storage._verify_delivery
    def changed(project,prepared):
        verifier(project,prepared)
        case.file('model.json').write_bytes(case.file('model.json').read_bytes()+b'\n')
    monkeypatch.setattr(storage,'_verify_delivery',changed)
    with pytest.raises(storage.StorageError):
        storage.apply_prepared(case.project,case.request_id,dict(entrypoint='generate',prepared_path=result['result']['prepared_ref']['path'],expected_current=None,plan_path=None))
    assert not (case.project/'.ai-sow-lite/current.json').exists()


def test_native_clipped_rows_get_actual_column_width_and_cjk_margin(tmp_path):
    project(tmp_path)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    # Reserve room for AC, the formula task list and the inline question using
    # actual template widths. No generated reference filler is needed.
    assert 90 <= w['01-需求故事'].row_dimensions[5].height <=409
    assert 100 <= w['01-需求故事'].row_dimensions[7].height <=409
    assert 180 <= w['02-任务清单'].row_dimensions[10].height <=409
    assert w['01-需求故事'].column_dimensions['E'].width==30
    assert w['02-任务清单'].column_dimensions['G'].width==36


def test_layout_overflow_keeps_full_cell_with_excel_row_height_cap(tmp_path):
    model,pending,decisions=bundle()
    model['stories'][0]['notes']='甲'*700
    project(tmp_path,model,pending,decisions)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    cell=w['01-需求故事']['E5']
    assert cell.value == '甲'*700
    assert not (tmp_path/'details.md').exists()
    assert w['01-需求故事'].row_dimensions[5].height<=409


def test_narrow_parent_title_is_not_truncated(tmp_path):
    m,p,d=bundle();m['epics'][0]['title']='窄列完整标题'*600
    project(tmp_path,m,p,d)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx')
    assert w['01-需求故事']['A5'].value == m['epics'][0]['title']
    assert not (tmp_path/'details.md').exists()
    assert w['01-需求故事'].row_dimensions[5].height<=409


@pytest.mark.office
def test_medium_task_list_expands_row_without_filling_notes_after_office(tmp_path):
    import shutil
    if not shutil.which('soffice') and not shutil.which('libreoffice'):
        pytest.skip('Real Office engine unavailable')
    from .support.excel import prepare_case
    case,payload,result=prepare_case(tmp_path/'project','medium-list')
    assert result['ok'],result
    output=result['result'];path=case.project/output['workbook_ref']['path']
    w=openpyxl.load_workbook(path);cached=openpyxl.load_workbook(path,data_only=True)
    ws=w['01-需求故事'];text=cached['01-需求故事']['H5'].value
    names=[f'测试任务长名称用于检验故事任务列表可读性{i}' for i in range(3)]
    assert all(name in text for name in names)
    assert isinstance(ws['H5'].value,ArrayFormula) and ws['H5'].value.ref=='H5'
    assert ws.row_dimensions[5].height>=243.75  # Office stores heights in 0.75pt steps.
    assert ws['E5'].value is None
    assert output['details_ref'] is None
    assert ws.row_dimensions[5].height<=409
    assert module().formula_cache_inventory(path)==module().formula_cache_inventory(path.with_name('sow.office-raw.xlsx'))


def test_medium_list_notes_and_pending_remain_complete_inline(tmp_path):
    from .support.excel import medium_list_model
    model,pending,decisions=bundle();model=medium_list_model(model)
    story=model['stories'][0];story['notes']='甲'*700
    item=pending['items'][0];item['targets']=[dict(object_id=story['id'],field='notes')]
    item['current_handling']='保留全部待确认处理说明。'*100
    result=project(tmp_path,model,pending,decisions)
    w=openpyxl.load_workbook(tmp_path/'projected.xlsx');notes=w['01-需求故事']['E5'].value
    assert result['details']==[]
    assert notes.startswith('待确认：')
    assert story['notes'] in notes and item['current_handling'] in notes
    assert all(w['02-任务清单'].cell(row,2).value==t['name'] for row,t in enumerate(model['tasks'],5))
    assert w['01-需求故事'].row_dimensions[5].height<=409
