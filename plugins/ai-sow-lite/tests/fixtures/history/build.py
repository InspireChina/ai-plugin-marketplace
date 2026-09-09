"""Regenerate synthetic physical-reading fixtures; no customer data or expected model."""
from pathlib import Path
from datetime import date, datetime, time, timedelta
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table


def build(directory):
    book = Workbook()
    sheet = book.active
    sheet.title = '历史范围'
    for row in [ ['范围', '工作说明', '限定'],
                 ['订单查询', '订单详情查询 API', '候选；服务及接口未记载'],
                 [None, '订单状态事件', '消费方未记载'],
                 [None, None, None],
                 ['订单查询', '归档订单查询', '仅归档数据，不能合并整条'],
                 ['退货', '退货处理', '已取消'],
                 ['对账', '对账能力', '未来范围，未交付'],
                 ['明确接口', 'GET /orders/{id}', '订单服务，只读当前订单'] ]:
        sheet.append(row)
    sheet.merge_cells('A2:A3')
    sheet['B3'].comment = Comment('仅状态变化通知，不含订单明细；参见附注 A2。', '合成夹具')
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions['C'].hidden = True
    sheet.add_table(Table(displayName='HistoryTable', ref='A1:C8'))
    notes = book.create_sheet('附注')
    notes.sheet_state = 'hidden'
    notes.append(['说明'])
    notes.append(['事件暂不包含补发；取消和未来范围不作为已交付实例。'])
    standard = book.create_sheet('标准目录')
    standard.append(['类型', '示例'])
    standard.append(['API', '目录与示例不代表历史交付'])
    typed = book.create_sheet('类型观察')
    typed.append([12.5, True, date(2026, 1, 2), datetime(2026, 1, 2, 3, 4, 5), time(3, 4, 5), '#DIV/0!', '=1+2', '=2+2'])
    typed.append([timedelta(days=1, seconds=30), '=SUM(A1,1)', '文字', '=0', '=FALSE()'])
    output = BytesIO()
    book.save(output)
    book.close()
    with ZipFile(output) as source, ZipFile(directory / 'sparse-history.xlsx', 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet4.xml':
                root = ET.fromstring(raw)
                ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                for cell in root.iter(ns + 'c'):
                    if cell.attrib['r'] in ('G1', 'D2', 'E2'):
                        cell.find(ns + 'v').text = '3' if cell.attrib['r'] == 'G1' else '0'
                        if cell.attrib['r'] == 'E2':
                            cell.set('t', 'b')
                raw = ET.tostring(root, encoding='utf-8')
            target.writestr(item, raw)
    (directory / 'prd-hld.md').write_bytes(b'\xef\xbb\xbf# PRD\r\nQuery orders.\r\n# HLD\r\nUse the order service.\r\n')


if __name__ == '__main__':
    build(Path(__file__).parent)
