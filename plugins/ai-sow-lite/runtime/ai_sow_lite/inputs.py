"""Immutable source copies and strict text/XLSX/prototype reading; no semantic analysis runner."""
from __future__ import annotations

import hashlib
from pathlib import Path
import stat
from uuid import uuid4

from .contracts import PLUGIN_ROOT, canonical_json_bytes
from .project import (StorageError, atomic_bytes, checked_json, ensure_request, file_ref,
                      initialize, request_area, safe_path, write_json)
from .validation import diagnostic


def input_index(project):
    index = checked_json(project, '.ai-sow-lite/inputs/index.json', 'input_index', '.ai-sow-lite/inputs')
    identities = [e['input_version_id'] for e in index['items']]
    if len(identities) != len(set(identities)):
        raise StorageError('VERSION_INCOMPATIBLE', '输入索引重复声明同一不可变版本。')
    return index


def _source_bytes(path):
    path = Path(path).absolute()
    # The explicitly supplied external directory may be an OS alias (e.g. /var).
    # The source file itself cannot be a link; project-owned paths remain stricter.
    if path.is_symlink():
        raise StorageError('INPUT_UNAVAILABLE', '来源不能使用符号链接；请提供实际文件。')
    if path.suffix.lower() not in ('.md', '.markdown', '.txt', '.xlsx'):
        raise StorageError('FORMAT_UNSUPPORTED', '只读取 Markdown/纯文本和 XLSX；其他非原型格式请提供可读文本或 XLSX。')
    if not stat.S_ISREG(path.stat().st_mode):
        raise StorageError('INPUT_UNAVAILABLE', '来源必须是普通文件。')
    if path.stat().st_size > 50 * 1024 * 1024:
        raise StorageError('RESULT_TOO_LARGE', '来源文件超过50 MiB读取上限；请提供较小的可读文件。')
    raw = path.read_bytes()
    if path.suffix.lower() == '.xlsx':
        if not raw.startswith(b'PK\x03\x04'):
            raise StorageError('FORMAT_UNSUPPORTED', 'XLSX 文件头无效，改名不能转换格式。')
        return raw
    if raw.startswith((b'%PDF-', b'PK\x03\x04', b'\xd0\xcf\x11\xe0')) or (b'\x00' in raw and not raw.startswith((b'\xff\xfe', b'\xfe\xff'))):
        raise StorageError('FORMAT_UNSUPPORTED', '文件字节不是受支持的文本，改扩展名不能改变格式。')
    return raw


def _decode_text(raw, entry):
    # Baseline lite-text-v1 registered utf-8 even with a BOM; preserve that meaning.
    # Newly registered BOM inputs explicitly use utf-8-sig. Neither rewrites originals.
    encoding = entry.get('encoding', 'utf-8')
    if encoding not in ('utf-8', 'utf-8-sig') or (encoding == 'utf-8-sig' and not raw.startswith(b'\xef\xbb\xbf')):
        raise StorageError('EVIDENCE_MISSING', '文本登记编码不受支持或与原件 BOM 不同。')
    return raw.decode(encoding, errors='strict')


def text_lines(project, entry):
    path = safe_path(project, entry['relative_path'], f".ai-sow-lite/inputs/originals/{entry['input_version_id']}")
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != entry['content_hash']:
        raise StorageError('EVIDENCE_MISSING', '原件与登记摘要不同。', entry['relative_path'])
    try:
        return _decode_text(raw, entry).splitlines(keepends=True)
    except UnicodeError:
        raise StorageError('INPUT_UNAVAILABLE', '文本无法严格按 UTF-8 解码；已保留原件，请提供可读 UTF-8 文本。', entry['relative_path']) from None


def _reading(project, entry):
    if entry['format'] == 'prototype':
        from ._prototype import reading
        return reading(project, entry)
    if entry['format'] == 'xlsx':
        return _xlsx_reading(project, entry, {'kind': 'workbook'})[0]
    identity = entry['input_version_id']
    area = f'.ai-sow-lite/inputs/originals/{identity}'
    locator_path = safe_path(project, area + '/reading-ref.json')
    if locator_path.exists():
        ref = checked_json(project, area + '/reading-ref.json', 'file_ref', area)
        path = safe_path(project, ref['path'], '.ai-sow-lite/inputs/readings')
        if file_ref(project, path) != ref:
            raise StorageError('EVIDENCE_MISSING', '读取记录摘要不一致。')
        lines = text_lines(project, entry)
        reading = checked_json(project, ref['path'], 'reading', '.ai-sow-lite/inputs/readings')
        if reading['input_version_id'] != identity or reading['content_hash'] != entry['content_hash'] or reading['adapter_version'] != 'lite-text-v1':
            raise StorageError('EVIDENCE_MISSING', '读取记录与实际来源或适配器不同。')
        for excerpt in reading['excerpts']:
            attachment = safe_path(project, excerpt['file_ref']['path'], str(Path(ref['path']).parent))
            if file_ref(project, attachment) != excerpt['file_ref'] or attachment.read_bytes() != ''.join(lines).encode('utf-8'):
                raise StorageError('EVIDENCE_MISSING', '已登记摘录附件与真实原件不同。')
        return ref
    lines = text_lines(project, entry)
    read_id = str(uuid4())
    reading_area = f'.ai-sow-lite/inputs/readings/{read_id}'
    excerpts = []
    if lines:
        excerpt_path = safe_path(project, reading_area + '/text.txt')
        atomic_bytes(excerpt_path, ''.join(lines).encode('utf-8'), immutable=True)
        excerpts.append(dict(locator=dict(kind='text_lines', start_line=1, end_line=len(lines)),
                             file_ref=file_ref(project, excerpt_path)))
    reading = dict(schema_version='1.0', read_id=read_id, input_version_id=identity,
                   content_hash=entry['content_hash'], adapter_version='lite-text-v1', options={}, excerpts=excerpts)
    ref = write_json(project, reading_area + '/reading.json', reading, immutable=True)
    write_json(project, area + '/reading-ref.json', ref, immutable=True)
    return ref


def _original_bytes(project, entry):
    path = safe_path(project, entry['relative_path'], f".ai-sow-lite/inputs/originals/{entry['input_version_id']}")
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != entry['content_hash']:
        raise StorageError('EVIDENCE_MISSING', '原件与登记摘要不同。')
    return raw


def _checked_xlsx_reading(project, entry, read_id):
    from . import _xlsx
    area = f'.ai-sow-lite/inputs/readings/{read_id}'
    reading = checked_json(project, area + '/reading.json', 'reading', area)
    if (reading['read_id'] != read_id or reading['input_version_id'] != entry['input_version_id'] or
            reading['content_hash'] != entry['content_hash'] or reading['adapter_version'] != _xlsx.READER_VERSION or
            reading['options'] != _xlsx.OPTIONS):
        raise StorageError('EVIDENCE_MISSING', 'XLSX 读取身份与实际来源、适配器或选项不同。')
    selection = reading.get('selection')
    if selection is None:
        raise StorageError('EVIDENCE_MISSING', 'XLSX reading 缺少 selection。')
    identity = dict(content_hash=entry['content_hash'], adapter_version=_xlsx.READER_VERSION, options=_xlsx.OPTIONS, selection=selection)
    key = hashlib.sha256(canonical_json_bytes(identity)).hexdigest()
    registration = checked_json(project, f".ai-sow-lite/inputs/originals/{entry['input_version_id']}/readings/{key}.json", 'file_ref')
    if registration != file_ref(project, safe_path(project, area + '/reading.json')):
        raise StorageError('EVIDENCE_MISSING', 'XLSX reading 未由此读取身份登记。')
    if selection['kind'] == 'xlsx_range':
        if len(reading['excerpts']) != 1 or any(reading['excerpts'][0][k] != selection[k] for k in ('sheet', 'range')):
            raise StorageError('EVIDENCE_MISSING', 'XLSX reading 附件定位与 selection 不同。')
    observed = _xlsx.read(_original_bytes(project, entry), selection)
    ref = reading['directory_ref'] if reading['selection']['kind'] == 'workbook' else reading['excerpts'][0]['file_ref']
    path = safe_path(project, ref['path'], area)
    if file_ref(project, path) != ref or path.read_bytes() != canonical_json_bytes(observed):
        raise StorageError('EVIDENCE_MISSING', 'XLSX 摘录与实际原件或选择范围不同。')
    return reading, observed


def _locator_error(message):
    error = StorageError('CANDIDATE_INVALID', message)
    error.diagnostics[0]['target']['field'] = 'locator'
    return error


def _text_directory(lines):
    """ATX headings are navigation only; fenced code is not a section boundary."""
    import re
    starts, fence = [], None
    for number, line in enumerate(lines, 1):
        marker = re.match(r'^ {0,3}(`{3,}|~{3,})', line)
        if marker:
            token = marker[1]
            if fence is None:
                fence = token
            elif token[0] == fence[0] and len(token) >= len(fence):
                fence = None
            continue
        heading = re.match(r'^ {0,3}(#{1,6})[ \t]+(.+?)[\r\n]*$', line) if fence is None else None
        if heading:
            starts.append((number, heading[2]))
    if lines and (not starts or starts[0][0] != 1):
        starts.insert(0, (1, None))
    return [dict(heading=title, locator=dict(kind='text_lines', start_line=start,
                 end_line=starts[index + 1][0] - 1 if index + 1 < len(starts) else len(lines)))
            for index, (start, title) in enumerate(starts)]


def source_excerpt(project, entry, locator, *, raw=None, uncovered=False):
    """Shared deterministic source verifier; returns bytes and physical dependencies."""
    from . import _xlsx, _prototype
    if locator['kind'] == 'observation':
        excerpt, _, dependencies = _prototype.observation_excerpt(project, entry, locator)
        return (None if uncovered else excerpt), dependencies
    if entry['format'] == 'prototype':
        _, dependencies, contents = _prototype.checked_package(project, entry)
        if locator['kind'] != 'text_lines':
            raise _locator_error('原型源码必须使用 text_lines。')
        lines = _prototype.source_lines(contents, locator)
        return (None if uncovered else ''.join(lines).encode('utf-8')), dependencies
    if locator['kind'] == 'text_lines':
        if entry['format'] != 'text' or locator.get('path') not in (None, entry['relative_path'], Path(entry['relative_path']).name):
            raise _locator_error('文本定位与已登记来源不同。')
        if locator['start_line'] > locator['end_line']:
            raise _locator_error('文本定位起止顺序不合法。')
        if uncovered:
            return None, []
        raw = _original_bytes(project, entry) if raw is None else raw
        try:
            lines = _decode_text(raw, entry).splitlines(keepends=True)
        except UnicodeError:
            raise StorageError('INPUT_UNAVAILABLE', '来源无法严格按 UTF-8 读取；请提供可读文本。') from None
        start, end = locator['start_line'], locator['end_line']
        if not 1 <= start <= end <= len(lines):
            raise StorageError('EVIDENCE_MISSING', '文本行范围超出实际来源。')
        return ''.join(lines[start - 1:end]).encode('utf-8'), []
    if locator['kind'] != 'xlsx_range':
        raise StorageError('OPERATION_UNSUPPORTED', '此来源定位适配尚未实现。')
    if entry['format'] != 'xlsx':
        raise StorageError('EVIDENCE_MISSING', 'XLSX 定位不能用于其他输入格式。')
    reading, observed = _checked_xlsx_reading(project, entry, locator['read_id'])
    selection = dict(kind='xlsx_range', sheet=locator['sheet'], range=_xlsx.canonical_range(locator['range']))
    if reading['selection'] != selection:
        if not uncovered or reading['selection'] != {'kind': 'workbook'}:
            raise StorageError('EVIDENCE_MISSING', '来源范围与不可变 reading selection 不同。')
        sheet = next((s for s in observed['sheets'] if s['sheet'] == selection['sheet']), None)
        if sheet is None or _xlsx.bounds(selection['range'])[2] > _xlsx.bounds(sheet['used_range'])[2] or _xlsx.bounds(selection['range'])[3] > _xlsx.bounds(sheet['used_range'])[3]:
            raise StorageError('EVIDENCE_MISSING', '未读区域不在已登记 XLSX 结构内。')
    area = f".ai-sow-lite/inputs/readings/{reading['read_id']}"
    ref = file_ref(project, safe_path(project, area + '/reading.json'))
    attachment = reading.get('directory_ref') or reading['excerpts'][0]['file_ref']
    return (None if uncovered else canonical_json_bytes(observed)), [ref, attachment]


def _xlsx_reading(project, entry, selection):
    from . import _xlsx
    identity = dict(content_hash=entry['content_hash'], adapter_version=_xlsx.READER_VERSION, options=_xlsx.OPTIONS, selection=selection)
    key = hashlib.sha256(canonical_json_bytes(identity)).hexdigest()
    cache = f".ai-sow-lite/inputs/originals/{entry['input_version_id']}/readings/{key}.json"
    if safe_path(project, cache).exists():
        ref = checked_json(project, cache, 'file_ref')
        path = safe_path(project, ref['path'], '.ai-sow-lite/inputs/readings')
        if file_ref(project, path) != ref:
            raise StorageError('EVIDENCE_MISSING', 'XLSX reading 摘要不同。')
        reading, observed = _checked_xlsx_reading(project, entry, path.parent.name)
        if reading['selection'] != selection:
            raise StorageError('EVIDENCE_MISSING', 'XLSX selection 与缓存不同。')
        return ref, reading, observed
    observed = _xlsx.read(_original_bytes(project, entry), selection)
    read_id = str(uuid4())
    area = f'.ai-sow-lite/inputs/readings/{read_id}'
    attachment = safe_path(project, area + '/excerpt.json')
    atomic_bytes(attachment, canonical_json_bytes(observed), immutable=True)
    reading = dict(schema_version='1.0', read_id=read_id, input_version_id=entry['input_version_id'],
                   **identity, excerpts=[])
    if selection['kind'] == 'workbook':
        reading['directory_ref'] = file_ref(project, attachment)
    else:
        reading['excerpts'] = [dict(sheet=selection['sheet'], range=selection['range'], file_ref=file_ref(project, attachment))]
    ref = write_json(project, area + '/reading.json', reading, immutable=True)
    write_json(project, cache, ref, immutable=True)
    return ref, reading, observed


def ingest_sources(project: Path, request_id: str, payload):
    project = Path(project).resolve()
    record = initialize(project, payload['project_type'], PLUGIN_ROOT / 'assets/sow-template.xlsx')
    ensure_request(project, request_id, payload['entrypoint'])
    index = input_index(project)
    refs, readings, failures = [], [], []
    for number, source in enumerate(payload['sources']):
        entry = None
        try:
            if Path(source['source_path']).is_dir():
                from ._prototype import ingest_package
                entry = ingest_package(project, index, source)
                write_json(project, '.ai-sow-lite/inputs/index.json', index)
                refs.append(entry)
                ref = _reading(project, entry)
                if ref not in readings:
                    readings.append(ref)
                continue
            raw = _source_bytes(source['source_path'])
            digest = hashlib.sha256(raw).hexdigest()
            source_format = 'xlsx' if Path(source['source_path']).suffix.lower() == '.xlsx' else 'text'
            for region in source['use_regions']:
                if region['material_type'] not in source['material_types'] or region['use'] not in source['uses']:
                    raise StorageError('CANDIDATE_INVALID', '用途区域必须属于所声明的角色与用途。')
                for locator in region['locators']:
                    if source_format == 'xlsx':
                        registered = next((e for e in index['items'] if e['content_hash'] == digest and
                                           source['input_id'] in (None, e['input_id'])), None)
                        if registered is None:
                            raise StorageError('EVIDENCE_MISSING', 'XLSX 用途定位须采用实际 reading；首次登记先省略 use_regions，再 inspect 所选区域。')
                        source_excerpt(project, registered, locator)
                    else:
                        # Validate the supplied basename before deduplication binds the stored name.
                        temporary_entry = dict(format='text', relative_path=Path(source['source_path']).name,
                                               encoding='utf-8-sig' if raw.startswith(b'\xef\xbb\xbf') else 'utf-8')
                        source_excerpt(project, temporary_entry, locator, raw=raw)
            logical = source['input_id']
            if logical is not None and not any(e['input_id'] == logical for e in index['items']):
                raise StorageError('INPUT_ID_CONFLICT', '指定 input_id 未登记；首次登记必须为 null。')
            matches = [e for e in index['items'] if e['format'] == source_format and e['content_hash'] == digest and
                       (logical is None or e['input_id'] == logical)]
            if matches:
                entry = matches[0]
                # Physical identity stays stable when explicitly adding roles/uses.
                for key in ('material_types', 'uses'):
                    entry[key].extend(v for v in source[key] if v not in entry[key])
                for region in source['use_regions']:
                    bound = dict(region, locators=[dict(locator, path=Path(entry['relative_path']).name)
                                 if 'path' in locator else dict(locator) for locator in region['locators']])
                    if bound not in entry['use_regions']:
                        entry['use_regions'].append(bound)
            else:
                identity = str(uuid4())
                relative = f'.ai-sow-lite/inputs/originals/{identity}/' + Path(source['source_path']).name
                path = safe_path(project, relative)
                atomic_bytes(path, raw, immutable=True)
                if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                    raise StorageError('IO_FAILED', '拷贝后摘要不一致；没有登记成功。')
                entry = dict(input_version_id=identity, input_id=logical or str(uuid4()), content_hash=digest,
                             relative_path=relative, format='xlsx' if Path(source['source_path']).suffix.lower() == '.xlsx' else 'text',
                             material_types=list(source['material_types']), uses=list(source['uses']),
                             use_regions=list(source['use_regions']))
                if entry['format'] == 'text':
                    entry['encoding'] = 'utf-8-sig' if raw.startswith(b'\xef\xbb\xbf') else 'utf-8'
                index['items'].append(entry)
            write_json(project, '.ai-sow-lite/inputs/index.json', index)
            refs.append(entry)
            ref = _reading(project, entry)
            if ref not in readings:
                readings.append(ref)
        except (StorageError, OSError) as error:
            if isinstance(error, StorageError):
                issues = error.diagnostics
            else:
                issues = [diagnostic('INPUT_UNAVAILABLE', message='输入不可读取或保存；保留其他有效材料。')]
            for issue in issues:
                issue['target']['field'] = f'sources[{number}]'
                if entry is not None:
                    issue['preserved_paths'] = [entry['relative_path']]
                failures.append(issue)
    checkpoint = safe_path(project, request_area(request_id, payload['entrypoint']) + '/checkpoint.json')
    return dict(project_id=record['project_id'], input_refs=refs, reading_refs=readings,
                failures=failures, checkpoint_ref=file_ref(project, checkpoint))


def _committed_topic(project, directory):
    """Return a fully committed topic's file_ref, or None when it is unwritten.

    Applies the same binding proof as a fresh registration: the record bytes must
    equal the canonical record rebuilt from the registration it names. Anything
    that contradicts its own source raises; only an interrupted write returns None.
    """
    relative = f'.ai-sow-lite/analysis/topics/{directory.name}/analysis.json'
    provenance_relative = f'.ai-sow-lite/analysis/topics/{directory.name}/registration-ref.json'
    record, provenance_path = safe_path(project, relative), safe_path(project, provenance_relative)
    if not provenance_path.exists():
        if record.exists():
            raise StorageError('EVIDENCE_MISSING', '主题缺少来源绑定，无法证明其登记出处。')
        return None
    if not record.exists():
        # Interrupted between the provenance binding and the record itself.
        return None
    provenance = checked_json(project, provenance_relative, 'file_ref')
    registration = safe_path(project, provenance['path'], '.ai-sow-lite/analysis/registrations')
    if not registration.exists() or file_ref(project, registration) != provenance:
        raise StorageError('EVIDENCE_MISSING', '主题绑定的来源原件缺失或字节已变化。')
    registered = checked_json(project, provenance['path'], 'analysis', '.ai-sow-lite/analysis/registrations')
    stored = checked_json(project, relative, 'analysis', '.ai-sow-lite/analysis/topics')
    if len(stored['topics']) != 1 or stored['topics'][0]['topic_version_id'] != directory.name:
        raise StorageError('EVIDENCE_MISSING', '主题文件与其目录身份不一致。')
    topic = stored['topics'][0]
    if topic not in registered['topics']:
        raise StorageError('EVIDENCE_MISSING', '主题内容不属于其声明的登记来源。')
    if not {r['sha256'] for r in stored['observations']} <= {r['sha256'] for r in registered['observations']}:
        raise StorageError('EVIDENCE_MISSING', '主题观察摘要不属于原始登记来源。')
    bound = dict(schema_version='1.0', topics=[topic], evidence=registered['evidence'],
                 observations=stored['observations'])
    if record.read_bytes() != canonical_json_bytes(bound):
        raise StorageError('EVIDENCE_MISSING', '主题原字节与登记来源不同。')
    return file_ref(project, record)


def recover_analysis_index(project):
    """Rebuild an index lost mid-registration. Write path only; never a query.

    Each surviving topic must re-prove its own origin byte for byte. Topics whose
    write was interrupted before the record exists are left for the caller to
    complete; corrupted ones raise instead of being adopted.
    """
    if safe_path(project, '.ai-sow-lite/analysis/index.json').exists():
        return
    topics = safe_path(project, '.ai-sow-lite/analysis/topics')
    if not topics.exists():
        return
    refs = [ref for ref in (_committed_topic(project, directory)
                            for directory in sorted(topics.iterdir()) if directory.is_dir()) if ref]
    if refs:
        write_json(project, '.ai-sow-lite/analysis/index.json', dict(schema_version='1.0', items=refs))


def _analysis_records(project):
    path = safe_path(project, '.ai-sow-lite/analysis/index.json')
    if not path.exists():
        # Read-only: a missing index is reported, never repaired here. Only a
        # directory holding an actual record counts as an existing analysis; one
        # left with just its provenance binding is an unfinished write.
        topics = safe_path(project, '.ai-sow-lite/analysis/topics')
        if topics.exists() and any((directory / 'analysis.json').exists()
                                   for directory in topics.iterdir() if directory.is_dir()):
            raise StorageError('EVIDENCE_MISSING', '已有分析文件但索引缺失，不能当作空分析。')
        return [], []
    index = checked_json(project, '.ai-sow-lite/analysis/index.json', 'analysis_index')
    records = []
    for ref in index['items']:
        path = safe_path(project, ref['path'], '.ai-sow-lite/analysis/topics')
        if file_ref(project, path) != ref:
            raise StorageError('EVIDENCE_MISSING', '分析记录实际字节与登记摘要不一致。')
        records.append(checked_json(project, ref['path'], 'analysis', '.ai-sow-lite/analysis/topics'))
    return index['items'], records


def ingest_analysis(project: Path, request_id: str, payload):
    from .validation import check_analysis
    project = Path(project).resolve()
    checked_json(project, '.ai-sow-lite/project.json', 'project')
    ensure_request(project, request_id, payload['entrypoint'])
    area = request_area(request_id, payload['entrypoint'])
    path = safe_path(project, payload['analysis_path'], area)
    raw = path.read_bytes()
    from .contracts import strict_json_loads
    analysis = strict_json_loads(raw)
    from .contracts import schema_validator
    from ._prototype import register_observations, topic_observations
    if list(schema_validator('artifacts', 'analysis').iter_errors(analysis)):
        raise StorageError('CANDIDATE_INVALID', '分析字段不符合当前 Schema。')
    adopted_inputs = {i for t in analysis['topics'] for i in t['input_version_ids']}
    entries = {e['input_version_id']: e for e in input_index(project)['items'] if e['input_version_id'] in adopted_inputs}
    analysis['observations'] = register_observations(project, analysis['observations'], entries, area)
    issues = check_analysis(project, analysis)
    if issues:
        error = StorageError('CANDIDATE_INVALID', '分析来源或结构无效。', payload['analysis_path'])
        error.diagnostics = issues
        raise error
    # Write path: an index lost to an interrupted registration is rebuilt here,
    # from re-proved sources, so the same request can be replayed to completion.
    recover_analysis_index(project)
    refs, old_records = _analysis_records(project)
    evidence, topics = {}, {}
    for record in old_records:
        for item in record['evidence']:
            if item['id'] in evidence and evidence[item['id']] != item:
                raise StorageError('IDENTITY_CONFLICT', '既有分析中的依据身份冲突。')
            evidence[item['id']] = item
        for item in record['topics']:
            topics[item['topic_version_id']] = item
    for collection, old, identity in [(analysis['evidence'], evidence, 'id'),
                                       (analysis['topics'], topics, 'topic_version_id')]:
        for item in collection:
            if item[identity] in old and item != old[item[identity]]:
                raise StorageError('IDENTITY_CONFLICT', '已登记的依据或主题版本不能覆盖不同含义。')
    digest = hashlib.sha256(raw).hexdigest()
    registration_relative = f'.ai-sow-lite/analysis/registrations/{digest}/analysis.json'
    candidate_ref = dict(path=registration_relative, sha256=digest)
    available_evidence, available_observations = set(), set()
    for topic in analysis['topics']:
        relative = f".ai-sow-lite/analysis/topics/{topic['topic_version_id']}/analysis.json"
        if safe_path(project, relative).exists():
            stored = checked_json(project, relative, 'analysis', '.ai-sow-lite/analysis/topics')
            if stored['topics'] != [topic]:
                raise StorageError('IDENTITY_CONFLICT', '已有主题目录不能绑定不同分析内容。')
            provenance = checked_json(project, (Path(relative).parent / 'registration-ref.json').as_posix(), 'file_ref')
            original = safe_path(project, provenance['path'], '.ai-sow-lite/analysis/registrations')
            if file_ref(project, original) != provenance:
                raise StorageError('EVIDENCE_MISSING', '已有主题的来源原字节与绑定摘要不同。')
            registered = checked_json(project, provenance['path'], 'analysis', '.ai-sow-lite/analysis/registrations')
            if not {r['sha256'] for r in stored['observations']} <= {r['sha256'] for r in registered['observations']}:
                raise StorageError('EVIDENCE_MISSING', '已有主题的观察摘要不属于原始登记来源。')
            bound_record = dict(schema_version='1.0', topics=[topic], evidence=registered['evidence'], observations=stored['observations'])
            if topic not in registered['topics'] or safe_path(project, relative).read_bytes() != canonical_json_bytes(bound_record):
                raise StorageError('EVIDENCE_MISSING', '已有主题原字节与登记来源不同。')
            if not any(ref['path'] == relative for ref in refs) and provenance != candidate_ref:
                raise StorageError('EVIDENCE_MISSING', '缺失索引项只能由相同原字节的候选完成登记。')
            available_evidence.update(e['id'] for e in stored['evidence'])
            available_observations.update((r['path'], r['sha256']) for r in stored['observations'])
        else:
            available_evidence.update(e['id'] for e in analysis['evidence'])
            available_observations.update((r['path'], r['sha256']) for r in topic_observations(project, analysis['observations'], topic))
    if not {e['id'] for e in analysis['evidence']} <= available_evidence:
        raise StorageError('IDENTITY_CONFLICT', '新依据需要新的主题版本承载；不能声称已登记到旧不可变主题。')
    if not {(r['path'], r['sha256']) for r in analysis['observations']} <= available_observations:
        raise StorageError('IDENTITY_CONFLICT', '新观察需要新的主题版本承载；不能声称已登记到旧不可变主题。')
    # Keep the actual candidate bytes as provenance; split topics for narrow adoption.
    registration = safe_path(project, registration_relative)
    atomic_bytes(registration, raw, immutable=True)
    for topic in analysis['topics']:
        relative = f".ai-sow-lite/analysis/topics/{topic['topic_version_id']}/analysis.json"
        existing = safe_path(project, relative)
        record = dict(schema_version='1.0', topics=[topic], evidence=analysis['evidence'],
                      observations=topic_observations(project, analysis['observations'], topic))
        if existing.exists():
            # Shared evidence can arrive in a different bundle; the immutable topic stays as first registered.
            if any(ref['path'] == relative for ref in refs):
                continue
            # Preflight proved this is the same candidate with only its index member missing.
            ref = file_ref(project, existing)
        else:
            # Bind provenance before the topic record: a topic that exists without
            # its registration-ref cannot be recovered, but the reverse can.
            write_json(project, (Path(relative).parent / 'registration-ref.json').as_posix(),
                       file_ref(project, registration), immutable=True)
            ref = write_json(project, relative, record, immutable=True)
        refs.append(ref)
    write_json(project, '.ai-sow-lite/analysis/index.json', dict(schema_version='1.0', items=refs))
    return dict(analysis_ref=file_ref(project, registration), evidence_ids=[e['id'] for e in analysis['evidence']],
                topic_version_ids=[t['topic_version_id'] for t in analysis['topics']])


def _selected(items, selector, fields):
    for field in fields:
        if field + 's' in selector:
            return [item for item in items if item[field] in selector[field + 's']]
    return items


def inspect_view(project: Path, payload):
    import base64
    from .contracts import semantic_digest, strict_json_loads
    from .project import recover_request
    project = Path(project).resolve()
    view, selector = payload['view'], payload['selector']
    if view == 'telemetry':
        from .telemetry import inspect_report
        return inspect_report(project, payload)
    if view not in ('inputs', 'regions', 'topics', 'current', 'request', 'objects', 'standards'):
        raise StorageError('OPERATION_UNSUPPORTED', '此视图尚未实现，不能返回模拟数据。')
    checked_json(project, '.ai-sow-lite/project.json', 'project')
    version, coverage = None, dict(view=view, selector=selector)
    if view == 'inputs':
        index = input_index(project)
        version = semantic_digest(index)
        items = _selected(index['items'], selector, ('input_version_id', 'input_id'))
    elif view == 'topics':
        refs, records = _analysis_records(project)
        version = semantic_digest(refs)
        items = _selected([t for r in records for t in r['topics']], selector, ('topic_version_id', 'topic_id'))
        if 'uses' in selector or 'historical_label' in selector:
            uses = selector.get('uses', ['as-is'])
            items = [t for t in items if set(t['uses']) & set(uses)]
            # An empty label query still depends on unanalysed members of this use.
            members = [e for e in input_index(project)['items'] if set(e['uses']) & set(uses)]
            source_members = [dict(input_version_id=e['input_version_id'], content_hash=e['content_hash'],
                                   uses=e['uses'], use_regions=e['use_regions']) for e in members]
            relevant_refs = [ref for ref, record in zip(refs, records) if any(set(t['uses']) & set(uses) for t in record['topics'])]
            scope = dict(uses=uses, source_members=source_members, analysis_refs=relevant_refs)
            version = semantic_digest(scope)
            coverage.update(scope_digest=version, source_members=source_members, analysis_refs=relevant_refs)
            if 'historical_label' in selector:
                items = [dict(topic_id=t['topic_id'], topic_version_id=t['topic_version_id'], historical_item=h)
                         for t in items for h in t['historical_items'] if selector['historical_label'] in h['label']]
    elif view == 'regions':
        entries = {entry['input_version_id']: entry for entry in input_index(project)['items']}
        entry = entries.get(selector['input_version_id'])
        if entry is None:
            raise StorageError('EVIDENCE_MISSING', '未登记指定输入版本。')
        if entry['format'] == 'prototype':
            from . import _prototype
            _, _, contents = _prototype.checked_package(project, entry)
            locator = selector.get('locator')
            version = dict(input_version_id=entry['input_version_id'], content_hash=entry['content_hash'],
                           adapter_version=_prototype.READER_VERSION)
            coverage['limitations'] = list(_prototype.LIMITATIONS)
            if locator is None:
                coverage['reading_ref'] = _reading(project, entry)
                items = []
                for resource in entry['resources']:
                    item = dict(resource, file_ref=_prototype.resource_ref(entry, resource), readable=False)
                    if Path(resource['path']).suffix.lower() in _prototype.TEXT_SUFFIXES:
                        try:
                            raw = contents[resource['path']]
                            lines = raw.decode('utf-8-sig' if raw.startswith(b'\xef\xbb\xbf') else 'utf-8').splitlines(keepends=True)
                            if b'\x00' in raw:
                                raise UnicodeError('binary')
                            item.update(readable=True, line_count=len(lines))
                            if lines:
                                item['locator'] = dict(kind='text_lines', path=resource['path'], start_line=1, end_line=len(lines))
                        except UnicodeError:
                            item['limitation'] = '无法按 UTF-8 读取；仅保留原件。'
                    items.append(item)
            elif locator['kind'] == 'observation':
                raw, record, dependencies = _prototype.observation_excerpt(project, entry, locator)
                version['observation_ref'] = dependencies[0]
                coverage.update(locator=locator, excerpt_hash=hashlib.sha256(raw).hexdigest(),
                                excerpt_ref=dependencies[0], attachment_refs=_prototype.attachment_refs(record))
                coverage['limitations'] = [record['limitations'], '仅校验提供的观察记录及附件，未执行或复放浏览器。']
                items = [record]
            elif locator['kind'] == 'text_lines':
                lines = _prototype.source_lines(contents, locator)
                coverage.update(locator=locator, excerpt_hash=hashlib.sha256(''.join(lines).encode('utf-8')).hexdigest(),
                                line_count=len(lines))
                items = [dict(line_number=n, character_offset=offset, text=line[offset:offset + 8192])
                         for n, line in enumerate(lines, locator['start_line']) for offset in range(0, len(line), 8192)]
            else:
                raise StorageError('CANDIDATE_INVALID', '原型区域需要 text_lines 或 observation。')
        elif entry['format'] == 'xlsx':
            from . import _xlsx
            selection = {'kind': 'workbook'}
            locator = selector.get('locator')
            if locator is not None:
                if locator['kind'] != 'xlsx_range':
                    raise StorageError('CANDIDATE_INVALID', 'XLSX 必须使用 xlsx_range。')
                _checked_xlsx_reading(project, entry, locator['read_id'])
                selection = dict(kind='xlsx_range', sheet=locator['sheet'], range=_xlsx.canonical_range(locator['range']))
            ref, reading, observed = _xlsx_reading(project, entry, selection)
            version = dict(input_version_id=entry['input_version_id'], content_hash=entry['content_hash'],
                           read_id=reading['read_id'], selection=selection, adapter_version=reading['adapter_version'], options=reading['options'])
            coverage.update(reading_ref=ref)
            if locator is None:
                items = observed['sheets']
                coverage['limitations'] = observed['limitations']
            else:
                excerpt = reading['excerpts'][0]['file_ref']
                coverage.update(locator=dict(selection, read_id=reading['read_id']), excerpt_ref=excerpt,
                                excerpt_hash=excerpt['sha256'], metadata=observed['metadata'], cell_count=len(observed['cells']))
                coverage['limitations'] = observed['limitations']
                items = [cell if len(canonical_json_bytes(cell)) <= 48 * 1024 else
                         dict(address=cell['address'], representation='attachment', excerpt_ref=excerpt,
                              message='单格正文超过分页上限；完整类型、值及附注保存在附件的同址 cell。')
                         for cell in observed['cells']]
        else:
            locator = selector.get('locator')
            if locator is not None and locator['kind'] != 'text_lines':
                raise StorageError('CANDIDATE_INVALID', '文本区域需要 text_lines。')
            lines = text_lines(project, entry)
            if locator is None:
                ref = _reading(project, entry)
                reading = checked_json(project, ref['path'], 'reading', '.ai-sow-lite/inputs/readings')
                version = dict(input_version_id=entry['input_version_id'], content_hash=entry['content_hash'],
                               adapter_version=reading['adapter_version'], read_id=reading['read_id'])
                items = _text_directory(lines)
                coverage.update(reading_ref=ref, line_count=len(lines), limitations=[])
            else:
                if locator.get('path') not in (None, entry['relative_path'], Path(entry['relative_path']).name):
                    raise StorageError('EVIDENCE_MISSING', '定位路径与已登记原件不同。')
                start, end = locator['start_line'], locator['end_line']
                if not 1 <= start <= end <= len(lines):
                    raise StorageError('CANDIDATE_INVALID', '行范围超出实际文本。')
                selected = lines[start - 1:end]
                version = dict(input_version_id=entry['input_version_id'], content_hash=entry['content_hash'])
                coverage['excerpt_hash'] = hashlib.sha256(''.join(selected).encode('utf-8')).hexdigest()
                coverage['line_count'] = len(selected)
                # Finite byte-bounded fragments let even a single long line be continued exactly.
                items = []
                for number, line in enumerate(selected, start):
                    offset = 0
                    while offset < len(line):
                        fragment = line[offset:offset + 8192]
                        items.append(dict(line_number=number, character_offset=offset, text=fragment))
                        offset += len(fragment)
    elif view == 'objects':
        version, items, selected_ref = _inspect_objects(project, selector)
        coverage.update(verification_scope='selected_file' if selected_ref else 'manifest', verified_file_ref=selected_ref)
    elif view == 'standards':
        version, items = _inspect_standards(project, selector)
    elif view == 'request':
        target = selector['request_id']
        recovered = recover_request(project, target)
        if recovered.get('diagnostics'):
            error = StorageError('VERSION_INCOMPATIBLE', '恢复事实不兼容。')
            error.diagnostics = recovered['diagnostics']
            raise error
        items = [recovered]
        version = semantic_digest(items)
    else:
        from .project import read_current_manifest
        pointer, manifest = read_current_manifest(project)
        coverage.update(verification_scope='manifest', verified_file_ref=None)
        version = pointer
        items = [] if pointer is None else [dict(current=pointer, manifest_ref=dict(
            path=f".ai-sow-lite/versions/{pointer['version_id']}/manifest.json", sha256=pointer['manifest_hash']))]
    binding = semantic_digest(dict(view=view, selector=selector, version=version))
    offset = 0
    if payload.get('cursor') is not None:
        try:
            decoded = strict_json_loads(base64.urlsafe_b64decode(payload['cursor'].encode('ascii')))
            if set(decoded) != {'binding', 'offset'} or decoded['binding'] != binding:
                raise ValueError('binding')
            offset = decoded['offset']
            if type(offset) is not int or not 0 < offset < len(items):
                raise ValueError('offset')
        except (ValueError, TypeError, KeyError, UnicodeError):
            raise StorageError('VERSION_INCOMPATIBLE', '游标与当前查询或来源版本不同；请用 cursor=null 明确读取新集合。') from None
    # CLI adds bounded observation status and envelope after this reader returns.
    body_limit = 64 * 1024 - 1024

    def page(selected):
        end = offset + len(selected)
        cursor = None if end == len(items) else base64.urlsafe_b64encode(canonical_json_bytes(dict(binding=binding, offset=end))).decode('ascii')
        return dict(selected_version=version, items=selected, matched_count=len(items), returned_count=len(selected),
                    remaining_count=len(items) - end, next_cursor=cursor,
                    coverage=dict(coverage, start_offset=offset, end_offset=end, complete=cursor is None), report_ref=None)

    def too_large():
        error = StorageError('RESULT_TOO_LARGE', '查询元数据或单项正文超过64 KiB返回上限；请缩小查询或读取已保存的附件。')
        if coverage.get('excerpt_ref'):
            error.diagnostics[0]['preserved_paths'] = [coverage['excerpt_ref']['path']]
        return error

    selected = []
    for item in items[offset:offset + payload.get('limit', 20)]:
        candidate = selected + [item]
        if len(canonical_json_bytes(page(candidate))) > body_limit:
            if not selected:
                raise too_large()
            break
        selected = candidate
    result = page(selected)
    # Zero matches still return scope/selector/limits, which must obey the same bound.
    if len(canonical_json_bytes(result)) > body_limit:
        raise too_large()
    return result


def _inspect_objects(project, selector):
    from .project import read_current_manifest, _read_manifest
    from .contracts import schema_validator, strict_json_loads
    pointer, manifest = read_current_manifest(project)
    requested = selector.get('version_id')
    seen = set()
    while requested is not None and manifest is not None and manifest['version_id'] != requested:
        if manifest['version_id'] in seen:
            raise StorageError('VERSION_INCOMPATIBLE', '历史基线链成环。')
        seen.add(manifest['version_id'])
        base = manifest['base_version_id']
        if base is None:
            manifest = None
            break
        ref = next(r for r in manifest['dependencies'] if r['path'] == f'.ai-sow-lite/versions/{base}/manifest.json')
        pointer = dict(version_id=base, manifest_hash=ref['sha256'])
        manifest = _read_manifest(project, pointer)
    if manifest is None:
        if requested is not None:
            raise StorageError('VERSION_INCOMPATIBLE', '指定版本不在 current 的有效历史链中。')
        return None, [], None
    collection = selector['collection']
    filename = {'pending_items': 'pending-items.json', 'decisions': 'decisions.json'}.get(collection, 'model.json')
    ref = next(r for r in manifest['files'] if Path(r['path']).name == filename)
    raw = safe_path(project, ref['path'], f".ai-sow-lite/versions/{manifest['version_id']}").read_bytes()
    if hashlib.sha256(raw).hexdigest() != ref['sha256']:
        raise StorageError('EVIDENCE_MISSING', '所选业务文件实际字节与 manifest 绑定摘要不同。', ref['path'])
    try:
        data = strict_json_loads(raw)
    except (ValueError, UnicodeError):
        raise StorageError('VERSION_INCOMPATIBLE', '所选业务文件不是有效 JSON。', ref['path']) from None
    if list(schema_validator(filename.removesuffix('.json')).iter_errors(data)):
        raise StorageError('VERSION_INCOMPATIBLE', '版本业务文件不符合 Schema，不能当作空集合。', ref['path'])
    if collection == 'acs':
        items = [dict(ac, story_id=story['id']) for story in data['stories'] for ac in story['acs']]
    elif collection in ('pending_items', 'decisions'):
        items = data['items']
    else:
        items = data[collection]
    if 'object_ids' in selector:
        items = [item for item in items if item['id'] in selector['object_ids']]
    elif 'title' in selector:
        items = [item for item in items if selector['title'] in (item['question'] if collection == 'pending_items'
                 else item.get('title', item.get('name', item.get('text', ''))))]
    elif 'status' in selector:
        items = [item for item in items if item['status'] == selector['status']]
    else:
        relation = selector['relation']
        fields = {'incoming': ['to_story_id'], 'outgoing': ['from_story_id'], 'both': ['from_story_id', 'to_story_id']}[relation['direction']]
        items = [item for item in items if any(item[field] == relation['object_id'] for field in fields)]
    return pointer, items, ref


def _inspect_standards(project, selector):
    from io import BytesIO
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    from zipfile import BadZipFile
    record = checked_json(project, '.ai-sow-lite/project.json', 'project')
    template = safe_path(project, f".ai-sow-lite/template/{record['template_hash']}/sow-template.xlsx")
    raw = template.read_bytes()
    if hashlib.sha256(raw).hexdigest() != record['template_hash']:
        raise StorageError('EVIDENCE_MISSING', '模板原字节与项目固定身份不同。')
    columns = ['工作类型 ID', '分类', '工作类型']
    if selector:
        columns += ['计量单位', '标准交付对象', '模式化完成标准', '说明', '主要计量维度',
                    'S（简单）', 'M（标准）', 'L（复杂）', 'X / 拆分条件', '模式适用说明', 'SIT适用', 'UAT适用']
    try:
        book = load_workbook(BytesIO(raw), data_only=False)
        try:
            matches = [(sheet, sheet.tables['TaskStandardTable']) for sheet in book if 'TaskStandardTable' in sheet.tables]
            if len(matches) != 1:
                raise ValueError('table')
            sheet, table = matches[0]
            left, top, right, bottom = range_boundaries(table.ref)
            rows = list(sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right))
            headers = [cell.value for cell in rows[0]]
            items = []
            for row in rows[1:]:
                cells = {name: row[headers.index(name)] for name in columns}
                if any(cell.data_type == 'f' or (cell.value is not None and not isinstance(cell.value, str)) for cell in cells.values()):
                    raise ValueError('qualitative columns')
                items.append({name: cell.value for name, cell in cells.items()})
        finally:
            book.close()
    except (ValueError, KeyError, BadZipFile):
        raise StorageError('VERSION_INCOMPATIBLE', '模板定性目录不兼容；不返回倍率或公式副本。') from None
    if 'work_type_ids' in selector:
        items = [item for item in items if item['工作类型 ID'] in selector['work_type_ids']]
    if 'work_type_names' in selector:
        items = [item for item in items if item['工作类型'] in selector['work_type_names']]
    return dict(template_hash=record['template_hash']), items
