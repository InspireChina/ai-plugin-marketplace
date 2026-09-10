"""D06 template projection and read-only verification; no estimation arithmetic.

Mechanical origins: D00 2fc8588 workbook.safe_text/copy_style/fill_table,
scan_workbook_integrity/dual_reopen. Lite keeps reserved rows, writes literal
strings without an added apostrophe. Controller Ruling3, under prior authorization
to reuse proven mechanics, permits two metadata fixes before the read-only seal.
"""
from __future__ import annotations

from copy import copy
import hashlib
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.formula import ArrayFormula

from .project import StorageError, atomic_bytes

TEMPLATE_HASH = '28d23be2b50e6abcb3abd81e97ebba3e9e696a8994e05dd9bda72987a1c8ab9e'
INLINE_TEMPLATE_HASH = '7b96f9d2d6f6f6175c4d99d875ee3cf0743df3d6884d64258271993432299919'
LEGACY_TEMPLATE_HASH = '6abc55d44bc66476a60c2251e18c0dfdb66709e07539c246dfdec3a0373f5332'
SUPPORTED_TEMPLATE_HASHES = (TEMPLATE_HASH, INLINE_TEMPLATE_HASH, LEGACY_TEMPLATE_HASH)
PROJECTOR_VERSION = 'lite-projection-v1'
# Render retry identity is separate from the unchanged projection data contract.
RENDER_IMPLEMENTATION_VERSION = 'lite-render-v9'
AC_COLUMN_WIDTH = 88
STORY_SHEET, TASK_SHEET = '01-需求故事', '02-任务清单'
SHEETS = (STORY_SHEET, TASK_SHEET, '03-工作量汇总', '90-估算标准')
XML_BAD = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def write_literal(cell, value: str) -> None:
    if not isinstance(value, str) or XML_BAD.search(value) or len(value.encode('utf-16-le')) // 2 > 32767:
        raise StorageError('WORKBOOK_INVALID', '文字无法完整写入单元格；请检查 XML 字符或 Excel 单元格容量。')
    cell.value = value
    cell.data_type = 's'


def name_key(value):
    return unicodedata.normalize('NFC', value).casefold().strip()


def _safe_name(value):
    return (bool(value.strip()) and value == value.strip() and not XML_BAD.search(value)
            and len(value.encode('utf-16-le')) // 2 <= 120
            and not any(c in value for c in '~*?\r\n\t') and value[0] not in '=<>\"'
            and not re.fullmatch(r'[+-]?(?:[0-9][0-9,.]*)(?:[eE][+-]?[0-9]+)?%?', value)
            and value.upper() not in ('TRUE','FALSE','#N/A','#VALUE!','#REF!','#NAME?','#NUM!','#NULL!','#DIV/0!'))


def allocate_aliases(items, field, prefix, previous=None):
    """Prior names are scoped by the caller to one baseline; original text is untouched."""
    previous = previous or {}
    names, used = {}, set()
    for obj in sorted(items, key=lambda x: x['id']):
        old = previous.get(obj['id'])
        if old and old['original'] == obj[field] and _safe_name(old['display_name']):
            key = name_key(old['display_name'])
            if key not in used:
                names[obj['id']] = old['display_name']; used.add(key)
    for obj in sorted(items, key=lambda x: x['id']):
        identity, title = obj['id'], obj[field]
        if identity in names:
            continue
        if _safe_name(title) and name_key(title) not in used:
            alias = title
        else:
            # UUID-derived suffix grows on collision, including common UUID prefixes.
            code = identity.replace('-', '')
            for length in range(8, len(code) + 1):
                suffix = code[-length:]
                alias = f'{title}〔{suffix}〕' if _safe_name(title) else f'{prefix}-{suffix}'
                if not _safe_name(alias):
                    alias = f'{prefix}-{suffix}'
                if name_key(alias) not in used:
                    break
            else:
                raise StorageError('WORKBOOK_INVALID', '无法分配唯一的安全显示名。')
        names[identity] = alias; used.add(name_key(alias))
    return names


def formula_text(value):
    return value.text if isinstance(value, ArrayFormula) else value


def _template(path):
    digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    if digest not in SUPPORTED_TEMPLATE_HASHES:
        raise StorageError('VERSION_INCOMPATIBLE', '模板字节或公式原型不属于本投影器支持的版本。')
    book = load_workbook(path)
    expected = {
        STORY_SHEET: ('SOWStoryTable', 'A4:J64', ['需求','子需求','故事','验收条件','备注','SIT适用','UAT适用','任务列表','故事人天','校验结果']),
        TASK_SHEET: ('TaskTable', 'A4:L204', ['所属故事','任务名称','工作类型名称','工作方式','复杂度','集成类型','备注','M档标准人天','复杂度系数','任务人天','SIT支持人天','校验结果'])}
    if tuple(book.sheetnames) != SHEETS:
        raise StorageError('VERSION_INCOMPATIBLE', '模板 Sheet 不兼容。')
    for sheet, (name, ref, headers) in expected.items():
        table = book[sheet].tables[name]
        if table.ref != ref or [c.name for c in table.tableColumns] != headers:
            raise StorageError('VERSION_INCOMPATIBLE', '模板 Table 或表头不兼容。')
        for col, metadata in enumerate(table.tableColumns, 1):
            cell = book[sheet].cell(5, col)
            if metadata.calculatedColumnFormula:
                if cell.data_type != 'f' or metadata.calculatedColumnFormula.text != formula_text(cell.value)[1:]:
                    raise StorageError('VERSION_INCOMPATIBLE', 'Table 计算列原型不一致。')
    if digest != TEMPLATE_HASH:
        # The user-authorized validation display update applies to the known
        # previous templates too. Their project bytes/hash and all estimate formulas
        # remain intact. Read the actual shipped prototypes, never reimplement them.
        current = _template(Path(__file__).resolve().parents[2] / 'assets/sow-template.xlsx')
        for sheet, column in [(STORY_SHEET, 10), (TASK_SHEET, 12)]:
            ws, source = book[sheet], current[sheet]
            ws['A2'].value = source['A2'].value
            for row in range(5, ws.max_row + 1):
                ws.cell(row, column).value = source.cell(row, column).value
            table, prototype = next(iter(ws.tables.values())), next(iter(source.tables.values()))
            table.tableColumns[column - 1].calculatedColumnFormula = copy(prototype.tableColumns[column - 1].calculatedColumnFormula)
        current.close()
    return book


def _extend(book, sheet, table_name, count):
    ws = book[sheet]; table = ws.tables[table_name]
    end = int(re.search(r'\d+$', table.ref).group())
    new_end = max(end, count + 4)
    if new_end > 1048576:
        raise StorageError('WORKBOOK_INVALID', '行数超出当前模板容量。')
    for row in range(end + 1, new_end + 1):
        ws.row_dimensions[row] = copy(ws.row_dimensions[5])
        ws.row_dimensions[row].index = row
        for source in ws[5]:
            cell = ws.cell(row, source.column)
            cell._style = copy(source._style)
            if source.data_type == 'f':
                translated = Translator(formula_text(source.value), origin=source.coordinate).translate_formula(cell.coordinate)
                cell.value = ArrayFormula(ref=cell.coordinate, text=translated) if isinstance(source.value, ArrayFormula) else translated
    if new_end != end:
        table.ref = re.sub(r'\d+$', str(new_end), table.ref)
        if table.autoFilter:
            table.autoFilter.ref = table.ref
    # This pinned template uses structured cross-table references, full-column
    # validation/conditional ranges, and no data-sheet print area. No A1 range
    # rewrite is needed or authorized; all those controls remain unchanged.


def _cell(sheet, coordinate, entry=None):
    value = dict(sheet=sheet, cell=coordinate)
    if entry is not None:
        value['entry'] = entry
    return value


def _block(text):
    # A fence longer than any source run preserves even embedded Markdown fences.
    fence = '`' * max(3, max((len(x) + 1 for x in re.findall(r'`+', text)), default=0))
    return f'{fence}text\n{text}\n{fence}\n'


def _projection(model, pending, decisions, version_id, evidence=None, previous=None, *, layout):
    """Project authored acceptance and open questions into the original tables."""
    stories = [s for e in model['epics'] for f in model['features'] if f['epic_id'] == e['id']
               for s in model['stories'] if s['feature_id'] == f['id']]
    tasks = [t for s in stories for t in model['tasks'] if t['story_id'] == s['id']]
    if len(stories) != len(model['stories']) or len(tasks) != len(model['tasks']):
        raise StorageError('WORKBOOK_INVALID', '对象父关系无法投影。')
    sn = allocate_aliases(stories, 'title', 'S', (previous or {}).get('stories'))
    tn = allocate_aliases(tasks, 'name', 'T', (previous or {}).get('tasks'))
    objects, inputs, notes = [], {}, {}
    story_rows = {s['id']: 5 + i for i, s in enumerate(stories)}
    # A known but undivided scope still needs a visible place for its question.
    # These rows carry only actual Epic/Feature titles, never invented Stories.
    scope_rows = []
    open_targets = {t['object_id'] for p in pending['items'] if p['status']=='open' for t in p['targets']}
    for epic in model['epics']:
        features = [f for f in model['features'] if f['epic_id'] == epic['id']]
        for feature in features:
            if not any(s['feature_id'] == feature['id'] for s in stories) and open_targets.intersection((epic['id'], feature['id'])):
                scope_rows.append((epic['id'], feature['id']))
        if not features and epic['id'] in open_targets:
            scope_rows.append((epic['id'], None))

    def text_value(value, identity, field):
        value = value.replace('\r\n', '\n').replace('\r', '\n')
        if XML_BAD.search(value) or len(value.encode('utf-16-le')) // 2 > 32767:
            error = StorageError('WORKBOOK_INVALID', '字段无法完整写入 Excel 单元格；请处理非法字符或精简该字段，保留验收边界。')
            error.diagnostics[0]['target'].update(object_id=identity, field=field)
            raise error
        return value

    def put(sheet, row, column, value, identity, field):
        if value not in ('', None):
            inputs[(sheet, f'{column}{row}')] = (text_value(value, identity, field), identity, field)

    note_sources = {}

    def note_source(key, identity, field, value):
        # A composed note still needs a repairable source field, including when
        # the row is only an Epic/Feature scope or the text is an alias original.
        size = len(text_value(value, identity, field))
        if key not in note_sources or size > note_sources[key][0]:
            note_sources[key] = (size, identity, field)

    def obj_record(obj, kind, sheet, table, rows, name, fields):
        objects.append(dict(object_id=obj['id'], kind=kind, sheet=sheet, table=table,
                            rows=rows, display_name=name, fields=fields))

    for kind, collection, column in [('epic','epics','A'),('feature','features','B')]:
        for obj in model[collection]:
            features = {f['id'] for f in model['features'] if f['epic_id'] == obj['id']} if kind == 'epic' else {obj['id']}
            rows = [story_rows[s['id']] for s in stories if s['feature_id'] in features]
            rows += [5 + len(stories) + i for i, (epic, feature) in enumerate(scope_rows)
                     if (epic if kind == 'epic' else feature) == obj['id']]
            cells = [_cell(STORY_SHEET, f'{column}{r}') for r in rows]
            for row in rows:
                put(STORY_SHEET, row, column, obj['title'], obj['id'], 'title')
            obj_record(obj, kind, STORY_SHEET, 'SOWStoryTable', rows, None, [dict(field='title',cells=cells)])
    for story in stories:
        row = story_rows[story['id']]
        for index, ac in enumerate(story['acs'], 1):
            obj_record(ac, 'ac', STORY_SHEET, 'SOWStoryTable', [row], None,
                       [dict(field='text',cells=[_cell(STORY_SHEET, f'D{row}', index)])])
        ac_text = '\n'.join(f'{i}. {ac["text"]}' for i, ac in enumerate(story['acs'], 1))
        put(STORY_SHEET, row, 'C', sn[story['id']], story['id'], 'title')
        put(STORY_SHEET, row, 'D', ac_text, story['id'], 'acs')
        fields = [dict(field=f,cells=[_cell(STORY_SHEET,f'{c}{row}')]) for f,c in [('title','C'),('acs','D'),('notes','E')]]
        obj_record(story, 'story', STORY_SHEET, 'SOWStoryTable', [row], sn[story['id']], fields)
    for row, task in enumerate(tasks, 5):
        columns = [('story_id','A'),('name','B'),('work_type_name','C'),('work_mode','D'),('complexity','E'),('integration_type','F'),('notes','G')]
        for field, col in columns[:-1]:
            value = sn[task['story_id']] if field == 'story_id' else tn[task['id']] if field == 'name' else task[field]
            put(TASK_SHEET, row, col, value, task['id'], field)
        obj_record(task, 'task', TASK_SHEET, 'TaskTable', [row], tn[task['id']],
                   [dict(field=f,cells=[_cell(TASK_SHEET,f'{c}{row}')]) for f,c in [*columns, ('classification_basis', 'G')]])
    by_id = {o['object_id']:o for o in objects}
    mappings = []
    for item in pending['items']:
        targets, affected = [], {}
        for target in item['targets']:
            obj = by_id.get(target['object_id'])
            cells = [c for f in obj['fields'] if target['field'] in (None, f['field']) for c in f['cells']] if obj else []
            targets.append(dict(target, cells=cells))
            if obj:
                for row in obj['rows']:
                    key = (obj['sheet'], row)
                    affected.setdefault(key, [])
                    if obj['kind'] == 'ac':
                        label = f"第 {obj['fields'][0]['cells'][0]['entry']} 条 AC"
                        if label not in affected[key]: affected[key].append(label)
        mappings.append(dict(pending_item_id=item['id'], path='pending-items.md',
                             anchor=f'pending-{item["id"]}', targets=targets))
        if item['status'] == 'open':
            if not affected:
                error = StorageError('WORKBOOK_INVALID', '待确认事项没有可展示的关联对象。')
                error.diagnostics[0]['target'].update(object_id=item['id'], field='targets')
                raise error
            for key, labels in affected.items():
                for field in ('question', 'current_handling'):
                    note_source(key, item['id'], field, item[field])
                context = ('（' + '、'.join(labels) + '）') if labels else ''
                text = f'待确认：{context}{item["question"]}\n当前处理：{item["current_handling"]}'
                if item['unestimated_work']: text += '\n未拆明工作，尚未估算。'
                notes.setdefault(key, []).append(text)
    # Pending questions stay first for the template's status formula. Project
    # existing classification reasons without inferring or rewriting them.
    for collection, sheet, names, name_field in [(stories,STORY_SHEET,sn,'title'), (tasks,TASK_SHEET,tn,'name')]:
        for obj in collection:
            key = (sheet, by_id[obj['id']]['rows'][0])
            if sheet == TASK_SHEET:
                reasons = {}
                for basis in obj['classification_basis']:
                    reasons.setdefault(basis['rationale'], set()).update(basis['fields'])
                for rationale, fields in reasons.items():
                    labels = [f'{label}（{obj[field]}）' for field, default, label in
                              [('work_mode', '新建', '工作方式'), ('complexity', 'M', '复杂度')]
                              if field in fields and obj[field] not in (None, default)]
                    if labels:
                        note_source(key, obj['id'], 'classification_basis', rationale)
                        notes.setdefault(key, []).append('、'.join(labels) + '判断原因：' + rationale)
            if obj.get('notes'):
                note_source(key, obj['id'], 'notes', obj['notes'])
                notes.setdefault(key, []).append(obj['notes'])
            if names[obj['id']] != obj[name_field]:
                note_source(key, obj['id'], name_field, obj[name_field])
                notes.setdefault(key, []).append('原名：' + obj[name_field])
    for (sheet, row), parts in notes.items():
        _, identity, field = note_sources[(sheet, row)]
        put(sheet, row, 'E' if sheet == STORY_SHEET else 'G', '\n'.join(parts), identity, field)
    pending_lines=[f'# 待确认事项\n\nversion_id: {version_id}\n']
    item_by_id={p['id']:p for p in pending['items']}
    for mapping in mappings:
        item=item_by_id[mapping['pending_item_id']]
        pending_lines.extend([f'<a id="{mapping["anchor"]}"></a>\n',f'## {item["id"]}\n',f'状态：{item["status"]}\n',
                              '问题：\n'+_block(item['question']),'当前处理：\n'+_block(item['current_handling'])])
        if item['unestimated_work']: pending_lines.append('未拆明业务工作：是\n')
        for target in mapping['targets']:
            positions='、'.join(f'{c["sheet"]}!{c["cell"]}' for c in target['cells']) or '无工作簿行'
            pending_lines.append(f'目标：{target["object_id"]} / {target["field"]}；{positions}\n')
        for identity in item['evidence_refs']:
            pending_lines.append(f'依据：{identity}\n'+(evidence or {}).get(identity,''))
        if item['resolution']:
            from .contracts import canonical_json_bytes
            pending_lines.append('处理记录：\n'+_block(canonical_json_bytes(item['resolution']).decode()))
        for decision in decisions['items']:
            if item['resolution'] and item['resolution'].get('decision_id')==decision['id']:
                from .contracts import canonical_json_bytes
                pending_lines.append('已采用答复：\n'+_block(canonical_json_bytes(decision).decode()))
    if not mappings: pending_lines.append('无待确认事项\n')
    docs={'pending-items.md':'\n'.join(pending_lines)}
    projection=dict(schema_version='1.0',version_id=version_id,projector_version=PROJECTOR_VERSION,
                    template_hash=TEMPLATE_HASH,objects=objects,pending_items=mappings,details=[])
    return projection,inputs,docs


def project_workbook(template, model, pending, decisions, version_id, directory, *, evidence=None, previous=None):
    """Write a pre-Office workbook and accompanying text; caller supplies checked data."""
    book,projection,docs=_expected_book(template,model,pending,decisions,version_id,evidence,previous)
    directory=Path(directory); directory.mkdir(parents=True,exist_ok=True)
    book.save(directory/'projected.xlsx'); book.close()
    for name,content in docs.items():
        atomic_bytes(directory/name,content.encode('utf-8'),immutable=True)
    return projection


# Only equivalent formula spellings observed in the selected engine are compared.
# Quoted strings are deliberately excluded (TRUE() may itself be business text).
def comparable_formula(value):
    parts = formula_text(value).split('"')
    for i in range(0,len(parts),2):
        parts[i] = re.sub(r'\b(TRUE|FALSE)\(\)',r'\1',parts[i])
    return '"'.join(parts)


def _xml_package(path):
    import posixpath
    import zipfile
    from xml.etree import ElementTree as ET
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(path) as z:
        if len(z.namelist()) != len(set(z.namelist())) or z.testzip():
            raise StorageError('WORKBOOK_INVALID','XLSX ZIP 成员重复或损坏。')
        parts = {n:z.read(n) for n in z.namelist()}
    for name, raw in parts.items():
        if name.endswith(('.xml','.rels')):
            ET.fromstring(raw)
        if name.startswith('xl/externalLinks/'):
            raise StorageError('WORKBOOK_INVALID','工作簿出现外部引用。')
    book = ET.fromstring(parts['xl/workbook.xml'])
    rels = {r.get('Id'):r.get('Target') for r in ET.fromstring(parts['xl/_rels/workbook.xml.rels'])}
    sheets = {}
    for s in book.find(ns+'sheets'):
        target=rels[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
        member=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        sheets[s.get('name')]=(member,ET.fromstring(parts[member]))
    return parts,sheets


def formula_cache_inventory(path):
    """Actual raw OOXML formula/cache records, including stored empty strings/errors."""
    from .contracts import canonical_json_bytes
    _,sheets=_xml_package(path)
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    records=[]
    for sheet,(_,root) in sheets.items():
        for c in root.iter(ns+'c'):
            f=c.find(ns+'f')
            if f is None: continue
            v=c.find(ns+'v')
            if v is None or (v.text is None and c.get('t','n') not in ('str','inlineStr')):
                raise StorageError('WORKBOOK_INVALID',f'公式缓存未保存：{sheet}!{c.get("r")}。')
            records.append([sheet,c.get('r'),dict(f.attrib),f.text,c.get('t','n'),v.text])
    return dict(formula_count=len(records),cache_count=len(records),
                formula_cache_hash=hashlib.sha256(canonical_json_bytes(records)).hexdigest())


def _color(color,book):
    if color is None or color.type=='auto': return None
    # The pinned template has no tinted colors. Do not discard an unexpected
    # tint while comparing an Office style to its original effective color.
    if color.tint: return (color.type,str(color.value),color.tint)
    if color.type=='rgb': return color.rgb[-6:].upper()
    if color.type=='indexed':
        from openpyxl.styles.colors import COLOR_INDEXED
        return COLOR_INDEXED[color.indexed][-6:].upper() if color.indexed<len(COLOR_INDEXED) else None
    if color.type=='theme' and book.loaded_theme:
        from xml.etree import ElementTree as ET
        ns='{http://schemas.openxmlformats.org/drawingml/2006/main}'
        scheme=ET.fromstring(book.loaded_theme).find('.//'+ns+'clrScheme')
        index=[1,0,3,2,4,5,6,7,8,9,10,11][color.theme]
        node=list(scheme[index])[0]
        return node.get('lastClr',node.get('val')).upper()
    return (color.type,color.value)


def _theme_typefaces(book,scheme):
    from xml.etree import ElementTree as ET
    ns='{http://schemas.openxmlformats.org/drawingml/2006/main}'
    root=ET.fromstring(book.loaded_theme)
    node=root.find('.//'+ns+scheme+'Font')
    if node is None: _fail('字体主题无法解析。')
    return tuple(sorted((c.tag,c.get('script',''),c.get('typeface','')) for c in node))


def _font_display(font,book):
    """Compare declared and effective theme typefaces, including CJK script fonts."""
    typefaces=_theme_typefaces(book,font.scheme) if font.scheme else ()
    return (font.name,font.scheme,typefaces,font.charset if font.charset is not None else 1,
            font.family or 0,font.vertAlign or 'baseline',bool(font.outline),bool(font.shadow),
            bool(font.condense),bool(font.extend))


def _style(cell,book):
    font,fill,border,align=cell.font,cell.fill,cell.border,cell.alignment
    return (_font_display(font,book),bool(font.bold),bool(font.italic),font.sz,font.u,bool(font.strike),_color(font.color,book),
            fill.patternType,_color(fill.fgColor,book) if fill.patternType else None,
            tuple((side.style,_color(side.color,book) if side.style else None) if side is not None else (None,None)
                  for side in [border.left,border.right,border.top,border.bottom]),
            align.horizontal or 'general',align.vertical or 'bottom',bool(align.wrapText),bool(align.shrinkToFit),
            align.textRotation or 0,align.indent or 0,cell.number_format,bool(cell.protection.locked),bool(cell.protection.hidden))


def _effective_cell(sheet,coordinate,root):
    """Resolve omitted blank cells through OOXML row/column style inheritance."""
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    cell=sheet[coordinate]
    node=root.find(f'.//{ns}c[@r="{coordinate}"]')
    if node is not None and node.get('s') is not None:
        return cell
    # Existing explicit style 0 is authoritative too. Row customFormat precedes column.
    row=root.find(f'{ns}sheetData/{ns}row[@r="{cell.row}"]')
    style=None
    if row is not None and row.get('customFormat') in ('1','true'):
        style=row.get('s')
    if style is None:
        cols=root.find(ns+'cols')
        if cols is not None:
            for col in cols:
                if int(col.get('min'))<=cell.column<=int(col.get('max')):
                    style=col.get('style'); break
    if style is not None:
        cell=copy(cell); cell._style=sheet.parent._cell_styles[int(style)]
    return cell


def _dv_rule(dv):
    value=dict(dv)
    value.pop('sqref',None)
    # operator is inapplicable for custom/list validation; LO spells the default.
    if dv.type in ('custom','list') and value.get('operator')=='between': value.pop('operator')
    second = None if dv.type in ('custom','list') and dv.formula2 in (None,'0') else dv.formula2
    return value,dv.formula1,second


def _fail(message):
    raise StorageError('WORKBOOK_INVALID',message)


def audit_workbook(path, expected, *, allow_omissions=False, caches=True):
    """Dual read of real bytes against a validated projection; never save either book.

    allow_omissions identifies ONLY the two authorized LibreOffice losses. All
    other structure, cells, formulas, protection and cache records must pass first.
    """
    from xml.etree import ElementTree as ET
    from .contracts import canonical_json_bytes
    import zipfile
    owned=not hasattr(expected,'sheetnames')
    source=load_workbook(expected) if owned else expected
    actual=cached=None
    try:
        _,xml_sheets=_xml_package(path)
        actual=load_workbook(path,data_only=False)
        cached=load_workbook(path,data_only=True) if caches else None
        if tuple(actual.sheetnames)!=SHEETS or actual.sheetnames!=source.sheetnames: _fail('Sheet 集合或顺序变化。')
        changes=[]; formulas=[]
        for original in source:
            ws=actual[original.title]; raw=xml_sheets[original.title][1]
            if set(ws.tables)!=set(original.tables): _fail(f'Table 集合变化：{ws.title}。')
            if dict(ws.protection)!=dict(original.protection) or ws.protection.password!=original.protection.password:
                _fail(f'工作表保护变化：{ws.title}。')
            if (ws.freeze_panes!=original.freeze_panes or str(ws.print_area)!=str(original.print_area)
                or ws.print_title_rows!=original.print_title_rows or ws.print_title_cols!=original.print_title_cols
                or str(ws.merged_cells)!=str(original.merged_cells) or ws.sheet_state!=original.sheet_state):
                _fail(f'冻结、打印、合并或可见性变化：{ws.title}。')
            for attr in ('orientation','paperSize','fitToHeight','fitToWidth'):
                if getattr(original.page_setup,attr) is not None and getattr(ws.page_setup,attr)!=getattr(original.page_setup,attr): _fail(f'打印设置变化：{ws.title} / {attr}。')
            for name in original.tables:
                table=original.tables[name]
                other=ws.tables[name]
                if (other.ref!=table.ref or other.name!=table.name or other.displayName!=table.displayName
                    or [c.name for c in other.tableColumns]!=[c.name for c in table.tableColumns]
                    or other.tableStyleInfo!=table.tableStyleInfo
                    or (other.autoFilter.ref.replace('$','') if other.autoFilter else None)!=(table.autoFilter.ref.replace('$','') if table.autoFilter else None)):
                    _fail(f'Table 身份、范围、表头、样式或筛选变化：{name}。')
                for column,(before,after) in enumerate(zip(table.tableColumns,other.tableColumns),1):
                    a,b=before.calculatedColumnFormula,after.calculatedColumnFormula
                    if a is not None and b is None and allow_omissions and name in ('SOWStoryTable','TaskTable'):
                        changes.append(dict(kind='table_formula',table=name,column=column))
                    elif (a is None)!=(b is None) or (a is not None and (a.array!=b.array or comparable_formula(a.text)!=comparable_formula(b.text))):
                        _fail(f'Table 计算列原型变化：{name} / {column}。')
            first=list(original.data_validations.dataValidation); second=list(ws.data_validations.dataValidation)
            if len(first)!=len(second): _fail(f'数据验证规则丢失：{ws.title}。')
            # Order has no meaning; match by actual rule and unique starting column.
            for dv in first:
                matches=[v for v in second if _dv_rule(v)==_dv_rule(dv)]
                if len(matches)!=1: _fail(f'数据验证规则变化：{ws.title}。')
                other=matches[0]
                if str(other.sqref)!=str(dv.sqref):
                    clipped=re.sub(r'1048576\b',str(original.max_row+1000),str(dv.sqref))
                    if not allow_omissions or clipped==str(dv.sqref) or str(other.sqref)!=clipped:
                        _fail(f'数据验证范围变化：{ws.title}。')
                    changes.append(dict(kind='validation_range',sheet=ws.title,before=str(other.sqref),after=str(dv.sqref)))
            def conditional(s):
                rules=[]
                for key,values in s.conditional_formatting._cf_rules.items():
                    for rule in sorted(values,key=lambda r:r.priority):
                        dxf=copy(rule.dxf)
                        color=None
                        if dxf and dxf.fill:
                            fill=dxf.fill
                            # Differential solid fills use fgColor or the equivalent
                            # background-only representation emitted by LibreOffice.
                            if fill.patternType=='solid': color=_color(fill.fgColor,s.parent)
                            elif fill.patternType is None and fill.bgColor.type=='rgb': color=_color(fill.bgColor,s.parent)
                            else: _fail('不支持的条件填充表示。')
                            dxf.fill=None
                        rules.append((str(key.sqref),rule.type,rule.formula,bool(rule.stopIfTrue),color,
                                      ET.tostring(dxf.to_tree()) if dxf else None))
                return rules
            if conditional(ws)!=conditional(original): _fail(f'条件格式变化：{ws.title}。')
            if ws.max_row!=original.max_row or ws.max_column!=original.max_column: _fail(f'工作表范围变化：{ws.title}。')
            for row in original:
                for cell in row:
                    if cell.__class__.__name__=='MergedCell': continue
                    other=ws[cell.coordinate]
                    def link(value):
                        return (value.target,value.location,value.tooltip) if value else None
                    if link(cell.hyperlink)!=link(other.hyperlink):
                        _fail(f'工作簿内跳转变化：{ws.title}!{cell.coordinate}。')
                    if other.hyperlink:
                        match = re.fullmatch(r"'([^']+)'!([A-Z]+[1-9][0-9]*)", other.hyperlink.location or '')
                        if other.hyperlink.target is not None or not match or match[1] not in actual.sheetnames:
                            _fail(f'工作簿内跳转目标无效：{ws.title}!{cell.coordinate}。')
                        target = actual[match[1]]
                        from openpyxl.utils.cell import coordinate_to_tuple
                        target_row, target_column = coordinate_to_tuple(match[2])
                        if target_row > target.max_row or target_column > target.max_column:
                            _fail(f'工作簿内跳转目标不存在：{ws.title}!{cell.coordinate}。')
                        value = target[match[2]].value
                        # Literal blank lines/spaces can be the first continuation
                        # of complete source text; only an unwritten cell is absent.
                        if value is None:
                            _fail(f'工作簿内跳转目标为空：{ws.title}!{cell.coordinate}。')
                    if cell.data_type=='f':
                        if other.data_type!='f' or comparable_formula(other.value)!=comparable_formula(cell.value):
                            _fail(f'公式变化：{ws.title}!{cell.coordinate}。')
                        if isinstance(cell.value,ArrayFormula)!=isinstance(other.value,ArrayFormula) or (isinstance(cell.value,ArrayFormula) and cell.value.ref!=other.value.ref):
                            _fail(f'数组公式类型或范围变化：{ws.title}!{cell.coordinate}。')
                        formulas.append([ws.title,cell.coordinate,comparable_formula(cell.value)])
                    elif cell.value!=other.value or (isinstance(cell.value,str) and other.data_type!='s'):
                        _fail(f'输入或原模板值变化：{ws.title}!{cell.coordinate}。')
                    if cached is not None and cell.data_type!='f' and cached[ws.title][cell.coordinate].value!=cell.value:
                        _fail(f'缓存视图输入不同：{ws.title}!{cell.coordinate}。')
                    effective=_effective_cell(ws,cell.coordinate,raw)
                    if cell.value is not None or cell.has_style:
                        before_style,after_style=_style(cell,source),_style(effective,actual)
                        # Controller's observed fallback exception is confined to
                        # the three fixed instructional cells. It is NOT theme
                        # equivalence and changes no workbook data or font metadata.
                        if (ws.title in SHEETS[:3] and cell.coordinate=='A2'
                            and (cell.font.name,cell.font.scheme)==('Calibri','minor')
                            and (effective.font.name,effective.font.scheme)==('Arial Unicode MS',None)
                            and _theme_typefaces(source,'minor')==_theme_typefaces(actual,'minor')):
                            after_style=(before_style[0][:3]+after_style[0][3:],*after_style[1:])
                        if after_style!=before_style:
                            _fail(f'单元格样式或保护变化：{ws.title}!{cell.coordinate}。')
                before_row=original.row_dimensions[row[0].row]
                after_row=ws.row_dimensions[row[0].row]
                if (bool(before_row.hidden),bool(before_row.collapsed),before_row.outlineLevel)!=(bool(after_row.hidden),bool(after_row.collapsed),after_row.outlineLevel):
                    _fail(f'行可见性或分组变化：{ws.title} / {row[0].row}。')
                height=before_row.height
                actual_height=after_row.height
                if height is not None and (actual_height not in (height, int(height / .75) * .75)):
                    _fail(f'行高变化：{ws.title} / {row[0].row}。')
            for key,dim in original.column_dimensions.items():
                col=dim.min or __import__('openpyxl').utils.column_index_from_string(key)
                matches=[x for x in ws.column_dimensions.values() if (x.min or col)<=col<=(x.max or col)]
                if len(matches)!=1 or abs(matches[0].width-dim.width)>.1 or bool(matches[0].hidden)!=bool(dim.hidden):
                    _fail(f'列宽或可见性变化：{ws.title} / {key}。')
        inventory=formula_cache_inventory(path) if caches else dict(formula_count=len(formulas))
        return dict(inventory,formula_hash=hashlib.sha256(canonical_json_bytes(formulas)).hexdigest(),compatibility_changes=changes)
    except (KeyError,IndexError,AttributeError,ValueError,TypeError,ET.ParseError,zipfile.BadZipFile) as error:
        if isinstance(error,StorageError): raise
        raise StorageError('WORKBOOK_INVALID','工作簿内容或 OOXML 无法完整复读。') from error
    finally:
        if owned: source.close()
        if actual is not None: actual.close()
        if cached is not None: cached.close()


def seal_office_output(source, raw_path, destination):
    """One narrow, pre-seal LO adapter. Raw input remains intact; cells never change."""
    from xml.etree import ElementTree as ET
    import zipfile
    if Path(destination).exists(): _fail('最终工作簿路径已存在，未覆盖。')
    report=audit_workbook(raw_path,source,allow_omissions=True)
    before=formula_cache_inventory(raw_path)
    raw_parts,sheets=_xml_package(raw_path)
    source_parts,_=_xml_package(source)
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    changes=report['compatibility_changes']
    replacements={}
    for member,payload in raw_parts.items():
        if member.startswith('xl/tables/') and member.endswith('.xml'):
            root=ET.fromstring(payload); name=root.get('displayName') or root.get('name')
            additions=[x for x in changes if x['kind']=='table_formula' and x['table']==name]
            if additions:
                originals=[ET.fromstring(value) for key,value in source_parts.items() if key.startswith('xl/tables/') and key.endswith('.xml')]
                original=next(x for x in originals if (x.get('displayName') or x.get('name'))==name)
                for change in additions:
                    index=change['column']-1
                    root.find(ns+'tableColumns')[index].append(copy(original.find(ns+'tableColumns')[index].find(ns+'calculatedColumnFormula')))
                replacements[member]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    for sheet,(member,root) in sheets.items():
        fixes=[x for x in changes if x['kind']=='validation_range' and x['sheet']==sheet]
        if fixes:
            for fix in fixes:
                nodes=[x for x in root.iter(ns+'dataValidation') if x.get('sqref')==fix['before']]
                if len(nodes)!=1: _fail('已核验的数据验证位置不唯一。')
                nodes[0].set('sqref',fix['after'])
            replacements[member]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    # ElementTree only reserializes the two permitted metadata-bearing parts;
    # prove sheetData (all inputs/formulas/caches/styles) is exactly equivalent.
    for sheet,(member,root) in sheets.items():
        if member in replacements:
            revised=ET.fromstring(replacements[member])
            if ET.tostring(root.find(ns+'sheetData'))!=ET.tostring(revised.find(ns+'sheetData')):
                _fail('兼容转换触及单元格，拒绝封存。')
    temporary=Path(destination).with_suffix('.sealing.xlsx')
    try:
        with zipfile.ZipFile(raw_path) as src,zipfile.ZipFile(temporary,'w') as dst:
            for info in src.infolist(): dst.writestr(info,replacements.get(info.filename,raw_parts[info.filename]))
        if formula_cache_inventory(temporary)!=before: _fail('兼容转换改变公式或缓存。')
        final=audit_workbook(temporary,source)
        atomic_bytes(Path(destination),temporary.read_bytes(),immutable=True)
        # Final path is only ever read after the seal.
        if formula_cache_inventory(destination)!=before: _fail('最终缓存复读不同。')
        return dict(final,compatibility_changes=changes)
    finally:
        temporary.unlink(missing_ok=True)


def _evidence_labels(project,candidate):
    from .project import checked_json
    index=checked_json(project,'.ai-sow-lite/inputs/index.json','input_index')
    sources={x['input_version_id']:x for x in index['items']}
    records={}
    for version in candidate['topic_version_ids']:
        analysis=checked_json(project,f'.ai-sow-lite/analysis/topics/{version}/analysis.json','analysis')
        records.update({e['id']:e for e in analysis['evidence']})
    labels={}
    for identity in candidate['evidence_ids']:
        visited=set(); stack=[identity]; lines=[]
        while stack:
            current=stack.pop()
            if current in visited or current not in records: continue
            visited.add(current); record=records[current]
            for ref in record['source_refs']:
                loc=ref['locator']
                source=sources[ref['input_version_id']]
                name=Path(source['relative_path']).name
                if loc['kind']=='text_lines':
                    if source['format']=='prototype': name=loc['path']
                    position=f'第 {loc["start_line"]}—{loc["end_line"]} 行'
                elif loc['kind']=='xlsx_range':
                    position=f'{loc["sheet"]}!{loc["range"]}'
                elif loc['kind']=='observation':
                    name='原型'
                    position='；'.join(str(value) for value in (
                        '观察 '+loc['observation_id'],loc.get('attachment'),loc.get('region')) if value)
                else:
                    raise StorageError('OPERATION_UNSUPPORTED','来源定位尚不支持工作簿显示。')
                lines.append(f'{name}；{position}；依据 {current}')
            if record['limitations']: lines.append('限制：'+record['limitations'])
            stack.extend(reversed(record['basis_refs']))
        labels[identity]='\n'.join(lines)+'\n'
    return labels


def _summary(model,pending,version,plan=None):
    lines=[f'# 交付摘要\n\nversion_id: {version}\n',
           f'本版范围：{len(model["epics"])} Epic、{len(model["features"])} Feature、{len(model["stories"])} Story、{len(model["tasks"])} Task。\n',
           ('本次生成修改版候选；' if plan is not None else '本次生成首版候选；')+'生效事实以 current 指针为准。\n',
           f'待确认：{sum(p["status"]=="open" for p in pending["items"])} 项，见 [待确认事项](pending-items.md)。\n',
           '工作簿：[sow.xlsx](sow.xlsx)。公式由 Office 计算保存；输入、结构、公式、缓存和正文已机械复读。\n']
    if plan is not None:
        import json
        lines.append(f'本次变化：方案 {plan["plan_id"]}，修订 {plan["revision"]}。\n')
        # A subset inherits the full plan's prose; only its selected concrete
        # changes describe this package. Never present deferred work as applied.
        if 'subset_of' not in plan:
            lines.append(_block(plan['change_summary']))
        for change in plan['changes']:
            lines.append(f'{change["op"]} / {change["collection"]} / {change["object_id"]} / {change["field"]}：\n'
                         +_block(json.dumps(change['before'],ensure_ascii=False)+' → '
                                 +json.dumps(change['after'],ensure_ascii=False)))
    for item in pending['items']:
        if item['status']=='open':
            label='未拆明业务工作' if item['unestimated_work'] else '当前处理'
            lines.append(f'{label}：{item["id"]}；见 pending-items.md#pending-{item["id"]}\n'+_block(item['current_handling']))
    return '\n'.join(lines)


def _bound_plan(project,checked):
    from .contracts import strict_json_loads
    from .project import safe_path
    ref=checked.get('plan_ref')
    if ref is None:
        return None
    raw=safe_path(project,ref['path']).read_bytes()
    if hashlib.sha256(raw).hexdigest()!=ref['sha256']:
        _fail('修改摘要的方案与完整检查绑定不同。')
    return strict_json_loads(raw)


def _legacy_prepared(project,prepared):
    """Only an existing successful pre-v4 attempt can use the old summary text."""
    from .contracts import load_json
    from .project import safe_path,file_ref
    area='/'.join(Path(prepared['candidate_ref']['path']).parts[:4])
    path=safe_path(project,str(Path(prepared['candidate_ref']['path']).with_name('render-attempt.json')),area)
    if not path.exists():
        path=safe_path(project,area+'/render-attempt.json',area)
    if not path.exists():
        return False
    attempt=load_json(path)
    ref=attempt.get('prepared_ref')
    if attempt.get('implementation_version') not in (None,'lite-render-v2','lite-render-v3') or not ref:
        return False
    saved=safe_path(project,ref['path'],area)
    return file_ref(project,saved)==ref and load_json(saved)==prepared


def _baseline_aliases(project,candidate,checked):
    """Read original names and display names from the same checked baseline bytes."""
    from .contracts import strict_json_loads
    from .project import safe_path
    if candidate['entrypoint']=='generate':
        return None
    refs={ref['path']:ref for ref in checked['dependencies']}
    area=f'.ai-sow-lite/versions/{candidate["base_version_id"]}'
    values=[]
    for name in ('model.json','projection.json'):
        ref=refs[area+'/'+name]
        raw=safe_path(project,ref['path']).read_bytes()
        if hashlib.sha256(raw).hexdigest()!=ref['sha256']:
            _fail('基线显示名依赖与完整检查绑定不同。')
        values.append(strict_json_loads(raw))
    model,projection=values
    if (projection['version_id']!=candidate['base_version_id']
        or projection['model_hash']!=refs[area+'/model.json']['sha256']):
        _fail('基线原名与显示名不是同一版模型。')
    previous={}
    for collection,kind,field in [('stories','story','title'),('tasks','task','name')]:
        names={obj['object_id']:obj['display_name'] for obj in projection['objects'] if obj['kind']==kind}
        if set(names)!={obj['id'] for obj in model[collection]}:
            _fail('基线显示名与对象集合不同。')
        previous[collection]={obj['id']:dict(original=obj[field],display_name=names[obj['id']])
                              for obj in model[collection]}
    return previous


def _expected_book(template,model,pending,decisions,version,evidence,previous=None):
    book=_template(template)
    book[STORY_SHEET].column_dimensions['D'].width=AC_COLUMN_WIDTH
    for row in book[STORY_SHEET].iter_rows(min_row=5,min_col=4,max_col=4):
        alignment=copy(row[0].alignment)
        alignment.vertical='top'
        alignment.wrap_text=True
        row[0].alignment=alignment
    projection,inputs,docs=_projection(model,pending,decisions,version,evidence,previous,layout=book)
    _extend(book,STORY_SHEET,'SOWStoryTable',max((r - 4 for obj in projection['objects'] if obj['sheet']==STORY_SHEET for r in obj['rows']),default=0))
    _extend(book,TASK_SHEET,'TaskTable',len(model['tasks']))
    _fill_inputs(book,inputs)
    _fit_task_lists(book,model,projection)
    projection['template_hash']=hashlib.sha256(Path(template).read_bytes()).hexdigest()
    return book,projection,docs


def _fit_task_lists(book, model, projection):
    # Names and qualitative labels reserve wrap space for the template's H
    # formula. This is layout only; the task table contains every actual Task.
    names = {o['object_id']:o['display_name'] for o in projection['objects'] if o['kind']=='task'}
    rows = {o['object_id']:o['rows'][0] for o in projection['objects'] if o['kind']=='story'}
    for story in model['stories']:
        row = rows[story['id']]
        lines = [' '.join([names[t['id']], *(t[f] or '' for f in ('work_type_name','work_mode','complexity','integration_type'))])
                 for t in model['tasks'] if t['story_id']==story['id']]
        _fit_height(book[STORY_SHEET][f'H{row}'], '\n'.join(lines), story['id'], None)


def _column_width(sheet,coordinate):
    from openpyxl.utils.cell import coordinate_from_string
    column,_=coordinate_from_string(coordinate)
    return sheet.column_dimensions[column].width or 8.43


def _required_height(value,width,*,acceptance=False):
    """Estimate visible lines; retain the proven margin for notes/task lists.

    AC uses the wider original column and counts a spare line where Excel wraps
    a long numbered paragraph at its space. This is display layout only.
    """
    import math
    if not value: return 0
    usable=max(1,width-2) if acceptance else max(1,(width-2)*.85)
    lines=0
    for paragraph in value.split('\n'):
        units=sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in paragraph)
        wrapped=max(1,math.ceil(units/usable))
        spare=int(units>usable and bool(re.search(r'\s',paragraph))) if acceptance else 1
        lines+=wrapped+spare
    if acceptance:
        content_height=lines*16
        # Optional bottom padding must not reject text whose full wrap reserve
        # fits the row. Content beyond the limit still reaches _fit_height.
        return min(content_height+12,409) if content_height<=409 else content_height+12
    return lines*18+28


def _fit_height(cell,value,identity,field):
    height=_required_height(value,_column_width(cell.parent,cell.coordinate),
                            acceptance=cell.parent.title==STORY_SHEET and cell.column==4)
    if height>409:
        error=StorageError('WORKBOOK_LAYOUT_OVERFLOW',
            f'{cell.parent.title}!{cell.coordinate} 的完整正文预计超过 409 点可见高度；保留候选，不能删减验收义务来凑布局。')
        error.diagnostics[0]['target'].update(object_id=identity,field=field)
        raise error
    dimension=cell.parent.row_dimensions[cell.row]
    dimension.height=max(dimension.height or 15,height)


def _fill_inputs(book,inputs):
    for (sheet,coordinate),(value,identity,field) in inputs.items():
        cell=book[sheet][coordinate]
        write_literal(cell,value)
        if cell.alignment.wrap_text:
            _fit_height(cell,value,identity,field)


def verify_prepared(project,prepared):
    """Rebuild expected mappings in memory; reread final bytes without saving/Office.

    Receipt is a local execution record, not a signed attestation. Success flags
    never substitute for actual mapping, formula, metadata and cache inspection.
    """
    from .contracts import canonical_json_bytes,file_sha256,load_json,schema_validator
    from .project import safe_path,checked_json,_verify_refs
    from .validation import check_candidate
    from .changes import checks_match
    expected=None
    try:
        project=Path(project).resolve()
        if list(schema_validator('artifacts','prepared').iter_errors(prepared)): _fail('准备记录结构无效。')
        _verify_refs(project,[prepared['candidate_ref'],prepared['check_ref'],prepared['verification_ref'],*prepared['files']])
        candidate=checked_json(project,prepared['candidate_ref']['path'],'candidate')
        checked=checked_json(project,prepared['check_ref']['path'],'check')
        current_check=check_candidate(project,safe_path(project,prepared['candidate_ref']['path']),'full',
                                      (checked.get('plan_ref') or {}).get('path'))
        if not current_check['valid_for_render'] or not checks_match(project,checked,current_check): _fail('候选或依赖在准备后变化。')
        files={Path(r['path']).name:r for r in prepared['files']}
        names={'model.json','pending-items.json','decisions.json','projection.json','sow.xlsx','summary.md','pending-items.md'}
        if len(files)!=len(prepared['files']) or set(files) not in (names,names|{'details.md'}): _fail('交付文件集合不完整或有意外文件。')
        bundles=[]
        for name,key in [('model.json','model_path'),('pending-items.json','pending_items_path'),('decisions.json','decisions_path')]:
            if file_sha256(safe_path(project,candidate[key]))!=files[name]['sha256']: _fail('交付 JSON 不是候选原字节。')
            bundles.append(load_json(safe_path(project,files[name]['path'])))
        model,pending,decisions=bundles
        if prepared['projection_version']!=PROJECTOR_VERSION or prepared['template_hash'] not in SUPPORTED_TEMPLATE_HASHES: _fail('准备记录模板或投影器不匹配。')
        template=safe_path(project,f".ai-sow-lite/template/{prepared['template_hash']}/sow-template.xlsx")
        expected,projection,docs=_expected_book(template,model,pending,decisions,prepared['version_id'],
            _evidence_labels(project,candidate),_baseline_aliases(project,candidate,checked))
        projection.update(model_hash=files['model.json']['sha256'],pending_items_hash=files['pending-items.json']['sha256'],
                          decisions_hash=files['decisions.json']['sha256'],workbook_hash=files['sow.xlsx']['sha256'])
        actual_projection=checked_json(project,files['projection.json']['path'],'projection')
        if projection!=actual_projection: _fail('投影身份、字段位置或摘要与实际候选不同。')
        docs['summary.md']=_summary(model,pending,prepared['version_id'],_bound_plan(project,checked))
        if ('details.md' in docs)!=('details.md' in files): _fail('完整正文文件集合不同。')
        for name,text in docs.items():
            actual=safe_path(project,files[name]['path']).read_bytes()
            if actual!=text.encode('utf-8'):
                if (name=='summary.md' and _legacy_prepared(project,prepared)
                    and actual==_summary(model,pending,prepared['version_id']).encode('utf-8')):
                    continue
                _fail(f'同版正文不完整：{name}。')
        verification=checked_json(project,prepared['verification_ref']['path'],'verification')
        receipt=verification['office']
        if (verification['version_id']!=prepared['version_id'] or verification['candidate_digest']!=checked['candidate_digest']
            or receipt['final_hash']!=files['sow.xlsx']['sha256'] or receipt['office_identity']!=prepared['office_identity']
            or receipt['exit_code']!=0 or receipt['office_identity']!='sha256:'+hashlib.sha256(canonical_json_bytes(receipt['engine'])).hexdigest()):
            _fail('Office 记录身份、输入或最终文件绑定不同。')
        # Work-only raw and pre-Office files are evidence for preparation, not
        # dependencies of immutable applied versions. Validate them here before apply.
        directory=safe_path(project,prepared['verification_ref']['path']).parent
        projected=directory/'projected.xlsx'; raw=directory/'sow.office-raw.xlsx'
        if file_sha256(projected)!=receipt['input_hash'] or file_sha256(raw)!=receipt['raw_hash']: _fail('Office 原始或投影输入字节变化。')
        audit_workbook(projected,expected,caches=False)
        raw_report=audit_workbook(raw,expected,allow_omissions=True)
        final_report=audit_workbook(safe_path(project,files['sow.xlsx']['path']),expected)
        if (dict(final_report,compatibility_changes=raw_report['compatibility_changes'])!=receipt['verification']
            or formula_cache_inventory(raw)!=formula_cache_inventory(safe_path(project,files['sow.xlsx']['path']))):
            _fail('实际公式、缓存或兼容变换与 Office 记录不同。')
        return dict(diagnostics=[])
    except (StorageError,OSError,ValueError,KeyError) as error:
        from .validation import diagnostic
        # Corrupt internal prepared data consistently fails at the workbook boundary.
        message=error.diagnostics[0]['message'] if isinstance(error,StorageError) else '准备包文件无法完整复读。'
        return dict(diagnostics=[diagnostic('WORKBOOK_INVALID',field='verification_ref',message=message)])
    finally:
        if expected is not None: expected.close()


def _render_result(project,prepared,path,pending_count):
    from .project import file_ref
    files={Path(r['path']).name:r for r in prepared['files']}
    return dict(prepared_ref=file_ref(project,path),version_id=prepared['version_id'],workbook_ref=files['sow.xlsx'],
                projection_ref=files['projection.json'],summary_ref=files['summary.md'],pending_items_ref=files['pending-items.md'],
                details_ref=files.get('details.md'),office_identity=prepared['office_identity'],
                verification_ref=prepared['verification_ref'],pending_count=pending_count)


def render_candidate(project: Path,request_id: str,payload):
    from uuid import uuid4
    from . import office
    from .contracts import canonical_json_bytes,file_sha256,load_json,schema_validator,semantic_digest
    from .project import (safe_path,checked_json,file_ref,write_json,request_area,ensure_request,
                          read_current,_check_cancelled,save_checkpoint)
    from .validation import check_candidate
    from .changes import checks_match
    project=Path(project).resolve()
    if list(schema_validator('protocol','render_payload').iter_errors(payload)):
        raise StorageError('PROTOCOL_INVALID','导出字段不符合合同。')
    candidate=checked_json(project,payload['candidate_path'],'candidate','.ai-sow-lite/work')
    entrypoint=candidate['entrypoint']
    area=request_area(request_id,entrypoint)
    candidate_path=safe_path(project,payload['candidate_path'],area)
    check_path=safe_path(project,payload['check_path'],area)
    checked=checked_json(project,payload['check_path'],'check',area)
    plan_path=(checked.get('plan_ref') or {}).get('path')
    report=check_candidate(project,candidate_path,'full',plan_path)
    if not report['valid_for_render'] or not checks_match(project,checked,report):
        raise StorageError('CANDIDATE_INVALID','候选或依赖与实际完整检查记录不一致。')
    expected=payload['expected_current']
    if (read_current(project)[0]!=expected or candidate['base_version_id']!=(expected['version_id'] if expected else None)
        or (entrypoint=='generate')!=(expected is None)):
        raise StorageError('BASE_STALE','当前完整指针、候选基线和调用期望不同。')
    if plan_path and not load_json(safe_path(project,plan_path,area))['changes']:
        raise StorageError('CANDIDATE_INVALID','没有实际变化；使用当前交付文件，无需再次导出。')
    checkpoint=ensure_request(project,request_id,entrypoint); _check_cancelled(project,request_id,entrypoint)
    # Clarify construction owns a finite immutable r1/r2/repair/subset slot.
    # Keep each successful preview there; changing slots is not a failed retry.
    # Generate retains its original request-wide receipt and retry behavior.
    attempt_path=(candidate_path.with_name('render-attempt.json') if entrypoint=='clarify'
                  else safe_path(project,area+'/render-attempt.json'))
    signature_inputs=dict(check=checked,payload=payload,projector_version=PROJECTOR_VERSION,
                          engine=office.selection_fingerprint())
    legacy_signature=semantic_digest(signature_inputs)
    v2_signature=semantic_digest(dict(signature_inputs,implementation_version='lite-render-v2'))
    v3_signature=semantic_digest(dict(signature_inputs,implementation_version='lite-render-v3'))
    signature=semantic_digest(dict(signature_inputs,implementation_version=RENDER_IMPLEMENTATION_VERSION))
    previous=load_json(attempt_path) if attempt_path.exists() else None
    if entrypoint=='clarify' and previous is None:
        legacy_path=safe_path(project,area+'/render-attempt.json')
        if legacy_path.exists():
            legacy=load_json(legacy_path)
            # A failed legacy receipt has no candidate identity. A changed engine
            # signature cannot prove an independent slot or erase its retry history.
            applicable=True
            if legacy.get('prepared_ref'):
                saved=checked_json(project,legacy['prepared_ref']['path'],'prepared',area)
                applicable=saved['candidate_ref']==file_ref(project,candidate_path)
            if applicable:
                previous=legacy
    reusable=False
    if entrypoint=='clarify' and previous and previous.get('prepared_ref'):
        old_path=safe_path(project,previous['prepared_ref']['path'],area)
        if file_ref(project,old_path)!=previous['prepared_ref']: _fail('已准备记录字节变化。')
        old_prepared=checked_json(project,previous['prepared_ref']['path'],'prepared',area)
        old_check=checked_json(project,old_prepared['check_ref']['path'],'check',area)
        # v4/v5 had appendices; v6 omitted reasons; v7 could clip text. Preserve the
        # old package and use the existing bounded retry for the current output.
        reusable=(previous.get('implementation_version') not in ('lite-render-v4','lite-render-v5','lite-render-v6','lite-render-v7')
                  and old_prepared['candidate_ref']==file_ref(project,candidate_path)
                  and old_prepared['expected_current']==expected
                  and old_prepared['template_hash']==candidate['template_hash']
                  and old_check.get('plan_digest')==report.get('plan_digest'))
    # Pre-fix successful text packages remain reusable only after full verification.
    if previous and (previous['signature'] in (signature,legacy_signature,v2_signature,v3_signature) or reusable) and previous.get('prepared_ref'):
        prepared_path=safe_path(project,previous['prepared_ref']['path'],area)
        if file_ref(project,prepared_path)!=previous['prepared_ref']: _fail('已准备记录字节变化。')
        prepared=checked_json(project,previous['prepared_ref']['path'],'prepared',area)
        verified=verify_prepared(project,prepared)
        if verified['diagnostics']:
            error=StorageError('WORKBOOK_INVALID','准备记录失效。');error.diagnostics=verified['diagnostics'];raise error
        return _render_result(project,prepared,prepared_path,report['unknowns_count'])
    if previous:
        if previous['signature']==signature:
            raise StorageError('LOOP_LIMIT_REACHED','同一导出输入与环境没有变化，不重复调用 Office。',preserved_paths=[area])
        if checkpoint['repair_batches']>=2 or checkpoint['operation_retries'].get('render',0)>=1:
            raise StorageError('LOOP_LIMIT_REACHED','导出返修次数已到上限。',preserved_paths=[area])
        checkpoint['repair_batches']+=1;checkpoint['operation_retries']['render']=1
        save_checkpoint(project,checkpoint)
    version=str(uuid4()); directory=safe_path(project,area+'/render-'+version);directory.mkdir()
    attempt=dict(signature=signature,prepared_ref=None,implementation_version=RENDER_IMPLEMENTATION_VERSION)
    write_json(project,attempt_path.relative_to(project).as_posix(),attempt)
    try:
        bundle=[]
        for name,key in [('model.json','model_path'),('pending-items.json','pending_items_path'),('decisions.json','decisions_path')]:
            source=safe_path(project,candidate[key],area)
            atomic_bytes(directory/name,source.read_bytes(),immutable=True)
            bundle.append(load_json(directory/name))
        model,pending,decisions=bundle
        template=safe_path(project,f".ai-sow-lite/template/{candidate['template_hash']}/sow-template.xlsx")
        projection=project_workbook(template,model,pending,decisions,version,directory,
            evidence=_evidence_labels(project,candidate),previous=_baseline_aliases(project,candidate,checked))
        audit_workbook(directory/'projected.xlsx',directory/'projected.xlsx',caches=False)
        receipt=office.recalculate(directory/'projected.xlsx',directory/'sow.xlsx')
        # A change during Office invalidates all caches; no second internal calculation.
        if not checks_match(project,report,check_candidate(project,candidate_path,'full',plan_path)):
            raise StorageError('CANDIDATE_INVALID','Office 处理期间候选或来源变化，准备记录失效。')
        _check_cancelled(project,request_id,entrypoint)
        if read_current(project)[0]!=payload['expected_current']: raise StorageError('BASE_STALE','导出期间当前版本变化。')
        projection.update(model_hash=file_sha256(directory/'model.json'),pending_items_hash=file_sha256(directory/'pending-items.json'),
                          decisions_hash=file_sha256(directory/'decisions.json'),workbook_hash=file_sha256(directory/'sow.xlsx'))
        atomic_bytes(directory/'projection.json',canonical_json_bytes(projection),immutable=True)
        atomic_bytes(directory/'summary.md',_summary(model,pending,version,_bound_plan(project,checked)).encode('utf-8'),immutable=True)
        verification=dict(schema_version='1.0',version_id=version,candidate_digest=report['candidate_digest'],office=receipt)
        atomic_bytes(directory/'verification.json',canonical_json_bytes(verification),immutable=True)
        names=['model.json','pending-items.json','decisions.json','projection.json','sow.xlsx','summary.md','pending-items.md']
        if (directory/'details.md').exists(): names.append('details.md')
        prepared=dict(schema_version='1.0',version_id=version,candidate_ref=file_ref(project,candidate_path),check_ref=file_ref(project,check_path),
                      expected_current=payload['expected_current'],files=[file_ref(project,directory/n) for n in names],template_hash=candidate['template_hash'],
                      projection_version=PROJECTOR_VERSION,office_identity=receipt['office_identity'],verification_ref=file_ref(project,directory/'verification.json'))
        verified=verify_prepared(project,prepared)
        if verified['diagnostics']:
            error=StorageError('WORKBOOK_INVALID','最终工作簿复读失败。');error.diagnostics=verified['diagnostics'];raise error
        path=directory/'prepared.json';atomic_bytes(path,canonical_json_bytes(prepared),immutable=True)
        attempt['prepared_ref']=file_ref(project,path);write_json(project,attempt_path.relative_to(project).as_posix(),attempt)
        return _render_result(project,prepared,path,report['unknowns_count'])
    except StorageError as error:
        for d in error.diagnostics:
            d['preserved_paths']=list(dict.fromkeys([*d['preserved_paths'],directory.relative_to(project).as_posix()]))
        write_json(project,(directory/'failure.json').relative_to(project).as_posix(),dict(diagnostics=error.diagnostics))
        raise
