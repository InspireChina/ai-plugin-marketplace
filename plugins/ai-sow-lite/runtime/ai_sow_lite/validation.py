"""Read-only D02 checks. Professional scope/classification stays with the Agent."""
from __future__ import annotations

import hashlib
from pathlib import Path, PurePosixPath
from uuid import UUID
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .contracts import schema_validator, semantic_digest, strict_json_loads

VALIDATOR_VERSION = "lite-check-v1"
CLASSIFICATION_FIELDS = ("work_type_name", "work_mode", "complexity", "integration_type")


def diagnostic(code, path=None, object_id=None, field=None, message="候选不符合合同。"):
    return dict(code=code, target=dict(path=path, object_id=object_id, field=field),
                message=message, preserved_paths=[])


def project_file(project: Path, value: str | Path, area: str) -> Path:
    """Resolve a persistent reference within its exact project-owned area."""
    supplied_root = project.absolute()
    project = project.resolve()
    if isinstance(value, Path) and value.is_absolute():
        # Normalize the explicitly supplied root alias only. Resolving the whole
        # candidate first would hide a child symlink crossing an ownership area.
        try:
            value = value.relative_to(supplied_root).as_posix()
        except ValueError:
            value = value.relative_to(project).as_posix()
    if not isinstance(value, str) or not value or "\\" in value:
        raise ValueError("invalid project reference")
    pure = PurePosixPath(value)
    if pure.is_absolute() or ".." in pure.parts or ":" in value or str(pure) != value:
        raise ValueError("invalid project reference")
    root = project / ".ai-sow-lite"
    root.resolve().relative_to(project)
    allowed = project / area
    # Both lexical ownership and resolved ownership matter (including symlinks).
    (project / value).relative_to(allowed)
    resolved = (project / value).resolve()
    resolved.relative_to(allowed.resolve())
    resolved.relative_to(project)
    if allowed.resolve() != allowed:
        raise ValueError("redirected project area")
    return resolved


class _Check:
    def __init__(self, project):
        self.project = project.resolve()
        self.diagnostics = []
        self.dependencies = {}

    def relative(self, path):
        try:
            return path.relative_to(self.project).as_posix()
        except ValueError:
            return None

    def add(self, field, message, *, path=None, object_id=None, code="CANDIDATE_INVALID"):
        self.diagnostics.append(diagnostic(code, self.relative(path) if isinstance(path, Path) else path,
                                           object_id, field, message))

    def bytes(self, value, area, field=None):
        try:
            path = project_file(self.project, value, area)
        except (ValueError, OSError, RuntimeError):
            self.add(field, "文件引用必须位于指定项目子目录，不能逃逸。")
            return None, None
        try:
            raw = path.read_bytes()
        except OSError:
            self.add(field, "引用文件不可读取；保留现有文件。", path=path, code="IO_FAILED")
            return path, None
        self.dependencies[self.relative(path)] = dict(path=self.relative(path), sha256=hashlib.sha256(raw).hexdigest())
        return path, raw

    def json(self, value, area, schema, definition=None, field=None):
        path, raw = self.bytes(value, area, field)
        if raw is None:
            return None
        try:
            data = strict_json_loads(raw)
        except (ValueError, UnicodeError):
            self.add(field, "文件必须为严格 UTF-8 JSON；拒绝重复 key、浮点数和非法字符。", path=path)
            return None
        errors = list(schema_validator(schema, definition).iter_errors(data))
        for error in errors:
            parts = list(error.absolute_path)
            target_field = next((part for part in reversed(parts) if isinstance(part, str)), field)
            if error.validator == "required" and isinstance(error.instance, dict):
                missing = [key for key in error.validator_value if key not in error.instance]
                target_field = missing[0] if missing else target_field
            object_id = None
            node = data
            for part in parts:
                if isinstance(node, dict) and isinstance(node.get("id"), str):
                    try:
                        if UUID(node["id"]).version == 4 and str(UUID(node["id"])) == node["id"]:
                            object_id = node["id"]
                    except ValueError:
                        pass
                node = node[part]
            code = "VERSION_INCOMPATIBLE" if target_field in {"schema_version", "projection_version"} else "CANDIDATE_INVALID"
            self.add(target_field, "字段缺失、类型、枚举或结构不符合当前 Schema。", path=path, object_id=object_id, code=code)
        return None if errors else data

    def report(self, scope, candidate_path=None, digest=None, unknowns=0):
        reference = self.dependencies.get(self.relative(candidate_path)) if candidate_path else None
        return dict(schema_version="1.0", validator_version=VALIDATOR_VERSION, scope=scope,
                    candidate_ref=reference, candidate_digest=digest,
                    dependencies=list(self.dependencies.values()), valid_for_render=scope == "full" and not self.diagnostics,
                    unknowns_count=unknowns, diagnostics=self.diagnostics)


def _standards(ctx, template_hash):
    relative = f".ai-sow-lite/template/{template_hash}/sow-template.xlsx"
    path, raw = ctx.bytes(relative, ".ai-sow-lite/template", "template_hash")
    if raw is None:
        return None
    if hashlib.sha256(raw).hexdigest() != template_hash:
        ctx.add("template_hash", "模板实际字节与候选绑定的身份不同。", path=path)
        return None
    try:
        from io import BytesIO
        workbook = load_workbook(BytesIO(raw), data_only=False)
        try:
            matches = [(sheet, sheet.tables['TaskStandardTable']) for sheet in workbook if 'TaskStandardTable' in sheet.tables]
            if len(matches) != 1:
                raise ValueError("catalog missing")
            sheet, table = matches[0]
            lo_col, lo_row, hi_col, hi_row = range_boundaries(table.ref)
            rows = list(sheet.iter_rows(min_row=lo_row, max_row=hi_row, min_col=lo_col, max_col=hi_col, values_only=True))
            headers = rows[0]
            required = ['工作类型 ID', '分类', '工作类型', '新建 PD', '调整 PD', '复用 PD']
            if not all(key in headers for key in required):
                raise ValueError("catalog incompatible")
            standards = {}
            for row in rows[1:]:
                record = dict(zip(headers, row))
                identity, name = record['工作类型 ID'], record['工作类型']
                if not isinstance(identity, str) or not isinstance(name, str) or name in standards:
                    raise ValueError("catalog identity invalid")
                # Only availability is read. No amount, multiplier or formula is copied/calculated.
                modes = [mode for mode, column in [('新建', '新建 PD'), ('调整', '调整 PD'), ('接入复用', '复用 PD')]
                         if isinstance(record[column], (int, float)) and not isinstance(record[column], bool)]
                standards[name] = dict(id=identity, modes=modes, integration=record['分类'] == '系统集成')
            return standards
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, TypeError, BadZipFile):
        ctx.add("template_hash", "模板标准目录与 Lite 合同不兼容。", path=path, code="VERSION_INCOMPATIBLE")
        return None


def _objects(ctx, model, model_path):
    objects, collections = {}, {}
    for collection in ('epics', 'features', 'stories', 'tasks', 'dependencies'):
        for obj in model[collection]:
            members = [(collection, obj)]
            if collection == 'stories':
                members += [('acs', ac) for ac in obj['acs']]
            for kind, member in members:
                identity = member['id']
                if identity in objects:
                    ctx.add('id', '对象 ID 在候选内必须唯一。', path=model_path, object_id=identity)
                else:
                    objects[identity], collections[identity] = member, kind
    for collection, parent_field, parent_kind in [('features', 'epic_id', 'epics'), ('stories', 'feature_id', 'features'), ('tasks', 'story_id', 'stories')]:
        for obj in model[collection]:
            if collections.get(obj[parent_field]) != parent_kind:
                ctx.add(parent_field, '父引用必须指向对应层级的当前对象。', path=model_path, object_id=obj['id'])
    edges = set()
    for dep in model['dependencies']:
        edge = (dep['from_story_id'], dep['to_story_id'], dep['kind'])
        if edge in edges:
            ctx.add('to_story_id', '同方向同种依赖不能重复。', path=model_path, object_id=dep['id'])
        edges.add(edge)
        for field in ('from_story_id', 'to_story_id'):
            if collections.get(dep[field]) != 'stories' or dep['from_story_id'] == dep['to_story_id']:
                ctx.add(field, '依赖端点须为两个不同的当前 Story。', path=model_path, object_id=dep['id'])
    return objects, collections


def _lineage_key(record):
    return record['from_version_id'], tuple(record['from_ids'])


def _history(ctx, model, base_version_id, current_objects):
    """Read the bounded ancestry needed by explicit lineage; never apply/recover."""
    lineage = model['lineage']
    if not lineage:
        return {}, set()
    current = ctx.json('.ai-sow-lite/current.json', '.ai-sow-lite', 'artifacts', 'current')
    if current is None:
        return None, None
    manifests = {}
    version, expected_hash = current['version_id'], current['manifest_hash']
    needed = {record['from_version_id'] for record in lineage}
    while version and version not in manifests:
        path = f'.ai-sow-lite/versions/{version}/manifest.json'
        manifest = ctx.json(path, '.ai-sow-lite/versions', 'artifacts', 'manifest')
        if manifest is None:
            return None, None
        if ctx.dependencies[path]['sha256'] != expected_hash or manifest['version_id'] != version:
            ctx.add('lineage', '历史 manifest 身份或摘要与已应用依赖链不符。')
            return None, None
        manifests[version] = manifest
        if needed <= set(manifests):
            break
        version = manifest['base_version_id']
        if version:
            parent_path = f'.ai-sow-lite/versions/{version}/manifest.json'
            parent_ref = next((r for r in manifest['dependencies'] if r['path'] == parent_path), None)
            if parent_ref is None:
                ctx.add('lineage', '历史前序 manifest 必须由当前版依赖摘要绑定。')
                return None, None
            expected_hash = parent_ref['sha256']
    if not needed <= set(manifests):
        ctx.add('from_version_id', 'Lineage 原版本不是可核实的已应用前序版本。')
        return None, None
    if base_version_id and base_version_id not in manifests:
        ctx.add('base_version_id', '候选基线不属于已应用历史。')
    # Only this current-to-explicit-origin interval is read. These snapshots locate
    # where an inherited replacement actually first appeared, not just any bare ID.
    versions = list(reversed(manifests))
    snapshots, objects_by_version = {}, {}
    for version in versions:
        path = f'.ai-sow-lite/versions/{version}/model.json'
        old = ctx.json(path, f'.ai-sow-lite/versions/{version}', 'model')
        file_ref = next((r for r in manifests[version]['files'] if r['path'] == path), None)
        if old is None:
            return None, None
        if file_ref is None or ctx.dependencies[path]['sha256'] != file_ref['sha256']:
            ctx.add('from_ids', '历史模型未匹配 manifest。')
            return None, None
        snapshots[version] = old
        objects_by_version[version], _ = _objects(ctx, old, path)
    ranks = {version: position for position, version in enumerate(versions)}
    candidate_rank = len(versions)
    keys, history, transitions = set(), {}, []
    for record in lineage:
        key = _lineage_key(record)
        source_version = record['from_version_id']
        source_rank = ranks[source_version]
        if key in keys:
            ctx.add('lineage', '历史复合键重复。')
        keys.add(key)
        for identity in record['from_ids']:
            obj = objects_by_version[source_version].get(identity)
            if obj is None:
                ctx.add('from_ids', '旧对象未出现在精确历史版本中。', object_id=identity)
            else:
                history[(source_version, identity)] = obj
        occurrences = [(ranks[v], saved) for v in versions for saved in snapshots[v]['lineage']
                       if _lineage_key(saved) == key]
        if any(saved != record for _, saved in occurrences):
            ctx.add('lineage', '继承的历史替换记录与已保存内容不一致，不能重写去向。')
            continue
        applied_rank = min((rank for rank, _ in occurrences), default=candidate_rank)
        if applied_rank <= source_rank:
            ctx.add('from_version_id', '替换必须发生在所引用原版本之后。')
            continue
        live_until_replaced = all(identity in objects_by_version[versions[rank]]
                                  for identity in record['from_ids']
                                  for rank in range(source_rank, applied_rank))
        if not live_until_replaced:
            ctx.add('from_ids', '原对象在此次替换之前已退出；不能借早期版本重新接续。')
            continue
        destination = current_objects if applied_rank == candidate_rank else objects_by_version[versions[applied_rank]]
        if any(identity not in destination for identity in record['to_ids']):
            ctx.add('to_ids', '去向必须存在于该次替换实际发生的版本；历史中间项需要已绑定替换记录。')
            continue
        transitions.append((record, source_rank, applied_rank))
    for record, _, applied_rank in transitions:
        for identity in record['to_ids']:
            # A surviving target must remain present; retired targets need a later
            # source coordinate and a later replacement, ending in current or [].
            if identity in current_objects and all(identity in objects_by_version[versions[rank]]
                                                   for rank in range(applied_rank, candidate_rank)):
                continue
            successor = any(identity in later['from_ids'] and later_source >= applied_rank and later_applied > applied_rank
                            and all(identity in objects_by_version[versions[rank]]
                                    for rank in range(applied_rank, later_applied))
                            for later, later_source, later_applied in transitions)
            if not successor:
                ctx.add('to_ids', '已退出的中间去向缺少时间更晚且可追溯的替换或明确删除。', object_id=identity)
    return history, keys


def _pending_relationships(ctx, pending, decisions, objects, kinds, history, lineage_keys):
    """Check each valid collection; None means unknown, never a legal empty file."""
    questions = {item['id']: item for item in pending['items']} if pending is not None else None
    applied = {item['id']: item for item in decisions['items']} if decisions is not None else None
    ids_seen = set(objects) if objects is not None else set()
    for collection in (pending, decisions):
        if collection is None:
            continue
        for obj in collection['items']:
            if obj['id'] in ids_seen:
                ctx.add('id', '问题、决定和模型对象的 ID 不能重复。', object_id=obj['id'])
            ids_seen.add(obj['id'])

    def target_valid(target, allow_history, owner):
        if objects is None:
            return  # model structure is invalid; no claim that this target is absent
        identity, field = target['object_id'], target['field']
        obj = objects.get(identity)
        if obj is None and allow_history:
            if history is None:
                return  # its bound history could not be read; absence is unknown
            matches = [value for (version, old_id), value in history.items() if old_id == identity]
            obj = next((value for value in matches if field is None or field in value), None)
        if obj is None or (field is not None and (field not in obj or field == 'id')):
            ctx.add('targets', '目标对象或字段不存在于允许的当前/历史集合。', object_id=owner)

    exact_open, work_gaps = (set(), set()) if pending is not None else (None, None)
    if pending is not None:
        for item in pending['items']:
            for target in item['targets']:
                target_valid(target, item['status'] != 'open', item['id'])
                if item['status'] == 'open':
                    exact_open.add((target['object_id'], target['field']))
                    if item['unestimated_work']:
                        if target['field'] is not None or (kinds is not None and kinds.get(target['object_id']) not in ('epics', 'features', 'stories')):
                            ctx.add('unestimated_work', '未拆明工作只能指向 Epic/Feature/Story 对象级目标。', object_id=item['id'])
                        work_gaps.add(target['object_id'])
            resolution = item['resolution']
            if item['status'] == 'resolved' and applied is not None:
                decision = applied.get(resolution['decision_id'])
                if decision is None or not all(target in decision['applies_to'] for target in item['targets']):
                    ctx.add('resolution', '已解决问题须有同版决定覆盖其全部目标。', object_id=item['id'])
            elif item['status'] == 'superseded':
                for replacement in resolution['replacement_item_ids']:
                    if replacement not in questions or replacement == item['id']:
                        ctx.add('replacement_item_ids', '替代问题不存在或自引用。', object_id=item['id'])
                if lineage_keys is not None:
                    for key in resolution['lineage_refs']:
                        if (key['from_version_id'], tuple(key['from_ids'])) not in lineage_keys:
                            ctx.add('lineage_refs', '问题历史复合键不能解析。', object_id=item['id'])
        _acyclic(ctx, {item['id']: item['resolution']['replacement_item_ids']
                      for item in pending['items'] if item['status'] == 'superseded'}, 'replacement_item_ids')
    if decisions is not None:
        for decision in decisions['items']:
            for target in decision['applies_to']:
                target_valid(target, True, decision['id'])
            for identity in decision.get('supersedes', []):
                if identity not in applied or identity == decision['id']:
                    ctx.add('supersedes', '替代决定不存在或自引用。', object_id=decision['id'])
        _acyclic(ctx, {item['id']: item.get('supersedes', []) for item in decisions['items']}, 'supersedes')
    return exact_open, work_gaps


def _model_classification(ctx, model, standards, exact_open, work_gaps):
    """Model-only rules always run; null/gap checks also need a valid pending file."""
    for collection, child_kind, parent in [('epics', 'features', 'epic_id'), ('features', 'stories', 'feature_id'), ('stories', 'tasks', 'story_id')]:
        parents = {child[parent] for child in model[child_kind]}
        for obj in model[collection]:
            if work_gaps is not None and obj['id'] not in parents and obj['id'] not in work_gaps:
                ctx.add(None, '空父项须有明确的 open 未拆明工作，或从本期范围移除。', object_id=obj['id'])
    for story in model['stories']:
        if exact_open is not None and not story['acs'] and (story['id'], 'acs') not in exact_open:
            ctx.add('acs', '空 AC 数组须有指向 Story.acs 的 open 问题。', object_id=story['id'])
    for task in model['tasks']:
        standard = standards.get(task['work_type_name']) if standards is not None else None
        if standards is not None and task['work_type_name'] is not None and standard is None:
            ctx.add('work_type_name', '工作类型必须精确匹配同版模板。', object_id=task['id'])
        if standard and task['work_mode'] is not None and task['work_mode'] not in standard['modes']:
            ctx.add('work_mode', '此工作方式在同版模板中不适用。', object_id=task['id'])
        na = 'integration_type' in task['not_applicable_fields']
        if na and (task['integration_type'] is not None or (standard and standard['integration'])):
            ctx.add('integration_type', '集成工作或已有集成类型不能标为不适用。', object_id=task['id'])
        if standard and not standard['integration'] and task['integration_type'] is not None:
            ctx.add('integration_type', '非集成工作不能填入集成类型。', object_id=task['id'])
        covered = set()
        for basis in task['classification_basis']:
            covered.update(basis['fields'])
            if standard and basis['standard_id'] != standard['id']:
                ctx.add('standard_id', '分类依据须定位该工作类型的同版标准行。', object_id=task['id'])
            if task['work_type_name'] is None and basis['standard_id'] is not None:
                ctx.add('standard_id', '未匹配类型不能编造标准身份。', object_id=task['id'])
        for field in CLASSIFICATION_FIELDS:
            applicable_na = field == 'integration_type' and na
            if task[field] is None and not applicable_na:
                if exact_open is not None and (task['id'], field) not in exact_open:
                    ctx.add(field, '未知分类须有精确字段的 open 问题。', object_id=task['id'])
            elif field not in covered:
                ctx.add(field, '有效分类和明确不适用均须有分类依据。', object_id=task['id'])


def _sources_and_evidence(ctx, candidate, model, pending, decisions, objects, analysis_records=None):
    model_valid = objects is not None
    index = ctx.json('.ai-sow-lite/inputs/index.json', '.ai-sow-lite/inputs', 'artifacts', 'input_index')
    inputs, raw_sources = {}, {}
    if index:
        for entry in index['items']:
            if entry['input_version_id'] in inputs:
                ctx.add('input_version_id', '输入版本重复登记。')
            inputs[entry['input_version_id']] = entry
    for identity in candidate['input_version_ids']:
        if index is None:
            break  # invalid registry: existence is unknown, not an empty index
        entry = inputs.get(identity)
        if not entry:
            ctx.add('input_version_ids', '候选输入版本未登记。', object_id=identity, code='EVIDENCE_MISSING')
            continue
        path, raw = ctx.bytes(entry['relative_path'], f'.ai-sow-lite/inputs/originals/{identity}', 'relative_path')
        if raw is not None:
            raw_sources[identity] = raw
            if hashlib.sha256(raw).hexdigest() != entry['content_hash']:
                ctx.add('content_hash', '原件字节与登记摘要不同。', path=path, code='EVIDENCE_MISSING')

    def source_valid(source, owner, *, topic_inputs=None, uncovered=False):
        identity = source['input_version_id']
        entry, raw = inputs.get(identity), raw_sources.get(identity)
        field = 'uncovered_regions' if uncovered else 'covered_regions' if topic_inputs is not None else 'source_refs'
        registered = entry is not None and identity in candidate['input_version_ids']
        if index is not None and not registered:
            ctx.add(field, '来源必须在输入索引登记并由候选采用。', object_id=owner, code='EVIDENCE_MISSING')
        if topic_inputs is not None and identity not in topic_inputs:
            ctx.add(field, '覆盖记录的来源必须属于该主题采用的输入集合。', object_id=owner)
        locator = source['locator']
        if locator['kind'] != 'text_lines':
            # XLSX readers belong to I2.1; prototype/observations to I4.
            ctx.add('locator', '仅支持 text_lines 原字节定位；此定位适配尚未实现。', object_id=owner, code='OPERATION_UNSUPPORTED')
            return
        if locator['start_line'] > locator['end_line']:
            ctx.add('locator', '文本范围的起始行不能大于结束行。', object_id=owner)
            return
        if not registered:
            return
        accepted_paths = (None, entry['relative_path'], PurePosixPath(entry['relative_path']).name)
        if entry['format'] != 'text' or locator.get('path') not in accepted_paths:
            ctx.add('locator', '单文件定位路径须省略或精确匹配登记相对路径/文件名。', object_id=owner)
            return
        if uncovered:
            # An unread region has no successful excerpt yet. Check its declared
            # identity/order, without manufacturing a read or an excerpt hash.
            return
        if raw is None:
            ctx.add(field, '采用的来源原件不可读取。', object_id=owner, code='EVIDENCE_MISSING')
            return
        try:
            text = raw.decode(entry.get('encoding', 'utf-8'), errors='strict')
            lines = text.splitlines(keepends=True)
            start, end = locator['start_line'], locator['end_line']
            if not 1 <= start <= end <= len(lines):
                raise ValueError('line range')
            excerpt = ''.join(lines[start - 1:end]).encode('utf-8')
            if hashlib.sha256(excerpt).hexdigest() != source['excerpt_hash']:
                raise ValueError('excerpt hash')
        except (UnicodeError, LookupError, ValueError):
            ctx.add('source_refs', '来源行范围、编码或摘录摘要无法匹配真实原件。', object_id=owner, code='EVIDENCE_MISSING')

    evidence, topics_by_version = {}, {}
    analyses_valid = True
    if not candidate['topic_version_ids']:
        ctx.add('topic_version_ids', '候选必须绑定实际分析记录；空模型也不例外。', code='EVIDENCE_MISSING')
    for version in candidate['topic_version_ids']:
        area = f'.ai-sow-lite/analysis/topics/{version}'
        analysis = (analysis_records[version] if analysis_records is not None else
                    ctx.json(area + '/analysis.json', area, 'artifacts', 'analysis'))
        if not analysis:
            analyses_valid = False
            continue
        if not any(topic['topic_version_id'] == version for topic in analysis['topics']):
            ctx.add('topic_version_id', '分析记录与不可变目录身份不同。')
        local_versions = set()
        for topic in analysis['topics']:
            topic_version = topic['topic_version_id']
            if topic_version not in candidate['topic_version_ids']:
                ctx.add('topic_version_id', '分析中的主题必须由候选明确采用。')
            if topic_version in local_versions:
                ctx.add('topic_version_id', '同一分析文件不能重复声明主题版本。', object_id=topic['topic_id'])
            local_versions.add(topic_version)
            previous = topics_by_version.get(topic_version)
            if previous is not None and previous != topic:
                ctx.add('topic_version_id', '不可变主题版本的内容或 topic_id 相互冲突。', object_id=topic['topic_id'])
            elif previous is None:
                topics_by_version[topic_version] = topic
        local_ids = set()
        for record in analysis['evidence']:
            if record['id'] in local_ids:
                ctx.add('id', '同一分析文件内的依据 ID 必须唯一。', object_id=record['id'])
            local_ids.add(record['id'])
            old = evidence.get(record['id'])
            if old is not None and old != record:
                ctx.add('id', '不同分析不能用相同依据 ID 覆盖不同内容。', object_id=record['id'])
            evidence[record['id']] = record
        if analysis['observations']:
            ctx.add('observations', '观察登记与附件校验将在 I4 实现。', code='OPERATION_UNSUPPORTED')
    topics = list(topics_by_version.values())
    adopted = set(candidate['evidence_ids'])
    for identity in adopted:
        if analyses_valid and identity not in evidence:
            ctx.add('evidence_ids', '采用的依据未出现在对应不可变分析中。', object_id=identity, code='EVIDENCE_MISSING')
    nodes = {identity: evidence[identity] for identity in adopted if identity in evidence}
    if decisions is not None:
        for decision in decisions['items']:
            if decision['id'] in nodes or (objects is not None and decision['id'] in objects):
                ctx.add('id', '依据/决定/模型身份必须互异。', object_id=decision['id'])
            nodes[decision['id']] = decision
    business_ids = set(objects) if objects is not None else set()
    if pending is not None:
        business_ids.update(item['id'] for item in pending['items'])
    for identity in adopted:
        if identity in business_ids:
            ctx.add('id', '依据身份不能与业务对象或问题重复。', object_id=identity)

    def references(refs, field, owner):
        for identity in refs:
            if analyses_valid and decisions is not None and identity not in nodes:
                ctx.add(field, '引用未属于候选采用的依据或同版决定。', object_id=owner, code='EVIDENCE_MISSING')

    collections = []
    if model is not None:
        collections.extend([list(objects.values()), model['lineage']])
    for bundle in (pending, decisions):
        if bundle is not None:
            collections.append(bundle['items'])
    for collection in collections:
        for obj in collection:
            references(obj['evidence_refs'], 'evidence_refs', obj.get('id'))
            for basis in obj.get('classification_basis', []):
                references(basis['evidence_refs'], 'classification_basis', obj.get('id'))
    for topic in topics:
        if not set(topic['input_version_ids']) <= set(candidate['input_version_ids']):
            ctx.add('input_version_ids', '主题采用了候选之外的来源。', object_id=topic['topic_id'])
        for identity in topic['related_object_ids']:
            if model_valid and identity not in objects:
                ctx.add('related_object_ids', '分析仍引用候选中不存在的业务对象。', object_id=identity)
        references(topic['evidence_refs'], 'evidence_refs', topic['topic_id'])
        for source in topic['covered_regions']:
            source_valid(source, topic['topic_id'], topic_inputs=topic['input_version_ids'])
        for source in topic['uncovered_regions']:
            source_valid(source, topic['topic_id'], topic_inputs=topic['input_version_ids'], uncovered=True)
        # parent_id has no cross-version coordinate: it resolves only inside
        # this immutable topic version, never against another topic's bare IDs.
        historical_ids = set()
        parents = {}
        for item in topic['historical_items']:
            if item['id'] in historical_ids:
                ctx.add('id', '同一主题版本中的历史条目 ID 必须唯一。', object_id=item['id'])
            historical_ids.add(item['id'])
            references(item['evidence_refs'], 'evidence_refs', item['id'])
            for fact in item.get('instance_facts', []):
                references(fact['evidence_refs'], 'instance_facts', item['id'])
            parents[item['id']] = [item['parent_id']] if 'parent_id' in item else []
        for identity, parent_ids in parents.items():
            if any(parent not in historical_ids for parent in parent_ids):
                ctx.add('parent_id', '历史父引用必须在同一主题版本中明确存在。', object_id=identity)
        _acyclic(ctx, parents, 'parent_id')
    graph = {}
    for identity, record in nodes.items():
        if identity in adopted:
            for source in record['source_refs']:
                source_valid(source, identity)
            edges = record['basis_refs']
        else:
            edges = record['evidence_refs']
        references(edges, 'basis_refs', identity)
        graph[identity] = edges
    _acyclic(ctx, graph, 'basis_refs')
    return evidence, topics


def check_analysis(project, analysis):
    """Share source/identity checks with candidates; no business objects are invented."""
    ctx = _Check(Path(project))
    for error in schema_validator('artifacts', 'analysis').iter_errors(analysis):
        ctx.add('analysis', '分析字段不符合当前 Schema。')
    if ctx.diagnostics:
        return ctx.diagnostics
    records = {topic['topic_version_id']: analysis for topic in analysis['topics']}
    candidate = dict(input_version_ids=list(dict.fromkeys(i for t in analysis['topics'] for i in t['input_version_ids'])),
                     topic_version_ids=list(records), evidence_ids=[e['id'] for e in analysis['evidence']])
    _sources_and_evidence(ctx, candidate, None, None, dict(items=[]), None, records)
    return ctx.diagnostics


def _acyclic(ctx, graph, field):
    # Iterative DFS keeps long valid chains from depending on Python recursion depth.
    done = set()
    for root in graph:
        if root in done:
            continue
        active = set()
        stack = [(root, False)]
        while stack:
            node, exiting = stack.pop()
            if exiting:
                active.discard(node)
                done.add(node)
                continue
            if node in active:
                ctx.add(field, '引用图包含循环。', object_id=node)
                continue
            if node in done or node not in graph:
                continue
            active.add(node)
            stack.append((node, True))
            stack.extend((child, False) for child in reversed(graph[node]))


def check_candidate(project: Path, candidate_path: Path, scope: str, plan_path: Path | None):
    """Return a report body, never the CLI result; no writes or inferred business values."""
    ctx = _Check(Path(project))
    if scope not in ('slice', 'full'):
        ctx.add('scope', 'scope 必须为 slice 或 full。')
        return ctx.report('full')
    if plan_path is not None:
        ctx.add('plan_path', '带计划的检查尚未实现。', code='OPERATION_UNSUPPORTED')
        return ctx.report(scope)
    try:
        path = project_file(Path(project), candidate_path, '.ai-sow-lite/work')
    except (ValueError, OSError, RuntimeError):
        ctx.add('candidate_path', '候选必须在显式项目的 work 目录中。')
        return ctx.report(scope)
    candidate = ctx.json(path, '.ai-sow-lite/work', 'artifacts', 'candidate')
    if not candidate:
        return ctx.report(scope, path)
    parts = path.relative_to(ctx.project).parts
    if len(parts) < 5 or parts[2] != candidate['entrypoint']:
        ctx.add('candidate_path', '候选路径必须包含入口和请求身份。')
        return ctx.report(scope, path)
    try:
        if UUID(parts[3]).version != 4:
            raise ValueError('request identity')
    except ValueError:
        ctx.add('candidate_path', 'work 请求目录必须为 UUID4。')
        return ctx.report(scope, path)
    work_area = '/'.join(parts[:4])
    project_record = ctx.json('.ai-sow-lite/project.json', '.ai-sow-lite', 'artifacts', 'project')
    if project_record and candidate['template_hash'] != project_record['template_hash']:
        ctx.add('template_hash', '候选模板身份与项目固定模板不同。')
    if candidate['entrypoint'] == 'generate' and candidate['base_version_id'] is not None:
        ctx.add('base_version_id', '首版 generate 的基线必须为 null。')
    if candidate['entrypoint'] == 'clarify':
        ctx.add('entrypoint', 'Clarify 基线和方案检查将在 I3 实现。', code='OPERATION_UNSUPPORTED')
    model = ctx.json(candidate['model_path'], work_area, 'model', field='model_path')
    pending = ctx.json(candidate['pending_items_path'], work_area, 'pending-items', field='pending_items_path')
    decisions = ctx.json(candidate['decisions_path'], work_area, 'decisions', field='decisions_path')
    standards = _standards(ctx, candidate['template_hash'])
    objects, kinds = _objects(ctx, model, candidate['model_path']) if model is not None else (None, None)
    history, keys = _history(ctx, model, candidate['base_version_id'], objects) if model is not None else (None, None)
    exact_open, work_gaps = _pending_relationships(ctx, pending, decisions, objects, kinds, history, keys)
    if model is not None:
        _model_classification(ctx, model, standards, exact_open, work_gaps)
        lineage_graph = {}
        for record in model['lineage']:
            for old in record['from_ids']:
                lineage_graph.setdefault(old, []).extend(record['to_ids'])
        _acyclic(ctx, lineage_graph, 'lineage')
    evidence, topics = _sources_and_evidence(ctx, candidate, model, pending, decisions, objects)
    unknowns = sum(item['status'] == 'open' for item in pending['items']) if pending is not None else 0
    if model is None or pending is None or decisions is None:
        return ctx.report(scope, path, unknowns=unknowns)
    digest = semantic_digest(dict(candidate=candidate, model=model, pending_items=pending,
                                  decisions=decisions, topics=topics,
                                  evidence=[evidence[i] for i in candidate['evidence_ids'] if i in evidence]))
    return ctx.report(scope, path, digest, unknowns)
