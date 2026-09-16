"""Bounded physical XLSX observations, never calculation or scope inference."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
import posixpath
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile

from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_datetime, is_timedelta_format
from openpyxl.utils.cell import range_boundaries, get_column_letter
from openpyxl.utils.datetime import from_excel, from_ISO8601, MAC_EPOCH, WINDOWS_EPOCH

from .project import StorageError

READER_VERSION = 'lite-xlsx-v1'
OPTIONS = {'data_only': False}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ENTRIES = 2048
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_SHEETS = 64
MAX_ROWS = 100000
MAX_COLUMNS = 512
MAX_DIMENSION_CELLS = 1000000
MAX_TOTAL_CELLS = 2000000
MAX_CELL_TEXT = 32767
MAX_TOTAL_TEXT = 10000000
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def _limit(message):
    return StorageError('RESULT_TOO_LARGE', message + '；未返回部分读取结果，请提供较小文件或明确区域的可读副本。')


def bounds(address):
    try:
        left, top, right, bottom = range_boundaries(address)
        if not all(type(v) is int for v in (left, top, right, bottom)) or not 1 <= left <= right <= MAX_COLUMNS or not 1 <= top <= bottom <= MAX_ROWS:
            raise ValueError('range')
        if (right - left + 1) * (bottom - top + 1) > MAX_DIMENSION_CELLS:
            raise ValueError('range size')
        return left, top, right, bottom
    except (ValueError, TypeError):
        raise StorageError('CANDIDATE_INVALID', 'XLSX 区域必须是有限 A1 范围（上限100000行、512列、1000000格）。') from None


def canonical_range(address):
    left, top, right, bottom = bounds(address)
    start, end = f'{get_column_letter(left)}{top}', f'{get_column_letter(right)}{bottom}'
    return start if start == end else f'{start}:{end}'


def _physical_bounds(address):
    """Validate OOXML coordinates before applying our smaller reading budget."""
    left, top, right, bottom = range_boundaries(address)
    if (not all(type(v) is int for v in (left, top, right, bottom)) or
            not 1 <= left <= right <= 16384 or not 1 <= top <= bottom <= 1048576):
        raise ValueError('physical range')
    return left, top, right, bottom


def _target(base, target):
    resolved = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(base), target))
    if resolved.startswith('../') or '\\' in resolved:
        raise ValueError('relationship')
    return resolved


def _preflight(archive, raw):
    entries = archive.infolist()
    if len(raw) > MAX_FILE_BYTES:
        raise _limit('XLSX 压缩文件超过50 MiB上限')
    if len(entries) > MAX_ENTRIES:
        raise _limit('XLSX ZIP 成员超过2048条上限')
    if len({e.filename for e in entries}) != len(entries):
        raise ValueError('duplicate ZIP members')
    total, compressed = sum(e.file_size for e in entries), sum(e.compress_size for e in entries)
    if total > MAX_UNCOMPRESSED_BYTES:
        raise _limit('XLSX ZIP 解压体积超过200 MiB上限')
    if total and (not compressed or total > MAX_COMPRESSION_RATIO * compressed):
        raise _limit('XLSX ZIP 压缩比超过100倍上限')
    if any(e.flag_bits & 1 for e in entries):
        raise ValueError('encrypted ZIP')


def _number(text):
    value = Decimal(text)
    if not value.is_finite():
        raise ValueError('nonfinite number')
    return dict(type='number', value=str(value))


def _typed(kind, text, shared):
    if text is None:
        return dict(type='blank', value=None)
    if kind == 's':
        index = int(text)
        if not 0 <= index < len(shared):
            raise ValueError('shared string')
        return dict(type='string', value=shared[index])
    if kind in ('inlineStr', 'str'):
        return dict(type='string', value=text)
    if kind == 'b':
        if text not in ('0', '1'):
            raise ValueError('boolean')
        return dict(type='boolean', value=text == '1')
    if kind == 'e':
        return dict(type='error', value=text)
    if kind == 'd':
        value = from_ISO8601(text)
        return dict(type=type(value).__name__, value=value.isoformat())
    if kind != 'n':
        raise ValueError('unsupported cell type')
    return _number(text)


def _temporal(value, number_format, epoch):
    if value['type'] != 'number' or not is_date_format(number_format):
        return value
    duration = is_timedelta_format(number_format)
    date_value = from_excel(float(value['value']), epoch, timedelta=duration)
    if duration:
        return dict(type='duration', value=format(Decimal(str(date_value.total_seconds())).normalize(), 'f'))
    kind = is_datetime(number_format)
    if kind == 'date' and hasattr(date_value, 'date'):
        date_value = date_value.date()
    elif kind == 'time' and hasattr(date_value, 'time'):
        date_value = date_value.time()
    return dict(type=type(date_value).__name__, value=date_value.isoformat())


def _body_text(node):
    """OOXML body t/r/t only; rPh phonetic annotations are not cell text."""
    if node is None:
        return ''
    parts = []
    for child in node:
        if child.tag == NS + 't':
            parts.append(child.text or '')
        elif child.tag == NS + 'r':
            parts.extend(t.text or '' for t in child.findall(NS + 't'))
    return ''.join(parts)


def read(raw, selection):
    """Directory or one exact region. Parsing limits never masquerade as completeness."""
    try:
        with ZipFile(BytesIO(raw)) as archive:
            _preflight(archive, raw)
            names = set(archive.namelist())

            def xml(part):
                payload = archive.read(part)
                declarations = payload.replace(b'\x00', b'')  # also inspect UTF-16/32 XML declarations
                if b'<!DOCTYPE' in declarations or b'<!ENTITY' in declarations:
                    raise ValueError('XML entity declaration')
                return ET.fromstring(payload)

            def relationships(part):
                path = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
                if path not in names:
                    return {}
                return {node.attrib['Id']: _target(part, node.attrib['Target']) for node in xml(path)
                        if node.attrib.get('TargetMode') != 'External'}

            shared = []
            if 'xl/sharedStrings.xml' in names:
                shared = [_body_text(node) for node in xml('xl/sharedStrings.xml')]
            formats = []
            if 'xl/styles.xml' in names:
                styles = xml('xl/styles.xml')
                custom = {int(n.attrib['numFmtId']): n.attrib['formatCode'] for n in styles.findall(NS + 'numFmts/' + NS + 'numFmt')}
                formats = [custom.get(int(n.attrib.get('numFmtId', 0)), BUILTIN_FORMATS.get(int(n.attrib.get('numFmtId', 0)), 'General'))
                           for n in styles.findall(NS + 'cellXfs/' + NS + 'xf')]
            workbook = xml('xl/workbook.xml')
            props = workbook.find(NS + 'workbookPr')
            epoch = MAC_EPOCH if props is not None and props.attrib.get('date1904') in ('1', 'true') else WINDOWS_EPOCH
            sheets = workbook.find(NS + 'sheets')
            if len(sheets) > MAX_SHEETS:
                raise _limit('XLSX Sheet 数量超过64个上限')
            targets = relationships('xl/workbook.xml')
            directory, selected, total_cells, total_text = [], None, 0, 0
            limitations = []
            for part in sorted(names):
                if part.endswith('.vml') or part.startswith(('xl/drawings/', 'xl/media/', 'xl/embeddings/', 'xl/charts/', 'xl/externalLinks/', 'xl/pivot', 'xl/threadedComments/', 'xl/richData/')) or 'vbaproject' in part.lower():
                    limitations.append(dict(part=part, coverage='unread', reason='图形、对象、外链或扩展内容未解释；不会刷新或执行。'))
                if part.endswith('.rels'):
                    for rel in xml(part):
                        if rel.attrib.get('TargetMode') == 'External':
                            limitations.append(dict(part=part, coverage='unread', reason='外部关系未访问。'))
            for sheet in sheets:
                name = sheet.attrib['name']
                part = targets[sheet.attrib[REL]]
                root = xml(part)
                dimension = root.find(NS + 'dimension')
                declared = dimension.attrib['ref'] if dimension is not None else 'A1'
                _, _, maxcol, maxrow = _physical_bounds(declared)
                nodes, content_col, content_row = {}, 1, 1
                try:
                    merges = [canonical_range(n.attrib['ref']) for n in root.findall(NS + 'mergeCells/' + NS + 'mergeCell')]
                    for n in root.findall(NS + 'sheetData/' + NS + 'row/' + NS + 'c'):
                        col, row, right, bottom = _physical_bounds(n.attrib['r'])
                        address = f'{get_column_letter(col)}{row}'
                        if col != right or row != bottom or address in nodes:
                            raise ValueError('cell address')
                        nodes[address] = n
                        maxcol, maxrow = max(maxcol, col), max(maxrow, row)
                        # An empty value, formula or inline string is still content.
                        # Only cells with none of these children are pure formatting.
                        if any(n.find(NS + tag) is not None for tag in ('v', 'f', 'is')):
                            content_col, content_row = max(content_col, col), max(content_row, row)
                    for merge in merges:
                        _, _, right, bottom = bounds(merge)
                        maxcol, maxrow = max(maxcol, right), max(maxrow, bottom)
                        content_col, content_row = max(content_col, right), max(content_row, bottom)
                except StorageError:
                    raise _limit(f'XLSX Sheet {name} 实际内容或合并区域超过读取上限') from None
                try:
                    used = canonical_range(f'A1:{get_column_letter(maxcol)}{maxrow}')
                    inflated = False
                except StorageError:
                    # Preserve byte-identical v1 observations for previously accepted
                    # sheets. Only sheets formerly rejected get a content-bound fallback.
                    maxcol, maxrow, inflated = content_col, content_row, True
                hidden_rows = [int(n.attrib['r']) for n in root.findall(NS + 'sheetData/' + NS + 'row') if n.attrib.get('hidden') in ('1', 'true')]
                hidden_columns = [dict(min=int(n.attrib['min']), max=int(n.attrib['max'])) for n in root.findall(NS + 'cols/' + NS + 'col') if n.attrib.get('hidden') in ('1', 'true')]
                links, comments, tables = relationships(part), {}, []
                for target in links.values():
                    if target not in names:
                        raise ValueError('missing relationship part')
                    if 'comment' in target.lower() and target.endswith('.xml'):
                        comment_root = xml(target)
                        authors = [n.text or '' for n in comment_root.findall(NS + 'authors/' + NS + 'author')]
                        for n in comment_root.findall(NS + 'commentList/' + NS + 'comment'):
                            comments[n.attrib['ref']] = dict(author=authors[int(n.attrib['authorId'])], text=''.join(t.text or '' for t in n.iter(NS + 't')))
                for n in root.findall(NS + 'tableParts/' + NS + 'tablePart'):
                    table = xml(links[n.attrib[REL]])
                    try:
                        tables.append(dict(name=table.attrib['displayName'], range=canonical_range(table.attrib['ref'])))
                    except StorageError:
                        raise _limit(f'XLSX Sheet {name} Table 区域超过读取上限') from None
                if inflated:
                    # Comments and table/merge structure must not disappear just
                    # because the associated cell has no value-bearing child.
                    for address in list(comments) + [table['range'] for table in tables]:
                        _, _, right, bottom = _physical_bounds(address)
                        maxcol, maxrow = max(maxcol, right), max(maxrow, bottom)
                    try:
                        used = canonical_range(f'A1:{get_column_letter(maxcol)}{maxrow}')
                    except StorageError:
                        raise _limit(f'XLSX Sheet {name} 实际内容、附注或结构范围超过读取上限') from None
                    limitations.append(dict(part=part, coverage='unread', reason=(
                        f'声明维度 {declared} 或纯样式空单元格范围超过读取上限；'
                        f'已按实际内容边界 {used} 读取，外围纯样式未解释，原件未改写。')))
                total_cells += maxcol * maxrow
                if total_cells > MAX_TOTAL_CELLS:
                    raise _limit('XLSX 全部 Sheet 单元格总数超过2000000格上限')
                info = dict(sheet=name, state=sheet.attrib.get('state', 'visible'), used_range=used, merges=merges,
                            tables=tables, hidden_rows=hidden_rows, hidden_columns=hidden_columns, comment_cells=sorted(comments))
                directory.append(info)
                # Validate all represented text even during inventory; never silently truncate.
                for n in nodes.values():
                    value = n.findtext(NS + 'v')
                    texts = [_body_text(n.find(NS + 'is')), n.findtext(NS + 'f') or '']
                    if n.attrib.get('t') == 's' and value is not None:
                        texts.append(_typed('s', value, shared)['value'])
                    if n.attrib.get('t') == 'str':
                        texts.append(value or '')
                    count = sum(len(t) for t in texts)
                    if count > MAX_CELL_TEXT:
                        raise _limit(f"XLSX Sheet {name} 单元格 {n.attrib['r']} 文本超过32767字符上限")
                    total_text += count
                total_text += sum(len(c['text']) + len(c['author']) for c in comments.values())
                if total_text > MAX_TOTAL_TEXT:
                    raise _limit('XLSX 文本总量超过10000000字符上限')
                if selection.get('sheet') != name:
                    continue
                left, top, right, bottom = bounds(selection['range'])
                if right > maxcol or bottom > maxrow:
                    raise StorageError('EVIDENCE_MISSING', 'XLSX 所选区域超出实际 Sheet 维度。')
                merge_bounds = [(m, bounds(m)) for m in merges]
                shared_formulas = {n.find(NS + 'f').attrib['si']: (a, n.findtext(NS + 'f')) for a, n in nodes.items()
                                   if n.find(NS + 'f') is not None and n.find(NS + 'f').attrib.get('t') == 'shared' and n.findtext(NS + 'f')}
                cells = []
                for row in range(top, bottom + 1):
                    for col in range(left, right + 1):
                        address = f'{get_column_letter(col)}{row}'
                        n = nodes.get(address)
                        raw_value, formula, cache, fmt = dict(type='blank', value=None), None, None, 'General'
                        if n is not None:
                            kind = n.attrib.get('t', 'n')
                            value = n.findtext(NS + 'v')
                            if kind == 'inlineStr':
                                value = _body_text(n.find(NS + 'is'))
                            # Empty <v/> is absent; 0 and false are explicitly present.
                            if value == '' and kind not in ('str', 'inlineStr'):
                                value = None
                            raw_value = _typed(kind, value, shared)
                            if 's' in n.attrib:
                                fmt = formats[int(n.attrib['s'])]
                            f = n.find(NS + 'f')
                            if f is not None:
                                text = f.text
                                if f.attrib.get('t') == 'shared' and not text:
                                    origin, master = shared_formulas[f.attrib['si']]
                                    text = Translator('=' + master, origin=origin).translate_formula(address)[1:]
                                formula = dict(text='=' + text if text else None, attributes=dict(f.attrib))
                                cache = dict(present=value is not None, value=_temporal(raw_value, fmt, epoch) if value is not None else None)
                        merge = next((m for m, b in merge_bounds if b[0] <= col <= b[2] and b[1] <= row <= b[3]), None)
                        cells.append(dict(address=address, value=_temporal(raw_value, fmt, epoch) if formula is None else dict(type='blank', value=None),
                                          raw_value=raw_value, formula=formula, cache=cache, number_format=fmt,
                                          merged_range=merge, merge_anchor=merge.split(':')[0] if merge else None,
                                          hidden_row=row in hidden_rows, hidden_column=any(c['min'] <= col <= c['max'] for c in hidden_columns),
                                          comment=comments.get(address)))
                selected = dict(sheet=name, range=canonical_range(selection['range']), cells=cells, metadata=info, limitations=limitations)
            if selection['kind'] == 'workbook':
                return dict(sheets=directory, limitations=limitations)
            if selected is None:
                raise StorageError('EVIDENCE_MISSING', 'XLSX Sheet 不存在。')
            return selected
    except StorageError:
        raise
    except (BadZipFile, KeyError, ValueError, TypeError, IndexError, OverflowError, InvalidOperation, ET.ParseError):
        raise StorageError('INPUT_UNAVAILABLE', 'XLSX 结构或内容无法读取；请提供有效 XLSX 或可读文本。') from None
