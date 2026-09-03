from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table


SUPPORT = Path(__file__).parent / "support"
if str(SUPPORT) not in sys.path:
    sys.path.insert(0, str(SUPPORT))

from compare_e2e import (  # noqa: E402
    compare_runs,
    delta_evidence_markdown,
    defect_dispositions,
    derive_manifest_checks,
    derive_workflow_evidence,
    load_summary,
    reported_effort,
    workbook_effort,
)


def write_summary(root: Path, value: dict[str, object]) -> None:
    root.mkdir(parents=True)
    (root / "e2e-run-summary.json").write_text(
        json.dumps(value, ensure_ascii=False), encoding="utf-8"
    )


def test_comparison_separates_expected_template_changes_from_regressions(
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline"
    current = tmp_path / "current"
    input_hashes = {"request.json": "a" * 64, "inputs/prd.md": "b" * 64}
    write_summary(
        baseline,
        {
            "inputSha256": input_hashes,
            "workflow": {"published": 2, "blocked": 1, "reused": 1},
            "objects": {"features": 9, "stories": 8, "tasks": 41},
            "granularity": {
                "maxFeatureLinksPerStory": 5,
                "maxTasksPerStory": 15,
                "maxStoryDays": 39,
                "averageStoryDays": 14.62,
            },
            "effort": {"directDays": 117, "sitDays": 0, "uatDays": 4, "totalDays": 121},
            "workbook": {"sheetCount": 12, "tableCount": 12, "formulaErrors": 0},
            "timingSeconds": 3269,
            "checks": {},
        },
    )
    write_summary(
        current,
        {
            "inputSha256": input_hashes,
            "workflow": {"published": 2, "blocked": 1, "reused": 1},
            "objects": {"features": 9, "stories": 8, "tasks": 41},
            "granularity": {
                "maxFeatureLinksPerStory": 1,
                "maxTasksPerStory": 4,
                "maxStoryDays": 9,
                "averageStoryDays": 6.58,
            },
            "effort": {"directDays": 117, "sitDays": 11, "uatDays": 4, "totalDays": 132},
            "workbook": {
                "sheetCount": 4,
                "tableCount": 5,
                "formulaErrors": 0,
                "packagingInvariant": True,
            },
            "timingSeconds": 120,
            "checks": {
                "emptyEstimateRejected": True,
                "realOfficeRoundtrip": True,
                "formulaAuthorityReread": True,
                "verifiedTrustBoundary": True,
                "lastKnownGood": True,
                "exactReplay": True,
                "earlyScopeConflict": True,
                "reviewSubjectInventory": True,
                "unambiguousCounts": True,
                "postRenderSemanticAudit": True,
                "packagingInvariant": True,
            },
        },
    )

    result = compare_runs(baseline, current)

    assert result["inputs"]["sameContent"] is True
    assert result["workbook"]["sheets"] == {"before": 12, "after": 4}
    assert result["granularity"]["changed"] == [
        "averageStoryDays",
        "maxFeatureLinksPerStory",
        "maxStoryDays",
        "maxTasksPerStory",
    ]
    assert result["defects"]["D1"]["status"] == "COVERED_BY_TEMPLATE"
    assert result["defects"]["D3"]["status"] == "FIXED"
    assert result["regressions"] == []


def test_comparison_reports_changed_input_and_failed_control(tmp_path: Path) -> None:
    baseline = tmp_path / "baseline"
    current = tmp_path / "current"
    common = {
        "workflow": {},
        "objects": {},
        "effort": {},
        "workbook": {
            "sheetCount": 4,
            "tableCount": 5,
            "formulaErrors": 0,
            "packagingInvariant": True,
        },
        "timingSeconds": 1,
    }
    write_summary(
        baseline,
        {**common, "inputSha256": {"request.json": "a" * 64}, "checks": {}},
    )
    write_summary(
        current,
        {
            **common,
            "inputSha256": {"request.json": "b" * 64},
            "checks": {
                "emptyEstimateRejected": False,
                "realOfficeRoundtrip": True,
                "formulaAuthorityReread": True,
                "verifiedTrustBoundary": True,
                "lastKnownGood": True,
                "exactReplay": True,
                "earlyScopeConflict": True,
                "reviewSubjectInventory": True,
                "unambiguousCounts": True,
                "postRenderSemanticAudit": True,
            },
        },
    )

    result = compare_runs(baseline, current)

    assert result["inputs"]["sameContent"] is False
    assert result["inputs"]["changed"] == ["request.json"]
    assert "D3" in result["regressions"]


def test_reported_effort_reads_legacy_e2e_summary_sentence() -> None:
    report = (
        "可见结果为直接开发 117.0 人天、UAT 支持 4.0 人天、"
        "SIT 支持 0.0 人天、总计 121.0 人天。"
    )

    assert reported_effort(report) == {
        "directDays": 117.0,
        "sitDays": 0.0,
        "uatDays": 4.0,
        "totalDays": 121.0,
    }


def test_workbook_effort_reads_current_four_sheet_total_label(tmp_path: Path) -> None:
    path = tmp_path / "sow.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["工作量项", "人天"])
    worksheet.append(["直接开发人天", 116.3])
    worksheet.append(["SIT支持人天", 11])
    worksheet.append(["UAT支持人天", 4])
    worksheet.append(["总开发人天", 131.3])
    worksheet.add_table(Table(displayName="ProjectSummaryTable", ref="A1:B5"))
    workbook.save(path)
    workbook.close()

    assert workbook_effort(path) == {
        "directDays": 116.3,
        "sitDays": 11.0,
        "uatDays": 4.0,
        "totalDays": 131.3,
    }


def test_load_summary_derives_missing_story_granularity_from_generation(
    tmp_path: Path,
) -> None:
    root = tmp_path / "run"
    delivery = {
        "stories": [
            {"storyId": "story-1", "featureIds": ["f1", "f2"]},
            {"storyId": "story-2", "featureId": "f3"},
        ],
        "tasks": [
            {"storyId": "story-1"},
            {"storyId": "story-1"},
            {"storyId": "story-2"},
        ],
    }
    write_summary(
        root,
        {
            "delivery": {"stories": [], "tasks": []},
            "granularity": {"maxTasksPerStory": 99},
            "workflow": {"published": 99, "blocked": 99, "reused": 99},
        },
    )
    generation = root / ".ai-sow/generations/000001"
    generation.mkdir(parents=True)
    write_json = lambda path, value: path.write_text(  # noqa: E731
        json.dumps(value), encoding="utf-8"
    )
    write_json(root / ".ai-sow/current.json", {"generationId": "000001"})
    write_json(generation / "manifest.json", {})
    data = generation / "data"
    data.mkdir()
    write_json(
        data / "scope.json",
        {"features": [], "integrations": [], "nfrs": []},
    )
    write_json(data / "delivery.json", delivery)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["故事人天"])
    worksheet.append([12.5])
    worksheet.append([5.5])
    worksheet.add_table(Table(displayName="SOWStoryTable", ref="A1:A3"))
    output = generation / "output"
    output.mkdir()
    workbook.save(output / "sow.xlsx")
    workbook.close()

    assert load_summary(root)["granularity"] == {
        "maxFeatureLinksPerStory": 2,
        "maxTasksPerStory": 2,
        "maxStoryDays": 12.5,
        "averageStoryDays": 9.0,
    }
    assert load_summary(root)["workflow"] == {
        "published": 1,
        "blocked": 0,
        "reused": 0,
    }


def test_blocked_and_replay_workflow_require_hash_bound_contract_evidence(
    tmp_path: Path,
) -> None:
    root = tmp_path / "run"
    evidence = root / "e2e-evidence"
    evidence.mkdir(parents=True)
    packet = {
        "runId": "run-000001-000001",
        "inputRevisionId": "000001",
        "artifacts": {
            "scope": {"sha256": "a" * 64},
            "delivery": {"sha256": "b" * 64},
        },
        "bundles": {
            "scope": {
                "features": [{"featureId": "feature-return"}],
            },
            "delivery": {"stories": []},
        },
    }
    packet_bytes = (
        json.dumps(packet, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    import hashlib

    packet_sha = hashlib.sha256(packet_bytes).hexdigest()
    (evidence / "blocked-review-packet.json").write_bytes(packet_bytes)
    invalid_review = {
        "contract": "ai-sow-final-review-v1",
        "runId": "run-000001-000001",
        "inputRevisionId": "000001",
        "outcome": "BLOCKED",
        "questions": ["本期是否包含换货？"],
    }
    (evidence / "blocked-final-review.json").write_text(
        json.dumps(invalid_review, ensure_ascii=False), encoding="utf-8"
    )
    (evidence / "blocked-result.json").write_text(
        json.dumps({"outcome": "BLOCKED", "diagnostics": []}), encoding="utf-8"
    )
    (evidence / "replay-result.json").write_text(
        json.dumps(
            {
                "outcome": "REUSED",
                "diagnostics": [],
                "generationId": "000001",
                "revisionId": "000001",
                "workbookPath": ".ai-sow/generations/000001/output/sow.xlsx",
            }
        ),
        encoding="utf-8",
    )
    manifest = {
        "generationId": "000001",
        "revisionId": "000001",
        "workbookPath": ".ai-sow/generations/000001/output/sow.xlsx",
    }

    assert derive_workflow_evidence(root, manifest) == {
        "published": 0,
        "blocked": 0,
        "reused": 0,
    }

    review = {
        "contract": "ai-sow-final-review-v1",
        "runId": "run-000001-000001",
        "inputRevisionId": "000001",
        "scopeSha256": "a" * 64,
        "deliverySha256": "b" * 64,
        "packetSha256": packet_sha,
        "decision": "BLOCKED",
        "notes": [],
        "questions": [
            {
                "blockingConditionId": "block-exchange-scope",
                "subjectIds": ["feature-return"],
                "summary": "换货范围影响交付边界。",
                "question": "请确认本期是否包含换货处理。",
            }
        ],
    }
    review_bytes = (
        json.dumps(review, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    (evidence / "blocked-final-review.json").write_bytes(review_bytes)
    (evidence / "blocked-result.json").write_text(
        json.dumps(
            {
                "outcome": "BLOCKED",
                "diagnostics": [],
                "reviewSha256": hashlib.sha256(review_bytes).hexdigest(),
                "questions": ["请确认本期是否包含换货处理。"],
            }
        ),
        encoding="utf-8",
    )

    assert derive_workflow_evidence(root, manifest)["blocked"] == 0

    packet["contract"] = "ai-sow-final-review-packet-v1"
    packet_bytes = (
        json.dumps(packet, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    (evidence / "blocked-review-packet.json").write_bytes(packet_bytes)
    review["packetSha256"] = hashlib.sha256(packet_bytes).hexdigest()
    review_bytes = (
        json.dumps(review, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    (evidence / "blocked-final-review.json").write_bytes(review_bytes)
    (evidence / "blocked-result.json").write_text(
        json.dumps(
            {
                "outcome": "BLOCKED",
                "diagnostics": [],
                "reviewSha256": hashlib.sha256(review_bytes).hexdigest(),
                "questions": ["请确认本期是否包含换货处理。"],
            }
        ),
        encoding="utf-8",
    )

    assert derive_workflow_evidence(root, manifest)["blocked"] == 1

    review["notes"] = ["不是合法 review note"]
    invalid_note_bytes = (
        json.dumps(review, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    (evidence / "blocked-final-review.json").write_bytes(invalid_note_bytes)
    (evidence / "blocked-result.json").write_text(
        json.dumps(
            {
                "outcome": "BLOCKED",
                "diagnostics": [],
                "reviewSha256": hashlib.sha256(invalid_note_bytes).hexdigest(),
                "questions": ["请确认本期是否包含换货处理。"],
            }
        ),
        encoding="utf-8",
    )

    assert derive_workflow_evidence(root, manifest)["blocked"] == 0


def test_replay_requires_input_result_current_and_manifest_hash_closure(
    tmp_path: Path,
) -> None:
    import hashlib

    root = tmp_path / "run"
    evidence = root / "e2e-evidence"
    generation = root / ".ai-sow/generations/000001"
    evidence.mkdir(parents=True)
    generation.mkdir(parents=True)
    (root / "request.json").write_text("{}\n", encoding="utf-8")
    manifest = {
        "generationId": "000001",
        "revisionId": "000001",
        "workbookPath": ".ai-sow/generations/000001/output/sow.xlsx",
    }
    manifest_bytes = (
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    (generation / "manifest.json").write_bytes(manifest_bytes)
    current = {
        "generationId": "000001",
        "revisionId": "000001",
        "generationManifestPath": ".ai-sow/generations/000001/manifest.json",
        "generationManifestSha256": hashlib.sha256(manifest_bytes).hexdigest(),
    }
    current_bytes = (
        json.dumps(current, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    (root / ".ai-sow/current.json").write_bytes(current_bytes)
    replay = {
        "outcome": "REUSED",
        "diagnostics": [],
        "generationId": "000001",
        "revisionId": "000001",
        "workbookPath": ".ai-sow/generations/000001/output/sow.xlsx",
    }
    replay_bytes = (
        json.dumps(replay, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")
    (evidence / "replay-result.json").write_bytes(replay_bytes)

    assert derive_workflow_evidence(root, manifest)["reused"] == 0

    replay_evidence = {
        "contract": "ai-sow-e2e-replay-evidence-v1",
        "inputSha256": {
            "request.json": hashlib.sha256(b"{}\n").hexdigest(),
        },
        "resultPath": "e2e-evidence/replay-result.json",
        "resultSha256": hashlib.sha256(replay_bytes).hexdigest(),
        "currentPath": ".ai-sow/current.json",
        "currentSha256": hashlib.sha256(current_bytes).hexdigest(),
        "generationManifestPath": ".ai-sow/generations/000001/manifest.json",
        "generationManifestSha256": hashlib.sha256(manifest_bytes).hexdigest(),
    }
    (evidence / "replay-evidence.json").write_text(
        json.dumps(replay_evidence, ensure_ascii=False), encoding="utf-8"
    )

    assert derive_workflow_evidence(root, manifest)["reused"] == 1

    replay_evidence["inputSha256"]["request.json"] = "0" * 64
    (evidence / "replay-evidence.json").write_text(
        json.dumps(replay_evidence, ensure_ascii=False), encoding="utf-8"
    )

    assert derive_workflow_evidence(root, manifest)["reused"] == 0


def test_unambiguous_counts_checks_all_five_manifest_count_dimensions() -> None:
    scope = {"features": [{"featureId": "feature-return"}]}
    delivery = {
        "stories": [{"storyId": "story-return"}],
        "acceptanceCriteria": [{"acceptanceCriterionId": "ac-return"}],
        "tasks": [{"taskId": "task-return"}],
    }
    correct = {
        collection: {
            "affected": 0,
            "recomputed": 1,
            "reused": 0,
            "deleted": 0,
            "final": 1,
        }
        for collection in (
            "features",
            "stories",
            "acceptanceCriteria",
            "tasks",
        )
    }
    manifest = {
        "changeCounts": json.loads(json.dumps(correct)),
        "impact": {"baselineGenerationId": None},
    }
    manifest["changeCounts"]["tasks"]["reused"] = 1

    checks = derive_manifest_checks(
        manifest,
        scope,
        delivery,
        {},
        {},
        scope_sha256="a" * 64,
        delivery_sha256="b" * 64,
        workbook_sha256="c" * 64,
        expected_change_counts=correct,
    )

    assert checks["unambiguousCounts"] is False


def test_comparison_reports_task_and_effort_changes_without_treating_baseline_as_truth(
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline"
    current = tmp_path / "current"
    common = {
        "inputSha256": {"request.json": "a" * 64},
        "workflow": {},
        "objects": {},
        "effort": {
            "directDays": 8,
            "sitDays": 1,
            "uatDays": 0.5,
            "totalDays": 9.5,
        },
        "workbook": {"sheetCount": 4, "tableCount": 5, "formulaErrors": 0},
        "timingSeconds": 1,
        "checks": {
            "emptyEstimateRejected": True,
            "realOfficeRoundtrip": True,
            "formulaAuthorityReread": True,
            "verifiedTrustBoundary": True,
            "lastKnownGood": True,
            "exactReplay": True,
            "earlyScopeConflict": True,
            "reviewSubjectInventory": True,
            "unambiguousCounts": True,
            "postRenderSemanticAudit": True,
        },
    }
    story = {"storyId": "story-1", "name": "退款"}
    task = {
        "storyId": "story-1",
        "name": "实现退款",
        "workMode": "新建",
        "complexity": "M",
    }
    write_summary(
        baseline,
        {
            **common,
            "delivery": {
                "stories": [story],
                "tasks": [{**task, "baseUnit": "BU-API"}],
            },
        },
    )
    write_summary(
        current,
        {
            **common,
            "effort": {
                "directDays": 10,
                "sitDays": 1,
                "uatDays": 0.5,
                "totalDays": 11.5,
            },
            "delivery": {
                "stories": [story],
                "tasks": [{**task, "baseUnit": "BU-INTEGRATION"}],
            },
        },
    )

    result = compare_runs(baseline, current)

    assert result["taskChanges"]["removed"] == [
        ["退款", "实现退款", "BU-API", "新建", "M"]
    ]
    assert result["taskChanges"]["added"] == [
        ["退款", "实现退款", "BU-INTEGRATION", "新建", "M"]
    ]
    assert result["taskSemantics"]["removed"] == [
        ["实现退款", "BU-API", "新建", "M"]
    ]
    assert result["taskSemantics"]["added"] == [
        ["实现退款", "BU-INTEGRATION", "新建", "M"]
    ]
    assert result["effort"]["changed"] == ["directDays", "totalDays"]
    assert "taskSemanticsUnchanged" not in result["acceptance"]
    assert "TASK_SEMANTICS_CHANGED" not in result["regressions"]


def test_comparison_itemizes_five_level_semantic_deltas(tmp_path: Path) -> None:
    baseline = tmp_path / "baseline"
    current = tmp_path / "current"
    common = {
        "inputSha256": {"request.json": "a" * 64},
        "workflow": {},
        "objects": {},
        "granularity": {
            "maxFeatureLinksPerStory": 1,
            "maxTasksPerStory": 1,
        },
        "effort": {},
        "workbook": {
            "sheetCount": 4,
            "tableCount": 5,
            "formulaErrors": 0,
            "packagingInvariant": True,
        },
        "timingSeconds": 1,
        "checks": {key: True for key in (
            "emptyEstimateRejected",
            "realOfficeRoundtrip",
            "formulaAuthorityReread",
            "verifiedTrustBoundary",
            "lastKnownGood",
            "exactReplay",
            "earlyScopeConflict",
            "reviewSubjectInventory",
            "unambiguousCounts",
            "postRenderSemanticAudit",
        )},
    }
    write_summary(
        baseline,
        {
            **common,
            "scope": {
                "epics": [{
                    "epicId": "epic-old",
                    "name": "退款中心",
                    "sourceRefs": [{
                        "sourceId": "prd-refund",
                        "anchorId": "anchor-old-epic",
                        "locator": "section:退款中心",
                    }],
                }],
                "features": [{
                    "featureId": "feature-old",
                    "epicId": "epic-old",
                    "name": "退款受理",
                    "sourceRefs": [{
                        "sourceId": "prd-refund",
                        "anchorId": "anchor-old-feature",
                        "locator": "section:退款受理",
                    }],
                }],
            },
            "delivery": {
                "stories": [{
                    "storyId": "story-old",
                    "featureId": "feature-old",
                    "name": "[退款] 客户提交申请",
                }],
                "acceptanceCriteria": [{
                    "acceptanceCriterionId": "ac-old",
                    "storyId": "story-old",
                    "name": "有效申请进入待审核",
                }],
                "tasks": [{
                    "taskId": "task-old",
                    "storyId": "story-old",
                    "name": "实现退款接口",
                    "baseUnit": "BU-API",
                    "workMode": "新建",
                    "complexity": "M",
                }],
            },
        },
    )
    write_summary(
        current,
        {
            **common,
            "scope": {
                "epics": [{
                    "epicId": "epic-new",
                    "name": "售后服务中心",
                    "sourceRefs": [{
                        "sourceId": "prd-return",
                        "anchorId": "anchor-new-epic",
                        "locator": "section:售后服务中心",
                    }],
                }],
                "features": [{
                    "featureId": "feature-new",
                    "epicId": "epic-new",
                    "name": "退款申请处理",
                    "scopeDecision": {"decision": "IN_SCOPE"},
                    "sourceRefs": [{
                        "sourceId": "prd-return",
                        "anchorId": "anchor-new-feature",
                        "locator": "section:退款申请处理",
                    }],
                }],
            },
            "delivery": {
                "stories": [{
                    "storyId": "story-new",
                    "featureId": "feature-new",
                    "name": "[退款申请] 客户提交退款",
                }],
                "acceptanceCriteria": [{
                    "acceptanceCriterionId": "ac-new",
                    "storyId": "story-new",
                    "name": "有效申请受理并返回单号",
                }],
                "tasks": [{
                    "taskId": "task-new",
                    "storyId": "story-new",
                    "name": "开发退款申请接口",
                    "baseUnit": "BU-INTEGRATION",
                    "workMode": "新建",
                    "complexity": "M",
                }],
            },
        },
    )

    result = compare_runs(baseline, current)

    assert result["semanticChanges"]["requirements"]["removed"] == [
        {"requirement": "退款中心"}
    ]
    assert result["semanticChanges"]["subRequirements"]["added"] == [{
        "requirement": "售后服务中心",
        "subRequirement": "退款申请处理",
        "scopeDecision": "IN_SCOPE",
    }]
    assert result["semanticChanges"]["stories"]["added"] == [{
        "subRequirement": "退款申请处理",
        "story": "[退款申请] 客户提交退款",
    }]
    assert result["semanticChanges"]["acceptanceCriteria"]["added"] == [{
        "story": "[退款申请] 客户提交退款",
        "acceptanceCriterion": "有效申请受理并返回单号",
    }]
    assert result["semanticChanges"]["tasks"]["added"] == [{
        "story": "[退款申请] 客户提交退款",
        "task": "开发退款申请接口",
        "baseUnit": "BU-INTEGRATION",
        "workMode": "新建",
        "complexity": "M",
    }]
    assert len(result["semanticEvidence"]) == 10
    assert {
        (item["collection"], item["change"])
        for item in result["semanticEvidence"]
    } == {
        (collection, change)
        for collection in (
            "requirements",
            "subRequirements",
            "stories",
            "acceptanceCriteria",
            "tasks",
        )
        for change in ("REMOVED", "ADDED")
    }
    assert all(item["objectId"] for item in result["semanticEvidence"])
    assert all(
        item["sourceRefs"] or item.get("baseUnit")
        for item in result["semanticEvidence"]
    )
    rationale = delta_evidence_markdown(result)
    assert "Story 人天不作为拆分正确性的完成门禁" in rationale
    assert all(item["objectId"] in rationale for item in result["semanticEvidence"])
    assert "prd-return#anchor-new-feature" in rationale
    assert "BU-INTEGRATION" in rationale
    assert all(
        item["classification"]
        in {
            "AUTHORING_STANDARD",
            "DEFECT_FIX",
            "EVIDENCE_REINTERPRETATION",
            "TEMPLATE_PROJECTION",
            "UNCHANGED",
        }
        for item in result["semanticEvidence"]
    )


def test_comparison_fails_story_and_packaging_acceptance_gates(tmp_path: Path) -> None:
    baseline = tmp_path / "baseline"
    current = tmp_path / "current"
    common = {
        "inputSha256": {"request.json": "a" * 64},
        "workflow": {},
        "objects": {},
        "effort": {},
        "workbook": {"sheetCount": 4, "tableCount": 5, "formulaErrors": 0},
        "timingSeconds": 1,
        "delivery": {"stories": [], "tasks": []},
        "checks": {
            "emptyEstimateRejected": True,
            "realOfficeRoundtrip": True,
            "formulaAuthorityReread": True,
            "verifiedTrustBoundary": True,
            "lastKnownGood": True,
            "exactReplay": True,
            "earlyScopeConflict": True,
            "reviewSubjectInventory": True,
            "unambiguousCounts": True,
            "postRenderSemanticAudit": True,
        },
    }
    write_summary(
        baseline,
        {
            **common,
            "granularity": {
                "maxFeatureLinksPerStory": 1,
                "maxTasksPerStory": 4,
                "maxStoryDays": 9,
                "averageStoryDays": 6,
            },
        },
    )
    write_summary(
        current,
        {
            **common,
            "granularity": {
                "maxFeatureLinksPerStory": 2,
                "maxTasksPerStory": 5,
                "maxStoryDays": 12,
                "averageStoryDays": 8,
            },
            "checks": {**common["checks"], "packagingInvariant": False},
        },
    )

    result = compare_runs(baseline, current)

    assert {
        "STORY_FEATURE_CARDINALITY_EXCEEDED",
        "STORY_TASK_LIMIT_EXCEEDED",
        "PACKAGING_INVARIANCE_UNVERIFIED",
    } <= set(result["regressions"])


def test_comparison_does_not_trust_self_reported_packaging_invariance(
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline"
    current = tmp_path / "current"
    common = {
        "inputSha256": {"request.json": "a" * 64},
        "workflow": {},
        "objects": {},
        "granularity": {
            "maxFeatureLinksPerStory": 1,
            "maxTasksPerStory": 4,
        },
        "effort": {},
        "delivery": {"stories": [], "tasks": []},
        "timingSeconds": 1,
        "checks": {"packagingInvariant": True},
    }
    write_summary(
        baseline,
        {
            **common,
            "workbook": {
                "sheetCount": 4,
                "tableCount": 5,
                "formulaErrors": 0,
                "packagingInvariant": True,
            },
        },
    )
    write_summary(
        current,
        {
            **common,
            "workbook": {
                "sheetCount": 4,
                "tableCount": 5,
                "formulaErrors": 0,
                "packagingInvariant": False,
            },
        },
    )

    result = compare_runs(baseline, current)

    assert result["acceptance"]["packagingInvariant"] is False
    assert "PACKAGING_INVARIANCE_UNVERIFIED" in result["regressions"]


def test_defect_disposition_prefers_derived_evidence_over_reported_boolean() -> None:
    current = {
        "workbook": {"sheetCount": 4, "tableCount": 5},
        "effort": {"sitDays": 1},
        "workflow": {"published": 1, "blocked": 1, "reused": 1},
        "checks": {
            "emptyEstimateRejected": True,
            "realOfficeRoundtrip": True,
            "formulaAuthorityReread": True,
            "verifiedTrustBoundary": True,
            "lastKnownGood": True,
            "exactReplay": True,
            "earlyScopeConflict": True,
            "reviewSubjectInventory": True,
            "unambiguousCounts": True,
            "postRenderSemanticAudit": True,
        },
        "derivedChecks": {
            "realOfficeRoundtrip": False,
            "formulaAuthorityReread": False,
            "verifiedTrustBoundary": False,
            "reviewSubjectInventory": False,
            "unambiguousCounts": False,
            "postRenderSemanticAudit": False,
        },
    }

    defects = defect_dispositions(current)

    for defect in ("D4", "D5", "D6", "D10", "D11", "D12"):
        assert defects[defect]["status"] == "REGRESSED"
    assert defects["D8"]["status"] == "FIXED"
    assert defects["D9"]["status"] == "FIXED"
