"""Reproducible project-template variants for projection/Office compatibility tests."""
from copy import copy
from pathlib import Path
import re

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.formula.translate import Translator
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import TableFormula


def compact_task_template(source: Path, output: Path, *, story_capacity: int = 1, task_capacity: int = 1) -> Path:
    """Adapt a name-entry template; financial prototypes still come from source."""
    book = openpyxl.load_workbook(source)
    sheet = book['02-任务清单']
    assert sheet['H4'].value == 'SIT计费点ID'
    dimensions = {key: copy(value) for key, value in sheet.column_dimensions.items()}
    for merged in list(sheet.merged_cells):
        sheet.unmerge_cells(str(merged))
        sheet.merge_cells(start_row=merged.min_row, end_row=merged.max_row,
                          start_column=merged.min_col, end_column=merged.max_col - 1)
    sheet.delete_cols(8)
    sheet.column_dimensions.clear()
    for key, dimension in dimensions.items():
        old = column_index_from_string(key)
        if old == 8:
            continue
        new = old - (old > 8)
        dimension.index = get_column_letter(new)
        dimension.min = dimension.max = new
        sheet.column_dimensions[dimension.index] = dimension
    sheet['G4'] = '集成类型'
    sheet['A2'] = ('工作类型按名称选择，任务名称全表唯一。仅集成类任务选择内部集成或外部集成，'
                   '其他任务留空；SIT 按该集成任务计取。灰色计算列由公式维护。')
    for row in sheet:
        for cell in row:
            if cell.data_type == 'f':
                cell.value = re.sub(r'(?<![A-Z0-9_])(\$?)([I-N])(\$?\d+)',
                    lambda m: m[1] + chr(ord(m[2]) - 1) + m[3], cell.value)
                cell.value = cell.value.replace('"NONE"', '""').replace(
                    '"INTERNAL"', '"内部集成"').replace('"EXTERNAL"', '"外部集成"')
    eligible = ('IFERROR(INDEX(TaskStandardTable[SIT支持资格],'
                'MATCH($C5,TaskStandardTable[工作类型ID],0)),"")="PER_INTEGRATION"')
    # Equality comparison treats wildcard-looking names as literal task names.
    checks = [
        ('COUNTA($A5:$B5,$D5:$H5)=0', '""'),
        ('OR($A5="",$B5="",$C5="",$E5="",$F5="")', '"必填项缺失"'),
        ('SUMPRODUCT(--(TaskTable[任务名称]=$B5))>1', '"任务名称重复"'),
        ('COUNTIF(SOWStoryTable[故事],$A5)<>1', '"所属故事未知"'),
        ('NOT(OR($F5="S",$F5="M",$F5="L"))', '"复杂度非法"'),
        ('NOT(OR($G5="",$G5="内部集成",$G5="外部集成"))', '"集成类型非法"'),
        (f'AND($G5<>"",NOT({eligible}))', '"非集成任务不应填写集成类型"'),
        (f'AND($G5="",{eligible})', '"集成类型缺失"'),
    ]
    validation = 'IF(ISNUMBER($I5),"通过","工作类型与工作方式不适用")'
    for condition, result in reversed(checks):
        validation = f'IF({condition},{result},{validation})'
    sheet['M5'] = '=' + validation
    sheet['L5'] = '=IF(COUNTA($A5:$B5,$D5:$H5)=0,"",' + sheet['L5'].value[1:] + ')'
    table = sheet.tables['TaskTable']
    table.ref = 'A4:M5'
    table.autoFilter.ref = table.ref
    table.tableColumns = [column for column in table.tableColumns if column.name != 'SIT计费点ID']
    for index, column in enumerate(table.tableColumns, 1):
        column.id = index
        column.name = sheet.cell(4, index).value
        cell = sheet.cell(5, index)
        column.calculatedColumnFormula = (TableFormula(attr_text=cell.value[1:])
            if cell.data_type == 'f' else None)
    # A pre-existing validation highlight was attached to the coefficient column.
    rules = [(str(region.sqref), list(sheet.conditional_formatting[region]))
             for region in sheet.conditional_formatting]
    sheet.conditional_formatting._cf_rules.clear()
    for region, entries in rules:
        for rule in entries:
            rule = copy(rule)
            rule.formula = [formula.replace('$K5', '$M5') for formula in rule.formula]
            sheet.conditional_formatting.add('M5:M1048576' if region == 'K5:K1048576' else region, rule)
    sheet['J5'].fill = copy(sheet['I5'].fill)
    standard = book['90-估算标准']
    option_column = get_column_letter(standard.max_column + 2)
    for row, value in ((5, '内部集成'), (6, '外部集成')):
        standard[f'{option_column}{row}'] = value
    standard.column_dimensions[option_column].hidden = True
    for name, rows in (('AiSowIntegrationTypes', '$5:$6'), ('AiSowNoIntegrationType', '$7:$7')):
        first, last = rows.split(':')
        book.defined_names.add(DefinedName(name, attr_text=
            f"'90-估算标准'!${option_column}{first}:${option_column}{last}"))
    for rule in sheet.data_validations.dataValidation:
        rule.errorStyle = 'stop'
        rule.showErrorMessage = True
        if str(rule.sqref) == 'G5:G1048576':
            rule.formula1 = ('IF(IFERROR(INDEX(INDIRECT("TaskStandardTable[SIT支持资格]"),'
                'MATCH($C5,INDIRECT("TaskStandardTable[工作类型ID]"),0)),"")="PER_INTEGRATION",'
                'AiSowIntegrationTypes,AiSowNoIntegrationType)')
            rule.promptTitle = '集成类型'
            rule.prompt = '仅集成类工作类型可选内部集成或外部集成；其他任务留空。'
            rule.errorTitle = '集成类型不适用'
            rule.error = '请先选择集成类工作类型，或清空本列。'
            rule.showInputMessage = True
    unique = DataValidation(type='custom', allow_blank=True, errorStyle='stop', showErrorMessage=True,
        formula1='COUNTIF($B:$B,"="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(B5,"~","~~"),"*","~*"),"?","~?"))=1')
    unique.errorTitle = '任务名称重复'
    unique.error = '任务名称必须全表唯一，请补充业务对象或集成方向。'
    unique.add('B5:B1048576')
    sheet.add_data_validation(unique)
    story_sheet = book['01-需求故事']
    story_unique = DataValidation(type='custom', allow_blank=True, errorStyle='stop', showErrorMessage=True,
        formula1='COUNTIF($C:$C,"="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(C5,"~","~~"),"*","~*"),"?","~?"))=1')
    story_unique.errorTitle = '故事名称重复'
    story_unique.error = '故事名称必须全表唯一，任务通过故事名称关联。'
    story_unique.add('C5:C1048576')
    story_sheet.add_data_validation(story_unique)
    for target, table_name, capacity in ((story_sheet, 'SOWStoryTable', story_capacity),
                                         (sheet, 'TaskTable', task_capacity)):
        target_table = target.tables[table_name]
        for row in range(6, 5 + capacity):
            target.row_dimensions[row].height = target.row_dimensions[5].height
            for prototype in target[5]:
                cell = target.cell(row, prototype.column)
                cell._style = copy(prototype._style)
                if prototype.data_type == 'f':
                    is_array = isinstance(prototype.value, ArrayFormula)
                    formula = prototype.value.text if is_array else prototype.value
                    translated = Translator(formula, origin=prototype.coordinate).translate_formula(cell.coordinate)
                    cell.value = ArrayFormula(ref=cell.coordinate, text=translated) if is_array else translated
        target_table.ref = f'A4:{get_column_letter(len(target_table.tableColumns))}{4 + capacity}'
        target_table.autoFilter.ref = target_table.ref
    for row in book['03-工作量汇总']:
        for cell in row:
            if cell.data_type == 'f':
                cell.value = cell.value.replace('TaskTable[SIT支持分类]', 'TaskTable[集成类型]').replace(
                    '"INTERNAL"', '"内部集成"').replace('"EXTERNAL"', '"外部集成"')
    book.save(output)
    book.close()
    return output
