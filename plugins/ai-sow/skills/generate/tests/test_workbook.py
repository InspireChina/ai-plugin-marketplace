"""Task12 projection and integrity boundaries; no full host setup."""
import json
import sys
from pathlib import Path
import openpyxl
import pytest

TEST_LAYER = 'unit'
ROOT = Path(__file__).parents[1]
LEGACY_TEMPLATE = ROOT / 'tests/fixtures/sow-template-legacy.xlsx'
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'tests'))
import workbook
from contracts import canonical_json_bytes, sha256_bytes


def render_model():
    from ir_samples import task_story_model
    from task_standard_catalog import catalog
    model = task_story_model()
    source = catalog(ROOT / 'assets/sow-template.xlsx')
    model['project']['templateSha256'] = source.template_sha256
    model['tasks'] = [{'taskId':'task-query','storyId':'story-query','name':'订单查询页面',
        'workTypeId':'FE-VIEW','rowSemanticSha256':source.by_work_type_id['FE-VIEW']['rowSemanticSha256'],
        'workMode':'新建','complexity':'M','actualMeasurementScope':'一张订单查询页面',
        'sourceRefs':model['stories'][0]['sourceRefs'],'integrationIds':[]}]
    return model


def name_entry_template(tmp_path):
    """Project-template fixture: the name is input and the stable key is calculated."""
    from copy import copy
    path = tmp_path / 'name-entry-template.xlsx'
    book = openpyxl.load_workbook(LEGACY_TEMPLATE)
    sheet = book['02-任务清单']
    input_style = copy(sheet['C5']._style)
    sheet['C5'].value = sheet['D5'].value.replace('工作类型名称', '__NAME__').replace(
        '工作类型ID', '工作类型名称').replace('__NAME__', '工作类型ID').replace('$C5', '$D5')
    sheet['C5']._style = copy(sheet['D5']._style)
    sheet['D5'].value = None
    sheet['D5']._style = input_style
    sheet['N5'].value = sheet['N5'].value.replace(
        'COUNTA($A5:$C5,$E5:$I5)', 'COUNTA($A5:$B5,$D5:$I5)')
    sheet.column_dimensions['C'].hidden = True
    sheet.column_dimensions['H'].hidden = True
    for rule in sheet.data_validations.dataValidation:
        if rule.formula1 == '=INDIRECT("TaskStandardTable[工作类型ID]")':
            rule.formula1 = '=INDIRECT("TaskStandardTable[工作类型名称]")'
            rule.sqref = 'D5:D1048576'
    book.save(path)
    book.close()
    return path


def test_name_entry_template_projects_editable_names_and_keeps_hidden_keys(tmp_path):
    model = render_model()
    template = name_entry_template(tmp_path)
    path = tmp_path / 'name-entry.xlsx'
    workbook.write_workbook(template, model, path)
    book = openpyxl.load_workbook(path)
    try:
        sheet = book['02-任务清单']
        assert sheet['D5'].value == '信息展示与查询页面'
        assert sheet['D5'].data_type == 's' and not sheet['D5'].protection.locked
        assert sheet['C5'].data_type == 'f' and sheet['C5'].protection.locked
        assert sheet.column_dimensions['C'].hidden and sheet.column_dimensions['H'].hidden
        assert any(str(rule.sqref) == 'D5:D1048576' and '工作类型名称' in rule.formula1
                   for rule in sheet.data_validations.dataValidation)
        columns = {col.name: col for col in sheet.tables['TaskTable'].tableColumns}
        assert columns['工作类型ID'].calculatedColumnFormula is not None
        assert columns['工作类型名称'].calculatedColumnFormula is None
    finally:
        book.close()


@pytest.mark.parametrize('broken', ['both', 'neither', 'missing_effort'])
def test_name_entry_does_not_relax_required_formula_roles(tmp_path, broken):
    template = name_entry_template(tmp_path)
    book = openpyxl.load_workbook(template)
    if broken == 'both':
        book['02-任务清单']['D5'] = '=C5'
    elif broken == 'neither':
        book['02-任务清单']['C5'] = None
    else:
        book['02-任务清单']['L5'] = None
    book.save(template)
    book.close()
    with pytest.raises(ValueError, match='formula prototype mismatch'):
        workbook.write_workbook(template, render_model(), tmp_path / 'invalid.xlsx')


@pytest.mark.e2e
def test_name_selection_office_recalculates_key_effort_and_blank_validation(tmp_path):
    from office_engine import recalculate_workbook, require_office_engine
    from task_standard_catalog import catalog
    model = render_model()
    template = name_entry_template(tmp_path)
    path = tmp_path / 'selected.xlsx'
    workbook.write_workbook(template, model, path)
    # Exercise the actual user input, without editing a hidden ID or copying a result.
    book = openpyxl.load_workbook(path)
    book['02-任务清单']['D5'] = '业务数据查询 API'
    book.save(path)
    book.close()
    engine = require_office_engine()
    selected = tmp_path / 'selected-calculated.xlsx'
    recalculate_workbook(path, selected, engine)
    standard = catalog(ROOT / 'assets/sow-template.xlsx')
    model['tasks'][0].update(workTypeId='FE-QUERY-API',
        rowSemanticSha256=standard.by_work_type_id['FE-QUERY-API']['rowSemanticSha256'])
    legacy = tmp_path / 'legacy.xlsx'
    legacy_calculated = tmp_path / 'legacy-calculated.xlsx'
    workbook.write_workbook(LEGACY_TEMPLATE, model, legacy)
    recalculate_workbook(legacy, legacy_calculated, engine)
    actual = openpyxl.load_workbook(selected, data_only=True)
    expected = openpyxl.load_workbook(legacy_calculated, data_only=True)
    try:
        assert workbook.table_records(actual, 'TaskTable') == workbook.table_records(expected, 'TaskTable')
        assert actual['02-任务清单']['C5'].value == 'FE-QUERY-API'
        assert actual['02-任务清单']['N5'].value == '通过'
    finally:
        actual.close(); expected.close()
    blank = openpyxl.load_workbook(template)
    blank['02-任务清单']['D5'] = None
    blank.save(tmp_path / 'blank.xlsx'); blank.close()
    recalculate_workbook(tmp_path / 'blank.xlsx', tmp_path / 'blank-calculated.xlsx', engine)
    blank = openpyxl.load_workbook(tmp_path / 'blank-calculated.xlsx', data_only=True)
    try:
        assert blank['02-任务清单']['N5'].value in (None, '')
    finally:
        blank.close()


@pytest.mark.e2e
def test_hidden_sit_keys_still_reject_duplicates_missing_and_ineligible_types(tmp_path):
    from copy import deepcopy
    from office_engine import recalculate_workbook, require_office_engine
    from task_standard_catalog import catalog
    model = render_model()
    standard = catalog(ROOT / 'assets/sow-template.xlsx')
    prototype = model['tasks'][0]
    model['tasks'] = []
    for i in range(4):
        task = deepcopy(prototype)
        task.update(taskId=f'task-{i}', name=f'独立集成方向{i}',
            workTypeId='IN-INTEGRATION', integrationIds=[f'integration-{i}'],
            rowSemanticSha256=standard.by_work_type_id['IN-INTEGRATION']['rowSemanticSha256'])
        model['tasks'].append(task)
        model['integrations'].append({'integrationId':f'integration-{i}', 'counterpartyBoundary':'INTERNAL'})
    path = tmp_path / 'invalid-points.xlsx'
    workbook.write_workbook(name_entry_template(tmp_path), model, path)
    book = openpyxl.load_workbook(path)
    sheet = book['02-任务清单']
    sheet['H6'] = sheet['H5'].value
    sheet['H7'] = None
    sheet['D8'] = '信息展示与查询页面'
    book.save(path); book.close()
    calculated = tmp_path / 'invalid-points-calculated.xlsx'
    recalculate_workbook(path, calculated, require_office_engine())
    book = openpyxl.load_workbook(calculated, data_only=True)
    try:
        assert [book['02-任务清单'][f'N{r}'].value for r in range(5,9)] == [
            'SIT计费点重复', 'SIT计费点重复', 'SIT计费点缺失', 'SIT支持资格不适用']
    finally:
        book.close()


@pytest.mark.e2e
@pytest.mark.parametrize('entry', ['legacy', 'compact'])
def test_catalog_selector_preserves_literal_prefix_names_through_office(tmp_path, entry):
    from copy import deepcopy
    from office_engine import recalculate_workbook, require_office_engine
    from task_standard_catalog import catalog
    from template_variants import compact_task_template
    template = (compact_task_template(name_entry_template(tmp_path), tmp_path / 'compact.xlsx')
                if entry == 'compact' else LEGACY_TEMPLATE)
    book = openpyxl.load_workbook(template)
    sheet, table = workbook.table_index(book)['TaskStandardTable']
    changed = dict(zip(['FE-VIEW', 'FE-EDIT', 'FE-QUERY-API', 'FE-COMMAND-API'], '=+-@'))
    for row in range(5, sheet.max_row + 1):
        work_type = sheet.cell(row, 3).value
        if work_type in changed:
            cell = sheet.cell(row, 4)
            cell.value = changed[work_type] + cell.value
            cell.data_type = 's'
    template = tmp_path / 'literal-catalog.xlsx'
    book.save(template); book.close()
    standard = catalog(template)
    model = render_model()
    prototype = model['tasks'][0]
    model['project']['templateSha256'] = standard.template_sha256
    model['tasks'] = []
    for index, work_type in enumerate(changed):
        task = deepcopy(prototype)
        task.update(taskId=f'task-{index}', name=f'独立交付对象{index}', workTypeId=work_type,
            rowSemanticSha256=standard.by_work_type_id[work_type]['rowSemanticSha256'])
        model['tasks'].append(task)
    projected = tmp_path / 'projected.xlsx'
    calculated = tmp_path / 'calculated.xlsx'
    workbook.write_workbook(template, model, projected)
    engine = require_office_engine()
    recalculate_workbook(projected, calculated, engine)
    audit = workbook.audit_calculated_workbook(calculated, template, model, engine)
    assert audit.trust_state == 'VERIFIED'
    book = openpyxl.load_workbook(calculated, data_only=True)
    try:
        for record in workbook.table_records(book, 'TaskTable'):
            assert record['工作类型名称'] == standard.by_work_type_id[record['工作类型ID']]['工作类型名称']
            assert record['校验结果'] == '通过'
    finally:
        book.close()


@pytest.mark.e2e
def test_compact_integration_types_use_unique_tasks_and_preserve_sit_totals(tmp_path):
    from copy import copy, deepcopy
    from office_engine import recalculate_workbook, require_office_engine
    from task_standard_catalog import catalog
    from template_variants import compact_task_template
    model = render_model()
    standard = catalog(ROOT / 'assets/sow-template.xlsx')
    prototype = model['tasks'][0]
    for index, (boundary, complexity) in enumerate([('INTERNAL', 'S'), ('EXTERNAL', 'M')]):
        task = deepcopy(prototype)
        task.update(taskId=f'integration-task-{index}', name=f'集成方向{index}',
            workTypeId='IN-INTEGRATION', complexity=complexity, integrationIds=[f'integration-{index}'],
            rowSemanticSha256=standard.by_work_type_id['IN-INTEGRATION']['rowSemanticSha256'])
        model['tasks'].append(task)
        model['integrations'].append({'integrationId':f'integration-{index}', 'counterpartyBoundary':boundary})
    template = compact_task_template(name_entry_template(tmp_path), tmp_path / 'compact.xlsx')
    engine = require_office_engine()
    results = []
    for name, source in [('legacy', LEGACY_TEMPLATE), ('compact', template)]:
        path = tmp_path / f'{name}-projected.xlsx'
        calculated = tmp_path / f'{name}-calculated.xlsx'
        workbook.write_workbook(source, model, path)
        recalculate_workbook(path, calculated, engine)
        results.append(workbook.audit_calculated_workbook(calculated, source, model, engine))
    assert results[0].direct_days == results[1].direct_days
    assert results[0].sit_days == results[1].sit_days
    assert results[0].total_days == results[1].total_days
    book = openpyxl.load_workbook(tmp_path / 'compact-calculated.xlsx', data_only=True)
    try:
        records = workbook.table_records(book, 'TaskTable')
        assert 'SIT计费点ID' not in records[0] and 'SIT支持分类' not in records[0]
        assert [row['集成类型'] for row in records] == [None, '内部集成', '外部集成']
        assert all(row['校验结果'] == '通过' for row in records)
        sheet = book['02-任务清单']
        assert all('J5' not in str(region.sqref) for region in sheet.conditional_formatting)
        assert copy(sheet['J5'].fill) == copy(sheet['I5'].fill)
        rules = sheet.data_validations.dataValidation
        assert any(str(rule.sqref) == 'B5:B1048576' and rule.errorStyle == 'stop' for rule in rules)
        assert any(str(rule.sqref) == 'G5:G1048576' and 'AiSowIntegrationTypes' in rule.formula1
                   and 'PER_INTEGRATION' in rule.formula1 and rule.errorStyle == 'stop' for rule in rules)
    finally:
        book.close()
    # Pasting can bypass Excel validation: recalculated checks must still reject it.
    invalid = openpyxl.load_workbook(tmp_path / 'compact-projected.xlsx')
    sheet = invalid['02-任务清单']
    sheet['G5'] = '内部集成'
    sheet['G6'] = None
    sheet['G7'] = 'INTERNAL'
    invalid.save(tmp_path / 'invalid.xlsx'); invalid.close()
    recalculate_workbook(tmp_path / 'invalid.xlsx', tmp_path / 'invalid-calculated.xlsx', engine)
    invalid = openpyxl.load_workbook(tmp_path / 'invalid-calculated.xlsx', data_only=True)
    assert [row['校验结果'] for row in workbook.table_records(invalid, 'TaskTable')] == [
        '非集成任务不应填写集成类型', '集成类型缺失', '集成类型非法']
    invalid.close()
    invalid = openpyxl.load_workbook(tmp_path / 'compact-projected.xlsx')
    invalid['02-任务清单']['B6'] = invalid['02-任务清单']['B7'].value
    invalid.save(tmp_path / 'duplicate.xlsx'); invalid.close()
    recalculate_workbook(tmp_path / 'duplicate.xlsx', tmp_path / 'duplicate-calculated.xlsx', engine)
    invalid = openpyxl.load_workbook(tmp_path / 'duplicate-calculated.xlsx', data_only=True)
    assert [row['校验结果'] for row in workbook.table_records(invalid, 'TaskTable')][1:] == ['任务名称重复'] * 2
    invalid.close()


@pytest.mark.parametrize('prefix', ['=', '+', '-', '@'])
def test_template_projection_preserves_authority_and_safe_text(tmp_path, prefix):
    model = render_model(); model['tasks'][0]['name'] = prefix+'plain text'
    path = tmp_path/'sow.xlsx'; template = ROOT/'assets/sow-template.xlsx'
    workbook.write_workbook(template, model, path)
    source = openpyxl.load_workbook(template); result = openpyxl.load_workbook(path)
    try:
        workbook.verify_static_authority(result, source)
        workbook.verify_worksheet_authority(result, source)
        assert result['02-任务清单']['B5'].data_type == 's'
        assert result['02-任务清单']['B5'].value == "'"+prefix+'plain text'
        for name in workbook.TABLES:
            sheet, table = workbook.table_index(result)[name]
            assert table.autoFilter.ref == table.ref
        assert any('TaskTable[' in workbook.formula_text(c.value) for s in result for row in s for c in row if c.data_type == 'f')
    finally:
        source.close(); result.close()


@pytest.mark.e2e
def test_manual_entry_uses_names_only_and_prepared_rows_recalculate(tmp_path):
    from copy import deepcopy
    from office_engine import recalculate_workbook, require_office_engine
    from task_standard_catalog import catalog
    template = ROOT / 'assets/sow-template.xlsx'
    book = openpyxl.load_workbook(template)
    stories = book['01-需求故事']
    for column, value in enumerate(['订单需求', '订单查询', '查询订单', '是', '可返回订单'], 1):
        assert not stories.cell(5, column).protection.locked
        stories.cell(5, column).value = value
    tasks = book['02-任务清单']
    standard = catalog(template)
    for row, work_type, boundary in ((5, 'FE-VIEW', None), (204, 'IN-INTEGRATION', '内部集成')):
        values = {'A': '查询订单', 'B': f'交付对象{row}',
                  'D': standard.by_work_type_id[work_type]['工作类型名称'],
                  'E': '新建', 'F': 'M', 'G': boundary}
        for column, value in values.items():
            assert not tasks[f'{column}{row}'].protection.locked
            tasks[f'{column}{row}'] = value
        for column in ['C', 'I', 'J', 'K', 'L', 'M']:
            assert tasks[f'{column}{row}'].protection.locked
            assert tasks[f'{column}{row}'].data_type == 'f'
    book.save(tmp_path / 'hand-filled.xlsx'); book.close()
    engine = require_office_engine()
    recalculate_workbook(tmp_path / 'hand-filled.xlsx', tmp_path / 'manual-calculated.xlsx', engine)
    actual = openpyxl.load_workbook(tmp_path / 'manual-calculated.xlsx', data_only=True)
    try:
        assert not workbook.formula_errors(actual)
        assert actual['02-任务清单']['C5'].value == 'FE-VIEW'
        assert actual['02-任务清单']['C204'].value == 'IN-INTEGRATION'
        assert actual['02-任务清单']['M5'].value == actual['02-任务清单']['M204'].value == '通过'
        assert actual['01-需求故事']['I5'].value == '通过'
        assert actual['02-任务清单']['M100'].value in (None, '')
        assert actual['02-任务清单']['L100'].value in (None, '')
        expected_summary = workbook.table_records(actual, 'ProjectSummaryTable')
    finally:
        actual.close()
    model = render_model()
    model['stories'][0].update(name='查询订单', uatApplicable=True)
    model['tasks'][0].update(name='交付对象5')
    task = deepcopy(model['tasks'][0])
    task.update(taskId='integration-task', name='交付对象204', workTypeId='IN-INTEGRATION',
                rowSemanticSha256=standard.by_work_type_id['IN-INTEGRATION']['rowSemanticSha256'],
                integrationIds=['integration-manual'])
    model['tasks'].append(task)
    model['integrations'] = [{'integrationId': 'integration-manual', 'counterpartyBoundary': 'INTERNAL'}]
    projected = tmp_path / 'generated.xlsx'
    workbook.write_workbook(template, model, projected)
    recalculate_workbook(projected, tmp_path / 'generated-calculated.xlsx', engine)
    generated = openpyxl.load_workbook(tmp_path / 'generated-calculated.xlsx', data_only=True)
    try:
        assert workbook.table_records(generated, 'ProjectSummaryTable') == expected_summary
        assert len(workbook.table_records(generated, 'TaskTable')) == 2
    finally:
        generated.close()


def test_compact_summary_preserves_template_without_appending_identity_ledger(tmp_path):
    from template_variants import compact_task_template
    template = compact_task_template(name_entry_template(tmp_path), tmp_path / 'compact.xlsx')
    path = tmp_path / 'business-sow.xlsx'
    workbook.write_workbook(template, render_model(), path)
    original = openpyxl.load_workbook(template)
    generated = openpyxl.load_workbook(path)
    try:
        before, after = original['03-工作量汇总'], generated['03-工作量汇总']
        assert before.max_row == after.max_row
        assert before.max_column == after.max_column
        assert before.print_area == after.print_area
        for row in before:
            for cell in row:
                assert cell.value == after[cell.coordinate].value
        assert before.column_dimensions['C'].width == after.column_dimensions['C'].width
        assert '实体 ID' not in [cell.value for row in after for cell in row]
    finally:
        original.close(); generated.close()


@pytest.mark.parametrize('error', ['#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!',"='missing sheet'!A1",'ZIP'])
def test_workbook_integrity_errors_block_verification(tmp_path, error):
    path = tmp_path/'bad.xlsx'; book = openpyxl.Workbook()
    book.active['A1'] = error
    book.save(path); book.close()
    if error == 'ZIP': path.write_bytes(b'PK corrupted')
    assert callable(getattr(workbook, 'scan_workbook_integrity', None)), 'complete ZIP/formula scanner missing'
    with pytest.raises(ValueError): workbook.scan_workbook_integrity(path)


@pytest.mark.integration
def test_template_projection_prior_adapter_round_trip(tmp_path):
    from prior_state import inventory_prior_workbook, materialize_prior_snapshot
    from test_prior_state import revision_for
    model = render_model(); path = tmp_path/'transferred.xlsx'
    workbook.write_workbook(LEGACY_TEMPLATE, model, path)
    before = path.read_bytes(); inventory = inventory_prior_workbook(path)
    ids = {'epic-orders','feature-query','story-query','ac-success','ac-empty','task-query','input-a','input-b'}
    entities = []
    for entity_id in sorted(ids):
        evidence = [e for e in inventory['evidence'] if any(c['value'] == entity_id for c in e['canonicalCellValues'])]
        assert evidence, 'visible whole-cell identity missing: '+entity_id
        entities.append({'localKey':entity_id,'sourceId':'prior','entityKind':'CONTRACT_ENTITY',
            'semanticSummary':entity_id,'deliveryStatus':'CURRENT_BY_CONTRACT','visiblePriorId':entity_id,
            'evidenceIds':[e['priorEvidenceId'] for e in evidence]})
    decision = {'entities':entities,'sourceRelations':[],'entitySupersessions':[],'unsupportedRegions':[]}
    snapshot = materialize_prior_snapshot([inventory], decision, input_revision_bytes=revision_for([inventory],['prior']))
    assert {e['entityId'] for e in snapshot['entities']} == ids
    values = [c['value'] for e in inventory['evidence'] for c in e['canonicalCellValues']]
    assert canonical_json_bytes(model['inputItems'][0]['sourceRefs'][0]).decode().strip() in values
    assert path.read_bytes() == before


def shared_asset_render_model():
    import copy
    model=render_model()
    other=copy.deepcopy(model['stories'][0]);other.update(storyId='story-acceptance',name='用户验收订单查询')
    model['stories'].append(other)
    model['acceptanceCriteria'][1]['storyId']=other['storyId']
    model['tasks'][0]['acceptanceCriterionIds']=['ac-success','ac-empty']
    return model


def test_shared_asset_projection_lists_coverage_but_keeps_one_charge_and_template_financial_formulas(tmp_path):
    model=shared_asset_render_model();template=ROOT/'assets/sow-template.xlsx';path=tmp_path/'shared.xlsx'
    workbook.write_workbook(template,model,path)
    b=openpyxl.load_workbook(path);source=openpyxl.load_workbook(template)
    try:
        tasks=workbook.table_records(b,'TaskTable');stories=workbook.table_records(b,'SOWStoryTable')
        assert len(tasks)==1 and len(stories)==2
        assert '共享' in stories[1]['备注'] and model['stories'][0]['name'] in stories[1]['备注']
        assert 'COUNTIFS(TaskTable[所属故事],' in stories[1]['校验结果']
        assert '订单查询页面' in workbook.formula_text(stories[1]['任务列表'])
        from openpyxl.formula.translate import Translator
        assert stories[1]['故事人天']==Translator(source['01-需求故事']['H5'].value,origin='H5').translate_formula('H6')
        contract=workbook.projection_contract(source)['TaskTable']
        for header in contract['formulas']:
            col=contract['headers'].index(header)+1
            assert workbook.comparable_formula(b['02-任务清单'].cell(5,col).value)==workbook.comparable_formula(source['02-任务清单'].cell(5,col).value)
        workbook.verify_static_authority(b,source)
    finally:b.close();source.close()


def test_distinct_work_types_qualify_shared_context_names_without_hiding_same_type_duplicates(tmp_path):
    import copy
    from task_standard_catalog import catalog
    model=render_model();original=copy.deepcopy(model)
    second=copy.deepcopy(model['tasks'][0]);second.update(taskId='task-query-api',workTypeId='FE-QUERY-API',
        actualMeasurementScope='一个查询 API')
    standard=catalog(ROOT/'assets/sow-template.xlsx')
    second['rowSemanticSha256']=standard.by_work_type_id[second['workTypeId']]['rowSemanticSha256']
    model['tasks'].append(second);before=copy.deepcopy(model)
    path=tmp_path/'different-assets.xlsx';workbook.write_workbook(ROOT/'assets/sow-template.xlsx',model,path)
    book=openpyxl.load_workbook(path)
    try:
        rows=workbook.table_records(book,'TaskTable')
        assert len(rows)==2 and len({row['任务名称'] for row in rows})==2
        for row,task in zip(rows,model['tasks'],strict=True):
            assert row['任务名称']==task['name']+'：'+standard.by_work_type_id[task['workTypeId']]['工作类型名称']
        assert model==before
    finally: book.close()
    model['tasks'][1]=copy.deepcopy(original['tasks'][0]);model['tasks'][1]['taskId']='another-same-type'
    with pytest.raises(ValueError,match='duplicated'): workbook.build_rows(model,standard)


@pytest.mark.e2e
def test_shared_asset_real_office_reread_keeps_all_stories_and_one_task(tmp_path):
    from office_engine import recalculate_workbook,require_office_engine
    model=shared_asset_render_model();template=ROOT/'assets/sow-template.xlsx'
    projected=tmp_path/'projected.xlsx';final=tmp_path/'final.xlsx';reference=tmp_path/'reference.xlsx'
    workbook.write_workbook(template,model,projected)
    engine=require_office_engine();recalculate_workbook(projected,final,engine);recalculate_workbook(projected,reference,engine)
    assert workbook.dual_reopen(final,projected)['cachedValueCount']>0
    audit=workbook.audit_calculated_workbook(final,template,model,engine,expected_layout_path=projected,reference_path=reference)
    assert audit.task_count==1 and audit.story_count==2
    b=openpyxl.load_workbook(final,data_only=True)
    try:
        rows=workbook.table_records(b,'SOWStoryTable')
        assert [row['校验结果'] for row in rows]==['通过','通过']
        assert all('订单查询页面' in row['任务列表'] for row in rows)
        assert rows[1]['故事人天']==0 and rows[0]['故事人天']>0
    finally:b.close()
