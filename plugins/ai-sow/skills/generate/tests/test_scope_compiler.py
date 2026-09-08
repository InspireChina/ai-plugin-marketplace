from __future__ import annotations

TEST_LAYER = "integration"

import copy
import json
import sys
import pytest
from pathlib import Path


SKILL_ROOT = Path(__file__).parents[1]
PLUGIN_ROOT = SKILL_ROOT.parents[1]
FIXTURES = SKILL_ROOT / "fixtures"
SCRIPTS = SKILL_ROOT / "scripts"
if str(SKILL_ROOT / "tests") not in sys.path:
    sys.path.insert(0, str(SKILL_ROOT / "tests"))
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))
if str(PLUGIN_ROOT) not in sys.path:
    sys.path.insert(0, str(PLUGIN_ROOT))

from contracts import canonical_json_bytes, sha256_bytes, InvalidActionResult  # noqa: E402
from ir_samples import scan_ir, audit_ir, scope_decision_ir, complete_scope_ir  # noqa: E402
import scope_compiler as scope_compiler_module  # noqa: E402
from sow_model import model_skeleton  # noqa: E402


# Deterministic Scope IR public boundary (Task 7).
def scope_ir_envelope(kind):
    from contracts import action_contract_binding
    _, digest = action_contract_binding(SKILL_ROOT, kind + "-v1")
    return {"actionContractId": kind + "-v1", "actionContractSha256": digest}


def scan_ir_packet():
    return {"workItems": [{"workItemId": "work-a", "payload": {
        "coverageRootId": "block-a", "sourceRole": "PRD", "evidenceIds": ["block-a"],
        "sourceBlock": {"blockId": "block-a", "content": "管理员可查询订单。"}}}], "contextRefs": []}


def test_source_scan_coverage_root_seals_only_facts_and_explicit_no_relevance():
    import pytest
    from contracts import normalize_action_result, InvalidActionResult
    envelope = scope_ir_envelope("SOURCE_SCAN")
    result = scan_ir()
    sealed = normalize_action_result(envelope, canonical_json_bytes(result), skill_root=SKILL_ROOT)
    scope_compiler_module.verify_source_scan(scan_ir_packet(), json.loads(sealed))
    assert json.loads(sealed)[0]["facts"][0]["qualifiers"] == ["仅管理员", "七日内"]
    absent = [{"coverageRootId": "block-a", "disposition": "NO_RELEVANT_FACT", "facts": [], "noRelevantReason": "仅含目录。"}]
    scope_compiler_module.verify_source_scan(scan_ir_packet(), absent)
    for invalid in [[], [{**result[0], "facts": []}], [{**absent[0], "noRelevantReason": ""}],
                    [{**result[0], "noRelevantReason": "多余"}], [{**result[0], "node": {}}]]:
        with pytest.raises((InvalidActionResult, ValueError)):
            normalized = normalize_action_result(envelope, canonical_json_bytes(invalid), skill_root=SKILL_ROOT)
            scope_compiler_module.verify_source_scan(scan_ir_packet(), json.loads(normalized))
    for invalid in [scan_ir("wrong"), result + result, scan_ir(key="duplicate")]:
        if invalid[0]["facts"][0]["localKey"] == "duplicate":
            invalid[0]["facts"].append(copy.deepcopy(invalid[0]["facts"][0]))
        with pytest.raises(InvalidActionResult):
            scope_compiler_module.verify_source_scan(scan_ir_packet(), invalid)


def audit_ir_packet():
    original = scan_ir_packet()["workItems"][0]["payload"]
    return {"workItems": [], "contextRefs": [
        {"refId": "source-block-a", "canonicalContent": {"kind": "SOURCE_BLOCK", **original},
         "contentSha256": sha256_bytes(canonical_json_bytes({"kind": "SOURCE_BLOCK", **original}))},
        {"refId": "dependency-result-logical-scan", "canonicalContent": {
            "kind": "DEPENDENCY_RESULT", "logicalWorkId": "logical-scan",
            "attemptRecordSha256": "a" * 64, "normalizedResult": scan_ir()}}]}


def test_independent_source_audit_requires_every_category_and_bound_fact_keys():
    import pytest
    from contracts import normalize_action_result, InvalidActionResult
    packet, result = audit_ir_packet(), audit_ir()
    sealed = normalize_action_result(scope_ir_envelope("SOURCE_AUDIT"), canonical_json_bytes(result), skill_root=SKILL_ROOT)
    scope_compiler_module.verify_source_audit(packet, json.loads(sealed))
    for mutation in ["missing_category", "duplicate_category", "wrong_fact", "wrong_root", "wrong_evidence", "missing_scan", "missing_original"]:
        candidate, context = copy.deepcopy(result), copy.deepcopy(packet)
        if mutation == "missing_category": candidate["checks"].pop()
        elif mutation == "duplicate_category": candidate["checks"][-1] = candidate["checks"][0]
        elif mutation == "wrong_fact": candidate["checks"][0]["relatedFactKeys"] = ["other-scan:fact"]
        elif mutation == "wrong_root": candidate["checks"][0]["coverageRootId"] = "other"
        elif mutation == "wrong_evidence": candidate["checks"][0]["evidenceIds"] = ["other"]
        elif mutation == "missing_scan": context["contextRefs"].pop()
        else: context["contextRefs"].pop(0)
        with pytest.raises((InvalidActionResult, ValueError)):
            scope_compiler_module.verify_source_audit(context, candidate)
    result["checks"][0].update(decision="MISSING", reason="来源阈值漏读。")
    with pytest.raises(scope_compiler_module.ScopeInputRequired):
        scope_compiler_module.verify_source_audit(packet, result)


def test_scope_decision_ir_is_narrow_and_has_controlled_boundary_members():
    import pytest
    from contracts import normalize_action_result, InvalidActionResult
    result = scope_decision_ir()
    for kind in ["SCOPE_SYNTHESIS", "SCOPE_PROPOSAL", "SCOPE_JOIN"]:
        envelope = scope_ir_envelope(kind)
        assert json.loads(normalize_action_result(envelope, canonical_json_bytes(result), skill_root=SKILL_ROOT)) == result
        for extra in ["node", "entityId", "sourceRefs", "sha256", "checkpoint", "effortPhase"]:
            invalid = copy.deepcopy(result)
            invalid["decisions"][0][extra] = {}
            with pytest.raises(InvalidActionResult):
                normalize_action_result(envelope, canonical_json_bytes(invalid), skill_root=SKILL_ROOT)
        for mutation in ["kind", "classification", "nested_node", "relation", "exclude_without_reason"]:
            invalid = copy.deepcopy(result)
            item = invalid["decisions"][0]
            if mutation == "kind": item["decisionKind"] = "ARBITRARY"
            elif mutation == "classification": item["boundaryEvidence"]["classification"] = "my identity text"
            elif mutation == "nested_node": item["boundaryEvidence"]["node"] = {"featureId": "forged"}
            elif mutation == "relation": item["relations"][0]["kind"] = "arbitrary"
            else:
                item["decisionKind"] = "EXCLUDE"
                item["boundaryEvidence"]["classification"] = "DISPOSITION"
            with pytest.raises(InvalidActionResult):
                normalize_action_result(envelope, canonical_json_bytes(invalid), skill_root=SKILL_ROOT)


def synthesis_packet():
    return {"workItems": [], "contextRefs": [
        {"refId": "dependency-result-" + key, "canonicalContent": {"kind": "DEPENDENCY_RESULT",
         "logicalWorkId": key, "attemptRecordSha256": digest * 64, "normalizedResult": value}}
        for key, digest, value in [("scan", "a", scan_ir()), ("audit", "b", audit_ir())]]}


def test_scope_decision_ir_closes_fact_handles_and_checks_relations():
    import pytest
    from contracts import InvalidActionResult
    packet, result = synthesis_packet(), complete_scope_ir()
    scope_compiler_module.verify_scope_decision(packet, result)
    for mutation in ["missing_fact", "duplicate_fact", "unknown_fact", "unknown_prior", "unknown_parent", "unknown_evidence"]:
        candidate = copy.deepcopy(result)
        item = candidate["decisions"][0]
        if mutation == "missing_fact": item["factIds"] = []
        elif mutation == "duplicate_fact": candidate["decisions"][1]["factIds"] = item["factIds"]
        elif mutation == "unknown_fact": item["factIds"] = ["other:fact"]
        elif mutation == "unknown_prior": item["priorEntityIds"] = ["other:entity"]
        elif mutation == "unknown_parent": item["relations"][0]["targetLocalKeys"] = ["missing"]
        else: item["boundaryEvidence"]["evidenceIds"] = ["other-source"]
        with pytest.raises(InvalidActionResult):
            scope_compiler_module.verify_scope_decision(packet, candidate)
    result["decisions"][0]["uncertainty"] = "业务意图冲突。"
    with pytest.raises(InvalidActionResult):
        scope_compiler_module.verify_scope_decision(packet, result)


def test_scope_relation_endpoint_failure_identifies_exact_root_and_path():
    result = complete_scope_ir()
    result['decisions'][0]['relations'][0]['kind'] = 'APPLIES_TO'
    with pytest.raises(InvalidActionResult) as captured:
        scope_compiler_module.verify_scope_decision(synthesis_packet(), result)
    diagnostic = captured.value.diagnostic
    assert diagnostic.code == 'SCOPE_RELATION_ENDPOINT_INVALID'
    assert diagnostic.path == '/decisions/0/relations/0'
    assert diagnostic.subject_ids == (result['decisions'][0]['localKey'],)


def owner_input_case(count=1, text_size=20, inventories=()):
    from test_contracts import valid_input_revision
    request = {'project': {'projectId': 'project-refund'}, 'mode': 'GREENFIELD',
        'responsibilityBoundaries': [{'responsibilityBoundaryId': 'boundary-vendor'}]}
    request["mode"] = "BROWNFIELD" if inventories else "GREENFIELD"
    revision = valid_input_revision()
    request["project"] = {**revision["project"], **request["project"]}
    revision["project"] = request["project"]
    revision["requestSha256"] = sha256_bytes(canonical_json_bytes(request))
    policy = json.loads((SKILL_ROOT / "contracts/delivery-policy-v1.json").read_bytes())
    revision["deliveryPolicySha256"] = sha256_bytes((SKILL_ROOT / "contracts/delivery-policy-v1.json").read_bytes())
    contents = {f"block-{i:03}": "订单能力" + "x" * text_size for i in range(count)}
    source_hash = sha256_bytes(canonical_json_bytes(contents))
    revision["sources"] = [{"sourceId": "prd-main", "role": "PRD", "status": "APPROVED", "path": "sources/prd.md", "rawSha256": source_hash,
        "parserId": "markdown-blocks", "parserVersion": "1", "blockIds": list(contents)}]
    revision["blocks"] = [{"blockId": key, "sourceId": "prd-main", "rawSha256": source_hash,
        "contentSha256": sha256_bytes(value.encode()), "locator": f"paragraph:{i + 1}", "primaryCoverageBlockId": key,
        "contextBlockIds": [], "structuralParentId": None, "extractionDisposition": "INCLUDED", "droppedContentCategories": []}
        for i, (key, value) in enumerate(contents.items())]
    for i, inventory in enumerate(inventories):
        revision["sources"].append({"sourceId": f"prior-{i}", "role": "PRIOR_SOW", "status": "APPLICABLE", "path": f"sources/prior-{i}.xlsx",
            "rawSha256": inventory["workbookSha256"], "parserId": "xlsx-inventory", "parserVersion": "1", "blockIds": []})
    revision["priorSowState"] = "PROVIDED" if inventories else "NOT_PROVIDED"
    revision["priorSowSha256s"] = sorted({inventory["workbookSha256"] for inventory in inventories})
    return request, canonical_json_bytes(revision), contents, policy


def test_scope_owner_e2e_prepares_authoritative_atomic_sources_and_prior_dates(tmp_path):
    import pytest
    from prior_state import inventory_prior_workbook
    from test_prior_state import workbook_fixture
    from stage_planner import plan_stage, materialize_packet
    from action_ledger import ActionLedger
    from test_stage_planner import policy
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path)
    inventory = inventory_prior_workbook(path)
    request, revision, contents, _ = owner_input_case(inventories=[inventory])
    before = path.read_bytes()
    items, contexts = scope_compiler_module.prepare_scope_inputs(revision, contents, request=request, prior_inventories=[inventory])
    descriptors = scope_compiler_module.build_scope_work_descriptors(items, contexts, policy())
    plan = plan_stage("SCOPE", items, contexts, descriptors, [], policy())
    for work in plan["works"]:
        if work["packetPlan"]["actionKind"] == "PRIOR_ANALYZE":
            packet = json.loads(materialize_packet(plan, work["logicalWorkId"], 1, items, contexts, [], ActionLedger()))
            from prior_state import validate_bound_prior_context
            validate_bound_prior_context("PRIOR_ANALYZE", packet, inventories=[inventory], input_revision_bytes=revision)
            assert any(ref["refId"] == "PROJECT_EFFECTIVE_START" for ref in packet["contextRefs"])
    assert path.read_bytes() == before


def test_large_prior_sheet_is_losslessly_partitioned_with_bound_headers(tmp_path):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table
    from prior_state import inventory_prior_workbook, validate_bound_prior_context
    from stage_planner import plan_stage, materialize_packet
    from test_stage_planner import policy
    from action_ledger import ActionLedger
    from dataclasses import replace

    path = tmp_path / 'large-prior.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['ID', '交付物', '状态'])
    for index in range(160):
        sheet.append([f'prior-{index:03}', '可验收的交付边界；' * 70, '已交付'])
    sheet.add_table(Table(displayName='PriorContract', ref='A1:C161'))
    workbook.save(path)
    before = path.read_bytes()
    inventory = inventory_prior_workbook(path)
    request, revision, contents, _ = owner_input_case(inventories=[inventory])
    items, contexts = scope_compiler_module.prepare_scope_inputs(
        revision, contents, request=request, prior_inventories=[inventory])
    planning = replace(policy(), model_context_limit_tokens=272000, max_planned_tokens=5000000)
    descriptors = scope_compiler_module.build_scope_work_descriptors(items, contexts, planning)
    plan = plan_stage('SCOPE', items, contexts, descriptors, [], planning)
    prior_items = [item for item in items if item.action_kind == 'PRIOR_ANALYZE']
    assert len(prior_items) > 1
    expected = {item['priorEvidenceId']: item for item in inventory['evidence']}
    collected = [item for work in prior_items for item in work.work_item_payload['evidence']]
    assert len(collected) == len(expected)
    assert {item['priorEvidenceId']: item for item in collected} == expected
    expected_cells = inventory['sheets'][0]['cells']
    actual_cells = [cell for item in collected for cell in item['canonicalCellValues']]
    assert sorted(actual_cells, key=lambda cell:cell['address']) == sorted(expected_cells, key=lambda cell:cell['address'])
    for item in prior_items:
        payload = item.work_item_payload
        assert payload['sheet'] == {key:value for key,value in inventory['sheets'][0].items() if key!='cells'}
        assert payload['headerEvidence'] == [expected[next(iter(expected))]]
        assert payload['evidenceIds'] == sorted(value['priorEvidenceId'] for value in payload['evidence'])
    for work in plan['works']:
        if work['packetPlan']['actionKind'] == 'PRIOR_ANALYZE':
            packet = json.loads(materialize_packet(plan, work['logicalWorkId'], 1, items, contexts, [], ActionLedger()))
            validate_bound_prior_context('PRIOR_ANALYZE', packet, inventories=[inventory], input_revision_bytes=revision)
    assert path.read_bytes() == before
    assert (items, contexts) == scope_compiler_module.prepare_scope_inputs(
        revision, contents, request=request, prior_inventories=[inventory])
    for mutate in ["content", "request", "inventory"]:
        bad_contents, bad_request, bad_inventories = copy.deepcopy(contents), copy.deepcopy(request), [inventory]
        if mutate == "content": bad_contents["block-000"] = "篡改"
        elif mutate == "request": bad_request["project"]["name"] = "另一个项目"
        else: bad_inventories = []
        with pytest.raises(ValueError):
            scope_compiler_module.prepare_scope_inputs(revision, bad_contents, request=bad_request, prior_inventories=bad_inventories)


def scope_owner_result(kind, packet):
    dependencies = [ref["canonicalContent"]["normalizedResult"] for ref in packet["contextRefs"] if ref["canonicalContent"].get("kind") == "DEPENDENCY_RESULT"]
    if kind == "SOURCE_SCAN":
        return [scan_ir(item["payload"]["coverageRootId"], "fact")[0] for item in packet["workItems"]]
    if kind == "SOURCE_AUDIT":
        return {"checks": [check for result in dependencies[0] for check in audit_ir(result["coverageRootId"], "fact")["checks"]]}
    if kind == "PRIOR_ANALYZE":
        return {"entities": [{"localKey": item["workItemId"] + ":entity", "sourceId": item["payload"]["sourceId"],
            "entityKind": "CONTRACT_ENTITY", "semanticSummary": "查询订单", "deliveryStatus": "CURRENT_BY_CONTRACT",
            "evidenceIds": item["payload"]["evidenceIds"]} for item in packet["workItems"]],
            "sourceRelations": [], "entitySupersessions": [], "unsupportedRegions": [], "unextractedEvidence": []}
    if kind == "PRIOR_CONSOLIDATE":
        result = {key: [item for dependency in dependencies for item in dependency[key]] for key in ["entities", "sourceRelations", "entitySupersessions", "unsupportedRegions"]}
        if len(result["entities"]) == 2:
            previous, current = sorted(result["entities"], key=lambda item: item["sourceId"])
            result["entitySupersessions"] = [{"predecessorLocalKeys": [previous["localKey"]], "successorLocalKeys": [current["localKey"]], "evidenceIds": current["evidenceIds"]}]
        return {key: result[key] for key in ("sourceRelations", "entitySupersessions")}
    if kind == "SCOPE_JOIN":
        children = [item for dependency in dependencies for item in dependency["decisions"] if item["decisionKind"] != "POLICY_INSTANCE"]
        return with_required_scope_policies({"decisions": children})
    result = {"decisions": []}
    for scan in [item for dependency in dependencies if isinstance(dependency, list) for item in dependency]:
        root = scan["coverageRootId"]
        feature = scope_decision_ir("feature-" + root, [root + ":" + item["localKey"] for item in scan["facts"]])["decisions"][0]
        feature["boundaryEvidence"].update(name="能力" + root, evidenceIds=[root])
        feature["relations"][0].update(targetLocalKeys=["epic-" + root], evidenceIds=[root])
        epic = copy.deepcopy(feature)
        epic.update(localKey="epic-" + root, decisionKind="EPIC", factIds=[], relations=[])
        epic["boundaryEvidence"]["name"] = "业务域" + root
        result["decisions"].extend([feature, epic])
    return with_required_scope_policies(result)


def with_required_scope_policies(result):
    if any(item["decisionKind"] == "POLICY_INSTANCE" for item in result["decisions"]):
        return result
    target = next(item for item in result["decisions"] if item["decisionKind"] == "FEATURE")
    for policy_id in sorted(scope_compiler_module.REQUIRED_POLICY_INCLUSIONS):
        result["decisions"].append({"localKey": policy_id, "decisionKind": "POLICY_INSTANCE", "factIds": [], "priorEntityIds": [],
            "boundaryEvidence": {"name": policy_id, "classification": policy_id, "evidenceIds": target["boundaryEvidence"]["evidenceIds"], "facetFacts": [], "observationKeys": []},
            "relations": [{"kind": "APPLIES_TO", "targetLocalKeys": [target["localKey"]], "evidenceIds": target["boundaryEvidence"]["evidenceIds"]}]})
    return result


def scope_owner_runtime(request, revision, contents, *, inventories=(), sizing=None):
    from test_stage_planner import policy
    from stage_planner import plan_stage
    from action_ledger import ActionLedger
    sizing = sizing or policy()
    items, contexts = scope_compiler_module.prepare_scope_inputs(revision, contents, request=request, prior_inventories=inventories)
    descriptors = scope_compiler_module.build_scope_work_descriptors(items, contexts, sizing)
    return [plan_stage("SCOPE", items, contexts, descriptors, [], sizing), items, contexts, ActionLedger(), sizing]


def seal_scope_work(runtime, logical_id, result_override=None, *, revision_number=1, allow_failure=False):
    from test_stage_planner import envelope_for_plan
    from stage_planner import materialize_packet, DependencyResultRef, run_budget_policy_value
    from action_ledger import issue, finish
    from models import ActionEnvelope
    from contracts import estimate_action_input_tokens
    from test_action_ledger import successful_completion
    plan, items, contexts, ledger, sizing = runtime
    work = next(work for work in plan["works"] if work["logicalWorkId"] == logical_id)
    refs = []
    for dependency in work["packetPlan"]["dependencyLogicalWorkIds"]:
        digest, record = next((digest, record) for digest, record in ledger.attempt_records.items() if record.logical_work_id == dependency and record.outcome == "SUCCEEDED")
        refs.append(DependencyResultRef(dependency, digest, ledger.normalized_results[record.normalized_result_sha256]))
    repair = None
    if revision_number == 2:
        from action_ledger import build_attempt_repair_context
        digest = next(digest for digest, record in ledger.attempt_records.items() if record.logical_work_id == logical_id and record.failure_kind == "INVALID_IR")
        repair = build_attempt_repair_context(logical_id, digest, ledger.attempt_records, ledger.raw_outputs, envelopes_by_sha256=ledger.envelopes_by_sha256)
    payload = materialize_packet(plan, logical_id, revision_number, items, contexts, refs, ledger, repair)
    packet = json.loads(payload)
    kind = work["packetPlan"]["actionKind"]
    result = scope_owner_result(kind, packet) if result_override is None else result_override
    value = {**envelope_for_plan(plan, logical_id).value, "packetSha256": sha256_bytes(payload)}
    value["revision"] = revision_number
    if revision_number == 2:
        value["actionId"] += "-r2"
    scope_context = next((json.loads(ref.canonical_content) for ref in contexts if json.loads(ref.canonical_content).get("kind") == "SCOPE_CONTEXT"), None)
    if scope_context:
        value["inputRevisionSha256"] = scope_context["inputRevisionSha256"]
    if ledger.envelopes_by_sha256:
        value["runId"] = next(iter(ledger.envelopes_by_sha256.values())).value["runId"]
    value["executionLimits"] = {"estimatedInputTokens": estimate_action_input_tokens(SKILL_ROOT, value["actionContractId"], payload,
        budget_policy=run_budget_policy_value(sizing), max_output_tokens=sizing.output_reserve_tokens),
        "maxOutputTokens": sizing.output_reserve_tokens, "maxHydrateTokens": sizing.hydrate_reserve_tokens}
    envelope = ActionEnvelope(value, "actions/" + value["actionId"] + "/envelope.json", sha256_bytes(canonical_json_bytes(value)))
    callback = None
    if kind in {"SOURCE_SCAN", "SOURCE_AUDIT", "SCOPE_SYNTHESIS", "SCOPE_PROPOSAL", "SCOPE_JOIN"}:
        scope_compiler_module.validate_bound_scope_context(kind, packet)
        callback = lambda normalized: scope_compiler_module.validate_bound_scope_result(kind, packet, normalized)
    ledger, record = finish(issue(ledger, envelope), envelope, successful_completion(canonical_json_bytes(result)), bound_result_validator=callback, packet_payload=payload)
    runtime[3] = ledger
    if not allow_failure:
        assert record.outcome == "SUCCEEDED", record.diagnostic
    return packet, record


def test_scope_materialize_once_per_revision_refuses_partial_groups_and_reuses_immutable_files(tmp_path):
    import pytest
    from runtime.project_io import ProjectFiles
    request, revision, contents, _ = owner_input_case()
    runtime = scope_owner_runtime(request, revision, contents)
    plan, items, contexts, ledger, sizing = runtime
    files = ProjectFiles.open(tmp_path)
    with pytest.raises(scope_compiler_module.ScopeInputRequired):
        scope_compiler_module.materialize_scope_candidate(revision, request, plan, items, contexts, ledger, sizing)
    assert not list(tmp_path.rglob("*candidate*"))
    for work in plan["works"]:
        seal_scope_work(runtime, work["logicalWorkId"])
    ledger = runtime[3]
    material = scope_compiler_module.materialize_scope_candidate(revision, request, plan, items, contexts, ledger, sizing)
    assert scope_compiler_module.validate_scope_candidate(material) == ()
    published = scope_compiler_module.publish_scope_candidate(files, "run-scope", material)
    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in tmp_path.rglob("*.json")}
    restored = scope_compiler_module.materialize_scope_candidate(revision, request, plan, items[::-1], contexts[::-1], ledger, sizing)
    assert scope_compiler_module.publish_scope_candidate(files, "run-scope", restored) == published
    assert {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in tmp_path.rglob("*.json")} == before
    candidate = json.loads(material.candidate_bytes)
    assert len(candidate["features"]) == len(candidate["inputItems"]) == 1
    assert candidate["stories"] == candidate["tasks"] == []
    assert all("featureId" not in item for item in scope_owner_result("SOURCE_SCAN", scan_ir_packet()))


def test_scope_assumptions_close_as_project_gates_without_delivery_obligations():
    request, revision, contents, _ = owner_input_case(count=2)
    runtime = scope_owner_runtime(request, revision, contents)
    for work in runtime[0]['works']:
        override = None
        if work['packetPlan']['actionKind'] == 'SOURCE_SCAN':
            override = [scan_ir(root, 'fact')[0] for root in sorted(contents)]
            override[1]['facts'][0].update(factKind='ASSUMPTION', statement='客户提供测试环境。')
        seal_scope_work(runtime, work['logicalWorkId'], override)
    material = scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], runtime[4])
    model = json.loads(material.candidate_bytes)
    identity = next(row['inputItemId'] for row in model['inputItems'] if row['kind'] == 'ASSUMPTION')
    gate = next(row for row in model['scopeClosure'] if row['inputItemId'] == identity)
    assert gate['disposition'] == 'PROJECT_GATE'
    assert gate['deliveryDisposition'] == 'PROJECT_LEVEL_ONLY'
    assert gate['assignedFeatureIds'] == []
    assert scope_compiler_module.validate_scope_candidate(material) == ()


def test_scope_owner_e2e_multi_prior_consumes_real_unique_root_and_full_supersession(tmp_path):
    import openpyxl
    from dataclasses import replace
    from test_stage_planner import policy
    from prior_state import inventory_prior_workbook, derive_effective_prior
    inventories = []
    for i in range(2):
        path = tmp_path / f"prior-{i}.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "合同"
        sheet.append(["ID", "交付物"])
        sheet.append([f"prior-visible-{i}", "替代旧版本并已交付" + "x" * 18000])
        workbook.save(path)
        workbook.close()
        inventories.append(inventory_prior_workbook(path))
    request, revision, contents, _ = owner_input_case(inventories=inventories)
    sizing = replace(policy(), model_context_limit_tokens=40000, output_reserve_tokens=4096,
                     hydrate_reserve_tokens=1024, safety_margin_tokens=1024)
    runtime = scope_owner_runtime(request, revision, contents, inventories=inventories, sizing=sizing)
    for work in runtime[0]["works"]:
        seal_scope_work(runtime, work["logicalWorkId"])
    assert sum(work["packetPlan"]["actionKind"] == "PRIOR_ANALYZE" for work in runtime[0]["works"]) == 2
    assert sum(work["packetPlan"]["actionKind"] == "PRIOR_CONSOLIDATE" for work in runtime[0]["works"]) == 1
    material = scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], sizing, prior_inventories=inventories)
    assert material.prior_state_bytes is not None
    snapshot = json.loads(material.prior_state_bytes)
    active = derive_effective_prior(snapshot)["activeEntityIds"]
    assert len(active) == 1
    assert next(item for item in snapshot["entities"] if item["entityId"] == active[0])["sourceId"] == "prior-1"
    assert len(snapshot["entities"]) == 2
    assert scope_compiler_module.validate_scope_candidate(material) == ()


def test_scope_owner_e2e_preserves_visible_prior_only_for_unchanged_unique_match(tmp_path):
    from test_prior_state import workbook_fixture
    from prior_state import inventory_prior_workbook
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path)
    inventory = inventory_prior_workbook(path)
    request, revision, contents, _ = owner_input_case(inventories=[inventory])
    runtime = scope_owner_runtime(request, revision, contents, inventories=[inventory])
    prior_key = None
    for work in runtime[0]["works"]:
        kind = work["packetPlan"]["actionKind"]
        if kind == "PRIOR_ANALYZE":
            item = next(item for item in runtime[1] if item.action_kind == "PRIOR_ANALYZE")
            prior_key = item.work_item_id + ":entity"
            result = {"entities": [{"localKey": prior_key, "sourceId": "prior-0", "entityKind": "CONTRACT_ENTITY", "semanticSummary": "订单查询",
                "deliveryStatus": "CURRENT_BY_CONTRACT", "evidenceIds": item.work_item_payload["evidenceIds"], "visiblePriorId": "prior-visible-1"}],
                "sourceRelations": [], "entitySupersessions": [], "unsupportedRegions": [], "unextractedEvidence": []}
        elif kind == "SCOPE_SYNTHESIS":
            result = complete_scope_ir()
            feature, epic = result["decisions"]
            feature["factIds"] = ["block-000:fact"]
            feature["priorEntityIds"] = [prior_key]
            for decision in result["decisions"]:
                decision["boundaryEvidence"]["evidenceIds"] = ["block-000"]
                for relation in decision["relations"]: relation["evidenceIds"] = ["block-000"]
            feature["relations"].extend([{ "kind": kind, "targetLocalKeys": ["feature-a"], "evidenceIds": ["block-000"]}
                                       for kind in ["REUSE_DEPENDENCY", "UNCHANGED_IDENTITY"]])
            result = with_required_scope_policies(result)
        else:
            result = None
        seal_scope_work(runtime, work["logicalWorkId"], result)
    material = scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], runtime[4], prior_inventories=[inventory])
    candidate, graph = json.loads(material.candidate_bytes), json.loads(material.change_graph_bytes)
    assert candidate["features"][0]["featureId"] == "prior-visible-1"
    assert graph["changeGroups"] == [{"kind": "REUSE_DEPENDENCY", "priorEntityIds": ["prior-visible-1"], "targetEntityIds": ["prior-visible-1"], "evidenceIds": ["block-000"]}]
    identity = next(row for row in material.review_obligations if row["kind"] == "PRIOR_IDENTITY")
    assert identity["targetEntityIds"] == ["prior-visible-1"]


def test_scope_decision_ir_bound_finish_rejects_invalid_refs_without_schema_reads(monkeypatch):
    import pytest
    from test_action_ledger import prepared_envelope, successful_completion
    from action_ledger import ActionLedger, issue, finish
    from models import ActionEnvelope
    from contracts import action_contract_binding
    kind = "SCOPE_SYNTHESIS"
    packet = synthesis_packet()
    scope_compiler_module.validate_bound_scope_context(kind, packet)
    result = complete_scope_ir()
    result["decisions"][0]["factIds"] = ["unrelated:fact"]
    _, digest = action_contract_binding(SKILL_ROOT, kind + "-v1")
    value = {**prepared_envelope().value, "actionContractId": kind + "-v1", "actionContractSha256": digest,
             "packetSha256": sha256_bytes(canonical_json_bytes(packet))}
    envelope = ActionEnvelope(value, "actions/scope/envelope.json", sha256_bytes(canonical_json_bytes(value)))
    def pure_validator(normalized):
        with monkeypatch.context() as guard:
            def denied(*args, **kwargs): raise AssertionError("callback attempted file read")
            guard.setattr(Path, "read_bytes", denied)
            guard.setattr(Path, "read_text", denied)
            scope_compiler_module.validate_bound_scope_result(kind, packet, normalized)
    ledger, record = finish(issue(ActionLedger(), envelope), envelope, successful_completion(canonical_json_bytes(result)), bound_result_validator=pure_validator)
    assert record.failure_kind == "INVALID_IR"
    assert record.normalized_result_sha256 is None
    assert ledger.raw_outputs[record.raw_sha256] == canonical_json_bytes(result)


def test_independent_source_audit_missing_is_execution_fact_but_not_scope_completion():
    from test_action_ledger import prepared_envelope, successful_completion
    from action_ledger import ActionLedger, issue, finish
    from models import ActionEnvelope
    from contracts import action_contract_binding
    import pytest
    packet, result = audit_ir_packet(), audit_ir()
    result["checks"][0].update(decision="MISSING", reason="遗漏阈值。")
    scope_compiler_module.validate_bound_scope_context("SOURCE_AUDIT", packet)
    _, digest = action_contract_binding(SKILL_ROOT, "SOURCE_AUDIT-v1")
    value = {**prepared_envelope().value, "actionContractId": "SOURCE_AUDIT-v1", "actionContractSha256": digest,
             "packetSha256": sha256_bytes(canonical_json_bytes(packet))}
    envelope = ActionEnvelope(value, "actions/audit/envelope.json", sha256_bytes(canonical_json_bytes(value)))
    ledger, record = finish(issue(ActionLedger(), envelope), envelope, successful_completion(canonical_json_bytes(result)),
        bound_result_validator=lambda normalized: scope_compiler_module.validate_bound_scope_result("SOURCE_AUDIT", packet, normalized))
    assert record.outcome == "SUCCEEDED"
    with pytest.raises(scope_compiler_module.ScopeInputRequired):
        scope_compiler_module.verify_source_audit(packet, json.loads(ledger.normalized_results[record.normalized_result_sha256]))


def scope_prototype_case(tmp_path):
    from test_orchestrator import start_demo, prototype_payload, submit_prototype
    from test_prototype_analysis import scenario_fixture, trace_fixture, observation_fixture
    import orchestrator
    from runtime.project_io import ProjectFiles
    action = start_demo(tmp_path, two_controls=True)["nextAction"]
    inventory = prototype_payload(tmp_path, action)["inventory"]
    run_id = action["runId"]
    for number in (1, 2):
        chosen = copy.deepcopy(inventory)
        chosen["interactions"] = [inventory["interactions"][number - 1]]
        scenario = scenario_fixture(chosen, round=number)
        submit_prototype(tmp_path, action, scenario)
        browser = orchestrator.run_mode(tmp_path, "resume")["nextAction"]
        submit_prototype(tmp_path, browser, trace_fixture(inventory, scenario))
        analyze = orchestrator.run_mode(tmp_path, "resume")["nextAction"]
        observation = observation_fixture(chosen, "same-key")
        if number == 2: observation["runtimeStatus"] = "CODE_ONLY"
        submit_prototype(tmp_path, analyze, {"observations": [observation]})
        resumed = orchestrator.run_mode(tmp_path, "resume")
        if number == 1:
            action = resumed["nextAction"]
    root = tmp_path / ".ai-sow/work/runs" / run_id
    prototype_ledger = max((json.loads(path.read_bytes()) for path in (root / "stages/SCOPE/prototype-ledgers").glob("*.json")), key=lambda value: len(value["rounds"]))
    return inventory, prototype_ledger, orchestrator._load_action_ledger(ProjectFiles.open(tmp_path), run_id)


def test_scope_owner_plan_topology_consumes_all_prototype_round_refs_and_final_ledger(tmp_path):
    import pytest
    from test_stage_planner import scope_plan_inputs, policy
    from stage_planner import plan_stage
    inventory, prototype_ledger, ledger = scope_prototype_case(tmp_path)
    prototype_contexts = scope_compiler_module.prepare_scope_prototype_contexts(inventory, prototype_ledger, ledger)
    references = [json.loads(ref.canonical_content) for ref in prototype_contexts]
    rounds = [ref for ref in references if ref["kind"] == "PROTOTYPE_OBSERVATION_REF"]
    assert [ref["round"] for ref in rounds] == [1, 2]
    assert rounds[0]["observationKeys"] != rounds[1]["observationKeys"]
    assert all("observations" not in ref and "normalizedResult" not in ref for ref in rounds)
    items, contexts = scope_plan_inputs()
    contexts.extend(prototype_contexts)
    descriptors = scope_compiler_module.build_scope_work_descriptors(items, contexts, policy())
    plan = plan_stage("SCOPE", items, contexts, descriptors, [], policy())
    root = next(work for work in plan["works"] if work["packetPlan"]["actionKind"] == "SCOPE_SYNTHESIS")
    assert {ref.ref_id for ref in prototype_contexts} <= {ref["refId"] for ref in root["packetPlan"]["contextRefs"]}
    with pytest.raises(ValueError):
        scope_compiler_module.build_scope_work_descriptors(items, contexts[:-1], policy())
    incomplete = copy.deepcopy(prototype_ledger)
    incomplete["sealed"] = False
    with pytest.raises(scope_compiler_module.ScopeInputRequired):
        scope_compiler_module.prepare_scope_prototype_contexts(inventory, incomplete, ledger)


def projected_scope_ir():
    result = complete_scope_ir()
    feature, epic = result["decisions"]
    epic.update(localKey="epic", factIds=[])
    feature.update(localKey="feature", factIds=["block-000:fact"])
    for item in (epic, feature):
        item["boundaryEvidence"].update(evidenceIds=["block-000"], classification="TECHNICAL")
    feature["relations"] = [{"kind": "PARENT", "targetLocalKeys": ["epic"], "evidenceIds": ["block-000"]},
        {"kind": "DESIGN", "targetLocalKeys": ["design"], "evidenceIds": ["block-001"]}]
    def add(key, kind, classification, indexes, facets=()):
        facts = [f"block-{i:03}:fact" for i in indexes]
        decision = {"localKey": key, "decisionKind": kind, "factIds": facts, "priorEntityIds": [],
            "boundaryEvidence": {"name": key, "classification": classification,
                "evidenceIds": [f"block-{i:03}" for i in indexes], "facetFacts": list(facets), "observationKeys": []},
            "relations": [{"kind": "APPLIES_TO", "targetLocalKeys": ["feature"], "evidenceIds": [f"block-{indexes[0]:03}"]}]}
        result["decisions"].append(decision)
        return decision
    add("design", "DESIGN_ITEM", "COMPONENT", [1])
    integration = add("integration", "INTEGRATION", "EXTERNAL", [2, 3, 4, 5],
        [{"role": role, "factId": f"block-{i:03}:fact"} for i, role in zip([2, 3, 4], ["DIRECTION", "METHOD", "PURPOSE"])])
    integration["boundaryEvidence"]["responsibilityBoundaryIds"] = ["boundary-app"]
    add("quality", "NFR", "PERFORMANCE", [6], [{"role": "TARGET", "factId": "block-006:fact"}])
    excluded = add("excluded", "EXCLUDE", "DISPOSITION", [7])
    excluded.update(relations=[], exclusionReason="本轮来源明确排除此项。")
    for i, policy_id in enumerate(["policy-sit-automation", "policy-uat-automation", "policy-go-live", "policy-data-migration"]):
        policy = add("policy-" + str(i), "POLICY_INSTANCE", policy_id, [8])
        if i < 3: policy["factIds"] = []
    return result


def test_scope_owner_e2e_projects_all_target_kinds_from_facts_policy_and_approved_design():
    request, revision_bytes, contents, policy = owner_input_case(count=9)
    revision = json.loads(revision_bytes)
    revision["sources"][0].update(role="HLD", status="APPROVED")
    request["responsibilityBoundaries"] = [{**request["responsibilityBoundaries"][0], "responsibilityBoundaryId": "boundary-app"}]
    revision["requestSha256"] = sha256_bytes(canonical_json_bytes(request))
    revision_bytes = canonical_json_bytes(revision)
    runtime = scope_owner_runtime(request, revision_bytes, contents)
    for work in runtime[0]["works"]:
        seal_scope_work(runtime, work["logicalWorkId"], projected_scope_ir() if work["packetPlan"]["actionKind"] == "SCOPE_SYNTHESIS" else None)
    material = scope_compiler_module.materialize_scope_candidate(revision_bytes, request, *runtime[:3], runtime[3], runtime[4])
    assert scope_compiler_module.validate_scope_candidate(material) == ()
    model = json.loads(material.candidate_bytes)
    feature = model["features"][0]
    assert feature["designRefs"] == [model["designItems"][0]["designItemId"]]
    assert model["designItems"][0]["status"] == "APPROVED"
    integration=model["integrations"][0]
    assert integration["responsibilityBoundaryIds"] == ["boundary-app"]
    assert integration["direction"] == scan_ir()[0]["facts"][0]["statement"]
    assert set(integration)=={
        "integrationId","name","featureIds","sourceRefs","direction","method",
        "purpose","responsibilityBoundaryIds","counterpartyBoundary"}
    assert model["nfrs"][0]["target"] == scan_ir()[0]["facts"][0]["statement"]
    assert {item["policyId"] for item in model["policyInstances"]} == {item["policyId"] for item in policy["policies"]}
    assert sum(item["disposition"] == "OUT_OF_SCOPE" for item in model["scopeClosure"]) == 1
    assert model["scopeAnnotations"][0]["text"] == "本轮来源明确排除此项。"




def test_scope_owner_e2e_57_items_consumes_actual_multi_group_attempts_with_two_join_levels():
    from dataclasses import replace
    from test_stage_planner import policy
    from stage_planner import materialize_packet, DependencyResultRef, _effective_success
    request, revision, contents, _ = owner_input_case(count=57, text_size=60000)
    sizing = replace(policy(), model_context_limit_tokens=100000, output_reserve_tokens=12000)
    runtime = scope_owner_runtime(request, revision, contents, sizing=sizing)
    kinds = [work["packetPlan"]["actionKind"] for work in runtime[0]["works"]]
    assert kinds.count("SOURCE_SCAN") == 57 and kinds.count("SCOPE_JOIN") > 2
    for work in runtime[0]["works"]:
        override = None
        if work["packetPlan"]["actionKind"] == "SCOPE_JOIN":
            refs = []
            for key in work["packetPlan"]["dependencyLogicalWorkIds"]:
                digest, record = _effective_success(runtime[3], key)
                refs.append(DependencyResultRef(key, digest, runtime[3].normalized_results[record.normalized_result_sha256]))
            packet = json.loads(materialize_packet(runtime[0], work["logicalWorkId"], 1, runtime[1], runtime[2], refs, runtime[3]))
            children = [item for ref in packet["contextRefs"] if ref["canonicalContent"].get("kind") == "DEPENDENCY_RESULT" for item in ref["canonicalContent"]["normalizedResult"]["decisions"]]
            feature = copy.deepcopy(next(item for item in children if item["decisionKind"] == "FEATURE"))
            epic = copy.deepcopy(next(item for item in children if item["decisionKind"] == "EPIC"))
            feature["factIds"] = sorted({handle for item in children for handle in item["factIds"]})
            feature["boundaryEvidence"]["evidenceIds"] = sorted({eid for item in children for eid in item["boundaryEvidence"]["evidenceIds"]})
            feature["relations"][0]["targetLocalKeys"] = [epic["localKey"]]
            override = with_required_scope_policies({"decisions": [feature, epic]})
        seal_scope_work(runtime, work["logicalWorkId"], override)
    material = scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], sizing)
    assert scope_compiler_module.validate_scope_candidate(material) == ()
    model = json.loads(material.candidate_bytes)
    assert len(model["inputItems"]) == len(model["scopeClosure"]) == 57
    assert {item["disposition"] for item in model["scopeClosure"]} == {"SCOPE_NODE"}
    assert model["stories"] == model["tasks"] == []
    restored = scope_compiler_module.materialize_scope_candidate(revision, request, runtime[0], runtime[1][::-1], runtime[2][::-1], runtime[3], sizing)
    assert restored.candidate_bytes == material.candidate_bytes


@pytest.mark.parametrize("defect", ["parent", "policy_target", "unconsumed_prior", "facet", "responsibility"])
def test_scope_decision_ir_rejects_invalid_structural_relations_before_seal(defect):
    packet, result = synthesis_packet(), complete_scope_ir()
    feature = result["decisions"][0]
    if defect == "parent": feature["relations"][0]["targetLocalKeys"] = [feature["localKey"]]
    elif defect == "policy_target": feature["relations"].append({"kind": "POLICY", "targetLocalKeys": ["epic-a"], "evidenceIds": ["block-a"]})
    elif defect == "unconsumed_prior":
        feature["relations"].append({"kind": "UNCHANGED_IDENTITY", "targetLocalKeys": [feature["localKey"]], "evidenceIds": ["block-a"]})
    elif defect == "facet": feature["boundaryEvidence"]["facetFacts"] = [{"role": "TARGET", "factId": "block-a:fact-a"}]
    else: feature["boundaryEvidence"]["responsibilityBoundaryIds"] = ["invented"]
    with pytest.raises(InvalidActionResult):
        scope_compiler_module.validate_bound_scope_result("SCOPE_SYNTHESIS", packet, canonical_json_bytes(result))


def test_scope_owner_e2e_requires_mandatory_policies_without_inventing_model_decisions():
    request, revision, contents, _ = owner_input_case()
    runtime = scope_owner_runtime(request, revision, contents)
    for work in runtime[0]["works"]:
        packet, record = seal_scope_work(runtime, work["logicalWorkId"])
    result = scope_owner_result("SCOPE_SYNTHESIS", packet)
    result["decisions"] = [item for item in result["decisions"] if item["decisionKind"] != "POLICY_INSTANCE"]
    # Build another genuine Attempt chain with the incomplete, otherwise valid IR.
    runtime = scope_owner_runtime(request, revision, contents)
    for work in runtime[0]["works"]:
        seal_scope_work(runtime, work["logicalWorkId"], result if work["packetPlan"]["actionKind"] == "SCOPE_SYNTHESIS" else None)
    with pytest.raises(scope_compiler_module.ScopeInputRequired, match="政策"):
        scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], runtime[4])


def test_scope_materialize_once_per_revision_recovers_real_revision_two_repair():
    request, revision, contents, _ = owner_input_case()
    runtime = scope_owner_runtime(request, revision, contents)
    frozen = canonical_json_bytes(runtime[0])
    root = runtime[0]["works"][-1]["logicalWorkId"]
    for work in runtime[0]["works"][:-1]:
        seal_scope_work(runtime, work["logicalWorkId"])
    _, failed = seal_scope_work(runtime, root, {"invalid": True}, allow_failure=True)
    assert failed.failure_kind == "INVALID_IR"
    with pytest.raises(scope_compiler_module.ScopeInputRequired):
        scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], runtime[4])
    packet, succeeded = seal_scope_work(runtime, root, revision_number=2)
    assert succeeded.revision == 2
    assert any(ref["canonicalContent"].get("kind") == "ATTEMPT_REPAIR" for ref in packet["contextRefs"])
    material = scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], runtime[4])
    assert scope_compiler_module.validate_scope_candidate(material) == ()
    assert canonical_json_bytes(runtime[0]) == frozen
    restored = scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], runtime[4])
    assert restored == material


@pytest.mark.parametrize("kind,schema,definition", [("scan", "fact-decision", "sourceScanResult"), ("audit", "source-audit", "sourceAudit"), ("scope", "scope-decision", "scopeDecision")])
def test_scope_owner_independent_validation_rereads_changed_schema(tmp_path, monkeypatch, kind, schema, definition):
    import shutil
    shutil.copytree(SKILL_ROOT / "contracts", tmp_path / "contracts")
    monkeypatch.setattr(scope_compiler_module, "SKILL_ROOT", tmp_path)
    verify, packet, result = {"scan": (scope_compiler_module.verify_source_scan, scan_ir_packet(), scan_ir()),
        "audit": (scope_compiler_module.verify_source_audit, audit_ir_packet(), audit_ir()),
        "scope": (scope_compiler_module.verify_scope_decision, synthesis_packet(), complete_scope_ir())}[kind]
    verify(packet, result)
    path = tmp_path / "contracts" / (schema + ".schema.json")
    value = json.loads(path.read_bytes())
    value["$defs"][definition]["not"] = {}
    path.write_bytes(canonical_json_bytes(value))
    with pytest.raises(InvalidActionResult):
        verify(packet, result)


def test_scope_owner_rejects_one_empty_prior_among_authorized_workbooks(tmp_path):
    from prior_state import inventory_prior_workbook
    from test_prior_state import workbook_fixture
    import openpyxl
    filled, empty = tmp_path / "filled.xlsx", tmp_path / "empty.xlsx"
    workbook_fixture(filled)
    workbook = openpyxl.Workbook()
    workbook.save(empty)
    workbook.close()
    inventories = [inventory_prior_workbook(path) for path in (filled, empty)]
    request, revision, contents, _ = owner_input_case(inventories=inventories)
    with pytest.raises(scope_compiler_module.ScopeInputRequired):
        scope_compiler_module.prepare_scope_inputs(revision, contents, request=request, prior_inventories=inventories)


def change_scope_packet_result():
    packet, result = synthesis_packet(), complete_scope_ir()
    packet["contextRefs"].append({"refId": "dependency-result-prior", "canonicalContent": {"kind": "DEPENDENCY_RESULT",
        "logicalWorkId": "prior", "attemptRecordSha256": "c" * 64, "normalizedResult": {
            "entities": [{"localKey": key, "sourceId": "prior", "entityKind": "CONTRACT_ENTITY", "semanticSummary": "合同能力",
                "deliveryStatus": "CURRENT_BY_CONTRACT", "evidenceIds": ["d" * 64]} for key in ["p1", "p2"]],
            "sourceRelations": [], "entitySupersessions": [], "unsupportedRegions": [], "unextractedEvidence": []}}})
    return packet, result


def assert_scope_invalid_before_seal(packet, result):
    from action_ledger import ActionLedger, issue, finish
    from test_action_ledger import prepared_envelope, successful_completion
    from models import ActionEnvelope
    scope_compiler_module.validate_bound_scope_context("SCOPE_SYNTHESIS", packet)
    value = {**prepared_envelope().value, **scope_ir_envelope("SCOPE_SYNTHESIS"), "packetSha256": sha256_bytes(canonical_json_bytes(packet))}
    envelope = ActionEnvelope(value, "actions/cardinality/envelope.json", sha256_bytes(canonical_json_bytes(value)))
    ledger, record = finish(issue(ActionLedger(), envelope), envelope, successful_completion(canonical_json_bytes(result)),
        bound_result_validator=lambda normalized: scope_compiler_module.validate_bound_scope_result("SCOPE_SYNTHESIS", packet, normalized))
    assert record.failure_kind == "INVALID_IR"
    assert record.normalized_result_sha256 is None
    assert ledger.raw_outputs[record.raw_sha256] == canonical_json_bytes(result)


@pytest.mark.parametrize("kind,priors,targets", [("ADJUST", ["p1", "p2"], ["feature-a"]),
    ("REUSE_DEPENDENCY", ["p1"], ["feature-a", "epic-a"]), ("SPLIT", ["p1"], ["feature-a"]),
    ("MERGE", ["p1"], ["feature-a"]), ("MERGE", ["p1", "p2"], ["feature-a", "epic-a"])])
def test_scope_change_cardinality_is_invalid_ir_before_attempt_seal(kind, priors, targets):
    packet, result = change_scope_packet_result()
    feature = result["decisions"][0]
    feature["priorEntityIds"] = priors
    feature["relations"].append({"kind": kind, "targetLocalKeys": targets, "evidenceIds": ["block-a"]})
    assert_scope_invalid_before_seal(packet, result)


@pytest.mark.parametrize("defect", ["prior_overlap", "target_overlap", "retire_without_prior", "retire_without_history", "retire_without_current"])
def test_scope_change_closure_is_invalid_ir_before_attempt_seal(defect):
    packet, result = change_scope_packet_result()
    feature, epic = result["decisions"]
    if defect in {"prior_overlap", "target_overlap"}:
        feature["priorEntityIds"] = ["p1"]
        feature["relations"].append({"kind": "ADJUST", "targetLocalKeys": ["feature-a"], "evidenceIds": ["block-a"]})
        epic["priorEntityIds"] = ["p1"] if defect == "prior_overlap" else ["p2"]
        epic["relations"] = [{"kind": "ADJUST", "targetLocalKeys": ["feature-a"] if defect == "target_overlap" else ["epic-a"], "evidenceIds": ["block-a"]}]
    else:
        feature.update(decisionKind="RETIRE", priorEntityIds=[] if defect == "retire_without_prior" else ["p1"], relations=[], exclusionReason="本轮明确移除。")
        feature["boundaryEvidence"].update(classification="DISPOSITION", evidenceIds=["d" * 64] if defect == "retire_without_current" else ["block-a"])
    assert_scope_invalid_before_seal(packet, result)


def test_scope_design_authority_is_visible_to_model_and_rejected_before_seal():
    request, revision_bytes, contents, _ = owner_input_case(count=9)
    revision = json.loads(revision_bytes)
    revision["sources"][0].update(role="HLD", status="REFERENCE_ONLY")
    request["responsibilityBoundaries"][0]["responsibilityBoundaryId"] = "boundary-app"
    revision["requestSha256"] = sha256_bytes(canonical_json_bytes(request))
    runtime = scope_owner_runtime(request, canonical_json_bytes(revision), contents)
    for work in runtime[0]["works"][:-1]:
        seal_scope_work(runtime, work["logicalWorkId"])
    packet, record = seal_scope_work(runtime, runtime[0]["works"][-1]["logicalWorkId"], projected_scope_ir(), allow_failure=True)
    context = next(ref["canonicalContent"] for ref in packet["contextRefs"] if ref["canonicalContent"].get("kind") == "SCOPE_CONTEXT")
    assert context["sourceDirectory"][0]["status"] == "REFERENCE_ONLY"
    assert record.failure_kind == "INVALID_IR"
    assert record.normalized_result_sha256 is None


def test_scope_owner_e2e_retire_preserves_dual_evidence_and_bound_review_obligation(tmp_path):
    from prior_state import inventory_prior_workbook
    from test_prior_state import workbook_fixture
    from change_graph import derive_change_views
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path)
    inventory = inventory_prior_workbook(path)
    request, revision, contents, _ = owner_input_case(count=2, inventories=[inventory])
    runtime = scope_owner_runtime(request, revision, contents, inventories=[inventory])
    prior_item = next(item for item in runtime[1] if item.action_kind == "PRIOR_ANALYZE")
    for work in runtime[0]["works"]:
        result = None
        if work["packetPlan"]["actionKind"] == "SCOPE_SYNTHESIS":
            scans = [scan_ir(f"block-{index:03}", "fact")[0] for index in range(2)]
            result = scope_owner_result("SCOPE_SYNTHESIS", {"contextRefs": [{"canonicalContent": {"kind": "DEPENDENCY_RESULT", "normalizedResult": scans}}]})
            result["decisions"] = [item for item in result["decisions"] if item["localKey"] != "epic-block-001"]
            retiring = next(item for item in result["decisions"] if item["localKey"] == "feature-block-001")
            retiring.update(decisionKind="RETIRE", priorEntityIds=[prior_item.work_item_id + ":entity"], relations=[], exclusionReason="本轮来源明确移除该历史承诺。")
            retiring["boundaryEvidence"].update(classification="DISPOSITION", evidenceIds=["block-001", *prior_item.work_item_payload["evidenceIds"]])
        seal_scope_work(runtime, work["logicalWorkId"], result)
    material = scope_compiler_module.materialize_scope_candidate(revision, request, *runtime[:3], runtime[3], runtime[4], prior_inventories=[inventory])
    assert scope_compiler_module.validate_scope_candidate(material) == ()
    graph, prior = json.loads(material.change_graph_bytes), json.loads(material.prior_state_bytes)
    retired = graph["retiredPrior"][0]
    views = derive_change_views(graph, prior, list(material.identity_by_local_key.values()))
    assert views.prior_fate[retired["priorEntityId"]] == "RETIRE"
    obligation = next(item for item in material.review_obligations if item.get("relationKind") == "RETIRE")
    assert obligation["priorEntityIds"] == [retired["priorEntityId"]]
    assert obligation["targetEntityIds"] == []
    assert obligation["evidenceIds"] == retired["evidenceIds"]
    assert obligation["priorStateSha256"] == sha256_bytes(material.prior_state_bytes)
    assert obligation["priorRootAttemptRecordSha256"] in runtime[3].attempt_records


@pytest.mark.parametrize("kind", ["DESIGN", "POLICY"])
def test_scope_explicit_design_policy_links_cannot_disagree(kind):
    packet, result = synthesis_packet(), complete_scope_ir()
    context = {"kind": "SCOPE_CONTEXT", "sourceDirectory": [{"sourceId": "source", "role": "HLD", "status": "APPROVED", "blockIds": ["block-a"]}], "responsibilityBoundaries": []}
    packet["contextRefs"].append({"refId": "scope-context", "canonicalContent": context, "contentSha256": sha256_bytes(canonical_json_bytes(context))})
    feature = result["decisions"][0]
    other = copy.deepcopy(feature)
    other.update(localKey="feature-b", factIds=[])
    related = copy.deepcopy(feature)
    related.update(localKey="related", decisionKind="DESIGN_ITEM" if kind == "DESIGN" else "POLICY_INSTANCE", factIds=[])
    related["boundaryEvidence"]["classification"] = "COMPONENT" if kind == "DESIGN" else "policy-go-live"
    related["relations"] = [{"kind": "APPLIES_TO", "targetLocalKeys": ["feature-a", "feature-b"], "evidenceIds": ["block-a"]}]
    feature["relations"].append({"kind": kind, "targetLocalKeys": ["related"], "evidenceIds": ["block-a"]})
    result["decisions"].extend([other, related])
    scope_compiler_module.verify_scope_decision(packet, result)
    related["relations"][0]["targetLocalKeys"] = ["feature-b"]
    assert_scope_invalid_before_seal(packet, result)
