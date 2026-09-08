from __future__ import annotations

TEST_LAYER = "unit"

import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest


SKILL_ROOT = Path(__file__).parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from source_readers import (  # noqa: E402
    SourceReadError,
    extract_document,
    extract_source_blocks,
)
import source_readers as source_readers_module  # noqa: E402


def test_numbered_scope_items_keep_distinct_evidence_and_continuation(tmp_path):
    source = tmp_path / 'scope.md'
    source.write_text('# Scope\n\nApproved.\n\n1. Deliver upload.\n   At most 100 rows.\n2. Customer provides storage.\n3. Exclude migration.\n')
    document = extract_source_blocks(source, source_role='PRD', parser_version='2')
    items = [row for row in document.blocks if '/paragraph:' in row['locator']]
    assert [row['content'] for row in items] == ['Approved.', '1. Deliver upload.\n   At most 100 rows.', '2. Customer provides storage.', '3. Exclude migration.']
    assert len({row['locator'] for row in items}) == 4
    anchors = extract_document(source, source_id='prd', role='PRD')
    assert len([row for row in anchors if row.kind == 'PARAGRAPH']) == 4


def test_html_elements_exposes_inline_dependencies_without_executing_them():
    html = '<script type="module">import "./app.js";</script><style>p{background:url("pixel.png")}</style><p style="background:url(other.png)"><img srcset="one.png 1x, two.png 2x"></p>'
    elements = source_readers_module.html_elements(html)
    assert [(item["tag"], item.get("content")) for item in elements[:2]] == [
        ("script", 'import "./app.js";'), ("style", 'p{background:url("pixel.png")}')]
    assert elements[2]["attributes"]["style"] == "background:url(other.png)"
    assert elements[3]["attributes"]["srcset"] == "one.png 1x, two.png 2x"
    assert elements[2]["selector"] == "p:nth-of-type(1)"


def synthetic_source(tmp_path: Path, suffix: str) -> Path:
    path = tmp_path / f"source{suffix}"
    if suffix == ".md":
        path.write_text(
            "# Refund\n\nUser submits a refund request and sees the result.\n",
            encoding="utf-8",
        )
    elif suffix == ".txt":
        path.write_text(
            "Refund request\n\nUser sees the processing result.\n",
            encoding="utf-8",
        )
    elif suffix in {".html", ".htm"}:
        path.write_text(
            '<button id="refund">Submit refund</button>\n'
            '<p id="status">Waiting</p>\n'
            '<script>refund.onclick = () => { status.textContent = "Submitted"; };</script>\n',
            encoding="utf-8",
        )
    elif suffix in {".ts", ".tsx"}:
        path.write_text(
            "export function RefundButton() {\n"
            "  const [status, setStatus] = useState('waiting');\n"
            "  return <button onClick={() => setStatus('submitted')}>Refund</button>;\n"
            "}\n",
            encoding="utf-8",
        )
    elif suffix == ".xlsx":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Refund"
        sheet.append(["Feature", "Outcome"])
        sheet.append(["Refund request", "User sees the processing result"])
        workbook.save(path)
    else:
        raise AssertionError(suffix)
    return path


@pytest.mark.parametrize(
    ("role", "suffix"),
    [
        ("PRD", ".md"),
        ("HLD", ".md"),
        ("PRIOR_SOW", ".xlsx"),
        ("SUPPLEMENT", ".md"),
        ("SUPPLEMENT", ".txt"),
        ("SUPPLEMENT", ".html"),
        ("SUPPLEMENT", ".htm"),
        ("SUPPLEMENT", ".ts"),
        ("SUPPLEMENT", ".tsx"),
        ("SUPPLEMENT", ".xlsx"),
    ],
)
def test_supported_role_format_pairs_produce_non_page_non_line_anchors(
    tmp_path: Path, role: str, suffix: str
) -> None:
    anchors = extract_document(
        synthetic_source(tmp_path, suffix), source_id="source-main", role=role
    )
    assert anchors
    assert all(
        "page=" not in anchor.locator and "line=" not in anchor.locator
        for anchor in anchors
    )
    assert all(anchor.source_id == "source-main" for anchor in anchors)


@pytest.mark.parametrize(
    ("role", "suffix"),
    [
        ("PRD", ".txt"),
        ("PRD", ".html"),
        ("PRD", ".xlsx"),
        ("HLD", ".ts"),
        ("HLD", ".xlsx"),
        ("PRIOR_SOW", ".md"),
        ("PRIOR_SOW", ".pdf"),
        ("SUPPLEMENT", ".pdf"),
        ("SUPPLEMENT", ".docx"),
    ],
)
def test_unsupported_role_format_pairs_are_rejected(
    tmp_path: Path, role: str, suffix: str
) -> None:
    source = tmp_path / f"source{suffix}"
    source.write_bytes(b"synthetic source")
    with pytest.raises(SourceReadError) as captured:
        extract_document(source, source_id="source-main", role=role)
    assert captured.value.code == "SOURCE_FORMAT_UNSUPPORTED"


@pytest.mark.parametrize("suffix", [".html", ".ts", ".tsx"])
def test_prototype_source_preserves_function_and_interaction_evidence(
    tmp_path: Path, suffix: str
) -> None:
    anchors = extract_document(
        synthetic_source(tmp_path, suffix),
        source_id="prototype-demo",
        role="SUPPLEMENT",
    )
    text = " ".join(anchor.normalized_text for anchor in anchors)
    assert "Refund" in text or "refund" in text
    assert "Submitted" in text or "submitted" in text
    assert "button" in text


def test_same_text_under_moved_heading_keeps_semantic_hash(tmp_path: Path) -> None:
    before = tmp_path / "before.md"
    after = tmp_path / "after.md"
    before.write_text("## 原章节\n\n退款必须可追踪\n", encoding="utf-8")
    after.write_text("## 新章节\n\n退款必须可追踪\n", encoding="utf-8")

    def semantic_anchor(path: Path):
        return next(
            anchor
            for anchor in extract_document(path, source_id="prd-main", role="PRD")
            if anchor.normalized_text == "退款必须可追踪"
        )

    assert semantic_anchor(before).sha256 == semantic_anchor(after).sha256


def test_duplicate_semantic_anchors_get_ordered_suffixes(tmp_path: Path) -> None:
    source = tmp_path / "duplicate.txt"
    source.write_text("相同结果\n\n相同结果\n", encoding="utf-8")
    anchors = extract_document(source, source_id="supplement-main", role="SUPPLEMENT")
    assert [anchor.anchor_id[-5:] for anchor in anchors] == ["-0001", "-0002"]
    assert anchors[0].sha256 == anchors[1].sha256


@pytest.mark.parametrize(
    ("name", "payload", "expected_code"),
    [
        ("blank.md", b" \n", "SOURCE_BLANK"),
        ("placeholder.md", b"# TODO\n\nTBD\n", "SOURCE_PLACEHOLDER_ONLY"),
        ("invalid.ts", b"\xff\xfe\x00", "SOURCE_UNREADABLE"),
        ("binary.html", b"valid\x00text", "SOURCE_UNREADABLE"),
    ],
)
def test_invalid_text_sources_return_stable_codes(
    tmp_path: Path, name: str, payload: bytes, expected_code: str
) -> None:
    source = tmp_path / name
    source.write_bytes(payload)
    role = "PRD" if source.suffix == ".md" else "SUPPLEMENT"
    with pytest.raises(SourceReadError) as captured:
        extract_document(source, source_id="source-main", role=role)
    assert captured.value.code == expected_code


def test_source_blocks_preserve_lossless_locators_structure_and_dropped_categories(
    tmp_path: Path,
) -> None:
    markdown = tmp_path / "design.md"
    markdown.write_text(
        "# 系统上下文\r\n\r\n订单由门户提交。\r\n\r\n"
        "## Integration\r\n\r\n| 来源 | 目标 |\r\n|---|---|\r\n| 门户 | 订单服务 |\r\n",
        encoding="utf-8",
        newline="",
    )

    document = extract_source_blocks(
        markdown,
        source_role="HLD",
        parser_version="1",
    )

    assert document.parser_id == "markdown-blocks"
    assert document.parser_version == "1"
    assert document.raw_sha256 == __import__("hashlib").sha256(
        markdown.read_bytes()
    ).hexdigest()
    assert document.blocks
    assert all("locator" in block and "content" in block for block in document.blocks)
    assert all("\r" not in str(block["content"]) for block in document.blocks)
    assert any(block["structuralParentId"] is not None for block in document.blocks)
    assert all(block["primaryCoverageBlockId"] for block in document.blocks)
    assert all("contextBlockIds" in block for block in document.blocks)
    assert all("droppedContentCategories" in block for block in document.blocks)

    prototype = tmp_path / "demo.html"
    prototype.write_text(
        "<main><button>提交</button><script>secret()</script><!-- note --></main>",
        encoding="utf-8",
    )
    demo = extract_source_blocks(
        prototype,
        source_role="DEMO",
        parser_version="1",
    )
    assert any(
        block["extractionDisposition"] == "DROPPED"
        and "SCRIPT" in block["droppedContentCategories"]
        for block in demo.blocks
    )
    assert any(
        block["extractionDisposition"] == "DROPPED"
        and "COMMENT" in block["droppedContentCategories"]
        for block in demo.blocks
    )


def test_xlsx_zip_expansion_limit_has_stable_failure_code(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = synthetic_source(tmp_path, ".xlsx")
    monkeypatch.setattr(source_readers_module, "MAX_XLSX_UNCOMPRESSED_BYTES", 100)

    with pytest.raises(SourceReadError) as captured:
        extract_source_blocks(source, source_role="PRIOR_SOW", parser_version="1")

    assert captured.value.code == "SOURCE_LIMIT_EXCEEDED"


def test_xlsx_declared_dimension_and_cell_text_limits_are_bounded(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    huge_dimension = tmp_path / "huge-dimension.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "有效内容"
    sheet.cell(
        row=source_readers_module.MAX_XLSX_ROWS_PER_SHEET + 1,
        column=1,
        value="越界",
    )
    workbook.save(huge_dimension)
    with pytest.raises(SourceReadError) as captured:
        extract_document(
            huge_dimension, source_id="prior-sow", role="PRIOR_SOW"
        )
    assert captured.value.code == "SOURCE_LIMIT_EXCEEDED"

    long_cell = synthetic_source(tmp_path, ".xlsx")
    monkeypatch.setattr(source_readers_module, "MAX_XLSX_CELL_TEXT_CHARS", 5)
    with pytest.raises(SourceReadError) as captured:
        extract_source_blocks(long_cell, source_role="PRIOR_SOW", parser_version="1")
    assert captured.value.code == "SOURCE_LIMIT_EXCEEDED"


EMPTY_OFFICE_DRAWING = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"></xdr:wsDr>'
)


def source_with_drawing(tmp_path, drawing):
    base = synthetic_source(tmp_path, ".xlsx")
    source = tmp_path / "with-drawing.xlsx"
    with zipfile.ZipFile(base) as original, zipfile.ZipFile(source, "w") as target:
        for item in original.infolist():
            target.writestr(item, original.read(item.filename))
        target.writestr("xl/drawings/drawing1.xml", drawing.encode("utf-8"))
    return base, source


def require_prior_surfaces_readable(inventory):
    from prior_state import _require_ready
    _require_ready({"unsupportedRegions": [], "sourceRelations": [],
                    "entities": [], "entitySupersessions": []}, [inventory])


@pytest.mark.parametrize("drawing", [
    EMPTY_OFFICE_DRAWING,
    EMPTY_OFFICE_DRAWING.replace("></xdr:wsDr>", ">\n  \t</xdr:wsDr>"),
    '<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"/>',
], ids=["office-299-bytes", "whitespace", "default-namespace"])
def test_empty_office_drawing_does_not_block_prior_or_change_cell_evidence(tmp_path, drawing):
    base, source = source_with_drawing(tmp_path, drawing)
    before = source.read_bytes()
    inventory = source_readers_module.inventory_xlsx(source)
    assert len(EMPTY_OFFICE_DRAWING.encode("utf-8")) == 299
    assert inventory["unsupportedSurfaces"] == []
    assert inventory["sheets"] == source_readers_module.inventory_xlsx(base)["sheets"]
    require_prior_surfaces_readable(inventory)
    documents = [extract_source_blocks(path, source_role="PRIOR_SOW", parser_version="2")
                 for path in (base, source)]
    assert [(block["locator"], block["content"]) for block in documents[0].blocks] == [
        (block["locator"], block["content"]) for block in documents[1].blocks]
    assert source.read_bytes() == before


@pytest.mark.parametrize("drawing", [
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>',
        '<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:ext cx="100000" cy="100000"/>'
        '<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="1" name="Scope note"/><xdr:cNvSpPr/>'
        '</xdr:nvSpPr><xdr:spPr/><xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r>'
        '<a:t>合同范围</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/>'
        '</xdr:absoluteAnchor></xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', '<xdr:oneCellAnchor/></xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', '<xdr:twoCellAnchor/></xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', '<xdr:unknown/></xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', '合同范围</xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', '\u00a0</xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('<xdr:wsDr ', '<xdr:wsDr unknown="value" '),
    EMPTY_OFFICE_DRAWING.replace('2006/spreadsheetDrawing', '2006/unknown'),
    '<wsDr/>',
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', ''),
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', '<!--未知内容--></xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('</xdr:wsDr>', '<?unknown content?></xdr:wsDr>'),
    EMPTY_OFFICE_DRAWING.replace('<xdr:wsDr ', '<!DOCTYPE xdr:wsDr><xdr:wsDr '),
    EMPTY_OFFICE_DRAWING.replace('encoding="UTF-8"', 'encoding="unknown-drawing-encoding"'),
], ids=["shape", "one-cell-anchor", "two-cell-anchor", "unknown-child", "text",
        "non-xml-whitespace", "attribute", "wrong-namespace", "missing-namespace", "malformed", "comment",
        "processing-instruction", "doctype", "unknown-encoding"])
def test_nonempty_unknown_or_malformed_drawing_still_blocks_prior(tmp_path, drawing):
    from prior_state import PriorInputRequired
    _base, source = source_with_drawing(tmp_path, drawing)
    before = source.read_bytes()
    inventory = source_readers_module.inventory_xlsx(source)
    assert inventory["unsupportedSurfaces"] == [{
        "part": "xl/drawings/drawing1.xml", "kind": "DRAWING", "coverage": "UNSUPPORTED",
    }]
    with pytest.raises(PriorInputRequired):
        require_prior_surfaces_readable(inventory)
    assert source.read_bytes() == before


def test_array_formula_source_blocks_preserve_formula_and_repeat_exactly(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    from contracts import canonical_json_bytes
    source=tmp_path/'prior-array.xlsx'
    workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='Stories'
    sheet.append(['Story','Formula']);sheet.append(['STORY-1',None])
    formula='=IF(A2="","",SUM(C2:D2))'
    sheet['B2']=ArrayFormula(ref='B2',text=formula)
    workbook.save(source);workbook.close()
    original=source.read_bytes()
    first=extract_source_blocks(source,source_role='PRIOR_SOW',parser_version='2')
    second=extract_source_blocks(source,source_role='PRIOR_SOW',parser_version='2')
    row=next(block for block in first.blocks if block['locator'].endswith('/row:000002'))
    assert row['content']=='STORY-1 | '+formula
    assert canonical_json_bytes(first.blocks)==canonical_json_bytes(second.blocks)
    assert source.read_bytes()==original
    inventory=source_readers_module.inventory_xlsx(source)
    cell=next(cell for cell in inventory['sheets'][0]['cells'] if cell['address']=='$B$2')
    assert cell['formula']==formula[1:]


def test_unreadable_array_formula_never_becomes_python_object_repr():
    from openpyxl.worksheet.formula import ArrayFormula
    with pytest.raises(SourceReadError,match='公式'):
        source_readers_module._xlsx_scalar(ArrayFormula(ref='B2',text=None))


def test_data_table_formula_source_rejects_unreadable_formula_without_object_repr(tmp_path):
    from openpyxl.worksheet.formula import DataTableFormula
    source=tmp_path/'prior-data-table.xlsx'
    workbook=openpyxl.Workbook();sheet=workbook.active
    sheet['A1']=1;sheet['B1']=DataTableFormula(ref='B1:C2',r1='A1')
    workbook.save(source);workbook.close()
    original=source.read_bytes()
    with pytest.raises(SourceReadError,match='公式') as error:
        extract_source_blocks(source,source_role='PRIOR_SOW',parser_version='2')
    assert error.value.code=='SOURCE_UNREADABLE'
    assert source.read_bytes()==original
