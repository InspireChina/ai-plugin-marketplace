from __future__ import annotations

import copy
import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


SKILL_ROOT = Path(__file__).parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
TEMPLATE = SKILL_ROOT / "assets" / "sow-template.xlsx"
sys.path.insert(0, str(SCRIPTS))

from task_standard_catalog import (  # noqa: E402
    CatalogContractError,
    catalog,
    compact_index,
    hydrate,
    row_semantic_sha256,
)


def _workbook_copy(tmp_path: Path, name: str) -> Path:
    target = tmp_path / name
    target.write_bytes(TEMPLATE.read_bytes())
    return target


def _mutate_cell(path: Path, address: str, value: object) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        workbook["90-估算标准"][address] = value
        workbook.save(path)
    finally:
        workbook.close()


def test_row_semantic_hash_canonicalizes_unicode_newline_decimal_null_and_neighbors() -> None:
    source = catalog(TEMPLATE).rows[0]
    left = copy.deepcopy(dict(source))
    right = copy.deepcopy(dict(source))
    left["分类"] = "Cafe\u0301"
    right["分类"] = "Café"
    left["说明"] = "  第一行\r\n第二行  "
    right["说明"] = "第一行\n第二行"
    left["新建M档人天"] = Decimal("1.000")
    right["新建M档人天"] = 1
    left["相邻工作类型IDs"] = ["FE-VIEW", "AN-DESIGN", "FE-VIEW"]
    right["相邻工作类型IDs"] = ["AN-DESIGN", "FE-VIEW"]
    left["接入复用M档人天"] = None
    right["接入复用M档人天"] = None

    assert row_semantic_sha256(left) == row_semantic_sha256(right)


def test_style_only_change_preserves_catalog_semantics_but_changes_template_hash(
    tmp_path: Path,
) -> None:
    changed_path = _workbook_copy(tmp_path, "style.xlsx")
    workbook = load_workbook(changed_path, data_only=False, read_only=False)
    try:
        workbook["90-估算标准"]["A1"].fill = PatternFill("solid", fgColor="7030A0")
        workbook.save(changed_path)
    finally:
        workbook.close()

    baseline = catalog(TEMPLATE)
    changed = catalog(changed_path)
    assert changed.template_sha256 != baseline.template_sha256
    assert changed.task_catalog_semantic_sha256 == baseline.task_catalog_semantic_sha256
    assert [row["rowSemanticSha256"] for row in changed.rows] == [
        row["rowSemanticSha256"] for row in baseline.rows
    ]


def test_one_row_semantic_change_changes_only_that_row_and_catalog_hash(
    tmp_path: Path,
) -> None:
    changed_path = _workbook_copy(tmp_path, "semantic.xlsx")
    _mutate_cell(changed_path, "I5", "语义变化后的说明")

    baseline = catalog(TEMPLATE)
    changed = catalog(changed_path)
    assert changed.task_catalog_semantic_sha256 != baseline.task_catalog_semantic_sha256
    changed_ids = {
        before["工作类型ID"]
        for before, after in zip(baseline.rows, changed.rows, strict=True)
        if before["rowSemanticSha256"] != after["rowSemanticSha256"]
    }
    assert changed_ids == {"AN-IMPACT"}


def test_compact_index_exposes_all_88_rows_without_full_rules() -> None:
    result = compact_index(catalog(TEMPLATE))
    assert len(result) == 88
    assert set(result[0]) == {
        "工作类型ID",
        "分类",
        "工作类型名称",
        "计量单位",
        "标准交付物",
        "适用模式",
        "相邻工作类型IDs",
        "SIT支持资格",
        "rowSemanticSha256",
    }
    assert {row["工作类型ID"] for row in result} == {
        row["工作类型ID"] for row in catalog(TEMPLATE).rows
    }
    for forbidden in ("S标准", "M标准", "L标准", "X/拆分条件", "说明"):
        assert all(forbidden not in row for row in result)


def test_hydration_includes_selected_neighbors_and_cross_category_challenger() -> None:
    source = catalog(TEMPLATE)
    result = hydrate(
        source,
        selected_work_type_ids=("FE-FLOW",),
        query="跨系统订单流程、接口映射与端到端业务结果",
    )

    assert result.selected_work_type_ids == ("FE-FLOW",)
    assert result.neighbor_work_type_ids == ("FE-EDIT", "FE-VIEW")
    assert result.challenger_work_type_ids
    selected_category = source.by_work_type_id["FE-FLOW"]["分类"]
    assert all(
        source.by_work_type_id[work_type_id]["分类"] != selected_category
        for work_type_id in result.challenger_work_type_ids
    )
    hydrated_ids = tuple(row["工作类型ID"] for row in result.rows)
    assert hydrated_ids[:3] == ("FE-FLOW", "FE-EDIT", "FE-VIEW")
    assert result.evidence["tokenizerVersion"] == "task-catalog-tokenizer-v1"
    scores = result.evidence["candidateScores"]
    assert scores
    assert [item["order"] for item in scores] == list(range(1, len(scores) + 1))
    assert list(scores) == sorted(
        scores, key=lambda item: (-item["score"], item["workTypeId"])
    )


@pytest.mark.parametrize(
    ("address", "value", "message"),
    [
        ("Y5", "UNKNOWN-TYPE", "UNKNOWN_NEIGHBOR"),
        ("Y5", "AN-IMPACT", "SELF_NEIGHBOR"),
        ("Y5", "AN-DESIGN,AN-DESIGN", "DUPLICATE_NEIGHBOR"),
        ("AB5", "MAYBE", "INVALID_SIT_ELIGIBILITY"),
        ("AB5", "PER_INTEGRATION", "INVALID_SIT_OWNER_TYPE"),
    ],
)
def test_unknown_neighbor_self_reference_duplicate_and_invalid_sit_eligibility_fail_closed(
    tmp_path: Path,
    address: str,
    value: str,
    message: str,
) -> None:
    changed_path = _workbook_copy(tmp_path, f"{message}.xlsx")
    _mutate_cell(changed_path, address, value)

    with pytest.raises(CatalogContractError, match=message) as raised:
        catalog(changed_path)
    assert raised.value.diagnostic.code == "CONTRACT_UNSUPPORTED"
