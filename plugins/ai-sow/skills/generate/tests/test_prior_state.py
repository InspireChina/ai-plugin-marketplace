from __future__ import annotations

TEST_LAYER = "integration"

import copy
import json
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

SKILL_ROOT = Path(__file__).parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
sys.path.insert(0, str(SKILL_ROOT / "tests"))
from contracts import canonical_json_bytes, sha256_bytes
from prior_state import inventory_prior_workbook


def revision_for(inventories, source_ids=None, date="2026-10-01"):
    from test_contracts import valid_input_revision
    value = valid_input_revision()
    value["project"]["plannedEffectiveDate"] = date
    value["sources"] = [{"sourceId": source_id, "role": "PRIOR_SOW", "status": "APPLICABLE", "path": f"sources/{source_id}.xlsx", "rawSha256": inventory["workbookSha256"], "parserId": "xlsx-inventory", "parserVersion": "1", "blockIds": []} for source_id, inventory in zip(source_ids or [f"source-{i}" for i in range(len(inventories))], inventories)]
    value["blocks"] = []
    value["priorSowState"] = "PROVIDED" if inventories else "NOT_PROVIDED"
    value["priorSowSha256s"] = sorted({inventory["workbookSha256"] for inventory in inventories})
    return canonical_json_bytes(value)


def decision_for(inventories, source_ids=None):
    return {"entities": [{"localKey": f"item-{i}:entity", "sourceId": source_id, "entityKind": "CONTRACT_ENTITY", "semanticSummary": "查询订单", "deliveryStatus": "CURRENT_BY_CONTRACT", "evidenceIds": [inventory["evidence"][1]["priorEvidenceId"]]} for i, (source_id, inventory) in enumerate(zip(source_ids or [f"source-{j}" for j in range(len(inventories))], inventories))], "sourceRelations": [], "entitySupersessions": [], "unsupportedRegions": []}


def test_prior_read_only_snapshot_binds_explicit_revision_sources_and_exact_shape(tmp_path):
    from prior_state import verify_prior_decision, materialize_prior_snapshot
    from contracts import validate_contract, load_schema_registry
    path = tmp_path / "arbitrary-name.xlsx"
    workbook_fixture(path)
    before = (path.stat().st_size, sha256_bytes(path.read_bytes()))
    inventory = inventory_prior_workbook(path)
    revision = revision_for([inventory], ["authorized-source"])
    decision = decision_for([inventory], ["authorized-source"])
    verify_prior_decision([inventory], decision, input_revision_bytes=revision)
    snapshot = materialize_prior_snapshot([inventory], decision, input_revision_bytes=revision)
    assert set(snapshot) == {"contractVersion", "inputRevisionSha256", "evidence", "entities", "sourceRelations", "entitySupersessions"}
    assert snapshot["contractVersion"] == "prior-state-snapshot-v1"
    assert snapshot["inputRevisionSha256"] == sha256_bytes(revision)
    assert len(snapshot["entities"]) == 1
    entity = snapshot["entities"][0]
    assert set(entity) == {"entityId", "sourceId", "entityKind", "semanticSummary", "deliveryStatus", "evidenceIds"}
    assert entity["sourceId"] == "authorized-source"
    assert entity["entityId"].startswith("prior-")
    assert snapshot["evidence"] == [{"sourceId": "authorized-source", **inventory["evidence"][1]}]
    assert validate_contract(snapshot, "prior-state-snapshot.schema.json", load_schema_registry(SKILL_ROOT)) == ()
    assert (path.stat().st_size, sha256_bytes(path.read_bytes())) == before


@pytest.mark.parametrize("mutation", ["DEMO", "SUPPLEMENT", "missing_inventory", "wrong_source", "wrong_evidence", "duplicate_local_key", "observation", "evidence_hash"])
def test_prior_downstream_boundary_rejects_unauthorized_source_and_demo_observation(tmp_path, mutation):
    from prior_state import verify_prior_decision, materialize_prior_snapshot
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path)
    inventories = [inventory_prior_workbook(path)]
    revision = json.loads(revision_for(inventories))
    decision = decision_for(inventories)
    if mutation in {"DEMO", "SUPPLEMENT"}:
        revision["sources"][0]["role"] = mutation
    elif mutation == "missing_inventory":
        inventories = []
    elif mutation == "wrong_source":
        decision["entities"][0]["sourceId"] = "unselected-source"
    elif mutation == "wrong_evidence":
        decision["entities"][0]["evidenceIds"] = ["f" * 64]
    elif mutation == "duplicate_local_key":
        decision["entities"].append(copy.deepcopy(decision["entities"][0]))
    elif mutation == "evidence_hash":
        inventories[0]["evidence"][1]["canonicalCellValuesSha256"] = "f" * 64
    else:
        from test_prototype_analysis import demo_files, observation_fixture, scenario_fixture, trace_fixture
        from prototype_analysis import inventory_demo_bundle, verify_prototype_observations
        demo_inventory = inventory_demo_bundle("demo/index.html", demo_files())
        observation = observation_fixture(demo_inventory)
        # A real, valid Demo observation is target-scope evidence, never prior state.
        decision = {"observations": [observation]}
        from contracts import validate_contract, load_schema_registry
        assert validate_contract(decision, "prototype-observation.schema.json", load_schema_registry(SKILL_ROOT)) == ()
        scenario = scenario_fixture(demo_inventory)
        assert verify_prototype_observations(demo_inventory, scenario, trace_fixture(demo_inventory, scenario), decision)["authority"] == "TARGET_SCOPE_ONLY"
    for operation in (verify_prior_decision, materialize_prior_snapshot):
        with pytest.raises(ValueError):
            operation(inventories, decision, input_revision_bytes=canonical_json_bytes(revision))


def test_effective_prior_projection_date_context_uses_exact_immutable_revision():
    from prior_state import build_project_effective_start_context
    from test_contracts import valid_input_revision
    revision = valid_input_revision()
    revision["project"]["plannedEffectiveDate"] = "2026-10-01"
    payload = canonical_json_bytes(revision)
    context = build_project_effective_start_context(payload)
    assert context.ref_id == "PROJECT_EFFECTIVE_START"
    assert json.loads(context.canonical_content) == {"kind": "PROJECT_EFFECTIVE_START", "inputRevisionSha256": sha256_bytes(payload), "plannedEffectiveDate": "2026-10-01"}
    for invalid in [b'{}\n', payload.rstrip(), payload.replace(b'2026-10-01', b'2026-02-30')]:
        with pytest.raises(ValueError):
            build_project_effective_start_context(invalid)


def workbook_fixture(path, layout="plain"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "合同"
    sheet.append(["ID", "交付物", "数量"])
    sheet.append(["prior-visible-1", "查询订单", 2])
    sheet.append(["prior-visible-1", "引用订单", "=C2+1"])
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions["C"].hidden = True
    if layout == "table":
        sheet.add_table(Table(displayName="PriorContracts", ref="A1:C3"))
    if layout == "merged":
        second = workbook.create_sheet("补充")
        second.merge_cells("A1:C1")
        second["A1"] = "交付说明"
        second["A2"] = "生产上线"
    workbook.save(path)
    workbook.close()


def prior_pair(tmp_path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    workbook_fixture(first)
    workbook_fixture(second, "merged")
    inventories = [inventory_prior_workbook(first), inventory_prior_workbook(second)]
    return inventories, revision_for(inventories), decision_for(inventories)

@pytest.mark.parametrize('defect', [None, 'other_sheet', 'other_source', 'no_owner_anchor', 'legacy'])
def test_prior_row_partition_does_not_split_one_authorized_entity(tmp_path, defect):
    from prior_state import build_project_effective_start_context, validate_bound_prior_result
    path=tmp_path/'prior.xlsx'
    workbook_fixture(path, 'merged')
    inventory=inventory_prior_workbook(path)
    revision=revision_for([inventory])
    context=build_project_effective_start_context(revision)
    own=inventory['evidence'][1]
    adjacent=inventory['evidence'][2 if defect!='other_sheet' else 3]
    items=[]
    for index,row in enumerate([own,adjacent]):
        items.append({'workItemId':f'item-{index}', 'payload':{
            'priorInputLayout':'ai-sow-prior-row-partition-v1',
            'sourceId':'source-0','sheet':{'sheet':row['sheet']},
            'evidenceIds':[row['priorEvidenceId']], 'evidence':[row]}})
    if defect=='other_source': items[1]['payload']['sourceId']='unselected-source'
    if defect=='legacy':
        for item in items:item['payload'].pop('priorInputLayout')
    packet={'workItems':items,'contextRefs':[{'refId':context.ref_id,
        'canonicalContent':json.loads(context.canonical_content),'contentSha256':sha256_bytes(context.canonical_content)}]}
    decision=decision_for([inventory])
    decision['entities'][0]['evidenceIds']=[own['priorEvidenceId'],adjacent['priorEvidenceId']]
    if defect=='no_owner_anchor':decision['entities'][0]['evidenceIds']=[adjacent['priorEvidenceId']]
    if defect:
        with pytest.raises(ValueError):
            validate_bound_prior_result('PRIOR_ANALYZE',packet,canonical_json_bytes(decision),
                inventories=[inventory],input_revision_bytes=revision)
    else:
        validate_bound_prior_result('PRIOR_ANALYZE',packet,canonical_json_bytes(decision),
            inventories=[inventory],input_revision_bytes=revision)



def analyze_packet(inventories, revision):
    from prior_state import build_project_effective_start_context
    from stage_planner import AtomicWorkItemDescriptor, make_planned_work, plan_stage, materialize_packet
    from action_ledger import ActionLedger
    from test_stage_planner import policy
    items = [AtomicWorkItemDescriptor(f"item-{i}", "PRIOR_ANALYZE", "PRIOR_SOW", inventory["workbookSha256"], i, {"sourceId": f"source-{i}", "evidenceIds": [item["priorEvidenceId"] for item in inventory["evidence"]]}) for i, inventory in enumerate(inventories)]
    contexts = [build_project_effective_start_context(revision)]
    descriptors = [make_planned_work("PRIOR_ANALYZE", items, contexts, [])]
    # These historical fixtures keep the v1 contract; v2 is exercised separately.
    from unittest.mock import patch
    with patch("stage_planner.current_action_contract_id", lambda kind: kind + "-v1"):
        plan = plan_stage("SCOPE", items, contexts, descriptors, [], policy())
    packet = materialize_packet(plan, plan["works"][0]["logicalWorkId"], 1, items, contexts, [], ActionLedger())
    return plan, json.loads(packet)


def prior_plan_case(tmp_path, count, *, consolidate=True, replacement=False, packed=False, unsupported_source=False):
    from prior_state import build_project_effective_start_context, validate_bound_prior_context, validate_bound_prior_result
    from stage_planner import AtomicWorkItemDescriptor, make_planned_work, plan_stage, materialize_packet, DependencyResultRef
    from action_ledger import ActionLedger, issue, finish, attempt_record_value
    from test_stage_planner import policy, envelope_for_plan
    from test_action_ledger import successful_completion
    paths, inventories, before = [], [], []
    for i, layout in enumerate(["plain", "table", "merged"][:count]):
        path = tmp_path / f"prior-{i}.xlsx"
        workbook_fixture(path, layout)
        paths.append(path)
        before.append((path.stat().st_size, sha256_bytes(path.read_bytes())))
        inventories.append(inventory_prior_workbook(path))
    revision = revision_for(inventories)
    decision = decision_for(inventories)
    if packed:
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "COMPLEMENTARY", "evidenceIds": [inventories[0]["evidence"][0]["priorEvidenceId"]]}]
    if unsupported_source:
        decision["entities"] = [item for item in decision["entities"] if item["sourceId"] != "source-1"]
        decision["sourceRelations"] = []
        decision["unsupportedRegions"] = [{"sourceId": "source-1", "regionId": "region-1", "reason": "等待输入"}]
    items = [AtomicWorkItemDescriptor(f"item-{i}", "PRIOR_ANALYZE", "PRIOR_SOW", inventory["workbookSha256"], i, {"sourceId": f"source-{i}", "evidenceIds": [item["priorEvidenceId"] for item in inventory["evidence"]]}) for i, inventory in enumerate(inventories)]
    contexts = [build_project_effective_start_context(revision)] if count else []
    descriptors = [make_planned_work("PRIOR_ANALYZE", [item], contexts, []) for item in items]
    if packed:
        descriptors = [make_planned_work("PRIOR_ANALYZE", items[:2], contexts, []), make_planned_work("PRIOR_ANALYZE", items[2:], contexts, [])]
    if count > 1 and consolidate:
        descriptors.append(make_planned_work("PRIOR_CONSOLIDATE", [], contexts, [item.work_key for item in descriptors]))
    # These historical fixtures keep the v1 contract; v2 is exercised separately.
    from unittest.mock import patch
    with patch("stage_planner.current_action_contract_id", lambda kind: kind + "-v1"):
        plan = plan_stage("SCOPE", items, contexts, descriptors, [], policy())
    ledger, refs, packets = ActionLedger(), {}, {}
    for work in plan["works"]:
        logical_id, packet_plan = work["logicalWorkId"], work["packetPlan"]
        deps = [refs[key] for key in packet_plan["dependencyLogicalWorkIds"]]
        payload = materialize_packet(plan, logical_id, 1, items, contexts, deps, ledger)
        packet = json.loads(payload)
        kind = packet_plan["actionKind"]
        result = copy.deepcopy(decision)
        if kind == "PRIOR_ANALYZE":
            namespaces = tuple(item["workItemId"] + ":" for item in packet["workItems"])
            result["entities"] = [entity for entity in result["entities"] if entity["localKey"].startswith(namespaces)]
            source_ids = {item["payload"]["sourceId"] for item in packet["workItems"]}
            result["sourceRelations"] = [relation for relation in result["sourceRelations"] if {relation["sourceAId"], relation["sourceBId"]} <= source_ids]
            result["unsupportedRegions"] = [region for region in result["unsupportedRegions"] if region["sourceId"] in source_ids]
        elif replacement:
            result["entitySupersessions"] = [{"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-1:entity"], "evidenceIds": decision["entities"][1]["evidenceIds"]}]
        validate_bound_prior_context(kind, packet, inventories=inventories, input_revision_bytes=revision)
        envelope = envelope_for_plan(plan, logical_id)
        ledger, record = finish(issue(ledger, envelope), envelope, successful_completion(canonical_json_bytes(result)), bound_result_validator=lambda raw: validate_bound_prior_result(kind, packet, raw, inventories=inventories, input_revision_bytes=revision))
        assert record.outcome == "SUCCEEDED"
        refs[logical_id] = DependencyResultRef(logical_id, sha256_bytes(canonical_json_bytes(attempt_record_value(record))), ledger.normalized_results[record.normalized_result_sha256])
        packets[logical_id] = packet
    assert [(path.stat().st_size, sha256_bytes(path.read_bytes())) for path in paths] == before
    return inventories, revision, plan, ledger, refs, packets, items, contexts


@pytest.mark.parametrize("count", [0, 1, 3])
def test_prior_read_only_zero_one_many_resolve_unique_effective_root(tmp_path, count):
    from prior_state import resolve_prior_root, materialize_prior_snapshot
    inventories, revision, plan, ledger, refs, *_ = prior_plan_case(tmp_path, count)
    root_ref = list(refs.values())[-1:] if count else []
    root = resolve_prior_root(plan, ledger, root_ref)
    assert root == (root_ref[0] if count else None)
    assert len(plan["works"]) == {0: 0, 1: 1, 3: 4}[count]
    if count:
        snapshot = materialize_prior_snapshot(inventories, json.loads(root.normalized_result), input_revision_bytes=revision)
        assert len(snapshot["entities"]) == count
    else:
        with pytest.raises(ValueError):
            materialize_prior_snapshot([], decision_for([]), input_revision_bytes=revision)
    assert {path.name for path in tmp_path.iterdir()} == {f"prior-{i}.xlsx" for i in range(count)}


@pytest.mark.parametrize("mutation", ["missing", "tampered", "logical", "extra", "no_consolidation", "failed_root", "superseded_root", "failed_leaf", "wrong_kind", "unknown_dependency", "analyze_root"])
def test_prior_multi_file_relations_root_rejects_unproven_or_wrong_topology(tmp_path, mutation):
    from dataclasses import replace
    from prior_state import resolve_prior_root
    from action_ledger import attempt_record_value
    from contracts import action_contract_binding
    inventories, revision, plan, ledger, refs, *_ = prior_plan_case(tmp_path, 3, consolidate=mutation != "no_consolidation", replacement=True)
    root = list(refs.values())[-1]
    provided = [root]
    if mutation == "missing":
        provided = []
    elif mutation == "tampered":
        provided = [replace(root, normalized_result=b"{}\n")]
    elif mutation == "logical":
        provided = [replace(root, logical_work_id="logical-wrong")]
    elif mutation == "extra":
        provided.append(list(refs.values())[0])
    elif mutation in {"failed_root", "superseded_root", "failed_leaf"}:
        selected = list(refs.values())[0] if mutation == "failed_leaf" else root
        record = ledger.attempt_records[selected.result_sha256]
        record = replace(record, outcome="SUPERSEDED" if mutation == "superseded_root" else "FAILED")
        records = {key: item for key, item in ledger.attempt_records.items() if key != selected.result_sha256}
        records[sha256_bytes(canonical_json_bytes(attempt_record_value(record)))] = record
        ledger = replace(ledger, attempt_records=records)
    elif mutation == "wrong_kind":
        record = ledger.attempt_records[root.result_sha256]
        old = ledger.envelopes_by_sha256[record.envelope_sha256]
        value = {**old.value, "actionContractId": "PRIOR_ANALYZE-v1"}
        value["actionContractSha256"] = action_contract_binding(SKILL_ROOT, value["actionContractId"])[1]
        envelope = replace(old, value=value, sha256=sha256_bytes(canonical_json_bytes(value)))
        record = replace(record, envelope_sha256=envelope.sha256)
        digest = sha256_bytes(canonical_json_bytes(attempt_record_value(record)))
        records = {key: item for key, item in ledger.attempt_records.items() if key != root.result_sha256}
        records[digest] = record
        envelopes = {key: item for key, item in ledger.envelopes_by_sha256.items() if key != old.sha256}
        envelopes[envelope.sha256] = envelope
        ledger = replace(ledger, envelopes_by_sha256=envelopes, attempt_records=records)
        provided = [replace(root, attempt_record_sha256=digest)]
    elif mutation == "unknown_dependency":
        plan["works"][-1]["packetPlan"]["dependencyLogicalWorkIds"].append("logical-unknown")
    elif mutation == "analyze_root":
        plan["works"][-1]["packetPlan"]["actionKind"] = "PRIOR_ANALYZE"
    with pytest.raises(ValueError):
        resolve_prior_root(plan, ledger, provided)


@pytest.mark.parametrize("mutation", ["lost_entity", "duplicate_entity", "lost_evidence", "lost_relation", "changed_summary", "intra_dependency_supersession", "unbound_relation_evidence", "unbound_supersession_evidence", "intra_waiting_source_relation"])
def test_prior_multi_file_relations_consolidation_preserves_dependencies_losslessly(tmp_path, mutation):
    from dataclasses import replace
    from prior_state import validate_bound_prior_context, validate_bound_prior_result
    from action_ledger import finish
    from test_action_ledger import successful_completion
    inventories, revision, plan, ledger, refs, packets, *_ = prior_plan_case(tmp_path, 3, packed=True, unsupported_source=mutation == "intra_waiting_source_relation")
    root = list(refs.values())[-1]
    packet = packets[root.logical_work_id]
    result = json.loads(root.normalized_result)
    if mutation == "lost_entity":
        result["entities"].pop()
    elif mutation == "duplicate_entity":
        result["entities"].append(copy.deepcopy(result["entities"][0]))
    elif mutation == "lost_evidence":
        result["entities"][0]["evidenceIds"] = [inventories[0]["evidence"][2]["priorEvidenceId"]]
    elif mutation == "lost_relation":
        result["sourceRelations"] = []
    elif mutation == "changed_summary":
        result["entities"][0]["semanticSummary"] = "改写dependency语义"
    elif mutation == "unbound_relation_evidence":
        result["sourceRelations"].append({"sourceAId": "source-0", "sourceBId": "source-2", "relation": "COMPLEMENTARY", "evidenceIds": [inventories[0]["evidence"][2]["priorEvidenceId"]]})
    elif mutation == "unbound_supersession_evidence":
        result["entitySupersessions"] = [{"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-2:entity"], "evidenceIds": [inventories[0]["evidence"][2]["priorEvidenceId"]]}]
    elif mutation == "intra_waiting_source_relation":
        result["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "COMPLEMENTARY", "evidenceIds": result["entities"][0]["evidenceIds"]}]
    else:
        result["entitySupersessions"] = [{"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-1:entity"], "evidenceIds": result["entities"][0]["evidenceIds"]}]
    validate_bound_prior_context("PRIOR_CONSOLIDATE", packet, inventories=inventories, input_revision_bytes=revision)
    root_record = ledger.attempt_records[root.result_sha256]
    envelope = ledger.envelopes_by_sha256[root_record.envelope_sha256]
    pending = replace(ledger, attempt_records={key: item for key, item in ledger.attempt_records.items() if key != root.result_sha256})
    completed, record = finish(pending, envelope, successful_completion(canonical_json_bytes(result)), bound_result_validator=lambda raw: validate_bound_prior_result("PRIOR_CONSOLIDATE", packet, raw, inventories=inventories, input_revision_bytes=revision))
    assert record.failure_kind == "INVALID_IR"
    assert record.normalized_result_sha256 is None


@pytest.mark.parametrize("mutation", ["valid", "namespace", "foreign_source", "late_invalid_after_wait", "waiting", "noncanonical", "unknown_region_source"])
def test_prior_multi_file_relations_preseal_bytes_validator_and_repair(tmp_path, monkeypatch, mutation):
    from dataclasses import replace
    from prior_state import validate_bound_prior_context, validate_bound_prior_result, verify_prior_decision
    from test_stage_planner import envelope_for_plan
    from test_action_ledger import successful_completion
    from action_ledger import ActionLedger, issue, finish, attempt_record_value, build_attempt_repair_context
    inventories, revision, decision = prior_pair(tmp_path)
    plan, packet = analyze_packet(inventories, revision)
    validate_bound_prior_context("PRIOR_ANALYZE", packet, inventories=inventories, input_revision_bytes=revision)
    if mutation == "namespace":
        decision["entities"][0]["localKey"] = "unissued:entity"
    elif mutation == "foreign_source":
        decision["entities"][0]["sourceId"] = "source-1"
        decision["entities"][0]["evidenceIds"] = decision["entities"][1]["evidenceIds"]
    elif mutation == "late_invalid_after_wait":
        decision["unsupportedRegions"] = [{"sourceId": "source-0", "regionId": "unknown-layout", "reason": "需补充"}]
        decision["entities"][1]["evidenceIds"] = ["f" * 64]
    elif mutation == "waiting":
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "CONFLICT", "evidenceIds": decision["entities"][0]["evidenceIds"]}]
    elif mutation == "noncanonical":
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "DUPLICATE", "evidenceIds": decision["entities"][0]["evidenceIds"]}]
        decision["entitySupersessions"] = [{"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-1:entity"], "evidenceIds": decision["entities"][1]["evidenceIds"]}]
    elif mutation == "unknown_region_source":
        decision["unsupportedRegions"] = [{"sourceId": "unknown-source", "regionId": "region-1", "reason": "需补充"}]
    envelope = envelope_for_plan(plan, plan["works"][0]["logicalWorkId"])
    ledger = issue(ActionLedger(), envelope)
    raw = canonical_json_bytes(decision)
    completion = replace(successful_completion(), raw_output=raw)
    def validator(normalized):
        with monkeypatch.context() as guard:
            def denied(*args, **kwargs):
                raise AssertionError("pre-seal callback must not read files")
            guard.setattr(Path, "open", denied)
            validate_bound_prior_result("PRIOR_ANALYZE", packet, normalized, inventories=inventories, input_revision_bytes=revision)
    ledger, record = finish(ledger, envelope, completion, bound_result_validator=validator)
    assert record.outcome == ("SUCCEEDED" if mutation in {"valid", "waiting"} else "FAILED")
    assert ledger.raw_outputs[record.raw_sha256] == raw
    if mutation == "waiting":
        from prior_state import materialize_prior_snapshot
        before = canonical_json_bytes(attempt_record_value(record))
        normalized = json.loads(ledger.normalized_results[record.normalized_result_sha256])
        for operation in (verify_prior_decision, materialize_prior_snapshot):
            with pytest.raises(ValueError) as failure:
                operation(inventories, normalized, input_revision_bytes=revision)
            assert failure.value.wait == "WAITING_INPUT"
        assert canonical_json_bytes(attempt_record_value(record)) == before
    elif mutation != "valid":
        assert record.failure_kind == "INVALID_IR"
        assert record.normalized_result_sha256 is None
        assert ledger.normalized_results == {}
        record_sha = sha256_bytes(canonical_json_bytes(attempt_record_value(record)))
        overlay = build_attempt_repair_context(record.logical_work_id, record_sha, ledger.attempt_records, ledger.raw_outputs, envelopes_by_sha256=ledger.envelopes_by_sha256)
        assert json.loads(overlay.canonical_content)["attemptRecordSha256"] == record_sha
        from models import ActionEnvelope
        value = {**envelope.value, "actionId": envelope.value["actionId"] + "-repair", "revision": 2, "attempt": 1}
        repair = ActionEnvelope(value, envelope.path, sha256_bytes(canonical_json_bytes(value)))
        updated = issue(ledger, repair)
        assert updated.envelopes_by_sha256[repair.sha256].value["logicalWorkId"] == record.logical_work_id


@pytest.mark.parametrize("mutation", ["missing", "duplicate", "date", "revision", "content_hash"])
def test_effective_prior_projection_every_packet_binds_exact_date(tmp_path, mutation):
    from prior_state import validate_bound_prior_context
    inventories, revision, _ = prior_pair(tmp_path)
    _, packet = analyze_packet(inventories, revision)
    validate_bound_prior_context("PRIOR_ANALYZE", packet, inventories=inventories, input_revision_bytes=revision)
    if mutation == "missing":
        packet["contextRefs"] = []
    elif mutation == "duplicate":
        packet["contextRefs"].append(copy.deepcopy(packet["contextRefs"][0]))
    elif mutation == "date":
        packet["contextRefs"][0]["canonicalContent"]["plannedEffectiveDate"] = "2028-01-01"
    elif mutation == "revision":
        packet["contextRefs"][0]["canonicalContent"]["inputRevisionSha256"] = "a" * 64
    else:
        packet["contextRefs"][0]["contentSha256"] = "a" * 64
    with pytest.raises(ValueError):
        validate_bound_prior_context("PRIOR_ANALYZE", packet, inventories=inventories, input_revision_bytes=revision)


@pytest.mark.parametrize("mutation", ["source", "evidence", "duplicate_item", "empty_items", "aliased_date"])
def test_prior_context_rejects_invalid_frozen_work_before_result_classification(tmp_path, mutation):
    from prior_state import validate_bound_prior_context, validate_bound_prior_result
    from contracts import InvalidActionResult
    inventories, revision, decision = prior_pair(tmp_path)
    _, packet = analyze_packet(inventories, revision)
    if mutation == "source":
        packet["workItems"][0]["payload"]["sourceId"] = "unknown-source"
    elif mutation == "evidence":
        packet["workItems"][0]["payload"]["evidenceIds"] = ["f" * 64]
    elif mutation == "duplicate_item":
        packet["workItems"].append(copy.deepcopy(packet["workItems"][0]))
    elif mutation == "empty_items":
        packet["workItems"] = []
    else:
        packet["contextRefs"].append({**copy.deepcopy(packet["contextRefs"][0]), "refId": "another-date"})
    for current_result in (False, True):
        with pytest.raises(ValueError) as failure:
            if current_result:
                validate_bound_prior_result("PRIOR_ANALYZE", packet, canonical_json_bytes(decision), inventories=inventories, input_revision_bytes=revision)
            else:
                validate_bound_prior_context("PRIOR_ANALYZE", packet, inventories=inventories, input_revision_bytes=revision)
        assert not isinstance(failure.value, InvalidActionResult)


@pytest.mark.parametrize("mutation", ["relation_source", "relation_evidence", "unsupported_source", "supersession_evidence"])
def test_prior_leaf_all_references_are_limited_to_issued_source_evidence(tmp_path, mutation):
    from prior_state import build_project_effective_start_context, validate_bound_prior_context, validate_bound_prior_result
    from stage_planner import AtomicWorkItemDescriptor, make_planned_work, plan_stage, materialize_packet
    from action_ledger import ActionLedger, issue, finish
    from test_stage_planner import policy, envelope_for_plan
    from test_action_ledger import successful_completion
    inventories, revision, decision = prior_pair(tmp_path)
    third = tmp_path / "third.xlsx"
    workbook_fixture(third, "table")
    inventories.append(inventory_prior_workbook(third))
    revision = revision_for(inventories)
    items = [AtomicWorkItemDescriptor(f"item-{i}", "PRIOR_ANALYZE", "PRIOR_SOW", inventory["workbookSha256"], i, {"sourceId": f"source-{i}", "evidenceIds": decision["entities"][i]["evidenceIds"]}) for i, inventory in enumerate(inventories[:2])]
    contexts = [build_project_effective_start_context(revision)]
    descriptors = [make_planned_work("PRIOR_ANALYZE", items, contexts, [])]
    # These historical fixtures keep the v1 contract; v2 is exercised separately.
    from unittest.mock import patch
    with patch("stage_planner.current_action_contract_id", lambda kind: kind + "-v1"):
        plan = plan_stage("SCOPE", items, contexts, descriptors, [], policy())
    logical_id = plan["works"][0]["logicalWorkId"]
    packet = json.loads(materialize_packet(plan, logical_id, 1, items, contexts, [], ActionLedger()))
    selected = decision["entities"][0]["evidenceIds"]
    unselected = [inventories[0]["evidence"][0]["priorEvidenceId"]]
    if mutation.startswith("relation_"):
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-2" if mutation == "relation_source" else "source-1", "relation": "CONFLICT", "evidenceIds": selected if mutation == "relation_source" else unselected}]
    elif mutation == "unsupported_source":
        decision["unsupportedRegions"] = [{"sourceId": "source-2", "regionId": "region-2", "reason": "该来源没有分配给本 leaf"}]
    else:
        decision["entitySupersessions"] = [{"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-1:entity"], "evidenceIds": unselected}]
    validate_bound_prior_context("PRIOR_ANALYZE", packet, inventories=inventories, input_revision_bytes=revision)
    envelope = envelope_for_plan(plan, logical_id)
    ledger, record = finish(issue(ActionLedger(), envelope), envelope, successful_completion(canonical_json_bytes(decision)), bound_result_validator=lambda raw: validate_bound_prior_result("PRIOR_ANALYZE", packet, raw, inventories=inventories, input_revision_bytes=revision))
    assert record.failure_kind == "INVALID_IR"
    assert record.normalized_result_sha256 is None
    assert ledger.normalized_results == {}


@pytest.mark.parametrize("status,active", [("CURRENT_BY_CONTRACT", True), ("EXCLUDED", False), ("CANCELLED", False), ("NOT_IMPLEMENTED", False), ("FUTURE", False)])
def test_current_by_contract_effective_prior_projection_excludes_explicit_noncurrent(tmp_path, status, active):
    from prior_state import materialize_prior_snapshot, derive_effective_prior
    inventories, revision, decision = prior_pair(tmp_path)
    decision["entities"][0]["deliveryStatus"] = status
    snapshot = materialize_prior_snapshot(inventories, decision, input_revision_bytes=revision)
    before = canonical_json_bytes(snapshot)
    view = derive_effective_prior(snapshot)
    entity = next(item for item in snapshot["entities"] if item["sourceId"] == "source-0")
    assert (entity["entityId"] in view["activeEntityIds"]) is active
    assert canonical_json_bytes(snapshot) == before


def test_effective_prior_projection_byte_identical_duplicate_keeps_audit_and_one_active(tmp_path):
    from prior_state import verify_prior_decision, materialize_prior_snapshot, derive_effective_prior
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path)
    inventory = inventory_prior_workbook(path)
    revision = revision_for([inventory, inventory])
    decision = decision_for([inventory, inventory])
    decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "DUPLICATE", "evidenceIds": decision["entities"][0]["evidenceIds"]}]
    verify_prior_decision([inventory], decision, input_revision_bytes=revision)
    snapshot = materialize_prior_snapshot([inventory], decision, input_revision_bytes=revision)
    assert len(snapshot["evidence"]) == len(snapshot["entities"]) == 2
    assert len({item["priorEvidenceId"] for item in snapshot["evidence"]}) == 1
    assert len({item["entityId"] for item in snapshot["entities"]}) == 2
    canonical = next(item for item in snapshot["entities"] if item["sourceId"] == "source-0")
    assert derive_effective_prior(snapshot) == {"activeEntityIds": [canonical["entityId"]]}
    decision["entities"].reverse()
    assert materialize_prior_snapshot([inventory], decision, input_revision_bytes=revision) == snapshot


@pytest.mark.parametrize("mode", ["unique", "repeated_reference", "duplicate", "ambiguous"])
def test_prior_identity_visible_id_retention_and_ambiguous_fallback(tmp_path, mode):
    from prior_state import verify_prior_decision, materialize_prior_snapshot, derive_effective_prior
    inventories, revision, decision = prior_pair(tmp_path)
    decision["entities"][0]["visiblePriorId"] = "prior-visible-1"
    if mode == "repeated_reference":
        decision["entities"][0]["evidenceIds"].append(inventories[0]["evidence"][2]["priorEvidenceId"])
    if mode in {"duplicate", "ambiguous"}:
        decision["entities"][1]["visiblePriorId"] = "prior-visible-1"
    if mode == "duplicate":
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "DUPLICATE", "evidenceIds": decision["entities"][0]["evidenceIds"]}]
    verify_prior_decision(inventories, decision, input_revision_bytes=revision)
    snapshot = materialize_prior_snapshot(inventories, decision, input_revision_bytes=revision)
    ids = {item["sourceId"]: item["entityId"] for item in snapshot["entities"]}
    assert (ids["source-0"] == "prior-visible-1") is (mode != "ambiguous")
    assert ids["source-1"] != "prior-visible-1"
    if mode == "duplicate":
        assert derive_effective_prior(snapshot) == {"activeEntityIds": ["prior-visible-1"]}


@pytest.mark.parametrize("visible", ["not-in-workbook", "visible", "bad id"])
def test_prior_identity_false_visible_cell_claim_is_invalid_ir(tmp_path, visible):
    from prior_state import verify_prior_decision, materialize_prior_snapshot
    from contracts import InvalidActionResult
    inventories, revision, decision = prior_pair(tmp_path)
    decision["entities"][0]["visiblePriorId"] = visible
    for operation in (verify_prior_decision, materialize_prior_snapshot):
        with pytest.raises(InvalidActionResult):
            operation(inventories, decision, input_revision_bytes=revision)


def test_prior_identity_anchors_stable_and_unexplained_collision_requires_input(tmp_path, monkeypatch):
    from prior_state import materialize_prior_snapshot
    inventories, revision, decision = prior_pair(tmp_path)
    snapshot = materialize_prior_snapshot(inventories, decision, input_revision_bytes=revision)
    before = {item["sourceId"]: item["entityId"] for item in snapshot["entities"]}
    expected = "prior-" + sha256_bytes(canonical_json_bytes({"idSchemaVersion": "prior-entity-id-v1", "entityKind": "CONTRACT_ENTITY", "sortedEvidenceAnchors": [{"sourceId": "source-0", "priorEvidenceId": decision["entities"][0]["evidenceIds"][0]}], "controlledDiscriminator": ["CONTRACT_ENTITY"]}))
    assert before["source-0"] == expected
    for item in decision["entities"]:
        item["semanticSummary"] = "名称润色"
        item["deliveryStatus"] = "FUTURE"
        item["localKey"] += "-renamed"
    changed_revision = revision_for(inventories, date="2027-01-01")
    with monkeypatch.context() as guard:
        def denied(*args, **kwargs):
            raise AssertionError("pure materializer must not read or write files")
        guard.setattr(Path, "open", denied)
        changed = materialize_prior_snapshot(inventories, decision, input_revision_bytes=changed_revision)
    assert {item["sourceId"]: item["entityId"] for item in changed["entities"]} == before
    decision["entities"].append({**decision["entities"][0], "localKey": "new:collision"})
    with pytest.raises(ValueError) as failure:
        materialize_prior_snapshot(inventories, decision, input_revision_bytes=revision)
    assert failure.value.wait == "WAITING_INPUT"


@pytest.mark.parametrize("date,status,active_source", [("2026-09-01", "FUTURE", "source-0"), ("2026-11-01", "CURRENT_BY_CONTRACT", "source-1")])
def test_effective_prior_projection_full_replacement_uses_new_revision_status(tmp_path, date, status, active_source):
    from prior_state import verify_prior_decision, materialize_prior_snapshot, derive_effective_prior
    inventories, _, decision = prior_pair(tmp_path)
    revision = revision_for(inventories, date=date)
    decision["entities"][1]["deliveryStatus"] = status
    decision["entitySupersessions"] = [{"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-1:entity"], "evidenceIds": decision["entities"][1]["evidenceIds"]}]
    verify_prior_decision(inventories, decision, input_revision_bytes=revision)
    snapshot = materialize_prior_snapshot(inventories, decision, input_revision_bytes=revision)
    ids = {item["sourceId"]: item["entityId"] for item in snapshot["entities"]}
    assert snapshot["entitySupersessions"] == [{"predecessorIds": [ids["source-0"]], "successorIds": [ids["source-1"]], "evidenceIds": decision["entities"][1]["evidenceIds"]}]
    assert derive_effective_prior(snapshot) == {"activeEntityIds": [ids[active_source]]}


@pytest.mark.parametrize("mutation", ["missing_endpoint", "unknown_evidence", "noncanonical", "PARTIAL", "affectedSlots", "source_SUPERSEDES"])
def test_prior_multi_file_relations_reject_invalid_full_replacement_ir(tmp_path, mutation):
    from prior_state import verify_prior_decision
    from contracts import InvalidActionResult
    inventories, revision, decision = prior_pair(tmp_path)
    relation = {"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-1:entity"], "evidenceIds": decision["entities"][1]["evidenceIds"]}
    decision["entitySupersessions"] = [relation]
    if mutation == "missing_endpoint":
        relation["predecessorLocalKeys"] = ["unknown:entity"]
    elif mutation == "unknown_evidence":
        relation["evidenceIds"] = ["f" * 64]
    elif mutation == "noncanonical":
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "DUPLICATE", "evidenceIds": decision["entities"][0]["evidenceIds"]}]
    elif mutation == "PARTIAL":
        relation["replacement"] = "PARTIAL"
    elif mutation == "affectedSlots":
        relation["affectedSlots"] = ["orders"]
    else:
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "SUPERSEDES", "evidenceIds": decision["entities"][0]["evidenceIds"]}]
    with pytest.raises(InvalidActionResult):
        verify_prior_decision(inventories, decision, input_revision_bytes=revision)


@pytest.mark.parametrize("mutation", ["cycle", "overlap", "mixed", "unknown_status", "conflict", "unsupported"])
def test_effective_prior_projection_unresolved_graph_requires_input(tmp_path, mutation):
    from prior_state import verify_prior_decision, materialize_prior_snapshot
    inventories, revision, decision = prior_pair(tmp_path)
    evidence = decision["entities"][1]["evidenceIds"]
    third = {**decision["entities"][1], "localKey": "item-1:third", "evidenceIds": [inventories[1]["evidence"][2]["priorEvidenceId"]]}
    decision["entities"].append(third)
    relation = {"predecessorLocalKeys": ["item-0:entity"], "successorLocalKeys": ["item-1:entity"], "evidenceIds": evidence}
    decision["entitySupersessions"] = [relation]
    if mutation == "cycle":
        decision["entitySupersessions"].append({**relation, "predecessorLocalKeys": ["item-1:entity"], "successorLocalKeys": ["item-0:entity"]})
    elif mutation == "overlap":
        decision["entitySupersessions"].append({**relation, "successorLocalKeys": ["item-1:third"]})
    elif mutation == "mixed":
        relation["successorLocalKeys"].append("item-1:third")
        third["deliveryStatus"] = "FUTURE"
    elif mutation == "unknown_status":
        decision["entities"][1]["deliveryStatus"] = "NOT_IMPLEMENTED"
    elif mutation == "conflict":
        decision["sourceRelations"] = [{"sourceAId": "source-0", "sourceBId": "source-1", "relation": "CONFLICT", "evidenceIds": evidence}]
    else:
        decision["unsupportedRegions"] = [{"sourceId": "source-1", "regionId": "region-1", "reason": "区域不可读取"}]
    for operation in (verify_prior_decision, materialize_prior_snapshot):
        with pytest.raises(ValueError) as failure:
            operation(inventories, decision, input_revision_bytes=revision)
        assert failure.value.wait == "WAITING_INPUT"


@pytest.mark.parametrize("mutation", ["reversed", "unknown_source", "empty_evidence", "unknown_evidence", "duplicate_pair"])
def test_prior_multi_file_relations_validate_declared_source_graph(tmp_path, mutation):
    from prior_state import verify_prior_decision
    inventories, revision, decision = prior_pair(tmp_path)
    relation = {"sourceAId": "source-0", "sourceBId": "source-1", "relation": "COMPLEMENTARY", "evidenceIds": [decision["entities"][0]["evidenceIds"][0]]}
    decision["sourceRelations"] = [relation]
    verify_prior_decision(inventories, decision, input_revision_bytes=revision)
    if mutation == "reversed":
        relation["sourceAId"], relation["sourceBId"] = relation["sourceBId"], relation["sourceAId"]
    elif mutation == "unknown_source":
        relation["sourceBId"] = "unknown-source"
    elif mutation == "empty_evidence":
        relation["evidenceIds"] = []
    elif mutation == "unknown_evidence":
        relation["evidenceIds"] = ["f" * 64]
    else:
        decision["sourceRelations"].append({**relation, "relation": "UNRELATED"})
    with pytest.raises(ValueError):
        verify_prior_decision(inventories, decision, input_revision_bytes=revision)


@pytest.mark.parametrize("layout", ["plain", "table", "merged"])
def test_workbook_inventory_preserves_cells_formula_cache_and_surfaces(tmp_path, layout):
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path, layout)
    before = (path.stat().st_size, sha256_bytes(path.read_bytes()))
    inventory = inventory_prior_workbook(path)
    assert inventory["workbookSha256"] == before[1]
    sheet = inventory["sheets"][0]
    assert sheet["sheet"] == "合同"
    assert sheet["usedRange"] == "$A$1:$C$3"
    assert len(sheet["cells"]) == 9
    formula = next(cell for cell in sheet["cells"] if cell["address"] == "$C$3")
    assert formula == {"address": "$C$3", "cellType": "n", "value": None, "formula": "C2+1", "cachedValue": None}
    assert sheet["hiddenRows"] == [3]
    assert sheet["hiddenColumns"] == [{"min": 3, "max": 3}]
    assert sheet["tables"] == ([{"name": "PriorContracts", "range": "$A$1:$C$3"}] if layout == "table" else [])
    if layout == "merged":
        assert inventory["sheets"][1]["merges"] == ["$A$1:$C$1"]
    assert inventory["unsupportedSurfaces"] == []
    assert (path.stat().st_size, sha256_bytes(path.read_bytes())) == before


def test_prior_evidence_id_uses_workbook_sheet_absolute_range_and_typed_values(tmp_path):
    import hashlib
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path)
    inventory = inventory_prior_workbook(path)
    evidence = inventory["evidence"][1]
    assert evidence["sheet"] == "合同"
    assert evidence["absoluteA1Range"] == "$A$2:$C$2"
    assert evidence["canonicalCellValues"] == inventory["sheets"][0]["cells"][3:6]
    basis = {"workbookSha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheet": "合同", "absoluteA1Range": "$A$2:$C$2", "canonicalCellValuesSha256": hashlib.sha256(canonical_json_bytes(evidence["canonicalCellValues"])).hexdigest()}
    assert evidence["priorEvidenceId"] == hashlib.sha256(canonical_json_bytes(basis)).hexdigest()
    assert inventory_prior_workbook(path)["evidence"] == inventory["evidence"]
    assert len({item["priorEvidenceId"] for item in inventory["evidence"]}) == 3


def test_workbook_inventory_reports_unsupported_surfaces_and_retains_cache(tmp_path):
    base = tmp_path / "base.xlsx"
    workbook_fixture(base)
    path = tmp_path / "surfaces.xlsx"
    with zipfile.ZipFile(base) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                from xml.etree import ElementTree as ET
                root = ET.fromstring(content)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                root.find(".//s:c[@r='C3']/s:v", ns).text = "3"
                content = ET.tostring(root)
            target.writestr(item, content)
        for name in ["xl/drawings/drawing1.xml", "xl/media/image1.png", "xl/embeddings/oleObject1.bin", "xl/vbaProject.bin", "xl/drawings/vmlDrawing1.vml"]:
            target.writestr(name, b"synthetic unsupported surface")
    before = (path.stat().st_size, sha256_bytes(path.read_bytes()))
    inventory = inventory_prior_workbook(path)
    assert len(inventory["unsupportedSurfaces"]) == 5
    assert {item["kind"] for item in inventory["unsupportedSurfaces"]} == {"DRAWING", "IMAGE", "EMBEDDED_OBJECT", "MACRO", "VML"}
    assert next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "$C$3")["cachedValue"] == 3
    assert (path.stat().st_size, sha256_bytes(path.read_bytes())) == before


@pytest.mark.parametrize("epoch", ["1900", "1904"])
def test_workbook_inventory_preserves_typed_temporal_values_and_formula_cache(tmp_path, epoch):
    from datetime import datetime, time
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
    from xml.etree import ElementTree as ET
    workbook = openpyxl.Workbook()
    if epoch == "1904":
        workbook.epoch = CALENDAR_MAC_1904
    sheet = workbook.active
    sheet["A1"] = datetime(2026, 10, 1)
    sheet["A1"].number_format = "yyyy-mm-dd"
    sheet["B1"] = datetime(2026, 10, 1, 12, 30)
    sheet["B1"].number_format = "yyyy-mm-dd hh:mm"
    sheet["C1"] = time(12, 30)
    sheet["C1"].number_format = "hh:mm"
    sheet["D1"] = 46296
    sheet["E1"] = "=A1"
    sheet["E1"].number_format = "yyyy-mm-dd"
    base = tmp_path / "base.xlsx"
    workbook.save(base)
    serial = to_excel(datetime(2026, 10, 1), workbook.epoch)
    workbook.close()
    path = tmp_path / "prior.xlsx"
    with zipfile.ZipFile(base) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                root.find(".//{*}c[@r='E1']/{*}v").text = str(serial)
                content = ET.tostring(root)
            target.writestr(item, content)
    before = (path.stat().st_size, sha256_bytes(path.read_bytes()))
    cells = {item["address"]: item for item in inventory_prior_workbook(path)["sheets"][0]["cells"]}
    assert (cells["$A$1"]["value"], cells["$A$1"]["temporalType"]) == ("2026-10-01", "date")
    assert (cells["$B$1"]["value"], cells["$B$1"]["temporalType"]) == ("2026-10-01T12:30:00", "datetime")
    assert (cells["$C$1"]["value"], cells["$C$1"]["temporalType"]) == ("12:30:00", "time")
    assert cells["$D$1"]["value"] == 46296 and "temporalType" not in cells["$D$1"]
    assert cells["$E$1"]["formula"] == "A1"
    assert (cells["$E$1"]["cachedValue"], cells["$E$1"]["temporalType"]) == ("2026-10-01", "date")
    assert (path.stat().st_size, sha256_bytes(path.read_bytes())) == before


def test_workbook_inventory_preserves_shared_formula_followers_and_caches(tmp_path):
    from xml.etree import ElementTree as ET
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([1, "结果"])
    sheet.append([2, "=A2+$A$1"])
    sheet.append([3, "=A3+$A$1"])
    base = tmp_path / "base.xlsx"
    workbook.save(base)
    workbook.close()
    path = tmp_path / "shared.xlsx"
    with zipfile.ZipFile(base) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                for address, cache in (("B2", "3"), ("B3", "4")):
                    cell = root.find(f".//{{*}}c[@r='{address}']")
                    formula = cell.find("{*}f")
                    formula.attrib.update({"t": "shared", "si": "0"})
                    if address == "B2":
                        formula.set("ref", "B2:B3")
                    else:
                        formula.text = None
                    cell.find("{*}v").text = cache
                content = ET.tostring(root)
            target.writestr(item, content)
    before = (path.stat().st_size, sha256_bytes(path.read_bytes()))
    inventory = inventory_prior_workbook(path)
    cells = {item["address"]: item for item in inventory["sheets"][0]["cells"]}
    assert (cells["$B$2"]["formula"], cells["$B$2"]["cachedValue"]) == ("A2+$A$1", 3)
    assert (cells["$B$3"]["formula"], cells["$B$3"]["cachedValue"]) == ("A3+$A$1", 4)
    assert (path.stat().st_size, sha256_bytes(path.read_bytes())) == before
