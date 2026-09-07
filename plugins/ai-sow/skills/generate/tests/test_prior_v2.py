from __future__ import annotations

TEST_LAYER = "integration"

import copy
import json
import sys
from pathlib import Path

import pytest

SKILL_ROOT = Path(__file__).parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
sys.path.insert(0, str(SKILL_ROOT / "tests"))

from test_prior_state import workbook_fixture, revision_for, decision_for
from contracts import canonical_json_bytes, sha256_bytes
from prior_state import inventory_prior_workbook, build_project_effective_start_context

def prior_case(tmp_path):
    path = tmp_path / "prior.xlsx"
    workbook_fixture(path, "merged")
    inventory = inventory_prior_workbook(path)
    revision = revision_for([inventory])
    date = build_project_effective_start_context(revision)
    packet = {"workItems": [{"workItemId": "item-0", "payload": {
        "sourceId": "source-0", "evidenceIds": [row["priorEvidenceId"] for row in inventory["evidence"]],
        "evidence": inventory["evidence"],
    }}], "contextRefs": [{"refId": date.ref_id, "canonicalContent": json.loads(date.canonical_content),
                           "contentSha256": sha256_bytes(date.canonical_content)}]}
    decision = decision_for([inventory])
    used = set(decision["entities"][0]["evidenceIds"])
    decision["unextractedEvidence"] = [{"sourceId": "source-0", "evidenceIds": [
        row["priorEvidenceId"] for row in inventory["evidence"] if row["priorEvidenceId"] not in used],
        "reason": "测试中未形成独立合同对象的表头和上下文"}]
    return inventory, revision, packet, decision


def normalize_v2(kind, result, packet=None):
    from contracts import action_contract_binding, normalize_action_result
    _, digest = action_contract_binding(SKILL_ROOT, kind + "-v2")
    raw_packet = canonical_json_bytes(packet) if packet is not None else None
    envelope = {"actionContractId": kind + "-v2", "actionContractSha256": digest,
                "packetSha256": sha256_bytes(raw_packet) if raw_packet else "0" * 64}
    return normalize_action_result(envelope, canonical_json_bytes(result), skill_root=SKILL_ROOT,
                                   packet_payload=raw_packet)


def dependency_packet(decisions):
    return {"workItems": [], "contextRefs": [{"refId": f"dependency-result-leaf-{index}", "canonicalContent": {
        "kind": "DEPENDENCY_RESULT", "logicalWorkId": f"leaf-{index}", "attemptRecordSha256": str(index) * 64,
        "normalizedResult": decision}} for index, decision in enumerate(decisions)]}


def test_narrow_consolidation_preserves_all_143_entities_across_two_levels():
    decisions = []
    for index, count in enumerate([21, 70, 37, 15, 0]):
        decisions.append({"entities": [{"localKey": f"item-{index}:{row}", "sourceId": f"source-{index}",
            "entityKind": "CONTRACT_ENTITY", "semanticSummary": f"合同交付 {row}",
            "deliveryStatus": "CURRENT_BY_CONTRACT", "evidenceIds": [f"row-{index}-{row}"]} for row in range(count)],
            "sourceRelations": [], "entitySupersessions": [], "unsupportedRegions": [],
            "unextractedEvidence": [{"sourceId": f"source-{index}", "evidenceIds": [f"header-{index}"], "reason": "表头"}]})
    additions = {"sourceRelations": [], "entitySupersessions": []}
    first = json.loads(normalize_v2("PRIOR_CONSOLIDATE", additions, dependency_packet(decisions[:2])))
    second = json.loads(normalize_v2("PRIOR_CONSOLIDATE", additions, dependency_packet(decisions[2:])))
    result = json.loads(normalize_v2("PRIOR_CONSOLIDATE", additions, dependency_packet([first, second])))
    assert len(result["entities"]) == 143
    assert len(result["unextractedEvidence"]) == 5
    assert {row["localKey"] for row in result["entities"]} == {
        f"item-{index}:{row}" for index, count in enumerate([21, 70, 37, 15, 0]) for row in range(count)}
    assert all(row in result["entities"] for decision in decisions for row in decision["entities"])


@pytest.mark.parametrize("defect", ["omission", "duplicate_reason", "reason_for_entity", "foreign_evidence"])
def test_analyze_v2_rejects_missing_or_misbound_coverage(tmp_path, defect):
    from prior_state import validate_bound_prior_result
    inventory, revision, packet, decision = prior_case(tmp_path)
    normalized = normalize_v2("PRIOR_ANALYZE", decision)
    validate_bound_prior_result("PRIOR_ANALYZE", packet, normalized, inventories=[inventory], input_revision_bytes=revision)
    if defect == "omission": decision["unextractedEvidence"][0]["evidenceIds"].pop()
    elif defect == "duplicate_reason": decision["unextractedEvidence"].append(copy.deepcopy(decision["unextractedEvidence"][0]))
    elif defect == "reason_for_entity": decision["unextractedEvidence"][0]["evidenceIds"].extend(decision["entities"][0]["evidenceIds"])
    else: decision["unextractedEvidence"][0]["evidenceIds"][0] = "f" * 64
    with pytest.raises(ValueError):
        validate_bound_prior_result("PRIOR_ANALYZE", packet, canonical_json_bytes(decision), inventories=[inventory], input_revision_bytes=revision)


def test_consolidation_rejects_entity_copy_and_duplicate_dependency_keys(tmp_path):
    _, _, _, decision = prior_case(tmp_path)
    second = copy.deepcopy(decision)
    with pytest.raises(ValueError):
        normalize_v2("PRIOR_CONSOLIDATE", {"sourceRelations": [], "entitySupersessions": []}, dependency_packet([decision, second]))
    with pytest.raises(ValueError):
        normalize_v2("PRIOR_CONSOLIDATE", decision, dependency_packet([decision, second]))


def test_consolidation_normalization_requires_exact_frozen_packet(tmp_path):
    from contracts import action_contract_binding, normalize_action_result
    _, _, _, decision = prior_case(tmp_path)
    packet = canonical_json_bytes(dependency_packet([decision, {**decision, "entities": []}]))
    _, digest = action_contract_binding(SKILL_ROOT, "PRIOR_CONSOLIDATE-v2")
    envelope = {"actionContractId": "PRIOR_CONSOLIDATE-v2", "actionContractSha256": digest, "packetSha256": "f" * 64}
    with pytest.raises(ValueError, match="packet"):
        normalize_action_result(envelope, canonical_json_bytes({"sourceRelations": [], "entitySupersessions": []}),
                                skill_root=SKILL_ROOT, packet_payload=packet)


def context_index(inventory, source_id="source-0"):
    return {"sourceId": source_id, "workbookSha256": inventory["workbookSha256"], "evidence": [
        {key: row[key] for key in ("priorEvidenceId", "sheet", "absoluteA1Range")}
        for row in inventory["evidence"]]}


def test_small_workbook_all_sheets_have_one_analyze_item(tmp_path):
    from scope_compiler import prepare_scope_inputs
    from test_scope_compiler import owner_input_case
    inventory, _, _, _ = prior_case(tmp_path)
    request, revision, contents, _ = owner_input_case(inventories=[inventory])
    items, _ = prepare_scope_inputs(revision, contents, request=request, prior_inventories=[inventory])
    prior = [item for item in items if item.action_kind == "PRIOR_ANALYZE"]
    assert len(prior) == 1
    assert {row["sheet"] for row in prior[0].work_item_payload["evidence"]} == {"合同", "补充"}


def test_prior_single_row_over_capacity_fails_without_truncating_source(tmp_path):
    import openpyxl
    from scope_compiler import prepare_scope_inputs
    from stage_planner import StagePlanningBlocked
    from test_scope_compiler import owner_input_case
    path = tmp_path / "large-row.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["合同限定" * 8000])
    book.save(path)
    before = path.read_bytes()
    inventory = inventory_prior_workbook(path)
    request, revision, contents, _ = owner_input_case(inventories=[inventory])
    with pytest.raises(StagePlanningBlocked):
        prepare_scope_inputs(revision, contents, request=request, prior_inventories=[inventory])
    assert inventory["evidence"][0]["canonicalCellValues"][0]["value"] == "合同限定" * 8000
    assert path.read_bytes() == before


@pytest.mark.parametrize("defect", [None, "no_primary", "foreign_source", "index_hash", "index_locator"])
def test_prior_cross_sheet_context_keeps_primary_anchor_and_source_binding(tmp_path, defect):
    from prior_state import validate_bound_prior_context, validate_bound_prior_result
    inventory, revision, packet, decision = prior_case(tmp_path)
    primary, footnote = inventory["evidence"][1], inventory["evidence"][-1]
    payload = packet["workItems"][0]["payload"]
    payload.update(evidenceIds=[primary["priorEvidenceId"]], evidence=[primary], priorContext=context_index(inventory))
    decision["entities"][0]["evidenceIds"] = [primary["priorEvidenceId"], footnote["priorEvidenceId"]]
    decision["entities"][0]["semanticSummary"] = "查询订单；补充合同约定生产上线。"
    decision["unextractedEvidence"] = []
    if defect == "no_primary": decision["entities"][0]["evidenceIds"] = [footnote["priorEvidenceId"]]
    elif defect == "foreign_source": payload["priorContext"]["sourceId"] = "other-source"
    elif defect == "index_hash": payload["priorContext"]["workbookSha256"] = "f" * 64
    elif defect == "index_locator": payload["priorContext"]["evidence"][-1]["sheet"] = "wrong-sheet"
    def validate():
        validate_bound_prior_context("PRIOR_ANALYZE", packet, inventories=[inventory], input_revision_bytes=revision)
        validate_bound_prior_result("PRIOR_ANALYZE", packet, normalize_v2("PRIOR_ANALYZE", decision),
                                    inventories=[inventory], input_revision_bytes=revision)
    if defect:
        with pytest.raises(ValueError): validate()
    else: validate()


def test_same_row_cell_anchors_keep_two_identities_through_change_graph(tmp_path):
    from prior_state import materialize_prior_snapshot, validate_bound_prior_result
    from scope_compiler import scope_change_graph
    inventory, revision, packet, decision = prior_case(tmp_path)
    first = decision["entities"][0]
    eid = first["evidenceIds"][0]
    first["cellAnchors"] = [{"evidenceId": eid, "address": "$A$2"}]
    second = {**copy.deepcopy(first), "localKey": "item-0:second", "semanticSummary": "另一个交付对象",
              "cellAnchors": [{"evidenceId": eid, "address": "$B$2"}]}
    decision["entities"].append(second)
    validate_bound_prior_result("PRIOR_ANALYZE", packet, normalize_v2("PRIOR_ANALYZE", decision),
                                inventories=[inventory], input_revision_bytes=revision)
    snapshot = materialize_prior_snapshot([inventory], decision, input_revision_bytes=revision)
    assert len({row["entityId"] for row in snapshot["entities"]}) == 2
    renamed = copy.deepcopy(decision)
    for index, entity in enumerate(renamed["entities"]): entity["localKey"] = f"different-{index}"
    assert materialize_prior_snapshot([inventory], renamed, input_revision_bytes=revision) == snapshot
    changes = {"decisions": [{"localKey": f"target-{index}", "decisionKind": "FEATURE", "priorEntityIds": [row["localKey"]],
        "relations": [{"kind": "REUSE_DEPENDENCY", "targetLocalKeys": [f"target-{index}"], "evidenceIds": [eid]}]}
        for index, row in enumerate(decision["entities"])]}
    graph = scope_change_graph(changes, {"target-0": "feature-0", "target-1": "feature-1"}, decision, snapshot)
    assert len({row["priorEntityIds"][0] for row in graph["changeGroups"]}) == 2
    first["cellAnchors"][0]["address"] = "$Z$999"
    with pytest.raises(ValueError): materialize_prior_snapshot([inventory], decision, input_revision_bytes=revision)


def test_public_prior_and_scope_review_hydrate_excluded_rows_and_preserve_clauses(tmp_path):
    import openpyxl
    import orchestrator
    from test_intake import write_next_request
    from test_orchestrator import write_budget_policy, submit_prototype
    from test_scope_compiler import scope_owner_result
    request_path = write_next_request(tmp_path, mode="BROWNFIELD", include_prior=True)
    source = tmp_path / "inputs/prior.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "合同"
    for row in [["项目合同范围"], ["退款申请，验收须在5秒内响应"], ["通用估算目录"], ["邮件发送任务，仅列目录，不在本项目"]]:
        book.active.append(row)
    book.create_sheet("条款").append(["客户提供短信账号，范围不含历史迁移"])
    book.save(source)
    original = source.read_bytes()
    request = json.loads(request_path.read_bytes())
    next(row for row in request["sources"] if row["role"] == "PRIOR_SOW")["expectedSha256"] = sha256_bytes(original)
    request_path.write_bytes(canonical_json_bytes(request))
    response = orchestrator.run_mode(tmp_path, "start", request=request_path.name, budget_policy=write_budget_policy(tmp_path))
    excluded_id = None
    for _ in range(30):
        assert response["outcome"] == "ACTIVE", response
        actions = response["nextAction"].get("actions", [response["nextAction"]])
        for action in actions:
            packet = json.loads((tmp_path / action["packetPath"]).read_bytes())
            kind = action["actionContractId"][:-3]
            if kind == "SOURCE_SCOPE":
                body = packet["workItems"][0]["payload"]
                assert "客户提供短信账号" in body["priorState"]["entities"][0]["semanticSummary"]
                assert "不含历史迁移" in body["priorState"]["entities"][0]["semanticSummary"]
                obligation = next(row for row in body["reviewObligations"] if row["kind"] == "PRIOR_EXTRACTION")
                assert "hydrate" in obligation["reviewInstruction"]
                assert "不能把估算目录" in obligation["reviewInstruction"]
                assert excluded_id in {eid for row in obligation["unextractedEvidence"] for eid in row["evidenceIds"]}
                hydrated = orchestrator.hydrate(tmp_path, action["actionId"], [excluded_id])
                assert hydrated["outcome"] == "HYDRATED", hydrated
                assert "仅列目录" in json.dumps(hydrated, ensure_ascii=False)
                assert orchestrator.hydrate(tmp_path, action["actionId"], ["f" * 64])["outcome"] == "BLOCKED"
                assert source.read_bytes() == original
                return
            result = scope_owner_result(kind, packet)
            if kind == "PRIOR_ANALYZE":
                item = packet["workItems"][0]
                rows = item["payload"]["evidence"]
                own = next(row for row in rows if row["sheet"] == "合同" and row["absoluteA1Range"] == "$A$2:$A$2")
                clause = next(row for row in rows if row["sheet"] == "条款")
                excluded_id = next(row["priorEvidenceId"] for row in rows if "仅列目录" in str(row["canonicalCellValues"]))
                hydrated = orchestrator.hydrate(tmp_path, action["actionId"], [clause["priorEvidenceId"]])
                assert hydrated["outcome"] == "HYDRATED", hydrated
                wire = json.loads(orchestrator.read_provider_request(tmp_path, action["actionId"]))
                assert json.loads(wire["messages"][-1]["content"]) == hydrated
                result["entities"] = [{"localKey": item["workItemId"] + ":refund", "sourceId": "prior-main",
                    "entityKind": "CONTRACT_ENTITY", "semanticSummary": "退款申请，验收5秒内响应；客户提供短信账号；范围不含历史迁移",
                    "deliveryStatus": "CURRENT_BY_CONTRACT", "evidenceIds": [own["priorEvidenceId"], clause["priorEvidenceId"]]}]
                result["unextractedEvidence"] = [{"sourceId": "prior-main", "evidenceIds": [row["priorEvidenceId"] for row in rows
                    if row not in (own, clause)], "reason": "标题和通用目录未形成项目交付"}]
            assert submit_prototype(tmp_path, action, result)["record"]["outcome"] == "SUCCEEDED"
        response = orchestrator.run_mode(tmp_path, "resume")
    pytest.fail("未到达 Scope Review")


def test_prior_v2_keeps_frozen_scope_review_contract_readable():
    from contracts import action_contract_binding, normalize_action_result
    contract, digest = action_contract_binding(SKILL_ROOT, "SOURCE_SCOPE-v1")
    # Published v1 identity: changing it invalidates an already verified Greenfield artifact.
    frozen_digest = "9581433587543cdff5558b2226c73766fad9c179b0323d6ce0d0421da99bbf1f"
    assert digest == frozen_digest
    envelope = {"actionContractId": "SOURCE_SCOPE-v1", "actionContractSha256": frozen_digest}
    assert normalize_action_result(envelope, b'{"decision":"PASS","findings":[]}', skill_root=SKILL_ROOT) == canonical_json_bytes({"decision": "PASS", "findings": []})


def test_prior_row_partition_keeps_large_shared_index_within_real_request_budget(tmp_path):
    import openpyxl
    from scope_compiler import prepare_scope_inputs, build_scope_work_descriptors
    from stage_planner import RunBudgetPolicy, DemoBudgetLimits, StagePlanningBlocked, plan_stage
    from test_scope_compiler import owner_input_case
    path = tmp_path / "many-rows.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "合同范围与估算目录"
    for index in range(430):
        book.active.append([f"第 {index} 行", "项目范围或完整目录说明" * 30])
    book.save(path)
    before = path.read_bytes()
    inventory = inventory_prior_workbook(path)
    request, revision, contents, _ = owner_input_case(inventories=[inventory])
    items, refs = prepare_scope_inputs(revision, contents, request=request, prior_inventories=[inventory])
    leaves = [item.work_item_payload for item in items if item.action_kind == "PRIOR_ANALYZE"]
    actual = [row for leaf in leaves for row in leaf["evidence"]]
    assert len(actual) == 430
    assert {canonical_json_bytes(row) for row in actual} == {canonical_json_bytes(row) for row in inventory["evidence"]}
    assert all(len(leaf["priorContext"]["evidence"]) == 430 for leaf in leaves)
    policy = RunBudgetPolicy("ai-sow-run-budget-policy-v1", "host-canonical-messages-v1", 272000,
        "utf8-bytes-v1", 5000000, 14400, 12000, 12000, 10000, 1000, 2, DemoBudgetLimits(3, 100, 30))
    works = build_scope_work_descriptors(items, refs, policy)
    assert plan_stage("SCOPE", items, refs, works, [], policy)["works"]
    from dataclasses import replace
    with pytest.raises(StagePlanningBlocked):
        build_scope_work_descriptors(items, refs, replace(policy, model_context_limit_tokens=70000))
    assert path.read_bytes() == before
