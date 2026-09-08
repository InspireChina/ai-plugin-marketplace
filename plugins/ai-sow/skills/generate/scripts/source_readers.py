from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, time
from html.parser import HTMLParser
from pathlib import Path
import zipfile
from io import BytesIO
import posixpath
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from contracts import canonical_json_bytes
from models import SourceAnchor, SourceDocument


SOURCE_ROLES = frozenset(
    {"PRD", "DEMO", "HLD", "ADR", "PRIOR_SOW", "SUPPLEMENT"}
)
ROLE_SUFFIXES = {
    "PRD": frozenset({".md"}),
    "HLD": frozenset({".md"}),
    "ADR": frozenset({".md"}),
    "PRIOR_SOW": frozenset({".xlsx"}),
}
PROTOTYPE_BINARY_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".gif", ".webp", ".ico"})
PROTOTYPE_SUFFIXES = frozenset({".html", ".htm", ".js", ".css", ".json", ".svg"}) | PROTOTYPE_BINARY_SUFFIXES
UNSUPPORTED_PARSED_SUFFIXES = frozenset(
    {
        ".doc",
        ".docm",
        ".docx",
        ".odt",
        ".pdf",
        ".ppt",
        ".pptm",
        ".pptx",
        ".rtf",
        ".xls",
        ".xlsb",
        ".xlsm",
    }
)
TEXT_MEDIA_TYPES = {
    ".html": "text/html",
    ".htm": "text/html",
    ".md": "text/markdown",
    ".ts": "text/typescript",
    ".tsx": "text/typescript",
}
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PLACEHOLDER_TERMS = frozenset(
    {
        "todo",
        "tbd",
        "placeholder",
        "待补充",
        "待填写",
        "待定",
        "填写",
        "示例",
        "n/a",
        "na",
        "无",
    }
)
MAX_SOURCE_FILE_BYTES = 50 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 2048
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 100
MAX_XLSX_SHEETS = 64
MAX_XLSX_ROWS_PER_SHEET = 100_000
MAX_XLSX_COLUMNS_PER_SHEET = 512
MAX_XLSX_DIMENSION_CELLS = 1_000_000
MAX_XLSX_TOTAL_CELLS = 2_000_000
MAX_XLSX_CELL_TEXT_CHARS = 32_767
MAX_XLSX_TOTAL_TEXT_CHARS = 10_000_000


class SourceReadError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def html_elements(text: str) -> list[dict[str, object]]:
    """Read element attributes and inline source without execution or URL access."""
    class Elements(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.elements = []
            self.stack = []
            self.counts = defaultdict(int)

        def handle_starttag(self, tag, attrs):
            parent = self.stack[-1][1] if self.stack else ""
            self.counts[(parent, tag)] += 1
            selector = f"{tag}:nth-of-type({self.counts[(parent, tag)]})"
            if parent:
                selector = parent + " > " + selector
            self.elements.append({"tag": tag, "attributes": dict(attrs), "line": self.getpos()[0], "selector": selector, "content": ""})
            if tag not in {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}:
                self.stack.append((tag, selector, self.elements[-1]))

        def handle_data(self, data):
            if self.stack and self.stack[-1][0] in {"script", "style"}:
                self.stack[-1][2]["content"] += data

        def handle_endtag(self, tag):
            for index in range(len(self.stack) - 1, -1, -1):
                if self.stack[index][0] == tag:
                    del self.stack[index:]
                    break

    parser = Elements()
    parser.feed(text)
    parser.close()
    return parser.elements


def _limit(message: str) -> SourceReadError:
    return SourceReadError("SOURCE_LIMIT_EXCEEDED", message)


def _preflight_xlsx(path: Path) -> None:
    try:
        snapshot = path.lstat()
        if snapshot.st_size > MAX_SOURCE_FILE_BYTES:
            raise _limit("XLSX 压缩文件超过读取上限。")
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
                raise _limit("XLSX ZIP 条目数超过读取上限。")
            uncompressed = sum(item.file_size for item in entries)
            compressed = sum(item.compress_size for item in entries)
            if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise _limit("XLSX 解压后体积超过读取上限。")
            if compressed == 0 and uncompressed > 0:
                raise _limit("XLSX 压缩比无法安全验证。")
            if compressed and uncompressed / compressed > MAX_XLSX_COMPRESSION_RATIO:
                raise _limit("XLSX 压缩比超过安全上限。")
    except SourceReadError:
        raise
    except (OSError, zipfile.BadZipFile) as error:
        raise SourceReadError("SOURCE_UNREADABLE", "XLSX 无法读取。") from error


def _bounded_xlsx_rows(path: Path, *, data_only: bool = False, temporal_types: bool = False):
    _preflight_xlsx(path)
    workbook = None
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            raise _limit("XLSX Sheet 数量超过读取上限。")
        total_cells = 0
        total_text_chars = 0
        for worksheet in workbook.worksheets:
            max_row = int(worksheet.max_row or 0)
            max_column = int(worksheet.max_column or 0)
            if (
                max_row > MAX_XLSX_ROWS_PER_SHEET
                or max_column > MAX_XLSX_COLUMNS_PER_SHEET
                or max_row * max_column > MAX_XLSX_DIMENSION_CELLS
            ):
                raise _limit(f"XLSX Sheet {worksheet.title} 的声明维度超过读取上限。")
            for row_number, row in enumerate(worksheet.iter_rows(values_only=False), 1):
                total_cells += len(row)
                if total_cells > MAX_XLSX_TOTAL_CELLS:
                    raise _limit("XLSX 单元格总数超过读取上限。")
                values: list[object] = []
                for cell in row:
                    value = cell.value
                    scalar = _xlsx_scalar(value)
                    temporal_type = None
                    if temporal_types and isinstance(value, (datetime, date, time)):
                        temporal_type = openpyxl.styles.numbers.is_datetime(cell.number_format)
                        if isinstance(value, datetime):
                            if temporal_type == "date":
                                scalar = value.date().isoformat()
                            elif temporal_type == "time":
                                scalar = value.time().isoformat()
                            else:
                                temporal_type = "datetime"
                        else:
                            temporal_type = "time" if isinstance(value, time) else "date"
                    text = "" if scalar is None else str(scalar)
                    if len(text) > MAX_XLSX_CELL_TEXT_CHARS:
                        raise _limit("XLSX 单元格文本超过读取上限。")
                    total_text_chars += len(text)
                    if total_text_chars > MAX_XLSX_TOTAL_TEXT_CHARS:
                        raise _limit("XLSX 文本总量超过读取上限。")
                    values.append((scalar, temporal_type) if temporal_types else scalar)
                yield worksheet.title, row_number, values
    except SourceReadError:
        raise
    except Exception as error:
        raise SourceReadError("SOURCE_UNREADABLE", "XLSX 无法读取。") from error
    finally:
        if workbook is not None:
            workbook.close()


def _normalize(text: object) -> str:
    return " ".join(unicodedata.normalize("NFC", str(text)).split())


def _empty_xlsx_drawing(payload: bytes) -> bool:
    """Recognize only an empty spreadsheet drawing container, never its contents."""
    class EmptyDrawingTree(ET.TreeBuilder):
        def comment(self, text):
            raise ValueError("Drawing contains a comment.")

        def pi(self, target, text):
            raise ValueError("Drawing contains a processing instruction.")

        def doctype(self, name, public_id, system_id):
            raise ValueError("Drawing contains a document type declaration.")

    try:
        root = ET.fromstring(payload, parser=ET.XMLParser(target=EmptyDrawingTree()))
    except (ET.ParseError, ValueError, LookupError):
        return False
    return (
        root.tag == "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}wsDr"
        and not root.attrib
        and len(root) == 0
        and not (root.text or "").strip(" \t\r\n")
    )


def inventory_xlsx(path: Path) -> dict[str, object]:
    """Read bounded OOXML metadata and typed cells without an Office save operation."""
    inspect_source_header(path, source_role="PRIOR_SOW")
    payload = path.read_bytes()
    # The bounded reader interprets the workbook's epoch and number formats;
    # data_only reads stored caches, never evaluates formulas.
    temporal_cells, formula_cells = {}, {}
    for data_only in (False, True):
        for sheet, row_number, values in _bounded_xlsx_rows(path, data_only=data_only, temporal_types=True):
            for column, (value, temporal_type) in enumerate(values, 1):
                address = f"${openpyxl.utils.get_column_letter(column)}${row_number}"
                if not data_only and isinstance(value, str) and value.startswith("="):
                    # The read-only parser expands shared references without calculation.
                    formula_cells[(sheet, address)] = value[1:]
                if temporal_type is not None:
                    temporal_cells[(data_only, sheet, address)] = (value, temporal_type)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    from openpyxl.utils.cell import absolute_coordinate

    def part_target(base, target):
        resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join(posixpath.dirname(base), target))
        if resolved.startswith("../"):
            raise SourceReadError("SOURCE_UNREADABLE", "XLSX relationship 越界。")
        return resolved

    with zipfile.ZipFile(BytesIO(payload)) as archive:
        names = set(archive.namelist())
        empty_drawings = {
            name for name in names
            if name.startswith("xl/drawings/") and name.endswith(".xml") and "/_rels/" not in name
            and _empty_xlsx_drawing(archive.read(name))
        }
        shared = []
        if "xl/sharedStrings.xml" in names:
            shared = ["".join(node.itertext()) for node in ET.fromstring(archive.read("xl/sharedStrings.xml")).findall("s:si", ns)]

        def relationships(base):
            relpath = posixpath.join(posixpath.dirname(base), "_rels", posixpath.basename(base) + ".rels")
            if relpath not in names:
                return {}
            return {item.attrib["Id"]: part_target(base, item.attrib["Target"]) for item in ET.fromstring(archive.read(relpath)) if item.attrib.get("TargetMode") != "External"}

        def scalar(node, kind):
            if node is None or node.text is None:
                return None
            if kind == "s":
                return shared[int(node.text)]
            if kind == "b":
                return node.text == "1"
            if kind == "n":
                number = float(node.text)
                if not __import__("math").isfinite(number):
                    raise SourceReadError("SOURCE_UNREADABLE", "XLSX 数值必须有限。")
                return int(number) if number.is_integer() else number
            return node.text

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet_targets = relationships("xl/workbook.xml")
        sheets = []
        for sheet in workbook.findall("s:sheets/s:sheet", ns):
            part = sheet_targets[sheet.attrib[f"{{{rel_ns}}}id"]]
            root = ET.fromstring(archive.read(part))
            cells = []
            for cell in root.findall("s:sheetData/s:row/s:c", ns):
                kind = cell.attrib.get("t", "n")
                value = scalar(cell.find("s:v", ns), kind)
                if kind == "inlineStr":
                    value = "".join(t.text or "" for t in cell.findall("s:is//s:t", ns))
                formula = cell.find("s:f", ns)
                if value is None and formula is None:
                    continue
                address = absolute_coordinate(cell.attrib["r"])
                record = {"address": address, "cellType": kind, "value": value if formula is None else None, "formula": None if formula is None else formula.text, "cachedValue": value if formula is not None else None}
                if formula is not None and formula.attrib.get("t") == "shared":
                    record["formula"] = formula_cells[(sheet.attrib["name"], address)]
                temporal = temporal_cells.get((formula is not None, sheet.attrib["name"], address))
                if temporal is not None:
                    record["cachedValue" if formula is not None else "value"], record["temporalType"] = temporal
                cells.append(record)
            targets = relationships(part)
            tables = []
            for table_ref in root.findall("s:tableParts/s:tablePart", ns):
                table = ET.fromstring(archive.read(targets[table_ref.attrib[f"{{{rel_ns}}}id"]]))
                tables.append({"name": table.attrib["displayName"], "range": absolute_coordinate(table.attrib["ref"])})
            dimension = root.find("s:dimension", ns)
            sheets.append({"sheet": sheet.attrib["name"], "state": sheet.attrib.get("state", "visible"), "usedRange": absolute_coordinate(dimension.attrib["ref"] if dimension is not None else "A1"), "cells": cells,
                           "merges": sorted(absolute_coordinate(item.attrib["ref"]) for item in root.findall("s:mergeCells/s:mergeCell", ns)),
                           "tables": sorted(tables, key=lambda item: item["name"]),
                           "hiddenRows": [int(item.attrib["r"]) for item in root.findall("s:sheetData/s:row", ns) if item.attrib.get("hidden") == "1"],
                           "hiddenColumns": [{"min": int(item.attrib["min"]), "max": int(item.attrib["max"])} for item in root.findall("s:cols/s:col", ns) if item.attrib.get("hidden") == "1"]})
    if path.read_bytes() != payload:
        raise SourceReadError("SOURCE_CHANGED", "读取期间来源发生变化。")
    surfaces = []
    for name in sorted(names):
        kind = None
        if "vbaproject" in name.lower():
            kind = "MACRO"
        elif name.endswith(".vml"):
            kind = "VML"
        elif name.startswith("xl/drawings/") and not "/_rels/" in name:
            if name not in empty_drawings:
                kind = "DRAWING"
        elif name.startswith("xl/media/"):
            kind = "IMAGE"
        elif name.startswith("xl/embeddings/"):
            kind = "EMBEDDED_OBJECT"
        elif name.startswith(("xl/charts/", "xl/externalLinks/")) and not "/_rels/" in name:
            kind = "UNSUPPORTED"
        if kind is not None:
            surfaces.append({"part": name, "kind": kind, "coverage": "UNSUPPORTED"})
    return {"workbookSha256": hashlib.sha256(payload).hexdigest(), "sizeBytes": len(payload), "sheets": sheets, "unsupportedSurfaces": surfaces}


def _markdown_anchors(text: str) -> list[tuple[str, str, str]]:
    anchors: list[tuple[str, str, str]] = []
    headings: list[str] = []
    paragraph: list[str] = []
    paragraph_index = 0
    table_index = 0

    def section() -> str:
        return "/".join(headings) if headings else "document"

    def flush_paragraph() -> None:
        nonlocal paragraph_index
        normalized = _normalize(" ".join(paragraph))
        paragraph.clear()
        if not normalized:
            return
        paragraph_index += 1
        anchors.append(
            ("PARAGRAPH", f"section:{section()}/paragraph:{paragraph_index:04d}", normalized)
        )

    for line in text.splitlines():
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = _normalize(heading.group(2))
            headings[level - 1 :] = [title]
            anchors.append(("HEADING", f"heading:{section()}", title))
            continue
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            cells = [_normalize(cell) for cell in stripped.strip("|").split("|")]
            if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            normalized = " | ".join(cell for cell in cells if cell)
            if normalized:
                table_index += 1
                anchors.append(
                    ("TABLE_ROW", f"section:{section()}/table-row:{table_index:04d}", normalized)
                )
            continue
        if not stripped:
            flush_paragraph()
        else:
            if re.match(r"^ {0,3}\d+[.)]\s+", line):
                flush_paragraph()
            paragraph.append(stripped)
    flush_paragraph()
    return anchors


def _text_anchors(text: str) -> list[tuple[str, str, str]]:
    paragraphs = [
        normalized
        for block in re.split(r"\n\s*\n", text)
        if (normalized := _normalize(block))
    ]
    return [
        ("PARAGRAPH", f"paragraph:{index:04d}", paragraph)
        for index, paragraph in enumerate(paragraphs, 1)
    ]


def _xlsx_anchors(path: Path) -> list[tuple[str, str, str]]:
    try:
        anchors: list[tuple[str, str, str]] = []
        semantic_rows: defaultdict[str, int] = defaultdict(int)
        for sheet_title, _row_number, row in _bounded_xlsx_rows(path):
            cells = [_normalize(value) for value in row if value is not None]
            normalized = " | ".join(cell for cell in cells if cell)
            if not normalized:
                continue
            semantic_rows[sheet_title] += 1
            anchors.append(
                (
                    "SHEET_ROW",
                    f"sheet:{sheet_title}/row:{semantic_rows[sheet_title]:04d}",
                    normalized,
                )
            )
        return anchors
    except SourceReadError:
        raise
    except Exception as error:
        raise SourceReadError("SOURCE_UNREADABLE", "XLSX 无法读取。") from error


def _source_kind(path: Path, role: str) -> str:
    if role not in SOURCE_ROLES:
        raise SourceReadError("SOURCE_ROLE_INVALID", "来源角色不受支持。")
    suffix = path.suffix.casefold()
    if role in ROLE_SUFFIXES:
        if suffix not in ROLE_SUFFIXES[role]:
            raise SourceReadError("SOURCE_FORMAT_UNSUPPORTED", "来源文件格式不受支持。")
        return "MARKDOWN" if suffix == ".md" else "XLSX"
    if role == "DEMO" and suffix not in PROTOTYPE_SUFFIXES:
        raise SourceReadError("SOURCE_FORMAT_UNSUPPORTED", "来源文件格式不受支持。")
    if role == "DEMO" and suffix in PROTOTYPE_BINARY_SUFFIXES:
        return "BINARY"
    if suffix == ".xlsx":
        return "XLSX"
    if suffix in UNSUPPORTED_PARSED_SUFFIXES:
        raise SourceReadError("SOURCE_FORMAT_UNSUPPORTED", "来源文件格式不受支持。")
    return "MARKDOWN" if suffix == ".md" else "TEXT"


def inspect_source_header(path: Path, *, source_role: str) -> str:
    """Run the bounded, non-semantic source check used by the cheap prepare gate."""
    kind = _source_kind(path, source_role)
    try:
        snapshot = path.lstat()
        if not path.is_file() or path.is_symlink():
            raise OSError("not a regular file")
        if snapshot.st_size > MAX_SOURCE_FILE_BYTES:
            raise _limit("来源文件超过读取上限。")
        if kind == "XLSX":
            with path.open("rb") as stream:
                header = stream.read(4)
            if snapshot.st_size < 4 or header != b"PK\x03\x04":
                raise SourceReadError("SOURCE_FORMAT_HEADER_INVALID", "XLSX 文件头无效。")
            _preflight_xlsx(path)
        elif kind != "BINARY":
            with path.open("rb") as stream:
                prefix = stream.read(4096)
            prefix.decode("utf-8")
    except SourceReadError:
        raise
    except UnicodeDecodeError as error:
        raise SourceReadError("SOURCE_FORMAT_HEADER_INVALID", "文本文件头不是 UTF-8。") from error
    except OSError as error:
        raise SourceReadError("SOURCE_UNREADABLE", "来源文件无法读取。") from error
    return kind


def _lossless_text(value: str) -> str:
    return unicodedata.normalize("NFC", value.replace("\r\n", "\n").replace("\r", "\n"))


class _BlockBuilder:
    def __init__(self, role: str, source_id: str) -> None:
        self.role = role
        self.source_id = source_id
        self.blocks: list[dict[str, object]] = []

    def add(
        self,
        *,
        locator: str,
        raw_content: str,
        parent_id: str | None,
        disposition: str = "INCLUDED",
        dropped: tuple[str, ...] = (),
        context_ids: tuple[str, ...] | None = None,
    ) -> str:
        content = _lossless_text(raw_content)
        raw_sha256 = hashlib.sha256(raw_content.encode("utf-8")).hexdigest()
        content_sha256 = hashlib.sha256(content.encode("utf-8")).hexdigest()
        identity = hashlib.sha256(
            canonical_json_bytes(
                [self.role, locator, raw_sha256, len(self.blocks) + 1]
            )
        ).hexdigest()
        block_id = f"block-{identity[:20]}"
        contexts = context_ids if context_ids is not None else ((parent_id,) if parent_id else ())
        self.blocks.append(
            {
                "blockId": block_id,
                "sourceId": self.source_id,
                "rawSha256": raw_sha256,
                "contentSha256": content_sha256,
                "locator": locator,
                "primaryCoverageBlockId": block_id,
                "contextBlockIds": list(contexts),
                "structuralParentId": parent_id,
                "extractionDisposition": disposition,
                "droppedContentCategories": list(dropped),
                "content": content,
            }
        )
        return block_id


def _markdown_blocks(text: str, builder: _BlockBuilder) -> None:
    headings: list[tuple[int, str, str]] = []
    paragraph: list[str] = []
    paragraph_index = 0
    table_index = 0
    gap_index = 0

    def section() -> str:
        return "/".join(item[1] for item in headings) if headings else "document"

    def parent_id() -> str | None:
        return headings[-1][2] if headings else None

    def flush_paragraph() -> None:
        nonlocal paragraph_index
        if not paragraph:
            return
        paragraph_index += 1
        builder.add(
            locator=f"section:{section()}/paragraph:{paragraph_index:04d}",
            raw_content="\n".join(paragraph),
            parent_id=parent_id(),
        )
        paragraph.clear()

    for line in _lossless_text(text).split("\n"):
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = unicodedata.normalize("NFC", heading.group(2).strip())
            while headings and headings[-1][0] >= level:
                headings.pop()
            parent = headings[-1][2] if headings else None
            block_id = builder.add(
                locator=f"heading:{'/'.join([*(item[1] for item in headings), title])}",
                raw_content=line,
                parent_id=parent,
            )
            headings.append((level, title, block_id))
            continue
        stripped = line.strip()
        if re.match(r"^ {0,3}\d+[.)]\s+", line):
            flush_paragraph()
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            table_index += 1
            disposition = (
                "CONTEXT_ONLY"
                if all(
                    re.fullmatch(r":?-{3,}:?", cell.strip())
                    for cell in stripped.strip("|").split("|")
                )
                else "INCLUDED"
            )
            builder.add(
                locator=f"section:{section()}/table-row:{table_index:04d}",
                raw_content=line,
                parent_id=parent_id(),
                disposition=disposition,
            )
            continue
        if not stripped:
            flush_paragraph()
            gap_index += 1
            builder.add(
                locator=f"section:{section()}/gap:{gap_index:04d}",
                raw_content=line,
                parent_id=parent_id(),
                disposition="DROPPED",
                dropped=("EMPTY",),
            )
        else:
            paragraph.append(line)
    flush_paragraph()


def _text_blocks(text: str, builder: _BlockBuilder) -> None:
    blocks = re.split(r"(\n\s*\n)", _lossless_text(text))
    content_index = 0
    gap_index = 0
    for block in blocks:
        if not block:
            continue
        if not block.strip():
            gap_index += 1
            builder.add(
                locator=f"gap:{gap_index:04d}",
                raw_content=block,
                parent_id=None,
                disposition="DROPPED",
                dropped=("EMPTY",),
            )
        else:
            content_index += 1
            builder.add(
                locator=f"paragraph:{content_index:04d}",
                raw_content=block,
                parent_id=None,
            )


class _LosslessHTMLParser(HTMLParser):
    def __init__(self, builder: _BlockBuilder) -> None:
        super().__init__(convert_charrefs=False)
        self.builder = builder
        self.stack: list[tuple[str, str, str]] = []
        self.counts: defaultdict[tuple[str, str], int] = defaultdict(int)

    def _parent(self) -> str | None:
        return self.stack[-1][1] if self.stack else None

    def _path(self, tag: str) -> str:
        parent_path = self.stack[-1][2] if self.stack else "html"
        key = (parent_path, tag)
        self.counts[key] += 1
        return f"{parent_path}/{tag}[{self.counts[key]}]"

    def handle_starttag(self, tag: str, attrs) -> None:
        path = self._path(tag)
        block_id = self.builder.add(
            locator=f"{path}/start",
            raw_content=self.get_starttag_text() or f"<{tag}>",
            parent_id=self._parent(),
            disposition="CONTEXT_ONLY",
        )
        self.stack.append((tag, block_id, path))

    def handle_startendtag(self, tag: str, attrs) -> None:
        path = self._path(tag)
        self.builder.add(
            locator=f"{path}/self",
            raw_content=self.get_starttag_text() or f"<{tag}/>",
            parent_id=self._parent(),
            disposition="CONTEXT_ONLY",
        )

    def handle_endtag(self, tag: str) -> None:
        current = self.stack[-1] if self.stack else None
        self.builder.add(
            locator=f"{current[2] if current else 'html'}/end:{tag}",
            raw_content=f"</{tag}>",
            parent_id=current[1] if current else None,
            disposition="CONTEXT_ONLY",
        )
        if current is not None:
            self.stack.pop()

    def handle_data(self, data: str) -> None:
        parent = self.stack[-1] if self.stack else None
        tag = parent[0].casefold() if parent else ""
        dropped = ("SCRIPT",) if tag == "script" else (("STYLE",) if tag == "style" else ())
        if not data.strip() and not dropped:
            dropped = ("EMPTY",)
        self.counts[((parent[2] if parent else "html"), "text")] += 1
        index = self.counts[((parent[2] if parent else "html"), "text")]
        self.builder.add(
            locator=f"{parent[2] if parent else 'html'}/text[{index}]",
            raw_content=data,
            parent_id=parent[1] if parent else None,
            disposition="DROPPED" if dropped else "INCLUDED",
            dropped=dropped,
        )

    def handle_comment(self, data: str) -> None:
        self.builder.add(
            locator=f"{self.stack[-1][2] if self.stack else 'html'}/comment:{len(self.builder.blocks)+1}",
            raw_content=f"<!--{data}-->",
            parent_id=self._parent(),
            disposition="DROPPED",
            dropped=("COMMENT",),
        )

    def handle_entityref(self, name: str) -> None:
        self.handle_data(f"&{name};")

    def handle_charref(self, name: str) -> None:
        self.handle_data(f"&#{name};")

    def handle_decl(self, decl: str) -> None:
        self.builder.add(
            locator="html/declaration",
            raw_content=f"<!{decl}>",
            parent_id=None,
            disposition="CONTEXT_ONLY",
        )


def _xlsx_scalar(value: object) -> object:
    if isinstance(value, DataTableFormula):
        raise SourceReadError("SOURCE_UNREADABLE", "XLSX 数据表公式没有可读取的原始公式文本。")
    if isinstance(value, ArrayFormula):
        if not isinstance(value.text, str) or not value.text.startswith("="):
            raise SourceReadError("SOURCE_UNREADABLE", "XLSX 数组公式缺少可读取的原始公式。")
        return value.text
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _xlsx_blocks(path: Path, builder: _BlockBuilder) -> None:
    try:
        sheet_ids: dict[str, str] = {}
        for sheet_title, row_number, values in _bounded_xlsx_rows(path):
            if sheet_title not in sheet_ids:
                sheet_ids[sheet_title] = builder.add(
                    locator=f"sheet:{sheet_title}",
                    raw_content=sheet_title,
                    parent_id=None,
                    disposition="CONTEXT_ONLY",
                )
            sheet_id = sheet_ids[sheet_title]
            if all(value is None or value == "" for value in values):
                continue
            text = " | ".join(str(value) for value in values if value not in {None, ""})
            builder.add(
                locator=f"sheet:{sheet_title}/row:{row_number:06d}",
                raw_content=text,
                parent_id=sheet_id,
                dropped=("STYLE",),
            )
        # Empty sheets intentionally produce no semantic block and are handled by SOURCE_BLANK.
    except SourceReadError:
        raise
    except Exception as error:
        raise SourceReadError("SOURCE_UNREADABLE", "XLSX 无法读取。") from error


def extract_source_blocks(
    path: Path,
    *,
    source_role: str,
    parser_version: str,
) -> SourceDocument:
    """Extract structure-preserving blocks while binding the untouched source bytes."""
    kind = inspect_source_header(path, source_role=source_role)
    raw_payload = path.read_bytes()
    source_id = re.sub(r"[^A-Za-z0-9._:-]+", "-", path.stem).strip("-") or "source"
    builder = _BlockBuilder(source_role, source_id)
    if kind == "BINARY":
        return SourceDocument(source_id=source_id, role=source_role,
                              raw_sha256=hashlib.sha256(raw_payload).hexdigest(),
                              parser_id="opaque-binary", parser_version=parser_version, blocks=())
    try:
        if kind == "XLSX":
            parser_id = "xlsx-rows"
            _xlsx_blocks(path, builder)
        else:
            text = _read_utf8_text(path)
            suffix = path.suffix.casefold()
            if suffix in {".html", ".htm"}:
                parser_id = "html-events"
                parser = _LosslessHTMLParser(builder)
                parser.feed(_lossless_text(text))
                parser.close()
            elif kind == "MARKDOWN":
                parser_id = "markdown-blocks"
                _markdown_blocks(text, builder)
            else:
                parser_id = "text-blocks"
                _text_blocks(text, builder)
    except SourceReadError:
        raise
    except (OSError, UnicodeDecodeError, UnicodeError) as error:
        raise SourceReadError("SOURCE_UNREADABLE", "来源文件无法读取。") from error

    semantic = [
        str(block["content"])
        for block in builder.blocks
        if block["extractionDisposition"] == "INCLUDED" and str(block["content"]).strip()
    ]
    if not semantic:
        raise SourceReadError("SOURCE_BLANK", "来源文件没有有效内容。")
    if _is_placeholder_only(semantic):
        raise SourceReadError("SOURCE_PLACEHOLDER_ONLY", "来源文件只有占位内容。")
    if _looks_like_unrelated_sample(semantic):
        raise SourceReadError("SOURCE_IRRELEVANT_SAMPLE", "来源文件是未填写的无关样例。")
    return SourceDocument(
        source_id=source_id,
        role=source_role,
        raw_sha256=hashlib.sha256(raw_payload).hexdigest(),
        parser_id=parser_id,
        parser_version=parser_version,
        blocks=tuple(builder.blocks),
    )


def source_media_type(path: Path, role: str) -> str:
    kind = _source_kind(path, role)
    if kind == "XLSX":
        return XLSX_MEDIA_TYPE
    return TEXT_MEDIA_TYPES.get(path.suffix.casefold(), "text/plain")


def _read_utf8_text(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    if any(
        unicodedata.category(character) == "Cc" and character not in "\n\r\t"
        for character in text
    ):
        raise SourceReadError("SOURCE_UNREADABLE", "文本来源包含二进制控制字符。")
    return text


def _is_placeholder_only(texts: list[str]) -> bool:
    tokens: list[str] = []
    for text in texts:
        simplified = re.sub(r"[^\w\u3400-\u9fff/]+", " ", text.casefold())
        tokens.extend(token for token in simplified.split() if token)
    return bool(tokens) and all(token in PLACEHOLDER_TERMS for token in tokens)


def _looks_like_unrelated_sample(texts: list[str]) -> bool:
    combined = " ".join(texts).casefold()
    has_sample_marker = any(marker in combined for marker in ("仅供参考", "示例模板", "sample template"))
    has_fill_marker = any(marker in combined for marker in ("请填写", "待填写", "replace this"))
    return has_sample_marker and has_fill_marker


def extract_document(
    path: Path,
    *,
    source_id: str,
    role: str,
) -> tuple[SourceAnchor, ...]:
    kind = _source_kind(path, role)
    if not path.is_file():
        raise SourceReadError("SOURCE_UNREADABLE", "来源文件无法读取。")
    if path.stat().st_size > MAX_SOURCE_FILE_BYTES:
        raise _limit("来源文件超过读取上限。")

    try:
        if kind in {"MARKDOWN", "TEXT"}:
            text = _read_utf8_text(path)
            raw = _markdown_anchors(text) if kind == "MARKDOWN" else _text_anchors(text)
        else:
            raw = _xlsx_anchors(path)
    except SourceReadError:
        raise
    except UnicodeDecodeError as error:
        raise SourceReadError("SOURCE_UNREADABLE", "文本来源不是有效 UTF-8。") from error
    except OSError as error:
        raise SourceReadError("SOURCE_UNREADABLE", "来源文件无法读取。") from error

    normalized = [(kind, locator, _normalize(text)) for kind, locator, text in raw]
    normalized = [item for item in normalized if item[2]]
    texts = [item[2] for item in normalized]
    if not texts:
        raise SourceReadError("SOURCE_BLANK", "来源文件没有有效内容。")
    if _is_placeholder_only(texts):
        raise SourceReadError("SOURCE_PLACEHOLDER_ONLY", "来源文件只有占位内容。")
    if _looks_like_unrelated_sample(texts):
        raise SourceReadError("SOURCE_IRRELEVANT_SAMPLE", "来源文件是未填写的无关样例。")

    identities = [
        hashlib.sha256(canonical_json_bytes([source_id, kind, text])).hexdigest()
        for kind, _, text in normalized
    ]
    totals = Counter(identities)
    occurrences: defaultdict[str, int] = defaultdict(int)
    anchors: list[SourceAnchor] = []
    for (kind, locator, text), identity in zip(normalized, identities, strict=True):
        occurrences[identity] += 1
        anchor_id = f"anchor-{identity[:16]}"
        if totals[identity] > 1:
            anchor_id += f"-{occurrences[identity]:04d}"
        anchors.append(
            SourceAnchor(
                anchor_id=anchor_id,
                source_id=source_id,
                kind=kind,  # type: ignore[arg-type]
                locator=locator,
                normalized_text=text,
                sha256=hashlib.sha256(text.encode("utf-8")).hexdigest(),
            )
        )
    return tuple(anchors)
