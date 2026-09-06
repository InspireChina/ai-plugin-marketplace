from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections.abc import Mapping, Sequence
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from models import CatalogHydration, Diagnostic, TaskStandardCatalog


TASK_STANDARD_HEADERS = (
    "序号",
    "分类",
    "工作类型ID",
    "工作类型名称",
    "计量单位",
    "标准交付物",
    "包含内容",
    "不包含内容",
    "说明",
    "新建适用",
    "新建M档人天",
    "新建完成标准",
    "调整适用",
    "调整M档人天",
    "调整完成标准",
    "接入复用适用",
    "接入复用M档人天",
    "接入复用完成标准",
    "主要计量维度",
    "S标准",
    "M标准",
    "L标准",
    "X/拆分条件",
    "模式适用说明",
    "相邻工作类型IDs",
    "相邻类型选择规则",
    "不建Task条件",
    "SIT支持资格",
    "标准版本",
    "参数状态",
    "来源版本",
)
SEMANTIC_HEADERS = TASK_STANDARD_HEADERS[1:28]
COMPACT_HEADERS = (
    "工作类型ID",
    "分类",
    "工作类型名称",
    "计量单位",
    "标准交付物",
    "适用模式",
    "相邻工作类型IDs",
    "SIT支持资格",
    "rowSemanticSha256",
)
MODE_COLUMNS = (
    ("新建", "新建适用", "新建M档人天", "新建完成标准"),
    ("调整", "调整适用", "调整M档人天", "调整完成标准"),
    ("接入复用", "接入复用适用", "接入复用M档人天", "接入复用完成标准"),
)
TOKENIZER_VERSION = "task-catalog-tokenizer-v1"
_SIT_OWNER_TYPES = frozenset({"IN-INTEGRATION", "IN-IDENTITY"})
_ASCII_TOKEN = re.compile(r"[a-z0-9]+(?:-[a-z0-9]+)*")
_HAN_RUN = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]+")


class CatalogContractError(ValueError):
    def __init__(self, reason: str, path: str = "", **details: object) -> None:
        self.diagnostic = Diagnostic(
            code="CONTRACT_UNSUPPORTED",
            message="Task 标准目录不满足冻结合同。",
            path=path,
            details={"reason": reason, **details},
        )
        super().__init__(f"CONTRACT_UNSUPPORTED:{reason}:{path}")


def _fail(reason: str, path: str = "", **details: object) -> None:
    raise CatalogContractError(reason, path, **details)


def _normalize_string(value: str) -> str:
    return unicodedata.normalize(
        "NFC", value.replace("\r\n", "\n").replace("\r", "\n").strip()
    )


def _decimal_token(value: int | float | Decimal) -> str:
    try:
        decimal = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError(f"invalid decimal: {value}") from error
    if not decimal.is_finite():
        raise ValueError(f"non-finite decimal: {value}")
    result = format(decimal, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return "0" if result in {"", "-0"} else result


def _canonical_json(value: object) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float, Decimal)):
        return _decimal_token(value)
    if isinstance(value, str):
        return json.dumps(
            _normalize_string(value), ensure_ascii=False, separators=(",", ":")
        )
    if isinstance(value, (list, tuple)):
        return "[" + ",".join(_canonical_json(item) for item in value) + "]"
    if isinstance(value, Mapping):
        pairs = sorted((str(key), item) for key, item in value.items())
        return "{" + ",".join(
            f"{_canonical_json(key)}:{_canonical_json(item)}" for key, item in pairs
        ) + "}"
    raise TypeError(f"unsupported canonical value: {type(value).__name__}")


def _sha256(value: object) -> str:
    return hashlib.sha256(_canonical_json(value).encode("utf-8")).hexdigest()


def _neighbor_values(value: object) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, str):
        values = value.split(",")
    elif isinstance(value, (list, tuple)):
        values = [str(item) for item in value]
    else:
        raise TypeError("相邻工作类型IDs 必须是逗号分隔文本、数组或空值")
    return tuple(sorted({_normalize_string(item) for item in values if item.strip()}))


def row_semantic_sha256(row: Mapping[str, object]) -> str:
    values: list[object] = []
    for header in SEMANTIC_HEADERS:
        value = row.get(header)
        values.append(_neighbor_values(value) if header == "相邻工作类型IDs" else value)
    return _sha256(values)


def _catalog_semantic_sha256(rows: Sequence[Mapping[str, object]]) -> str:
    entries = [
        {
            "workTypeId": str(row["工作类型ID"]),
            "rowSemanticSha256": str(row["rowSemanticSha256"]),
        }
        for row in sorted(rows, key=lambda item: str(item["工作类型ID"]))
    ]
    return _sha256(entries)


def _table_rows(template_path: Path) -> tuple[list[str], list[list[object]], set[str]]:
    workbook = load_workbook(template_path, data_only=False, read_only=False)
    try:
        matches = [
            (worksheet, worksheet.tables["TaskStandardTable"])
            for worksheet in workbook.worksheets
            if "TaskStandardTable" in worksheet.tables
        ]
        if len(matches) != 1:
            _fail("TASK_STANDARD_TABLE_COUNT", "TaskStandardTable", count=len(matches))
        worksheet, table = matches[0]
        if table.ref != "A4:AE92":
            _fail("TASK_STANDARD_TABLE_RANGE", "TaskStandardTable", actual=table.ref)
        headers = [column.name for column in table.tableColumns]
        min_column, min_row, max_column, max_row = range_boundaries(table.ref)
        values = [
            [
                worksheet.cell(row, column).value
                for column in range(min_column, max_column + 1)
            ]
            for row in range(min_row + 1, max_row + 1)
        ]
        table_names = {
            table_name
            for candidate in workbook.worksheets
            for table_name in candidate.tables
        }
        return headers, values, table_names
    finally:
        workbook.close()


def _nonempty_text(value: object) -> bool:
    return isinstance(value, str) and bool(value.strip())


def _validate_rows(raw_rows: list[list[object]]) -> tuple[Mapping[str, object], ...]:
    if len(raw_rows) != 88:
        _fail("ROW_COUNT", "TaskStandardTable", actual=len(raw_rows))
    rows: list[dict[str, object]] = [
        dict(zip(TASK_STANDARD_HEADERS, values, strict=True)) for values in raw_rows
    ]
    sequences = [row["序号"] for row in rows]
    if sequences != list(range(1, 89)):
        _fail("NON_CONTIGUOUS_SEQUENCE", "TaskStandardTable/序号")
    ids = [row["工作类型ID"] for row in rows]
    if not all(_nonempty_text(work_type_id) for work_type_id in ids):
        _fail("EMPTY_WORK_TYPE_ID", "TaskStandardTable/工作类型ID")
    if len(set(ids)) != 88:
        _fail("DUPLICATE_WORK_TYPE_ID", "TaskStandardTable/工作类型ID")
    known_ids = {str(work_type_id) for work_type_id in ids}

    mode_count = 0
    reuse_count = 0
    required_text = (
        "分类",
        "工作类型名称",
        "计量单位",
        "标准交付物",
        "主要计量维度",
        "S标准",
        "M标准",
        "L标准",
        "X/拆分条件",
        "模式适用说明",
        "相邻类型选择规则",
        "不建Task条件",
        "标准版本",
        "参数状态",
        "来源版本",
    )
    for row in rows:
        work_type_id = str(row["工作类型ID"])
        for header in required_text:
            if not _nonempty_text(row[header]):
                _fail("REQUIRED_TEXT_EMPTY", f"TaskStandardTable/{work_type_id}/{header}")
        for mode, applicable_header, effort_header, completion_header in MODE_COLUMNS:
            applicable = row[applicable_header]
            if type(applicable) is not bool:
                _fail(
                    "MODE_APPLICABILITY_NOT_BOOLEAN",
                    f"TaskStandardTable/{work_type_id}/{applicable_header}",
                )
            effort = row[effort_header]
            completion = row[completion_header]
            if applicable:
                if (
                    isinstance(effort, bool)
                    or not isinstance(effort, (int, float, Decimal))
                    or Decimal(str(effort)) < Decimal("1.0")
                ):
                    _fail(
                        "MODE_EFFORT_INVALID",
                        f"TaskStandardTable/{work_type_id}/{effort_header}",
                    )
                if not _nonempty_text(completion):
                    _fail(
                        "MODE_COMPLETION_EMPTY",
                        f"TaskStandardTable/{work_type_id}/{completion_header}",
                    )
                mode_count += 1
                reuse_count += int(mode == "接入复用")
            elif effort is not None or completion is not None:
                _fail(
                    "INAPPLICABLE_MODE_HAS_VALUES",
                    f"TaskStandardTable/{work_type_id}/{mode}",
                )

        raw_neighbors = row["相邻工作类型IDs"]
        if raw_neighbors is None:
            neighbor_parts: list[str] = []
        elif isinstance(raw_neighbors, str):
            neighbor_parts = [part.strip() for part in raw_neighbors.split(",") if part.strip()]
        else:
            _fail(
                "NEIGHBOR_FORMAT",
                f"TaskStandardTable/{work_type_id}/相邻工作类型IDs",
            )
        if len(neighbor_parts) != len(set(neighbor_parts)):
            _fail(
                "DUPLICATE_NEIGHBOR",
                f"TaskStandardTable/{work_type_id}/相邻工作类型IDs",
            )
        if work_type_id in neighbor_parts:
            _fail(
                "SELF_NEIGHBOR",
                f"TaskStandardTable/{work_type_id}/相邻工作类型IDs",
            )
        unknown = sorted(set(neighbor_parts) - known_ids)
        if unknown:
            _fail(
                "UNKNOWN_NEIGHBOR",
                f"TaskStandardTable/{work_type_id}/相邻工作类型IDs",
                unknown=unknown,
            )
        row["相邻工作类型IDs"] = tuple(sorted(neighbor_parts))

        eligibility = row["SIT支持资格"]
        if eligibility not in {"NONE", "PER_INTEGRATION"}:
            _fail(
                "INVALID_SIT_ELIGIBILITY",
                f"TaskStandardTable/{work_type_id}/SIT支持资格",
            )
        if eligibility == "PER_INTEGRATION" and work_type_id not in _SIT_OWNER_TYPES:
            _fail(
                "INVALID_SIT_OWNER_TYPE",
                f"TaskStandardTable/{work_type_id}/SIT支持资格",
            )

    if mode_count != 191:
        _fail("MODE_COUNT", "TaskStandardTable", actual=mode_count)
    if reuse_count != 26:
        _fail("REUSE_TYPE_COUNT", "TaskStandardTable", actual=reuse_count)
    if {
        str(row["工作类型ID"])
        for row in rows
        if row["SIT支持资格"] == "PER_INTEGRATION"
    } != _SIT_OWNER_TYPES:
        _fail("SIT_OWNER_SET", "TaskStandardTable/SIT支持资格")

    for row in rows:
        row["rowSemanticSha256"] = row_semantic_sha256(row)
    return tuple(rows)


def catalog(template_path: Path) -> TaskStandardCatalog:
    path = Path(template_path)
    headers, raw_rows, table_names = _table_rows(path)
    if tuple(headers) != TASK_STANDARD_HEADERS:
        _fail("HEADER_CONTRACT", "TaskStandardTable", actual=headers)
    if "BaseUnitCatalogTable" in table_names:
        _fail("LEGACY_CATALOG_PRESENT", "BaseUnitCatalogTable")
    rows = _validate_rows(raw_rows)
    by_work_type_id = {str(row["工作类型ID"]): row for row in rows}
    return TaskStandardCatalog(
        template_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        semantic_sha256=_catalog_semantic_sha256(rows),
        rows=rows,
        by_work_type_id=by_work_type_id,
    )


def compact_index(source: TaskStandardCatalog) -> tuple[Mapping[str, object], ...]:
    compact: list[Mapping[str, object]] = []
    for row in source.rows:
        modes = tuple(
            mode for mode, applicable, _, _ in MODE_COLUMNS if row[applicable] is True
        )
        item: dict[str, object] = {
            "工作类型ID": row["工作类型ID"],
            "分类": row["分类"],
            "工作类型名称": row["工作类型名称"],
            "计量单位": row["计量单位"],
            "标准交付物": row["标准交付物"],
            "适用模式": modes,
            "相邻工作类型IDs": tuple(row["相邻工作类型IDs"]),
            "SIT支持资格": row["SIT支持资格"],
            "rowSemanticSha256": row["rowSemanticSha256"],
        }
        if tuple(item) != COMPACT_HEADERS:
            raise AssertionError("compact index field order drift")
        compact.append(item)
    return tuple(compact)


def decision_catalog(source: TaskStandardCatalog) -> tuple[Mapping[str, object], ...]:
    """Project this revision's selection rules, never calculation parameters."""
    return tuple({
        'workTypeId': row['工作类型ID'], 'name': row['工作类型名称'],
        'category': row['分类'], 'unit': row['计量单位'],
        'deliverable': row['标准交付物'], 'includes': row['包含内容'],
        'excludes': row['不包含内容'], 'measurement': row['主要计量维度'],
        'modes': [mode for mode, applicable, _, _ in MODE_COLUMNS if row[applicable]],
        'modeRules': {mode: row[completion] for mode, applicable, _, completion in MODE_COLUMNS if row[applicable]},
        'complexityRules': {key: row[key+'标准'] for key in ('S', 'M', 'L')},
        'splitRule': row['X/拆分条件'], 'modeSelection': row['模式适用说明'],
        'neighbors': list(row['相邻工作类型IDs']), 'neighborRule': row['相邻类型选择规则'],
        'noTaskRule': row['不建Task条件'], 'sitEligibility': row['SIT支持资格'],
        'rowSemanticSha256': row['rowSemanticSha256'],
    } for row in sorted(source.rows, key=lambda row: row['序号']))


def _tokenize(value: str) -> tuple[str, ...]:
    normalized = _normalize_string(value).casefold()
    tokens = set(_ASCII_TOKEN.findall(normalized))
    for run in _HAN_RUN.findall(normalized):
        tokens.update(run)
        tokens.update(run[index : index + 1] for index in range(len(run)))
        tokens.update(run[index : index + 2] for index in range(max(0, len(run) - 1)))
    return tuple(sorted(token for token in tokens if token))


def hydrate(
    source: TaskStandardCatalog,
    selected_work_type_ids: Sequence[str],
    query: str,
    *,
    excluded_categories: Sequence[str] = (),
) -> CatalogHydration:
    selected = tuple(dict.fromkeys(str(item) for item in selected_work_type_ids))
    if not selected:
        _fail("HYDRATION_SELECTION_EMPTY", "hydrate/selectedWorkTypeIds")
    unknown = sorted(set(selected) - set(source.by_work_type_id))
    if unknown:
        _fail("HYDRATION_SELECTION_UNKNOWN", "hydrate/selectedWorkTypeIds", unknown=unknown)
    query_tokens = _tokenize(query)
    if not query_tokens:
        _fail("HYDRATION_QUERY_EMPTY", "hydrate/query")

    neighbor_ids = tuple(
        sorted(
            {
                str(neighbor)
                for work_type_id in selected
                for neighbor in source.by_work_type_id[work_type_id]["相邻工作类型IDs"]
                if neighbor not in selected
            }
        )
    )
    selected_categories = {
        str(source.by_work_type_id[work_type_id]["分类"]) for work_type_id in selected
    }
    blocked_categories = selected_categories | {str(item) for item in excluded_categories}
    query_set = set(query_tokens)
    weighted_fields = (("工作类型名称", 5), ("标准交付物", 3), ("计量单位", 2))
    candidates: list[dict[str, object]] = []
    for row in source.rows:
        work_type_id = str(row["工作类型ID"])
        if work_type_id in selected or str(row["分类"]) in blocked_categories:
            continue
        field_scores: dict[str, int] = {}
        score = 0
        for field, weight in weighted_fields:
            matches = len(query_set & set(_tokenize(str(row[field]))))
            field_scores[field] = matches * weight
            score += field_scores[field]
        candidates.append(
            {
                "workTypeId": work_type_id,
                "category": row["分类"],
                "score": score,
                "fieldScores": field_scores,
            }
        )
    candidates.sort(key=lambda item: (-int(item["score"]), str(item["workTypeId"])))
    if not candidates:
        _fail("HYDRATION_CHALLENGER_UNAVAILABLE", "hydrate/query")
    candidate_scores = tuple(
        {**item, "order": order} for order, item in enumerate(candidates, 1)
    )
    challenger_ids = (str(candidate_scores[0]["workTypeId"]),)
    hydration_order = tuple(dict.fromkeys((*selected, *neighbor_ids, *challenger_ids)))
    rows = tuple(source.by_work_type_id[work_type_id] for work_type_id in hydration_order)
    evidence: Mapping[str, object] = {
        "tokenizerVersion": TOKENIZER_VERSION,
        "query": _normalize_string(query),
        "queryTokens": query_tokens,
        "excludedCategories": tuple(sorted(blocked_categories)),
        "candidateScores": candidate_scores,
        "hydrationOrder": hydration_order,
    }
    return CatalogHydration(
        requested_work_type_ids=selected,
        query_terms=query_tokens,
        candidate_work_type_ids=tuple(
            str(item["workTypeId"]) for item in candidate_scores
        ),
        neighbor_work_type_ids=neighbor_ids,
        challenger_work_type_ids=challenger_ids,
        rows=rows,
        semantic_sha256=source.task_catalog_semantic_sha256,
        evidence=evidence,
    )
