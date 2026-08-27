from __future__ import annotations

import argparse
import json
import sys
from io import BytesIO
from pathlib import Path
from typing import Any


PLUGIN_VERSION = "0.1.0"
SOW_STANDARD_VERSION = "1.3"
MANAGED_DIRECTORIES = (
    ".ai-sow/templates",
    ".ai-sow/inputs",
    ".ai-sow/work",
    ".ai-sow/reviews",
    ".ai-sow/data",
    ".ai-sow/validation",
    ".ai-sow/outputs",
)
PROJECT_PATH = ".ai-sow/project.json"
TEMPLATE_PATH = ".ai-sow/templates/sow-template.xlsx"

# Windows 控制台默认使用本地代码页（如 cp936），会把中文结构化输出写成非 UTF-8 字节。
# 调用方按 UTF-8 读取 stdout/stderr，这里显式固定编码，与 POSIX 行为保持一致。
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

PLUGIN_ROOT = Path(__file__).resolve().parents[3]
if str(PLUGIN_ROOT) not in sys.path:
    sys.path.insert(0, str(PLUGIN_ROOT))

from runtime.project_io import (
    WINDOWS_MAX_PATH,
    ProjectFiles,
    ProjectIOError,
    managed_path_budget,
)

# 本插件已知最深的受管相对路径：staging 前缀 + SOW package 目录 + 最长的 sources 叶子文件。
# 路径布局属于本 Skill 的知识，不放进共享 runtime。
DEEPEST_MANAGED_RELATIVE_PATH = (
    ".ai-sow/.stage-0123456789ab"
    "/outputs/sow-sha256-" + "0" * 64
    + "/sources/data/analyze-requirement/requirements.json"
)


class BlockedError(ValueError):
    """A setup condition that requires user or environment action."""


def emit(outcome: str, summary: str, **details: Any) -> None:
    print(json.dumps({"outcome": outcome, "summary": summary, **details}, ensure_ascii=False))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Initialize a minimal AI SOW project")
    parser.add_argument("--project-root", required=True, type=Path)
    parser.add_argument("--project-id", required=True)
    parser.add_argument("--name", required=True)
    return parser.parse_args()


def project_bytes(project: dict[str, object]) -> bytes:
    return (json.dumps(project, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def verify_template_round_trip(payload: bytes, openpyxl: Any) -> None:
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    try:
        tables = {
            name: worksheet.tables[name]
            for worksheet in workbook.worksheets
            for name in worksheet.tables
        }
        required_tables = {"BaseUnitCatalogTable", "ProjectParameterTable"}
        if not required_tables.issubset(tables):
            raise BlockedError("项目模板缺少基础单元或项目参数 Table")
        catalog = tables["BaseUnitCatalogTable"]
        _, min_row, _, max_row = openpyxl.utils.range_boundaries(catalog.ref)
        if max_row - min_row != 37:
            raise BlockedError("项目模板的基础单元目录必须包含 37 项")
        saved = BytesIO()
        workbook.save(saved)
    finally:
        workbook.close()
    reopened = openpyxl.load_workbook(BytesIO(saved.getvalue()), data_only=False)
    reopened.close()


def validate_project(project: object, validator: Any) -> dict[str, object]:
    errors = sorted(validator.iter_errors(project), key=lambda item: list(item.path))
    if errors:
        raise BlockedError(
            "项目元数据不符合 Project Schema："
            + "；".join(error.message for error in errors)
        )
    assert isinstance(project, dict)
    return project


def verify_existing_project(
    files: ProjectFiles,
    requested: dict[str, object],
    validator: Any,
    openpyxl: Any,
) -> None:
    existing = validate_project(files.read_json(PROJECT_PATH), validator)
    if existing != requested:
        raise BlockedError("已登记项目身份或版本与本次请求不一致")
    for relative_path in MANAGED_DIRECTORIES:
        try:
            files.resolve(relative_path, expect="dir")
        except ProjectIOError as error:
            raise BlockedError(f"现有 AI SOW 项目不完整：{relative_path}") from error
    try:
        current_template = files.read_bytes(TEMPLATE_PATH)
    except ProjectIOError as error:
        raise BlockedError(f"现有 AI SOW 项目不完整：{TEMPLATE_PATH}") from error
    try:
        verify_template_round_trip(current_template, openpyxl)
    except BlockedError:
        raise
    except Exception as error:
        raise BlockedError("现有项目模板无法完成 XLSX round-trip") from error


def initialize_fresh_project(
    files: ProjectFiles,
    project: dict[str, object],
    template_bytes: bytes,
) -> None:
    try:
        files.resolve(".ai-sow", expect="any")
    except ProjectIOError as error:
        if error.code != "PROJECT_PATH_MISSING":
            raise
    else:
        raise BlockedError("目标包含未登记或不完整的 .ai-sow 受管内容")

    for relative_path in MANAGED_DIRECTORIES:
        files.ensure_dir(relative_path)
    files.publish_new(TEMPLATE_PATH, template_bytes)
    files.publish_new(PROJECT_PATH, project_bytes(project))


def main() -> int:
    args = parse_args()
    try:
        import openpyxl
        from jsonschema import Draft202012Validator
    except ImportError as error:
        emit(
            "NEEDS_INPUT",
            f"缺少 Python 依赖：{error.name}",
            nextStep="重新调用 setup；setup 会自动修复插件隔离环境，用户无需执行 uv 命令。",
        )
        return 2

    skill_root = Path(__file__).resolve().parents[1]
    asset_path = skill_root / "assets" / "sow-template.xlsx"
    schema_path = skill_root / "contracts" / "project.schema.json"

    budget = managed_path_budget(args.project_root)
    required = len(DEEPEST_MANAGED_RELATIVE_PATH)
    if budget is not None and budget < required:
        emit(
            "BLOCKED",
            "项目根目录过长：未启用长路径支持时无法创建本插件的受管输出路径",
            diagnostics=[
                {
                    "code": "WINDOWS_LONG_PATH_REQUIRED",
                    "message": (
                        f"当前项目根目录还剩 {budget} 个字符，最深的受管路径需要 {required} 个。"
                        "继续初始化会在生成阶段以 WinError 206 失败。"
                    ),
                    "path": ".",
                }
            ],
            nextStep=(
                "两个可选方案，任选其一后重新调用 setup："
                f"（1）把项目移动到长度不超过 {WINDOWS_MAX_PATH - required - 1} 个字符的路径；"
                "（2）启用 Windows 长路径支持——这会修改本机系统策略且需要管理员权限，"
                "必须先向用户说明影响并获得明确同意，再运行 "
                "skills/setup/scripts/enable_long_paths.ps1 -Apply。"
            ),
        )
        return 2

    try:
        template_bytes = asset_path.read_bytes()
        verify_template_round_trip(template_bytes, openpyxl)
        schema = json.loads(schema_path.read_text(encoding="utf-8"))
        validator = Draft202012Validator(schema)
        requested = validate_project(
            {
                "projectId": args.project_id,
                "name": args.name,
                "pluginVersion": PLUGIN_VERSION,
                "sowStandardVersion": SOW_STANDARD_VERSION,
            },
            validator,
        )
        files = ProjectFiles.open(args.project_root)
        try:
            files.resolve(PROJECT_PATH, expect="file")
        except ProjectIOError as error:
            if error.code != "PROJECT_PATH_MISSING":
                raise
            initialize_fresh_project(files, requested, template_bytes)
        else:
            verify_existing_project(files, requested, validator, openpyxl)

        emit(
            "OK",
            "AI SOW 项目外壳已通过验证",
            outputs=[PROJECT_PATH, TEMPLATE_PATH],
            nextStep="显式调用 analyze-requirement。",
        )
        return 0
    except (BlockedError, ProjectIOError) as error:
        emit("BLOCKED", str(error))
        return 2
    except Exception as error:
        emit("ERROR", str(error))
        return 3


if __name__ == "__main__":
    sys.exit(main())
